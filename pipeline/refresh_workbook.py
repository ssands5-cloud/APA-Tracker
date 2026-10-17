"""Validate the generated workbook in a new hidden Excel instance, never the user's session."""
from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

CAPTAIN_SHEETS = ("Match Day", "Command Center", "War Room", "Lineup Lab", "Scouting Cards", "Captain Packet")


def preflight():
    """Fail before login/scraping if the required Windows validators cannot run."""
    import sys
    if sys.platform != "win32":
        raise RuntimeError("complete refresh requires Windows and desktop Excel for workbook recalculation")
    try:
        import pythoncom
        import pywintypes
        import win32com.client.dynamic
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RuntimeError("refresh dependencies missing; install requirements-refresh.txt and Playwright Chromium") from exc
    pywintypes.IID("Excel.Application")


def validate_structure(path: Path):
    # Never save through openpyxl: that can discard the dynamic print area.
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        missing = set(CAPTAIN_SHEETS) - set(book.sheetnames)
        if missing:
            raise RuntimeError("workbook is missing captain sheets: " + ", ".join(sorted(missing)))
        if book.vba_archive or book._external_links:
            raise RuntimeError("unexpected workbook macros or external links")
        return {"captain_sheets": list(CAPTAIN_SHEETS)}
    finally:
        book.close()


def recalculate_and_validate(path: Path):
    import pythoncom
    import win32com.client.dynamic
    import pywintypes

    validate_structure(path)
    excel = book = None
    pythoncom.CoInitialize()
    try:
        instance = pythoncom.CoCreateInstance(
            pywintypes.IID("Excel.Application"), None,
            pythoncom.CLSCTX_LOCAL_SERVER, pythoncom.IID_IDispatch,
        )
        excel = win32com.client.dynamic.Dispatch(instance)
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3
        book = excel.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, ReadOnly=False)
        if book.ReadOnly:
            raise RuntimeError("generated workbook opened read-only; recalculation could not be saved")
        excel.CalculateFullRebuild()
        errors = {}
        for name in CAPTAIN_SHEETS:
            sheet = book.Worksheets(name)
            try:
                cells = sheet.UsedRange.SpecialCells(-4123, 16)
            except pywintypes.com_error as exc:
                # Excel's documented no-cells-found case, not a blanket COM exemption.
                if not exc.excepinfo or exc.excepinfo[5] != -2146827284:
                    raise
            else:
                errors[name] = str(cells.Address)
        if errors:
            raise RuntimeError("captain-sheet formula errors: " + repr(errors))
        if book.LinkSources(1):
            raise RuntimeError("workbook has external Excel links")
        book.Save()
        return {"native_excel_recalculated": True, "captain_sheet_formula_errors": 0,
                "captain_sheets": list(CAPTAIN_SHEETS)}
    finally:
        if book is not None:
            book.Close(SaveChanges=False)
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()


def validate_html(path: Path):
    from playwright.sync_api import sync_playwright
    errors = []
    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        try:
            page = browser.new_page(viewport={"width": 390, "height": 844})
            page.on("pageerror", lambda exc: errors.append(type(exc).__name__ + ": " + str(exc)))
            page.goto(Path(path).resolve().as_uri(), timeout=120000)
            page.wait_for_load_state("load")
            if not page.locator("#start-here").count():
                raise RuntimeError("Ultimate Coach HTML did not initialize its start panel")
            if errors:
                raise RuntimeError("Ultimate Coach browser errors: " + repr(errors))
            if page.evaluate("document.documentElement.scrollWidth>innerWidth"):
                raise RuntimeError("Ultimate Coach HTML has page-level phone overflow")
            return {"browser_started": True, "browser_page_errors": 0, "phone_width": 390}
        finally:
            browser.close()
