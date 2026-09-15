"""Tests for ui/export_excel_team_strength.py."""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.team_strength import build_report
from ui.export_excel_team_strength import (
    MATCHES_SHEET,
    PLAYERS_SHEET,
    SUMMARY_SHEET,
    build_workbook,
    write_workbook,
)


def player(external_id, name="P", won=0, played=0, skill=5, player_id=None):
    return {
        "player_id": player_id or hash(external_id) % 1000,
        "player_external_id": external_id, "player_name": name,
        "skill_level": skill, "matches_won": won, "matches_played": played,
    }


def match(match_id, points_for, points_against, is_home=True, **kwargs):
    row = {"match_id": match_id, "points_for": points_for, "points_against": points_against, "is_home": is_home}
    row.update(kwargs)
    return row


class TestBuildWorkbook:
    def test_all_three_sheets_present(self):
        report = build_report(
            "T1", "Chalk It Up", "Fall 2026",
            [player("A", won=5, played=10)], [match("M1", 8, 2)],
        )
        workbook = build_workbook(report)
        assert SUMMARY_SHEET in workbook.sheetnames
        assert PLAYERS_SHEET in workbook.sheetnames
        assert MATCHES_SHEET in workbook.sheetnames

    def test_players_sheet_row_count_matches_roster(self):
        report = build_report(
            "T1", "Chalk It Up", "Fall 2026",
            [player("A", won=5, played=10), player("B")], [],
        )
        sheet = build_workbook(report)[PLAYERS_SHEET]
        assert sheet.max_row == 3  # header + 2 players

    def test_fifth_depth_row_is_flagged(self):
        players = [
            player("A", won=10, played=10), player("B", won=8, played=10),
            player("C", won=6, played=10), player("D", won=4, played=10),
            player("E", won=2, played=10),
        ]
        report = build_report("T1", "Chalk It Up", "Fall 2026", players, [])
        sheet = build_workbook(report)[PLAYERS_SHEET]
        headers = [c.value for c in sheet[1]]
        flag_col = headers.index("Is Fifth Depth Row") + 1
        id_col = headers.index("Player External ID") + 1
        flags = {
            row[id_col - 1].value: row[flag_col - 1].value
            for row in sheet.iter_rows(min_row=2)
        }
        assert flags["E"] is True
        assert flags["A"] is False

    def test_matches_sheet_carries_every_eligible_match(self):
        report = build_report(
            "T1", "Chalk It Up", "Fall 2026", [],
            [match("M1", 8, 2), match("M2", 6, 4)],
        )
        sheet = build_workbook(report)[MATCHES_SHEET]
        assert sheet.max_row == 3  # header + 2 matches

    def test_write_workbook_produces_a_real_readable_file(self, tmp_path):
        report = build_report("T1", "Chalk It Up", "Fall 2026", [], [])
        out_path = write_workbook(report, tmp_path / "team_strength.xlsx")

        assert out_path.exists()
        reopened = load_workbook(out_path)
        assert SUMMARY_SHEET in reopened.sheetnames
