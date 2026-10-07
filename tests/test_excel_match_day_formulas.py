"""End-to-end Match Day / Match Night / Coach Dashboard formula results,
computed by tests/excel_formula_eval.py (no Excel, COM, pywin32 or macros).

These run the workbook's real formulas against its real lookup sheets, so a
wrong column, a broken key, or a blank-reads-as-0 cell fails here."""

from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from analytics.ultimate_coach_match_day import build_match_day_section
from tests.excel_formula_eval import Workbook, serial
from tests.test_export_excel_ultimate_coach import FALCONS_8, SHARKS_8, _payload, _team_match
from ui.export_excel_ultimate_coach import write_workbook

MD = "Match Day"


def _book(tmp_path, team_matches, *, viewer="1001", payload=None, card=None):
    payload = payload or _payload()
    payload["match_day"] = build_match_day_section(team_matches, payload["players"])
    path = write_workbook(payload, tmp_path / "uc.xlsx", viewer_member_external_id=viewer, viewer_card_number=card)
    wb = load_workbook(path)
    return Workbook(wb), wb


def _find_row(wb, predicate, sheet=MD):
    for row in wb[sheet].iter_rows():
        if predicate(row[0].value):
            return row[0].row
    raise AssertionError("row not found")


def _layout(wb):
    status = _find_row(wb, lambda v: isinstance(v, str) and v.startswith('=IF(B6="","Step 1'))
    header = _find_row(wb, lambda v: v == "# · Date")
    matchup = _find_row(wb, lambda v: v == "Matchup preview")
    return {"status": status, "slot1": header + 1, "matchup": matchup, "roster1": matchup + 5}


def _fixture_row(book, row):
    return {col: book.display(MD, f"{col}{row}") for col in "ABCDEFGHIJK"}


def _choose(book, *, day=None, team=None, fmt=None, compare=None, viewer=None):
    if viewer is not None:
        book.set(MD, "B7", viewer)
    if day is not None:
        book.set(MD, "B6", serial(day) if isinstance(day, date) else day)
    if team is not None:
        book.set(MD, "B8", team)
    if fmt is not None:
        book.set(MD, "B9", fmt)
    if compare is not None:
        book.set(MD, "B10", compare)


def test_single_fixture_is_found_listed_and_compared(tmp_path):
    book, wb = _book(tmp_path, [_team_match()], card="80000001")
    rows = _layout(wb)
    assert book.display(MD, "B7") == "Ann Archer (APA record ID 1001)"
    assert book.display(MD, "B13") == "Ann Archer (APA record ID 1001)"
    assert book.display(MD, "B14").startswith("2 verified players share this name")
    assert "Card #80000001" in book.display(MD, "B15")
    assert book.display(MD, "B16") == "1 current team scope (each division listed separately)."
    assert book.display(MD, "B18") == SHARKS_8
    assert book.display(MD, "B19") == ""

    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    assert book.display(MD, "C18") == "1 fixture(s) on this date"
    assert book.display(MD, f"A{rows['status']}") == (
        f"1 fixture found for {SHARKS_8} on Sunday, Oct 11, 2026 (8-Ball & 9-Ball). "
    )
    first = _fixture_row(book, rows["slot1"])
    assert first == {
        "A": "1 · Sun Oct 11, 2026", "B": "11:00 AM MDT", "C": "Home", "D": "8-Ball Open", "E": "Falcons",
        "F": "Spring 2026", "G": "No data", "H": "UNPLAYED", "I": "Not recorded",
        "J": "Current roster captured", "K": "2026-10-11T11:00:00-06:00",
    }
    assert set(_fixture_row(book, rows["slot1"] + 1).values()) == {""}  # empty slot is blank, never 0

    matchup = rows["matchup"]
    assert book.display(MD, f"A{matchup + 1}") == f"Fixture #1: {SHARKS_8} (Home) vs Falcons · Sun Oct 11, 2026 · 11:00 AM MDT"
    assert book.display(MD, f"A{matchup + 2}") == f"Opponent roster scope: {FALCONS_8}"
    our = [book.display(MD, f"{c}{rows['roster1']}") for c in "ABCD"]
    assert our == ["Bea Baker", "1002", "5*", "2-4"]
    assert [book.display(MD, f"{c}{rows['roster1'] + 1}") for c in "ABCD"] == ["Ann Archer", "1001", "4", "10-5"]
    assert [book.display(MD, f"{c}{rows['roster1']}") for c in "GHIJ"] == ["Cam Cole", "1003", "3", "3-2"]
    assert book.display(MD, f"G{rows['roster1'] + 1}") == ""


def test_utc_evening_fixture_is_on_the_denver_date_not_the_utc_date(tmp_path):
    book, wb = _book(tmp_path, [_team_match(match_date="2026-08-30T01:00:00Z")])
    rows = _layout(wb)
    _choose(book, day=date(2026, 8, 30), team=SHARKS_8)
    assert book.display(MD, f"A{rows['status']}").startswith("No scheduled match found")
    _choose(book, day=date(2026, 8, 29))
    assert book.display(MD, f"B{rows['slot1']}") == "7:00 PM MDT"
    assert book.display(MD, f"A{rows['slot1']}") == "1 · Sat Aug 29, 2026"
    assert book.display(MD, f"K{rows['slot1']}") == "2026-08-30T01:00:00Z"


def test_bye_says_no_opponent_and_offers_no_opponent_roster(tmp_path):
    book, wb = _book(tmp_path, [_team_match(is_bye=True, away_team_id="13082714", away_team_name="BYE")])
    rows = _layout(wb)
    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    fixture = _fixture_row(book, rows["slot1"])
    assert fixture["E"] == "Bye — no opponent"
    assert fixture["J"] == "Not applicable (bye)"
    assert 0 not in fixture.values()
    assert book.display(MD, f"A{rows['matchup'] + 2}") == "Opponent roster: Not applicable (bye)."
    assert book.display(MD, f"G{rows['matchup'] + 3}") == "OPPONENT — Bye — no opponent · Not applicable (bye)"
    assert [book.display(MD, f"{c}{rows['roster1']}") for c in "GHIJ"] == ["", "", "", ""]


def test_missing_opponent_name_status_and_venue_show_no_data_and_real_zero_scores_show(tmp_path):
    book, wb = _book(tmp_path, [_team_match(away_team_name="", status=None, is_scored=True, home_score=0, away_score=3)])
    rows = _layout(wb)
    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    fixture = _fixture_row(book, rows["slot1"])
    assert (fixture["E"], fixture["G"], fixture["H"], fixture["I"]) == ("No data", "No data", "No data", "0 – 3")


def test_every_fixture_on_a_date_is_listed_in_kickoff_order_and_compare_picks_the_chosen_one(tmp_path):
    book, wb = _book(tmp_path, [
        _team_match(match_id=1, match_external_id="1", match_date="2026-10-11T19:00:00-06:00"),
        _team_match(match_id=2, match_external_id="2", match_date="2026-10-11T11:00:00-06:00",
                    format="NINE", format_raw="9-Ball Open"),
    ])
    rows = _layout(wb)
    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    status = book.display(MD, f"A{rows['status']}")
    assert status.startswith("2 fixtures found") and "none is picked automatically" in status
    assert book.display(MD, f"B{rows['slot1']}") == "11:00 AM MDT"
    assert book.display(MD, f"B{rows['slot1'] + 1}") == "7:00 PM MDT"
    _choose(book, compare=2)
    assert "7:00 PM MDT" in book.display(MD, f"A{rows['matchup'] + 1}")
    _choose(book, compare=3)
    assert book.display(MD, f"A{rows['matchup'] + 1}") == "Fixture #3 is not in the list above — choose 1 to 2 in step 5."
    # Format filter: one recorded format only.
    _choose(book, compare=1, fmt="9-Ball Open")
    assert book.display(MD, f"A{rows['status']}").startswith("1 fixture found")
    _choose(book, fmt="Masters")
    assert book.display(MD, f"A{rows['status']}").startswith("No scheduled match found")


def test_typed_text_date_and_numeric_record_id_are_both_accepted(tmp_path):
    book, wb = _book(tmp_path, [_team_match()], viewer=None)
    rows = _layout(wb)
    assert book.display(MD, f"A{rows['status']}") == "Step 2: enter your APA record ID." or book.display(MD, "B6") is None
    _choose(book, viewer=1001.0, day="10/11/2026", team=SHARKS_8)
    assert book.display(MD, "B13") == "Ann Archer (APA record ID 1001)"
    assert book.display(MD, f"A{rows['status']}").startswith("1 fixture found")
    _choose(book, day="not a date")
    assert book.display(MD, f"A{rows['status']}") == "That date was not recognized — type it like 10/11/2026."


def test_unknown_record_id_and_someone_elses_team_are_refused_not_guessed(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    rows = _layout(wb)
    _choose(book, viewer="99999", day=date(2026, 10, 11), team=SHARKS_8)
    assert book.display(MD, "B13") == "No verified player has APA record ID 99999 in this build."
    assert book.display(MD, "B18") == ""
    assert book.display(MD, f"A{rows['status']}") == "That APA record ID is not a verified player in this build."
    # Ann's record with a team she is not on: never resolved.
    _choose(book, viewer="1001", team=FALCONS_8)
    assert book.display(MD, f"A{rows['status']}") == "Step 3: choose one of your teams."
    assert set(_fixture_row(book, rows["slot1"]).values()) == {""}


def test_viewer_with_two_scopes_sees_both_separately(tmp_path):
    book, wb = _book(tmp_path, [_team_match()], viewer="1002")  # Bea: Sharks 8-Ball and Sharks 9-Ball
    assert book.display(MD, "B16") == "2 current team scopes (each division listed separately)."
    assert {book.display(MD, "B18"), book.display(MD, "B19")} == {SHARKS_8, "Sharks · Spring 2026 · 9-Ball"}
    assert book.display(MD, "A18") == "Team 1" and book.display(MD, "A20") == ""


def test_match_night_lists_both_current_rosters(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    book.set("Match Night", "B5", SHARKS_8)
    book.set("Match Night", "C5", FALCONS_8)
    assert book.display("Match Night", "A20") == f"Our team: {SHARKS_8}"
    assert [book.display("Match Night", f"{c}22") for c in "ABCD"] == ["Bea Baker", "1002", "5*", "2-4"]
    assert book.display("Match Night", "A23") == "Ann Archer"
    opp_title = _find_row(wb, lambda v: isinstance(v, str) and v.startswith('=IF(C5="","Opponent'), sheet="Match Night")
    assert book.display("Match Night", f"A{opp_title + 2}") == "Cam Cole"


def test_coach_dashboard_pair_lookup_and_missing_sl(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    dash = "Coach Dashboard"
    book.set(dash, "B4", "Ann Archer (APA record ID 1001)")
    book.set(dash, "B5", "Cam Cole (APA record ID 1003)")
    book.set(dash, "B6", "8-Ball")
    assert [book.display(dash, r) for r in ("B21", "B22", "B23", "B24")] == [1, 1, 2, "50.0%"]
    book.set(dash, "B5", "Bea Baker (APA record ID 1002)")
    assert book.display(dash, "B21") == "0"
    assert "not evidence of a tie" in book.display(dash, "B25")
    assert book.display(dash, "B18") == "—"  # blank SL is "not captured", never 0


PM = "Print Matchup"


def _print_cells(book, row, cols="ABCDE"):
    return [book.display(PM, f"{c}{row}") for c in cols]


def test_identity_picker_shows_names_but_resolves_by_record_id(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    # The other "Ann Archer" (record 1004, no teams) is a different identity.
    _choose(book, viewer="Ann Archer (APA record ID 1004)")
    assert book.display(MD, "B13") == "Ann Archer (APA record ID 1004)"
    assert book.display(MD, "B16") == "No current team captured for this player."
    assert book.display(MD, "B18") == ""
    # A bare name is never treated as an identity.
    _choose(book, viewer="Ann Archer")
    assert book.display(MD, "B13") == "No verified player has APA record ID Ann Archer in this build."
    # Picking the label of record 1001 resolves to her teams.
    _choose(book, viewer="Ann Archer (APA record ID 1001)")
    assert book.display(MD, "B18") == SHARKS_8


def test_print_matchup_shows_fixture_and_clearly_labeled_rosters_for_both_teams(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    assert book.display(PM, "A1") == "Ultimate Coach — Matchup"
    assert book.display(PM, "A2") == "Sun Oct 11, 2026 · 11:00 AM MDT (America/Denver)"
    assert book.display(PM, "A3") == (
        "8-Ball Open · Spring 2026 · Venue: No data · Status: UNPLAYED · Score (home–away): Not recorded"
    )
    assert book.display(PM, "A4").startswith("Source timestamp: 2026-10-11T11:00:00-06:00 · fixture #1 of 1 on Match Day")
    assert (book.display(PM, "A6"), book.display(PM, "G6")) == ("OUR TEAM", "OPPONENT")
    assert book.display(PM, "A7") == f"{SHARKS_8} — Home"
    assert book.display(PM, "G7") == f"{FALCONS_8} — Away"
    assert book.display(PM, "A8") == "2 rostered · full-roster SL total 9"
    assert book.display(PM, "G8") == "1 rostered · full-roster SL total 3"
    assert _print_cells(book, 10) == [1, "Bea Baker", "1002", "5*", "2-4"]
    assert _print_cells(book, 11) == [2, "Ann Archer", "1001", "4", "10-5"]
    assert _print_cells(book, 12) == ["", "", "", "", ""]
    assert _print_cells(book, 10, "GHIJK") == [1, "Cam Cole", "1003", "3", "3-2"]
    assert _print_cells(book, 11, "GHIJK") == ["", "", "", "", ""]
    last = int(wb[PM].print_area.rsplit("$", 1)[1])
    shown = [book.display(PM, f"{c}{r}") for r in range(1, last + 1) for c in "ABCDEGHIJK"]
    assert 0 not in shown and "#N/A" not in shown and "#REF!" not in shown


def test_print_matchup_states_why_there_is_no_opponent_roster(tmp_path):
    book, wb = _book(tmp_path, [
        _team_match(match_id=1, match_external_id="1", is_bye=True, away_team_id="x", away_team_name="BYE"),
        _team_match(match_id=2, match_external_id="2", away_team_id="ghosts", away_team_name="Ghosts",
                    match_date="2026-10-18T11:00:00-06:00"),
    ])
    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    assert book.display(PM, "G7") == "Bye — no opponent"
    assert book.display(PM, "G8") == "Opponent roster: Not applicable (bye)."
    assert book.display(PM, "H10") == "No opponent roster to list — see above."
    assert _print_cells(book, 10)[1] == "Bea Baker"  # our roster still printed
    _choose(book, day=date(2026, 10, 18))
    assert book.display(PM, "G7") == "Ghosts"
    assert book.display(PM, "G8") == "Opponent roster: No current roster captured."
    assert book.display(PM, "H10") == "No opponent roster to list — see above."


def test_print_matchup_keeps_no_data_disclosures_for_missing_roster_values(tmp_path):
    payload = _payload()
    payload["players"].append({
        "id": 9, "external_id": "1009", "name": "Nia Null", "current_skill_level": None,
        "current_matches_won": None, "current_matches_played": None,
        "team_history": [{"team_external_id": "falcons-a", "team_name": "Falcons", "division_id": "d2",
                          "session_name": "Spring 2026", "format": "EIGHT", "is_current": True,
                          "skill_level": None, "matches_won": None, "matches_played": None}],
    })
    book, wb = _book(tmp_path, [_team_match()], payload=payload)
    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    rows = [_print_cells(book, r, "GHIJK") for r in (10, 11)]
    assert [1, "Cam Cole", "1003", "3", "3-2"] in rows
    assert [2, "Nia Null", "1009", "No data", "No data"] in rows
    assert book.display(PM, "G8") == (
        "2 rostered · SL total 3 from 1 players with a captured SL (missing excluded, not counted as 0)"
    )


def test_print_matchup_without_a_selection_explains_what_to_do(tmp_path):
    book, wb = _book(tmp_path, [_team_match()], viewer=None)
    assert book.display(PM, "A2").startswith("Choose a date, your team and a fixture on the Match Day sheet")
    assert book.display(PM, "A7") == "Not chosen — pick your team on Match Day (step 3)."
    assert book.display(PM, "G7") == "Not chosen — pick a fixture on Match Day."
    assert _print_cells(book, 10) == ["", "", "", "", ""]


PAIR_COMPARISON = {
    "Players rostered": ("2", "1"),
    "Players with a captured SL": ("2 of 2", "1 of 1"),
    "Skill total (captured SLs only)": ("9", "3"),
    "Recorded 8-Ball games (all opponents)": ("2", "2"),
    "Players with no recorded 8-Ball games": ("1", "0"),
}


def _comparison_rows(book, sheet, top, label_col, ours_col, theirs_col):
    out = {}
    for r in range(top + 3, top + 8):
        out[book.display(sheet, f"{label_col}{r}")] = (book.display(sheet, f"{ours_col}{r}"),
                                                        book.display(sheet, f"{theirs_col}{r}"))
    between = [book.display(sheet, f"{ours_col}{r}") for r in range(top + 9, top + 12)]
    return out, between


def _ranking_block(book, sheet, top, block, slots, cols):
    start = top + 4 + (block - 1) * (3 + slots + 1)
    header = book.display(sheet, f"A{start}")
    note = book.display(sheet, f"A{start + 1}")
    rows = [[book.display(sheet, f"{c}{start + 2 + r}") for c in cols] for r in range(1, slots + 1)]
    return header, note, rows


def _section_row(wb, sheet, title):
    return next(c.row for c in wb[sheet]["A"] if c.value == title)


RANK_TITLE = "Evidence ranking vs each opponent (not win odds)"
ANN_ROW = ["1", "Ann Archer (APA record ID 1001)", "4", "1-1 (2 meetings)", "No shared opponents", "Direct record"]
BEA_ROW = ["—", "Bea Baker (APA record ID 1002)", "5*", "No direct meetings", "No shared opponents",
           "No direct or shared-opponent evidence"]


def test_print_matchup_includes_team_comparison_and_evidence_ranking(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    _choose(book, day=date(2026, 10, 11), team=SHARKS_8)
    top = _section_row(wb, PM, "Team comparison")
    assert book.display(PM, f"A{top + 1}").startswith("Format: 8-Ball · recorded facts")
    assert book.display(PM, f"C{top + 2}") == f"Our team: {SHARKS_8}"
    assert book.display(PM, f"G{top + 2}") == f"Opponent: {FALCONS_8}"
    metrics, between = _comparison_rows(book, PM, top, "A", "C", "G")
    assert metrics == PAIR_COMPARISON
    assert between == ["2 games · our players 1-1", "1 of 1", "1 direct · 0 shared-opponent only · 1 no evidence (2 total)"]

    rank_top = _section_row(wb, PM, RANK_TITLE)
    slots = 8
    header, note, rows = _ranking_block(book, PM, rank_top, 1, slots, "ABCDGI")
    assert header == "vs Cam Cole (APA record ID 1003) · SL 3 · 2 recorded games in 8-Ball"
    assert note == "1 of 2 of our players have direct or shared-opponent evidence against this opponent."
    assert rows[0] == ANN_ROW and rows[1] == BEA_ROW
    assert set(rows[2]) == {""}
    header2, _, rows2 = _ranking_block(book, PM, rank_top, 2, slots, "ABCDGI")
    assert header2 == "" and all(set(r) == {""} for r in rows2)  # Falcons has one rostered player
    last = int(wb[PM].print_area.rsplit("$", 1)[1])
    shown = [book.display(PM, f"{c}{r}") for r in range(1, last + 1) for c in "ABCDEGHIJK"]
    assert 0 not in shown and "#N/A" not in shown and "#REF!" not in shown


def test_match_night_shows_comparison_and_ranking_for_a_scheduled_pairing(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    book.set("Match Night", "B5", SHARKS_8)
    book.set("Match Night", "C5", FALCONS_8)
    top = _section_row(wb, "Match Night", "Team comparison")
    assert book.display("Match Night", f"B{top + 2}") == f"Our team: {SHARKS_8}"
    metrics, between = _comparison_rows(book, "Match Night", top, "A", "B", "C")
    assert metrics == PAIR_COMPARISON
    assert between[0] == "2 games · our players 1-1"
    rank_top = _section_row(wb, "Match Night", RANK_TITLE)
    header, note, rows = _ranking_block(book, "Match Night", rank_top, 1, 8, "ABCDEF")
    assert header.startswith("vs Cam Cole (APA record ID 1003)")
    assert rows[0] == ANN_ROW and rows[1] == BEA_ROW


def test_match_night_explains_when_a_pairing_has_no_precomputed_evidence(tmp_path):
    book, wb = _book(tmp_path, [_team_match()])
    book.set("Match Night", "B5", "Sharks · Spring 2026 · 9-Ball")  # no fixture vs Falcons in 9-Ball
    book.set("Match Night", "C5", FALCONS_8)
    top = _section_row(wb, "Match Night", "Team comparison")
    assert book.display("Match Night", f"A{top + 1}").startswith("Team comparison and evidence ranking are precomputed")
    rank_top = _section_row(wb, "Match Night", RANK_TITLE)
    assert book.display("Match Night", f"A{rank_top + 3}").startswith("Team comparison and evidence ranking are precomputed")
    header, _, rows = _ranking_block(book, "Match Night", rank_top, 1, 8, "ABCDEF")
    assert header == "" and set(rows[0]) == {""}
    # Ranked rows never show a bare 0 from an empty lookup.
    assert 0 not in [book.display("Match Night", f"{c}{r}") for r in range(top, rank_top + 20) for c in "ABCDEF"]
