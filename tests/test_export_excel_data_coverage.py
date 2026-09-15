"""Tests for ui/export_excel_data_coverage.py."""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.data_coverage import build_report
from analytics.pairing_evidence import EvidenceLabel, PairingEvidence, PairingEvidenceMatrix
from ui.export_excel_data_coverage import SHEET_NAME, build_workbook, write_workbook


def _pairing(player_id, opponent_id, label, **kwargs):
    defaults = dict(
        player_external_id=f"P-{player_id}", player_name=f"Player {player_id}",
        player_skill_level=5, opponent_external_id=f"OPP-{opponent_id}",
        opponent_name=f"Opponent {opponent_id}", opponent_skill_level=4,
        format="8-Ball Open", session_name="Fall 2026",
        observed_win_rate=None, direct_evidence_count=0,
        modeled_win_probability=None, model_source=None,
    )
    defaults.update(kwargs)
    return PairingEvidence(player_id=player_id, opponent_id=opponent_id,
                            evidence_label=label, **defaults)


def _matrix(pairings):
    counts = {label.value: 0 for label in EvidenceLabel}
    for p in pairings:
        counts[p.evidence_label.value] += 1
    counts["total_feasible_pairings"] = len(pairings)
    return PairingEvidenceMatrix(
        our_team_external_id="OUR", opponent_team_external_id="THEIRS",
        format="8-Ball Open", session_name="Fall 2026",
        expected_pairings=tuple((p.player_id, p.opponent_id) for p in pairings),
        pairings=tuple(pairings), counts=counts,
        our_roster_available=True, opponent_roster_available=True,
    )


class TestBuildWorkbook:
    def test_produces_both_sheets_even_with_zero_pairings(self):
        report = build_report(_matrix([]))
        workbook = build_workbook(report)

        assert SHEET_NAME in workbook.sheetnames
        assert "Sample_Sizes" in workbook.sheetnames

    def test_evidence_counts_appear_in_the_sheet(self):
        report = build_report(_matrix([
            _pairing(1, 10, EvidenceLabel.DIRECT, direct_evidence_count=1),
            _pairing(2, 11, EvidenceLabel.UNKNOWN, player_skill_level=None),
        ]))
        workbook = build_workbook(report)
        sheet = workbook[SHEET_NAME]
        values = [cell.value for row in sheet.iter_rows() for cell in row]

        assert "DIRECT" in values
        assert 1 in values

    def test_missing_skill_levels_are_listed_by_name(self):
        report = build_report(_matrix([
            _pairing(1, 10, EvidenceLabel.UNKNOWN, player_skill_level=None, player_name="Bob"),
        ]))
        workbook = build_workbook(report)
        sheet = workbook[SHEET_NAME]
        values = [cell.value for row in sheet.iter_rows() for cell in row]

        assert "Bob" in values

    def test_write_workbook_produces_a_real_readable_file(self, tmp_path):
        report = build_report(_matrix([_pairing(1, 10, EvidenceLabel.DIRECT, direct_evidence_count=1)]))
        out_path = write_workbook(report, tmp_path / "data_coverage.xlsx")

        assert out_path.exists()
        reopened = load_workbook(out_path)
        assert SHEET_NAME in reopened.sheetnames
