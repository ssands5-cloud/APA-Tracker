"""Tests for the Head-to-Head Advantage Engine.

Covers the analytics, the database table, the builder and the demo tab.
Fixtures are built from real ORM objects so a column rename breaks these
rather than silently emptying a field.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from analytics.head_to_head import (
    MAX_WIN_PROBABILITY,
    MIN_WIN_PROBABILITY,
    average_skill_level_delta,
    evaluate,
    expected_balls,
    expected_points,
    skill_level_trend,
    win_loss,
    win_probability,
)
from database.ingest import ingest_h2h_advantage
from database.models import (
    Base,
    Match,
    Player,
    PlayerH2HAdvantage,
    PlayerHeadToHead,
    PlayerMatchup,
)
from scripts.build_head_to_head import build_rows
from ui.tabs.matchups import classify, load_rows, render


def game(result="W", own=5, opp=5, points=2.0, balls=None, fmt="8-Ball Open",
         session="Fall 2026"):
    """One head-to-head game, detached from any session."""
    return PlayerHeadToHead(
        own_skill_level=own, opponent_skill_level=opp, result=result,
        points_earned=points, nine_ball_points=balls, format=fmt,
        session_name=session,
    )


class TestRecord:
    def test_wins_and_losses_count_only_recognised_results(self):
        rows = [game("W"), game("L"), game("W"), game("?")]
        assert win_loss(rows) == (2, 1)

    def test_an_empty_pairing_is_zero_zero(self):
        assert win_loss([]) == (0, 0)


class TestSkillLevelDelta:
    def test_it_is_opponent_minus_own(self):
        """Positive means the player is giving up skill level -- the same
        sign convention PlayerMatchup.sl_delta uses."""
        assert average_skill_level_delta([game(own=4, opp=6)]) == 2.0
        assert average_skill_level_delta([game(own=6, opp=4)]) == -2.0

    def test_missing_levels_yield_none_not_zero(self):
        """No skill data is not the same fact as an even matchup."""
        assert average_skill_level_delta([game(own=None, opp=None)]) is None


class TestTrend:
    def test_first_versus_last_not_every_wobble(self):
        assert skill_level_trend([game(own=4), game(own=6), game(own=4)]) == "stable"
        assert skill_level_trend([game(own=4), game(own=5)]) == "up"
        assert skill_level_trend([game(own=5), game(own=4)]) == "down"

    def test_no_readings_is_no_data(self):
        assert skill_level_trend([game(own=None)]) == "no data"


class TestWinProbability:
    def test_a_skill_advantage_raises_it_above_even(self):
        assert win_probability([game("W", own=6, opp=4)]) > 0.5

    def test_a_skill_deficit_lowers_it_below_even(self):
        assert win_probability([game("L", own=4, opp=6)]) < 0.5

    def test_the_sign_convention_is_not_inverted(self):
        """The single easiest way to get this model backwards."""
        strong = win_probability([game("W", own=7, opp=3)])
        weak = win_probability([game("L", own=3, opp=7)])
        assert strong > weak

    def test_one_game_cannot_assert_certainty(self):
        """Laplace smoothing plus reliability damping: a 1-0 record is
        evidence, not proof."""
        p = win_probability([game("W", own=5, opp=5)])
        assert 0.5 < p < 0.75

    def test_more_evidence_moves_it_further(self):
        few = win_probability([game("W", own=5, opp=5)])
        many = win_probability([game("W", own=5, opp=5) for _ in range(10)])
        assert many > few

    def test_it_never_reaches_zero_or_one(self):
        """Pool has upsets; a stated 0% or 100% would overclaim."""
        p = win_probability([game("W", own=7, opp=2) for _ in range(50)])
        assert p <= MAX_WIN_PROBABILITY
        p = win_probability([game("L", own=2, opp=7) for _ in range(50)])
        assert p >= MIN_WIN_PROBABILITY

    def test_no_evidence_at_all_is_none(self):
        """No games and no skill levels: no probability, which is different
        from a 50/50 one."""
        assert win_probability([]) is None
        assert win_probability([game("?", own=None, opp=None)]) is None


class TestExpectedValues:
    def test_expected_points_is_eight_ball_only(self):
        assert expected_points([game(points=3.0)]) == 3.0
        assert expected_points([game(points=3.0, fmt="9-Ball Open")]) is None

    def test_expected_balls_is_nine_ball_only(self):
        assert expected_balls([game(balls=30, fmt="9-Ball Open")]) == 30.0
        assert expected_balls([game(balls=30, fmt="8-Ball Open")]) is None

    def test_balls_and_points_are_different_measurements(self):
        """nine_ball_points is a BALL count; points_earned carries match
        points. Reading one as the other would be silently wrong."""
        rows = [game(points=10.0, balls=45, fmt="9-Ball Open")]
        assert expected_balls(rows) == 45.0
        assert expected_points(rows) is None

    def test_a_thin_pairing_is_pulled_toward_the_baseline(self):
        """One big night must not become the expectation forever."""
        assert expected_points([game(points=6.0)], baseline=2.0) == 3.0

    def test_plenty_of_evidence_overrides_the_baseline(self):
        value = expected_points([game(points=6.0) for _ in range(10)], baseline=2.0)
        assert value > 5.0

    def test_without_a_baseline_it_is_the_plain_mean(self):
        assert expected_points([game(points=4.0), game(points=2.0)]) == 3.0


class TestEvaluate:
    def test_it_fills_every_stored_field(self):
        result = evaluate([game("W", own=6, opp=4), game("L", own=6, opp=4)],
                          player_id="P1", opponent_id="P2")
        assert (result.total_matches, result.wins, result.losses) == (2, 1, 1)
        assert result.sl_delta == -2.0
        assert result.matchup_score is not None
        assert 0 <= result.matchup_score <= 100
        assert result.win_probability is not None
        assert result.format == "8-Ball Open"

    def test_wins_plus_losses_always_equals_total(self):
        result = evaluate([game("W"), game("L"), game("?")], "P1", "P2")
        assert result.wins + result.losses == result.total_matches


@pytest.fixture
def db(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'h2h.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session


@pytest.fixture
def seeded(db):
    """Two players with a real pairing across two matches."""
    alice = Player(external_id="P1", name="Alice", skill_level=6)
    bob = Player(external_id="P2", name="Bob", skill_level=4)
    db.add_all([alice, bob])
    db.flush()

    for week, result in ((1, "W"), (2, "W")):
        match = Match(external_id=f"M{week}", week=week,
                      home_team_id="T1", away_team_id="T2")
        db.add(match)
        db.flush()
        db.add(PlayerHeadToHead(
            player_id=alice.id, opponent_id=bob.id, match_id=match.id,
            own_skill_level=6, opponent_skill_level=4, result=result,
            points_earned=3.0, format="8-Ball Open", session_name="Fall 2026",
        ))
    db.commit()
    return db


class TestDatabaseTable:
    def test_the_table_has_exactly_the_agreed_columns(self):
        """Innings and defensive shots are absent by decision, not oversight."""
        columns = {c.name for c in PlayerH2HAdvantage.__table__.columns}
        assert columns == {
            "id", "player_id", "opponent_id", "total_matches", "wins", "losses",
            "sl_delta", "trend_modifier", "matchup_score", "win_probability",
            "expected_points", "expected_balls", "format", "session_name",
        }
        assert "avg_innings" not in columns
        assert "avg_defense" not in columns

    def test_it_upserts_rather_than_duplicating(self, seeded):
        rows = build_rows(seeded)
        ingest_h2h_advantage(seeded, rows)
        ingest_h2h_advantage(seeded, rows)
        assert seeded.query(PlayerH2HAdvantage).count() == len(rows)

    def test_an_unknown_player_is_skipped_not_guessed(self, seeded):
        written = ingest_h2h_advantage(seeded, [{
            "player_id": "GHOST", "opponent_id": "P2", "total_matches": 1,
            "wins": 1, "losses": 0, "sl_delta": 0.0, "trend_modifier": 0,
            "matchup_score": 50, "win_probability": 0.5,
            "expected_points": None, "expected_balls": None,
            "format": "8-Ball Open", "session_name": "Fall 2026",
        }])
        assert written == 0

    def test_it_does_not_touch_player_matchups(self, seeded):
        """The Matchup Advantage Engine owns that table."""
        before = seeded.query(PlayerMatchup).count()
        ingest_h2h_advantage(seeded, build_rows(seeded))
        assert seeded.query(PlayerMatchup).count() == before


class TestBuilder:
    def test_it_builds_one_row_per_pairing(self, seeded):
        rows = build_rows(seeded)
        assert len(rows) == 1
        row = rows[0]
        assert (row["player_id"], row["opponent_id"]) == ("P1", "P2")
        assert (row["wins"], row["losses"], row["total_matches"]) == (2, 0, 2)
        assert row["format"] == "8-Ball Open"

    def test_an_empty_database_builds_nothing_rather_than_failing(self, db):
        assert build_rows(db) == []


class TestDemoTab:
    def test_it_highlights_only_when_score_and_probability_agree(self):
        """Either signal alone would promote a fluke or a mismatch."""
        assert classify({"matchup_score": 75, "win_probability": 0.8}) == "recommended"
        assert classify({"matchup_score": 75, "win_probability": 0.3}) == "neutral"
        assert classify({"matchup_score": 30, "win_probability": 0.2}) == "avoid"
        assert classify({"matchup_score": 30, "win_probability": 0.9}) == "neutral"

    def test_a_pairing_with_no_probability_is_neutral(self):
        assert classify({"matchup_score": 90, "win_probability": None}) == "neutral"

    def test_it_renders_rows_with_the_recommendation_class(self, seeded):
        ingest_h2h_advantage(seeded, build_rows(seeded))
        html = render(load_rows(seeded))
        assert "Bob" in html
        assert "<table" in html
        assert 'class="recommended"' in html or 'class="neutral"' in html

    def test_it_carries_no_external_resources(self, seeded):
        """It has to open at a venue with no internet."""
        ingest_h2h_advantage(seeded, build_rows(seeded))
        html = render(load_rows(seeded))
        assert "http://" not in html
        assert "https://" not in html

    def test_an_empty_tab_explains_what_to_run(self):
        html = render([])
        assert "build_head_to_head" in html

    def test_missing_values_render_as_a_dash_not_zero(self, seeded):
        ingest_h2h_advantage(seeded, build_rows(seeded))
        html = render(load_rows(seeded))
        # This pairing is 8-ball, so Expected Balls has no value.
        assert "—" in html


class TestExcelSheet:
    def test_the_sheet_exists_with_every_field(self, seeded, tmp_path):
        from ui.export_excel import export_to_excel

        ingest_h2h_advantage(seeded, build_rows(seeded))
        path = export_to_excel(seeded, {"export": {
            "excel_output_path": str(tmp_path / "wb.xlsx")}})

        sheet = load_workbook(path)["Head-to-Head"]
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        assert header == ["Player", "Opponent", "Format", "Session", "Games",
                          "Wins", "Losses", "SL Delta", "Trend Modifier",
                          "Matchup Score", "Win Probability", "Expected Points",
                          "Expected Balls"]

    def test_headers_freeze_and_the_score_is_colour_coded(self, seeded, tmp_path):
        from ui.export_excel import export_to_excel

        ingest_h2h_advantage(seeded, build_rows(seeded))
        path = export_to_excel(seeded, {"export": {
            "excel_output_path": str(tmp_path / "wb.xlsx")}})

        sheet = load_workbook(path)["Head-to-Head"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None
        assert sum(len(r.rules) for r in sheet.conditional_formatting) == 2

    def test_the_existing_matchups_sheet_still_ships(self, seeded, tmp_path):
        """Adding a sheet must not disturb the Matchup Advantage Engine's."""
        from ui.export_excel import export_to_excel

        path = export_to_excel(seeded, {"export": {
            "excel_output_path": str(tmp_path / "wb.xlsx")}})
        assert "Matchups" in load_workbook(path).sheetnames
