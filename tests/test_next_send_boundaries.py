"""Presentation boundaries; ranked evidence remains identical to Python."""
from pathlib import Path
import pytest
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
from ui.export_excel_ultimate_coach import write_workbook
from tests.excel_formula_eval import Workbook
from tests.test_excel_war_room_formulas import _payload
from tests.test_ultimate_coach_war_room_browser import _page, _python_war_room, OURS, THEIRS, CAM


@pytest.mark.parametrize('width',[390,1280])
def test_next_send_boundaries_before_list_and_scoped_context(tmp_path,width):
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page,errors=_page(tmp_path,browser);page.set_viewport_size({'width':width,'height':900})
            assert 'options against' in page.locator('#next-send h3').inner_text().lower()
            method=page.locator('#next-send .ns-method')
            assert 'Not an overall player ranking' in method.inner_text()
            assert method.evaluate('(e)=>!!(e.compareDocumentPosition(document.querySelector("#next-send .ns-list")) & Node.DOCUMENT_POSITION_FOLLOWING)')
            assert 'Direct-record order 1' in page.locator('#next-send .ns-m').first.inner_text()
            assert '2 meetings' in page.locator('#next-send .ns-list .ns-sample').first.inner_text()
            assert 'small sample' in page.locator('#next-send .ns-list .ns-sample').first.inner_text()
            context=page.locator('#next-send .ns-broader')
            assert 'Not used in this ordering' in context.inner_text()
            assert 'Fall 2026' in context.inner_text() and '8-Ball' in context.inner_text()
            assert 'Captured team/session record' in context.inner_text()
            assert 'not a lifetime record' in context.inner_text()
            assert 'No verified direct/shared comparison' in page.locator('#next-send .ns-method').inner_text()
            py=_python_war_room();js=page.evaluate(f"window.__ucWarRoomPair({OURS!r},{THEIRS!r},'EIGHT')")
            assert js['best']==[[r['player'] for r in b['rows'] if r['category'] in ('G','E','I')] for b in py['blocks']]
            assert js['matrix']==[[[c['category'],c['cell'],c['explanation'],c['reason'],c['captain']] for c in row['cells']] for row in py['matrix']]
            assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
            assert errors==[]
        finally:browser.close()


@pytest.mark.parametrize('width',[390,1280])
def test_one_win_and_eight_wins_keep_order_and_show_the_evidence_difference(tmp_path,width):
    from ui.ultimate_coach import render
    from tests.test_excel_war_room_formulas import _games
    payload=_payload()
    payload['evidence']=[r for r in payload['evidence'] if {r['player_id'],r['opponent_id']} not in ({1,10},{2,10})]
    payload['evidence']+=_games(1,10,'W')+_games(2,10,'WWWWWWWWLL')
    path=tmp_path/'one-vs-ten.html'
    path.write_text(render(payload,built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page=browser.new_page(viewport={'width':width,'height':900});page.set_default_timeout(2500);page.goto(path.as_uri())
            rows=page.locator('#next-send .ns-medal')
            assert rows.nth(0).locator('[data-send-our]').get_attribute('data-send-our')=='1'
            assert rows.nth(1).locator('[data-send-our]').get_attribute('data-send-our')=='2'
            assert '1-0' in rows.nth(0).inner_text() and '1 meeting · small sample' in rows.nth(0).inner_text()
            assert '8-2' in rows.nth(1).inner_text() and '10 meetings' in rows.nth(1).inner_text()
            assert 'small sample' not in rows.nth(1).inner_text()
            assert 'does not mean better supported' in page.locator('.ns-method').inner_text()
            assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
        finally:browser.close()


def test_workbook_neutral_order_and_broader_context(tmp_path):
    path=write_workbook(_payload(),tmp_path/'uc.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001')
    wb=load_workbook(path);book=Workbook(wb);book.set('Command Center','C6',CAM)
    assert book.display('Command Center','B7')=='NEXT SEND - RECORDED MATCHUP OPTIONS'
    text=book.display('Command Center','B8')
    assert 'Not an overall player ranking' in text
    assert 'Direct-record order 1' in text
    assert 'MEETINGS' in text and 'SMALL SAMPLE' in text
    broader=book.value('Engine',book.name('wr_NsBroader').ref)
    assert 'Not used in this ordering' in broader
    assert 'Fall 2026' in broader and '8-Ball' in broader
    assert 'not a lifetime record' in broader
    assert 'Ann Archer' in broader and 'Bea Baker' in broader and CAM in broader
    wb.close()


def test_broader_zero_missing_and_strong_records_never_reorder(tmp_path):
    from ui.ultimate_coach import render
    payload=_payload()
    results=[]
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            for label,won,played in [('zero',0,0),('strong',50,65)]:
                payload['players'][0]['team_history'][0]['matches_won']=won
                payload['players'][0]['team_history'][0]['matches_played']=played
                path=tmp_path/(label+'.html')
                path.write_text(render(payload,built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
                page=browser.new_page(viewport={'width':390,'height':900});page.goto(path.as_uri())
                results.append(page.evaluate(f"window.__ucWarRoomPair({OURS!r},{THEIRS!r},'EIGHT')"))
                text=page.locator('.ns-broader').inner_text()
                assert ('0-0' if label=='zero' else '50-15') in text
                assert 'Not captured' in text
                assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
                page.close()
        finally:browser.close()
    # Only scouting broader-record display changes; all matrix and sends are identical.
    assert results[0]['matrix']==results[1]['matrix']
    assert results[0]['best']==results[1]['best']


def test_one_direct_win_stays_first_while_unranked_broader_record_is_visible(tmp_path):
    from ui.ultimate_coach import render
    from tests.test_excel_war_room_formulas import _games
    payload=_payload()
    payload['evidence']=[r for r in payload['evidence'] if {r['player_id'],r['opponent_id']} not in ({1,10},{2,10})]
    old=_games(1,10,'W',date='2025-09-20T19:00:00-06:00')
    for r in old:r['session_name']='Fall 2025'
    payload['evidence']+=old
    for p in payload['players']:
        for h in p['team_history']:
            if p['id']==1:h.update(matches_won=3,matches_played=7)
            if p['id']==2:h.update(matches_won=50,matches_played=65)
            if p['id']==10:h.update(matches_won=8,matches_played=8)
    path=tmp_path/'early-direct.html';path.write_text(render(payload,built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page=browser.new_page(viewport={'width':390,'height':900});page.goto(path.as_uri())
            assert page.locator('#next-send [data-send-our]').first.get_attribute('data-send-our')=='1'
            assert page.locator('#next-send .ns-list .ns-sample').first.inner_text()=='1 meeting · small sample'
            broader=page.locator('.ns-broader')
            assert '3-4' in broader.inner_text() and '8-0' in broader.inner_text()
            b=broader.locator('li').filter(has_text='Bea Baker')
            assert '50-15' in b.inner_text() and 'no verified direct meetings' in b.inner_text()
            assert b.is_visible() and 'Not used in this ordering' in broader.inner_text()
        finally:browser.close()


@pytest.mark.parametrize('width',[320,390,1280])
def test_decision_labels_explain_samples_without_claiming_a_practical_best(tmp_path,width):
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page,errors=_page(tmp_path,browser)
            page.set_viewport_size({'width':width,'height':900})
            assert 'does not mean better supported' in page.locator('.ns-method').inner_text()
            plan=page.locator('#match-plan')
            # The lead pick carries its confidence level and states plainly that levels are evidence, not odds.
            assert plan.locator('.mp-lead .mp-level').count()==1
            assert 'not odds' in plan.inner_text() and '%' not in plan.inner_text()
            risks=page.locator('#wr-risks')
            assert 'Threats — winning direct records' in risks.inner_text()
            assert 'Avoid records — more losses than wins' in risks.inner_text()
            assert 'A single loss is a caution, not an instruction to rule out a player' in risks.inner_text()
            assert 'A single win is not an established threat' in risks.inner_text()
            assert 'recorded losses' in page.locator('#next-send .ns-more summary').inner_text()
            before=page.evaluate(f"window.__ucWarRoomPair({OURS!r},{THEIRS!r},'EIGHT')")
            page.locator('#ll-cap').fill('1')
            warnings=page.locator('#next-send .ns-budget.bad')
            assert warnings.count()>0
            assert warnings.first.evaluate('e=>parseFloat(getComputedStyle(e).borderLeftWidth)>=3')
            assert page.evaluate(f"window.__ucWarRoomPair({OURS!r},{THEIRS!r},'EIGHT')")==before
            page.locator('#ll-cap').fill('')
            assert page.locator('#next-send .ns-budget.bad').count()==0
            assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
            assert errors==[]
        finally:browser.close()


@pytest.mark.parametrize('width',[390,1280])
def test_thin_leader_headline_and_deeper_record_are_visible_without_reordering(tmp_path,width):
    from ui.ultimate_coach import render
    from tests.test_excel_war_room_formulas import _games
    payload=_payload()
    payload['evidence']=[r for r in payload['evidence'] if {r['player_id'],r['opponent_id']} not in ({1,10},{2,10})]
    payload['evidence']+=_games(1,10,'W')+_games(2,10,'WWWWWWWWLL')
    path=tmp_path/'thin-lead.html';path.write_text(render(payload,built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page=browser.new_page(viewport={'width':width,'height':900});page.set_default_timeout(2500);page.goto(path.as_uri())
            head=page.locator('#next-send .ns-evidence-head')
            assert 'Direct-record order #1' in head.inner_text()
            assert '1-0' in head.inner_text() and '1 meeting' in head.inner_text()
            assert '⚠ Small sample' in head.inner_text()
            deeper=page.locator('#next-send .ns-deeper')
            assert 'Bea Baker' in deeper.inner_text() and '8-2 in 10 meetings' in deeper.inner_text()
            assert deeper.is_visible() and 'not a new order' in deeper.inner_text()
            assert page.locator('#next-send [data-send-our]').first.get_attribute('data-send-our')=='1'
            assert '1 meeting, small sample' in page.locator('#scouting-cards .quick-read').first.inner_text()
            assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
        finally:browser.close()


@pytest.mark.parametrize('width',[390,1280])
@pytest.mark.parametrize('record,level',[('W','Thin'),('WWWWWW','Solid')])   # 6-0, but Cam's SL changed since: one level down
def test_match_plan_answers_carry_the_depth_of_their_evidence(tmp_path,width,record,level):
    from ui.ultimate_coach import render
    from tests.test_excel_war_room_formulas import _games,_player,_hist
    payload=_payload()
    payload['players'] += [_player(i,str(2000+i),'Fixture Opponent '+str(i),[_hist('falcons-a','d1','EIGHT',3)]) for i in (20,21)]
    payload['evidence']=sum((_games(1,i,record) for i in (10,11,20)),[])
    path=tmp_path/'blind-total.html';path.write_text(render(payload,built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page=browser.new_page(viewport={'width':width,'height':900});page.set_default_timeout(2500);page.goto(path.as_uri())
            plan=page.locator('#match-plan')
            # The same pairing counts, backed by 1 meeting or 6 each: the answer's level must show the difference.
            answer=plan.locator('.mp-resp li',has_text='Cam Cole →')
            assert 'Ann Archer' in answer.inner_text() and level in answer.locator('.mp-level').inner_text()
            if record!='W': assert '6 meetings' in answer.inner_text() and 'skill levels changed' in answer.inner_text()
            assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
        finally:browser.close()


@pytest.mark.parametrize('width',[390,1280])
def test_single_meeting_cautions_and_future_save_are_explicit(tmp_path,width):
    from ui.ultimate_coach import render
    from tests.test_excel_war_room_formulas import _games
    payload=_payload()
    payload['evidence']=[r for r in payload['evidence'] if {r['player_id'],r['opponent_id']} not in ({2,10},{3,11})]
    payload['evidence']+=_games(2,10,'L')+_games(3,11,'L')+_games(1,11,'W')
    path=tmp_path/'single-save.html';path.write_text(render(payload,built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page=browser.new_page(viewport={'width':width,'height':900});page.set_default_timeout(2500);page.goto(path.as_uri())
            target=page.locator('#next-send [data-inspect-our="1"][data-inspect-opp="11"]')
            assert '⚠ Only 1 meeting' in target.locator('strong').inner_text()
            assert 'Lost once' in page.locator('#wr-risks').inner_text()
            # Exclude our favorable Eve meeting to leave one live opponent win.
            page.select_option('#lineup-lab select[data-plan="avail"][data-pid="1"]','Unavailable')
            assert 'Beat us once' in page.locator('#tonight .decide-threats').inner_text()
            assert '1 meeting' in page.locator('#tonight .decide-threats').inner_text()
            quick=page.locator('#scouting-cards .quick-read').all_text_contents()
            assert any('Losing direct records: Dee Diaz (0-1, 1 meeting, small sample)' in q for q in quick)
            assert not page.evaluate('document.documentElement.scrollWidth>innerWidth')
        finally:browser.close()
