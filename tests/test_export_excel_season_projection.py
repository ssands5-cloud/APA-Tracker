"""Tests for ui/export_excel_season_projection.py."""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.season_projection import project_remaining_schedule, team_volatility_curve
from ui.export_excel_season_projection import (
    HISTORY_SHEET,
    MATCHES_SHEET,
    SUMMARY_SHEET,
    build_workbook,
    write_workbook,
)


def _projection():
    matches = [
        {"match_id": "M1", "opponent_team_id": "T2", "opponent_team_name": "Corner Pockets", "week": 5},
    ]
    return project_remaining_schedule(matches, our_win_rate=0.6, opponent_win_rates={"T2": 0.4})


class TestBuildWorkbook:
    def test_produces_summary_and_matches_sheets(self):
        workbook = build_workbook(_projection(), [], "T1", "Chalk It Up", "Fall 2026", 5, 2, "2026-09-15")

        assert SUMMARY_SHEET in workbook.sheetnames
        assert MATCHES_SHEET in workbook.sheetnames

    def test_history_sheet_omitted_when_no_standings_history_exists(self):
        workbook = build_workbook(_projection(), [], "T1", "Chalk It Up", "Fall 2026", 5, 2, None)

        assert HISTORY_SHEET not in workbook.sheetnames

    def test_history_sheet_present_when_real_points_exist(self):
        points = team_volatility_curve([{"captured_at": "2026-09-01", "wins": 1, "losses": 0}])
        workbook = build_workbook(_projection(), points, "T1", "Chalk It Up", "Fall 2026", 5, 2, None)

        assert HISTORY_SHEET in workbook.sheetnames

    def test_matches_sheet_carries_every_real_remaining_match(self):
        workbook = build_workbook(_projection(), [], "T1", "Chalk It Up", "Fall 2026", 5, 2, None)
        sheet = workbook[MATCHES_SHEET]
        assert sheet.max_row == 2  # header + 1 match

    def test_write_workbook_produces_a_real_readable_file(self, tmp_path):
        out_path = write_workbook(
            _projection(), [], "T1", "Chalk It Up", "Fall 2026", 5, 2, "2026-09-15",
            tmp_path / "season_projection.xlsx",
        )

        assert out_path.exists()
        reopened = load_workbook(out_path)
        assert SUMMARY_SHEET in reopened.sheetnames
