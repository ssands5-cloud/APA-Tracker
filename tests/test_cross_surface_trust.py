"""The same recorded loss cautions and live flags must agree across outputs."""
import re
import pytest
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from analytics.ultimate_coach_war_room import next_send,live_threats,lineup_budget
from tests.test_excel_war_room_formulas import _payload,_games,SHARKS,FALCONS
from tests.test_ultimate_coach_war_room_browser import _python_war_room_for
from tests.excel_formula_eval import Workbook
from ui.export_excel_ultimate_coach import write_workbook
from ui.ultimate_coach import render


def loss_payload():
    p=_payload(crowded=True)
    p['evidence']=(_games(1,10,'L')+_games(2,10,'LLLLLL')+_games(3,10,'WWLLL')
                   +_games(4,10,'WLLLLLLLL')+_games(5,10,'WW')+_games(6,10,'WL'))
    return p


def test_loss_cautions_have_one_order_but_positive_sends_are_unchanged():
    w=_python_war_room_for(loss_payload());ns=next_send(w,0)
    assert [r['member']['id'] for r in ns['ranked']]==[5,6]
    assert [r['member']['id'] for r in ns['avoid']]==[4,2,3,1]
    assert [r['member']['id'] for r in w['blocks'][0]['rows'] if r['category']=='R']==[4,2,3,1]
    assert [r['our']['id'] for r in w['concerning']]==[4,2,3,1]
    text=w['cards'][0]['quick_read']
    assert text.index('1-8, 9 meetings')<text.index('0-6, 6 meetings')<text.index('2-3, 5 meetings')<text.index('0-1, 1 meeting')


def test_workbook_next_send_uses_the_same_loss_order(tmp_path):
    wb=load_workbook(write_workbook(loss_payload(),tmp_path/'loss.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb);book.set('Command Center','C6','Cam Cole')
    text=book.value('Engine',book.name('wr_NsAvoidLine').ref)
    assert text.index('1-8')<text.index('0-6')<text.index('2-3')<text.index('0-1')
    wb.close()


def named_values(book,name):
    return [book.value(sheet,cell) for sheet,cell in book.name(name)]


def test_packet_preserves_the_same_future_save_warning_and_single_meeting_caution(tmp_path):
    payload=_payload()
    payload['evidence']=[r for r in payload['evidence'] if {r['player_id'],r['opponent_id']}!={1,10}]+_games(1,10,'W')
    wb=load_workbook(write_workbook(payload,tmp_path/'save.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb)
    text='\n'.join(str(book.display('Captain Packet',c.coordinate) or '') for row in wb['Captain Packet'] for c in row if c.column==3)
    assert 'consider saving' in text and 'Cam Cole' in text and 'Only 1 meeting' in text
    assert '1-0 in 1 meeting' in text
    # Spending or retiring the protected pairing removes its warning.
    book.set('Lineup Lab','C23','Played')
    text='\n'.join(str(book.display('Captain Packet',c.coordinate) or '') for row in wb['Captain Packet'] for c in row if c.column==3)
    assert 'consider saving' not in text
    wb.close()


def test_packet_shows_every_loss_caution_beside_its_opponent_not_only_global_top_three(tmp_path):
    wb=load_workbook(write_workbook(loss_payload(),tmp_path/'cautions.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb)
    head=next(c.row for c in wb['Captain Packet']['A'] if isinstance(c.value,str) and c.value.startswith('Recorded options preview'))
    text=book.display('Captain Packet',f'C{head+1}')
    assert 'Loss cautions' in text
    assert text.index('1-8')<text.index('0-6')<text.index('2-3')<text.index('0-1')
    assert 'Only 1 meeting' in text
    # A later opponent block must not leak through the unused roster slots.
    second=book.display('Captain Packet',f'C{head+2}')
    assert '1-8' not in second and '0-6' not in second
    book.set('Lineup Lab','C12','Unavailable')
    text=book.display('Captain Packet',f'C{head+1}')
    assert '0-1 (1 meeting)' not in text
    wb.close()


def test_packet_opponent_caution_survives_other_opponents_higher_global_losses(tmp_path):
    payload=_payload(crowded=True)
    payload['evidence']=_games(1,10,'LLLLL')+_games(2,10,'LLLL')+_games(3,10,'LLL')+_games(4,11,'LL')
    wb=load_workbook(write_workbook(payload,tmp_path/'local-caution.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb)
    head=next(c.row for c in wb['Captain Packet']['A'] if isinstance(c.value,str) and c.value.startswith('Recorded options preview'))
    line=book.display('Captain Packet',f'C{head+2}')
    assert '0-2 (2 meetings)' in line and 'SMALL SAMPLE' in line
    assert all(record not in line for record in ('0-5','0-4','0-3'))
    wb.close()


def test_paper_surfaces_expose_deeper_records_when_a_thin_leader_hides_them(tmp_path):
    payload=_payload();payload['evidence']=_games(1,10,'W')+_games(3,10,'WWWWWWWWWL')+_games(2,10,'WWWWWWWWLL')
    wb=load_workbook(write_workbook(payload,tmp_path/'deeper.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb);book.set('Command Center','C6','Cam Cole')
    head=next(c.row for c in wb['Captain Packet']['A'] if isinstance(c.value,str) and c.value.startswith('Recorded options preview'))
    packet=book.display('Captain Packet',f'C{head+1}')
    assert 'Deeper direct records' in packet and '8-2 (10 meetings)' in packet
    visible='\n'.join(str(book.display('Command Center',c.coordinate) or '') for row in wb['Command Center'] for c in row if c.column==2)
    assert 'Deeper direct records' in visible and '8-2 (10 meetings)' in visible
    assert packet.index('1-0 (1 meeting)')<packet.index('9-1 (10 meetings)')<packet.index('8-2 (10 meetings)')
    book.set('Lineup Lab','C13','Unavailable')
    assert '8-2 (10 meetings)' not in book.display('Captain Packet',f'C{head+1}')
    wb.close()


def test_paper_skill_cap_reference_warns_at_the_same_named_candidate_without_reranking(tmp_path):
    payload=_payload()
    for p in payload['players']:
        if p['id']==3:
            for history in p['team_history']:history['skill_level']=3
    wb=load_workbook(write_workbook(payload,tmp_path/'budget.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb);book.set('Command Center','C6','Cam Cole')
    book.set('Lineup Lab','C7',7)
    cc=book.value('Engine',book.name('wr_NsCard').ref)
    assert 'Reference comparison: this choice exceeds your entered skill budget' in cc
    head=next(c.row for c in wb['Captain Packet']['A'] if isinstance(c.value,str) and c.value.startswith('Recorded options preview'))
    packet=book.display('Captain Packet',f'C{head+1}')
    assert 'Reference comparison: this choice exceeds your entered skill budget' in packet
    assert 'not a match-legality check' in packet
    book.set('Lineup Lab','C7',23)
    assert 'this choice exceeds' not in book.value('Engine',book.name('wr_NsCard').ref)
    book.set('Lineup Lab','C7','')
    assert 'Reference budget' not in book.display('Captain Packet',f'C{head+1}')
    wb.close()


@pytest.mark.parametrize('state',['known','higher','missing','played','unavailable','no_players','no_opponents','unknown_played','equal_skills'])
def test_workbook_candidate_budget_uses_the_existing_phone_reference_arithmetic(tmp_path,state):
    payload=_payload()
    for p in payload['players']:
        if p['id']==3 and state not in ('missing','unknown_played'):
            for h in p['team_history']:h['skill_level']=3
        if p['id'] in (1,2,3) and state=='equal_skills':
            for h in p['team_history']:h['skill_level']=4
    wb=load_workbook(write_workbook(payload,tmp_path/'arithmetic.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb);cap=23 if state=='higher' else 7;book.set('Lineup Lab','C7',cap)
    remaining={1,2,3};played={};slots=2
    if state in ('played','unknown_played'):
        remaining.remove(3);book.set('Lineup Lab','D14','Played');book.set('Lineup Lab','C24','Played');slots=1;played={3:0 if state=='unknown_played' else 3}
    if state=='unavailable':remaining.remove(2);book.set('Lineup Lab','C13','Unavailable')
    if state=='no_players':
        remaining=set()
        for r in (12,13,14):book.set('Lineup Lab',f'C{r}','Unavailable')
    if state=='no_opponents':
        slots=0
        for r in (23,24):book.set('Lineup Lab',f'C{r}','Played')
    war=_python_war_room_for(payload)
    actual=named_values(book,'wr_CandidateBudget')[:3]
    unknown=(1 if state in ('missing','unknown_played') else 0)
    for i,pid in enumerate((1,2,3)):
        b=lineup_budget(war,cap=cap,slots=slots+len(played),played=played,remaining=remaining,sending=pid)
        expected=pid in remaining and slots>0 and b['sending_keeps_legal'] is False
        assert bool(actual[i])==expected,(state,pid,actual,b)
        if expected:assert ('incomplete' in actual[i])==bool(unknown)
    scope=book.value('Engine',book.name('wr_BudgetScope').ref)
    if state in ('missing','unknown_played','no_players'):assert 'No choice is cleared' in scope
    if not slots:assert scope==''
    wb.close()


@pytest.mark.parametrize('state',['initial','unavailable','played','opponent_played','none_left','live_not_winning'])
def test_workbook_live_flags_match_the_phone_engine_for_the_same_marks(tmp_path,state):
    payload=_payload()
    if state=='live_not_winning':
        payload['evidence']=_games(1,10,'LLLL')+_games(2,10,'W')+_games(3,10,'W')
    wb=load_workbook(write_workbook(payload,tmp_path/'live.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb);book.set('Lineup Lab','C5',SHARKS);book.set('Lineup Lab','C6',FALCONS)
    book.set('Lineup Lab','C9',book.value('Engine',book.name('uc_PlanKey').ref))
    remaining={1,2,3};unplayed={10,11}
    if state in ('unavailable','played'):
        remaining.remove(3);book.set('Lineup Lab','C14' if state=='unavailable' else 'D14','Unavailable' if state=='unavailable' else 'Played')
    if state=='opponent_played':unplayed.remove(11);book.set('Lineup Lab','C24','Played')
    if state=='none_left':
        remaining=set()
        for row in (12,13,14):book.set('Lineup Lab',f'D{row}','Played')
    if state=='live_not_winning':remaining.remove(1);book.set('Lineup Lab','C12','Unavailable')
    expected=[r for r in live_threats(_python_war_room_for(payload),remaining=remaining) if r['opponent']['id'] in unplayed]
    rows=list(zip(named_values(book,'wr_OppPids'),named_values(book,'wr_LiveThreatOrder'),named_values(book,'wr_LiveThreatUnknown'),named_values(book,'wr_LiveThreatText')))
    actual=sorted((order,pid,bool(unknown),text) for pid,order,unknown,text in rows if order!='')
    assert [(pid,unknown) for _,pid,unknown,_ in actual]==[(r['opponent']['id'],r['unknown']) for r in expected]
    if state in ('unavailable','played'):assert 'unknown vs who is left, not cleared' in actual[0][3]
    # Verify the visible views consume the live outputs, not just a hidden calculation.
    for sheet in ('War Room','Captain Packet'):
        displayed=[book.display(sheet,c.coordinate) for row in wb[sheet] for c in row
                   if isinstance(c.value,str) and c.value.startswith('=') and 'LiveThreatText' in c.value]
        for _,_,_,text in actual[:3]:assert any(text in line for line in displayed),(sheet,actual,displayed)
        assert all('Dangerous:' not in line for line in displayed)
    assert 'ignores tonight' in wb['Scouting Cards']['A5'].value
    wb.close()


@pytest.mark.parametrize('width',[390,1280])
def test_html_loss_lists_follow_top_risks_and_keep_positive_order(tmp_path,width):
    path=tmp_path/'parity.html';path.write_text(render(loss_payload(),built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page=browser.new_page(viewport={'width':width,'height':900});errors=[];page.on('pageerror',lambda e:errors.append(str(e)));page.goto(path.as_uri())
            page.locator('#next-send .ns-more summary').click()
            assert page.locator('#next-send [data-send-our]').evaluate_all('es=>es.map(e=>e.dataset.sendOur)')==['5','6']
            for text in (page.locator('#next-send').inner_text(),page.locator('#scouting-cards .quick-read').first.inner_text(),page.locator('#wr-risks').inner_text()):
                assert text.index('1-8')<text.index('0-6')<text.index('2-3')<text.index('0-1')
            assert not errors and not page.evaluate('document.documentElement.scrollWidth>innerWidth')
        finally:browser.close()


def test_workbook_single_loss_caution_uses_the_phone_and_packet_wording(tmp_path):
    wb=load_workbook(write_workbook(loss_payload(),tmp_path/'avoid-words.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb);book.set('Command Center','C6','Cam Cole')
    line=book.value('Engine',book.name('wr_NsAvoidLine').ref)
    assert line.startswith('⚠ Avoid records — recorded losses, not exclusions: ')
    assert '0-1 (1 meeting) · SMALL SAMPLE · Lost once · Only 1 meeting' in line
    assert '⚠ Avoid: ' not in line and line.count('Lost once')==1
    wb.close()


def test_packet_evidence_basis_keeps_direction_beside_small_sample(tmp_path):
    wb=load_workbook(write_workbook(loss_payload(),tmp_path/'basis.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb);ws=wb['Captain Packet']
    basis=[book.display('Captain Packet',f'K{c.row}') for c in ws['K']
           if isinstance(c.value,str) and 'Concerning direct' in c.value and 'SMALL SAMPLE' in c.value]
    shown=[b for b in basis if b]
    assert 'Favorable direct · SMALL SAMPLE' in shown and 'Concerning direct · SMALL SAMPLE' in shown
    assert 'Concerning direct' in shown and 'SMALL SAMPLE' not in shown
    wb.close()


def test_packet_blind_send_shows_the_phone_small_sample_pairing_line(tmp_path):
    payload=_payload()
    payload['evidence']=_games(1,10,'W')+_games(1,11,'WW')+_games(2,10,'L')
    expected='3 direct meetings total across 2 opponents · 2 small-sample pairings'
    path=tmp_path/'blind.html';path.write_text(render(payload,built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'),encoding='utf-8')
    with sync_playwright() as pw:
        browser=pw.chromium.launch()
        try:
            page=browser.new_page(viewport={'width':390,'height':900});page.goto(path.as_uri())
            assert expected in page.locator('.blind-send .bs-evidence').inner_text()
        finally:browser.close()
    wb=load_workbook(write_workbook(payload,tmp_path/'blind.xlsx',built_at='2026-10-07 18:00 UTC',viewer_member_external_id='1001'))
    book=Workbook(wb)
    row=next(c.row for c in wb['Captain Packet']['A'] if c.value=='IF WE PUT UP FIRST')
    assert expected in book.display('Captain Packet',f'C{row}')
    wb.close()
