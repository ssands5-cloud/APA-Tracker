"""Tests for ui/export_excel_opponent_volatility.py."""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.opponent_volatility import build_profile
from ui.export_excel_opponent_volatility import (
    PLAYERS_SHEET,
    SUMMARY_SHEET,
    build_workbook,
    write_workbook,
)


def roster_player(external_id, name="P", player_id=None):
    resolved_id = player_id if player_id is not None else hash(external_id) % 1000
    return {"player_id": resolved_id, "player_external_id": external_id, "player_name": name}


class TestBuildWorkbook:
    def test_both_sheets_present(self):
        roster = [roster_player("A", player_id=1)]
        profile = build_profile("T-OPP", "Corner Pockets", "Fall 2026", roster, {})
        workbook = build_workbook(profile)
        assert SUMMARY_SHEET in workbook.sheetnames
        assert PLAYERS_SHEET in workbook.sheetnames

    def test_players_sheet_includes_every_roster_player_including_nulls(self):
        roster = [roster_player("A", player_id=1), roster_player("B", player_id=2)]
        trends = {1: {"volatility": 0.2, "sample_size": 5}}
        profile = build_profile("T-OPP", "Corner Pockets", "Fall 2026", roster, trends)
        sheet = build_workbook(profile)[PLAYERS_SHEET]
        assert sheet.max_row == 3  # header + 2 players

    def test_write_workbook_produces_a_real_readable_file(self, tmp_path):
        profile = build_profile("T-OPP", "Corner Pockets", "Fall 2026", [], {})
        out_path = write_workbook(profile, tmp_path / "opponent_volatility.xlsx")

        assert out_path.exists()
        reopened = load_workbook(out_path)
        assert SUMMARY_SHEET in reopened.sheetnames
