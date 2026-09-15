"""Tests for ui/export_excel_player_vs_player.py.

Builds real PlayerVsPlayerExportRow fixtures via
analytics.player_vs_player_matrix.build_matrix_export (the same real
combination path the builder script will use) and checks the resulting
workbook: no empty placeholder sheet, deterministic fixed column widths,
UNKNOWN rows present, and the standalone (non-patching) posture.
"""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.pairing_evidence import EvidenceLabel, PairingEvidence, PairingEvidenceMatrix
from analytics.player_vs_player_matrix import build_matrix_export
from database.models import PlayerHeadToHead
from ui.export_excel_player_vs_player import (
    HISTORY_COLUMNS,
    HISTORY_SHEET,
    SUMMARY_COLUMNS,
    SUMMARY_SHEET,
    build_workbook,
    write_workbook,
)


def _pairing(player_id, opponent_id, label, **kwargs):
    defaults = dict(
        player_external_id=f"P-{player_id}", player_name=f"Player {player_id}",
        player_skill_level=5, opponent_external_id=f"OPP-{opponent_id}",
        opponent_name=f"Opponent {opponent_id}", opponent_skill_level=4,
        format="8-Ball Open", session_name="Fall 2026",
        observed_win_rate=None, direct_evidence_count=0,
        modeled_win_probability=None, model_source=None,
    )
    defaults.update(kwargs)
    return PairingEvidence(player_id=player_id, opponent_id=opponent_id,
                            evidence_label=label, **defaults)


def _matrix(pairings):
    counts = {label.value: 0 for label in EvidenceLabel}
    for p in pairings:
        counts[p.evidence_label.value] += 1
    counts["total_feasible_pairings"] = len(pairings)
    return PairingEvidenceMatrix(
        our_team_external_id="OUR", opponent_team_external_id="THEIRS",
        format="8-Ball Open", session_name="Fall 2026",
        expected_pairings=tuple((p.player_id, p.opponent_id) for p in pairings),
        pairings=tuple(pairings), counts=counts,
        our_roster_available=True, opponent_roster_available=True,
    )


def _game(player_id=1, opponent_id=10, result="W", match_id=1):
    return PlayerHeadToHead(
        player_id=player_id, opponent_id=opponent_id, match_id=match_id, result=result,
        own_skill_level=5, opponent_skill_level=4,
        format="8-Ball Open", session_name="Fall 2026",
    )


class TestNoEmptyPlaceholder:
    def test_zero_feasible_pairings_produces_no_workbook_at_all(self):
        assert build_workbook([]) is None

    def test_write_workbook_writes_nothing_when_there_are_no_rows(self, tmp_path):
        out_path = tmp_path / "player_vs_player.xlsx"
        result = write_workbook([], out_path)

        assert result is None
        assert not out_path.exists()

    def test_no_real_games_anywhere_omits_the_history_sheet_entirely(self):
        rows = build_matrix_export(
            _matrix([_pairing(1, 10, EvidenceLabel.UNKNOWN, player_skill_level=None)]),
            histories={},
        )
        workbook = build_workbook(rows)

        assert SUMMARY_SHEET in workbook.sheetnames
        assert HISTORY_SHEET not in workbook.sheetnames


class TestSummarySheet:
    def test_every_pairing_appears_exactly_once_evidence_labels_included(self):
        rows = build_matrix_export(
            _matrix([
                _pairing(1, 10, EvidenceLabel.DIRECT, direct_evidence_count=1, observed_win_rate=1.0),
                _pairing(1, 11, EvidenceLabel.INDIRECT, modeled_win_probability=0.6,
                         model_source="analytics.head_to_head:validated-skill-only"),
                _pairing(2, 10, EvidenceLabel.UNKNOWN, player_skill_level=None),
            ]),
            histories={(1, 10): [_game(match_id=1)]},
        )
        workbook = build_workbook(rows)
        sheet = workbook[SUMMARY_SHEET]

        assert sheet.max_row == 4  # header + 3 rows
        header = [cell.value for cell in sheet[1]]
        assert header == list(SUMMARY_COLUMNS)
        evidence_col = SUMMARY_COLUMNS.index("Evidence Label") + 1
        labels = {sheet.cell(row=r, column=evidence_col).value for r in range(2, 5)}
        assert labels == {"DIRECT", "INDIRECT", "UNKNOWN"}

    def test_column_widths_are_fixed_not_content_derived(self):
        rows = build_matrix_export(
            _matrix([_pairing(
                1, 10, EvidenceLabel.DIRECT, direct_evidence_count=1, observed_win_rate=1.0,
                player_name="A Very Long Real Player Name That Would Widen An Auto-Sized Column",
            )]),
            histories={(1, 10): [_game(match_id=1)]},
        )
        workbook = build_workbook(rows)
        sheet = workbook[SUMMARY_SHEET]

        player_col_letter = sheet.cell(row=1, column=SUMMARY_COLUMNS.index("Player") + 1).column_letter
        assert sheet.column_dimensions[player_col_letter].width == 20

    def test_observed_win_rate_is_null_outside_direct(self):
        rows = build_matrix_export(
            _matrix([_pairing(1, 10, EvidenceLabel.INDIRECT, modeled_win_probability=0.6,
                               model_source="analytics.head_to_head:validated-skill-only")]),
            histories={},
        )
        workbook = build_workbook(rows)
        sheet = workbook[SUMMARY_SHEET]
        rate_col = SUMMARY_COLUMNS.index("Observed Win Rate") + 1
        assert sheet.cell(row=2, column=rate_col).value is None


class TestHistorySheet:
    def test_real_games_appear_with_the_documented_columns(self, tmp_path):
        rows = build_matrix_export(
            _matrix([_pairing(1, 10, EvidenceLabel.DIRECT, direct_evidence_count=1, observed_win_rate=1.0)]),
            histories={(1, 10): [_game(match_id=1), _game(match_id=2, result="L")]},
        )
        out_path = write_workbook(rows, tmp_path / "pvp.xlsx")

        workbook = load_workbook(out_path)
        assert HISTORY_SHEET in workbook.sheetnames
        sheet = workbook[HISTORY_SHEET]
        header = [cell.value for cell in sheet[1]]
        assert header == list(HISTORY_COLUMNS)
        assert sheet.max_row == 3  # header + 2 games


class TestStandaloneWorkbook:
    def test_write_workbook_produces_a_real_readable_file(self, tmp_path):
        rows = build_matrix_export(
            _matrix([_pairing(1, 10, EvidenceLabel.DIRECT, direct_evidence_count=1, observed_win_rate=1.0)]),
            histories={(1, 10): [_game(match_id=1)]},
        )
        out_path = write_workbook(rows, tmp_path / "player_vs_player.xlsx")

        assert out_path is not None
        assert out_path.exists()
        reopened = load_workbook(out_path)
        assert SUMMARY_SHEET in reopened.sheetnames
