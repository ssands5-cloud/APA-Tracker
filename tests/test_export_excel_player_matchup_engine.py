"""Tests for ui/export_excel_player_matchup_engine.py."""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.pairing_evidence import EvidenceLabel, PairingEvidence
from analytics.player_matchup_engine import build_player_matchup_report
from ui.export_excel_player_matchup_engine import COLUMNS, write_workbook


def _pairing(**overrides) -> PairingEvidence:
    base = dict(
        player_id=1, player_external_id="P1", player_name="Ann", player_skill_level=5,
        opponent_id=2, opponent_external_id="P2", opponent_name="Bob", opponent_skill_level=4,
        format="8-Ball Open", session_name="Fall 2026",
        evidence_label=EvidenceLabel.DIRECT, observed_win_rate=0.75, direct_evidence_count=4,
        direct_wins=3, direct_losses=1,
        modeled_win_probability=0.7, model_source="analytics.head_to_head",
    )
    base.update(overrides)
    return PairingEvidence(**base)


def _build(**kwargs):
    return build_player_matchup_report(_pairing(), "OUR1", "OPP1", **kwargs)


def _col(name: str) -> int:
    return COLUMNS.index(name) + 1


class TestWriteWorkbook:
    def test_a_real_row_per_report(self, tmp_path):
        path = write_workbook([_build()], tmp_path / "player_matchup_engine.xlsx")
        workbook = load_workbook(path)
        sheet = workbook["Player Matchup Engine"]
        assert sheet.max_row == 2  # header + one report
        assert sheet.cell(row=2, column=_col("Our Team")).value == "OUR1"
        assert sheet.cell(row=2, column=_col("Opponent Team")).value == "OPP1"
        assert sheet.cell(row=2, column=_col("Player")).value == "Ann"
        assert sheet.cell(row=2, column=_col("Opponent")).value == "Bob"

    def test_direct_matches_column_is_renamed_from_the_old_direct_games(self):
        """GPT audit P1: direct_evidence_count counts distinct MATCHES, not
        games -- the old "Direct Games" header mislabeled it."""
        assert "Direct Games" not in COLUMNS
        assert "Direct Matches" in COLUMNS

    def test_direct_win_loss_record_is_a_real_column(self, tmp_path):
        path = write_workbook([_build()], tmp_path / "out.xlsx")
        sheet = load_workbook(path)["Player Matchup Engine"]
        assert sheet.cell(row=2, column=_col("Direct W-L")).value == "3-1"

    def test_headers_are_frozen_and_filterable(self, tmp_path):
        path = write_workbook([_build()], tmp_path / "out.xlsx")
        sheet = load_workbook(path)["Player Matchup Engine"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None

    def test_an_empty_report_list_still_produces_an_openable_workbook(self, tmp_path):
        path = write_workbook([], tmp_path / "empty.xlsx")
        workbook = load_workbook(path)
        assert workbook["Player Matchup Engine"].max_row == 1  # header only
