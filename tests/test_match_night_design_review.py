"""Small exhaustive design counterexamples, not an M3 engine implementation.

Numbers below are hypothetical proposal utilities, never win estimates. These
tests expose assumptions; passing does not validate match-winning performance.
"""
from itertools import permutations
from math import comb, factorial, perm
import pytest

from analytics.ultimate_coach_war_room import lineup_budget, next_send
from tests.test_recommendation_quality_risks import member, room


def two_player_boards(values):
    # Enumerate only this toy 2x2 audit table; no application solver is added.
    return [tuple(values[i][j] for i, j in enumerate(order)) for order in permutations(range(2))]


def test_m301_best_static_and_safest_static_board_can_both_be_unenforceable():
    values = ((2, -1), (-1, 2))
    boards = two_player_boards(values)
    assert max(map(sum, boards)) == 4
    safest_static = max(boards, key=lambda b: (min(b), sum(b)))
    assert safest_static == (2, 2)
    # Whichever player we lead, their responding captain can choose the cross.
    for lead in range(2):
        worst_reply = min(values[lead][j] + values[1-lead][1-j] for j in range(2))
        assert worst_reply == -2


def test_m302_local_first_send_can_consume_the_only_counter():
    w = room([(1, 10, 'WWW'), (2, 10, 'WWL'), (1, 11, 'WWWWW'), (2, 11, 'LLL')],
             theirs=[member(10, 'Opponent X'), member(11, 'Opponent Y')])
    answer = next_send(w, 0)
    assert answer['ranked'][0]['member']['id'] == 1
    assert 'consider saving' in answer['ranked'][0]['save']
    categories = { (r['member']['id'], b['opponent']['id']): r['category']
                   for b in w['blocks'] for r in b['rows'] }
    assert [categories[pair] for pair in ((1, 10), (2, 11))] == ['G', 'R']
    assert [categories[pair] for pair in ((2, 10), (1, 11))] == ['G', 'G']


def test_m303_ordered_confidence_labels_do_not_determine_additive_utility():
    # All three scales preserve Strong > Solid > Unknown, but choose differently.
    for strong, expected in ((1.5, 'two_solids'), (2, 'tie'), (3, 'strong_unknown')):
        a, b = strong + 0, 1 + 1
        selected = 'tie' if a == b else 'strong_unknown' if a > b else 'two_solids'
        assert selected == expected


def test_m304_unknown_zero_can_look_safer_than_recorded_caution():
    unknown_board, documented_board = (0, 0), (2, -1)
    assert min(unknown_board) > min(documented_board)
    assert sum(unknown_board) < sum(documented_board)
    # This reverses solely because of the chosen objective, not new evidence.


def test_m305_total_and_weakest_link_choose_different_concession_plans():
    concentrated, balanced = (2, 2, -1), (1, 1, 0.5)
    assert sum(concentrated) > sum(balanced)
    assert min(concentrated) < min(balanced)
    # "Must win" alone cannot choose the objective without match-score context.


def test_m306_full_unconstrained_board_count_is_not_single_roster_permutations():
    assert perm(8, 5) == 6720
    assert comb(8, 5) ** 2 * factorial(5) == 376320
    assert comb(8, 5) ** 2 * factorial(5) > perm(8, 5)


def test_m307_remaining_budget_can_remove_an_attractive_send():
    w = room([], ours=[member(1, 'Player A', sl=7), member(2, 'Player B', sl=6), member(3, 'Player C', sl=2)])
    a = lineup_budget(w, cap=8, slots=2, sending=1)
    b = lineup_budget(w, cap=8, slots=2, sending=2)
    assert a['feasible'] is True  # B+C fits this supplied arithmetic constraint.
    assert a['sending_keeps_legal'] is False  # A+C does not.
    assert b['sending_keeps_legal'] is True
    assert 'not a match-legality check' in a['note']


def test_m308_partial_skill_data_cannot_clear_a_board_from_feasible_field_alone():
    w = room([], ours=[member(1, 'Player A', sl=7), member(2, 'Player B', sl=6), member(3, 'Player C', sl=None)])
    result = lineup_budget(w, cap=13, slots=2)
    assert result['feasible'] is True
    assert result['complete'] is False and result['unknown_sl'] == 1
    assert 'no lineup cleared' in result['note']


def test_m309_final_pairings_do_not_identify_the_put_up_information_path():
    ours_lead = [('our lead', 'A'), ('their response', 'X'), ('their lead', 'Y'), ('our response', 'B')]
    theirs_lead = [('their lead', 'X'), ('our response', 'A'), ('our lead', 'B'), ('their response', 'Y')]
    def recorded_pairs(actions):
        return {tuple(next(player for role, player in actions[i:i+2] if role.startswith(side))
                      for side in ('our', 'their')) for i in range(0, len(actions), 2)}
    assert ours_lead != theirs_lead
    assert recorded_pairs(ours_lead) == recorded_pairs(theirs_lead) == {('A', 'X'), ('B', 'Y')}
    # Same recorded pairings admit different first-player and response decisions.


def test_m310_actual_result_does_not_supply_unplayed_alternative_outcome():
    observed = {('A', 'X'): 'W', ('B', 'Y'): 'L'}
    alternative = [('B', 'X'), ('A', 'Y')]
    assert all(pair not in observed for pair in alternative)
    # A retrospective replay cannot score this alternative as two real wins.


def test_m311_one_player_cannot_fulfil_two_simultaneous_hold_back_promises():
    w = room([(1, 10, 'WWW'), (2, 10, 'WWL'), (1, 11, 'WWW'), (1, 12, 'WWW'),
              (2, 11, 'L'), (2, 12, 'L'), (3, 11, 'L'), (3, 12, 'L')],
             ours=[member(1, 'Player A'), member(2, 'Player B'), member(3, 'Player C')],
             theirs=[member(10, 'Opponent X'), member(11, 'Opponent Y'), member(12, 'Opponent Z')])
    warning = next_send(w, 0)['ranked'][0]['save']
    assert 'Opponent Y' in warning and 'Opponent Z' in warning
    greens = {b['opponent']['id']: {r['member']['id'] for r in b['rows'] if r['category']=='G'}
              for b in w['blocks'][1:]}
    attainable = max(sum(chosen[i] in greens[opp] for i, opp in enumerate((11, 12)))
                     for chosen in permutations((1, 2, 3), 2))
    assert attainable == 1


def test_m312_implemented_plan_can_choose_differently_when_opponent_budget_is_known():
    from analytics.match_night_engine import match_plan
    from tests.test_match_night_engine import night
    grid = [['Caution', 'Caution', 'Thin'],
            ['Thin', 'Solid', 'Strong'],
            ['Solid', 'Strong', 'Unknown']]
    ours, theirs, values = night([('A', 4), ('B', 4), ('C', 4)],
                                [('X', 7), ('Y', 7), ('Z', 2)], grid)
    plan = match_plan(ours, theirs, values, games_left=2, cap=8)
    v = [[values[(o['id'], t['id'])]['points'] for t in theirs] for o in ours]

    def lead_values(enforce_opponent_budget):
        # Explicit two-round sequence: our lead, their response, their lead,
        # our best remaining response. Only the opponent's known 9-SL budget
        # differs. Values are existing implementation utility, not predicted wins.
        return [min(v[i][j] + min(max(v[h][k] for h in range(3) if h != i)
                    for k in range(3) if k != j and
                    (not enforce_opponent_budget or theirs[j]['sl'] + theirs[k]['sl'] <= 9))
                    for j in range(3)) for i in range(3)]

    unrestricted = lead_values(False)
    constrained = lead_values(True)
    assert unrestricted == [0, 2, 1]
    assert constrained == [3, 2, 1]
    assert plan['lead']['our'] == 'B' and plan['total_if_we_lead'] == 2
    assert constrained.index(max(constrained)) == 0  # A under known opponent budget.
    assert theirs[0]['sl'] + theirs[1]['sl'] > 9


def m313_budget_excluded_hold_case():
    from tests.test_match_night_engine import night
    return night([('A', 4), ('B', 4), ('C', 4)],
                 [('X', 7), ('Y', 2), ('Z', 2)],
                 [['Caution', 'Solid', 'Unknown'],
                  ['Unknown', 'Strong', 'Solid'],
                  ['Solid', 'Strong', 'Thin']])


def test_m313_hold_reason_excludes_budget_excluded_opponent():
    """Hold instructions use the same feasible opponents as the response search."""
    from analytics.match_night_engine import match_plan, plan_text
    ours, theirs, values = m313_budget_excluded_hold_case()
    plan = match_plan(ours, theirs, values, games_left=2, cap=8, opp_cap=8)
    assert theirs[0]['sl'] + min(p['sl'] for p in theirs[1:]) > 8
    assert {r['their'] for r in plan['responses']} == {'Y', 'Z'}
    for field in ('plan_if_we_lead', 'plan_if_they_lead', 'best_board', 'safest_board'):
        assert all(p['their'] != 'X' for p in plan[field])
    assert plan['lead']['our'] == 'B'
    assert plan['holds'] == [{'our': 'C', 'why': 'the plan is weaker without them', 'covers': []}]
    assert 'answer to X' not in plan_text(plan)['holds']
    unrestricted = match_plan(ours, theirs, values, games_left=2, cap=8)
    assert any('X' in h['covers'] for h in unrestricted['holds'])


def test_m313_phone_engine_excludes_budget_excluded_hold_reason(tmp_path):
    """Actual browser engine mirrors the corrected feasible hold coverage."""
    from playwright.sync_api import sync_playwright
    from analytics.match_night_engine import match_plan
    from tests.test_excel_war_room_formulas import _payload
    from ui.ultimate_coach import render
    ours, theirs, values = m313_budget_excluded_hold_case()
    expected = match_plan(ours, theirs, values, games_left=2, cap=8, opp_cap=8)
    path = tmp_path / 'm313.html'
    path.write_text(render(_payload(), built_at='2026-10-20 12:00 UTC',
                           viewer_member_external_id='1001'), encoding='utf-8')
    js_values = {f'{a}|{b}': v for (a, b), v in values.items()}
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        try:
            for width in (390, 1280):
                page = browser.new_page(viewport={'width': width, 'height': 850})
                errors = []
                page.on('pageerror', lambda e: errors.append(str(e)))
                page.goto(path.as_uri())
                actual = page.evaluate('a => window.__ucMatchPlan(...a)',
                                       [ours, theirs, js_values, 2, 8, 0, 8, 0])
                assert actual == expected
                assert actual['holds'][0]['covers'] == []
                assert 'answer to X' not in actual['holds'][0]['why']
                assert errors == []
                page.close()
        finally:
            browser.close()



@pytest.mark.parametrize('recommended,ann_record,bea_record', [
    ('Ann Archer', 'WWWWW', 'LLL'), ('Bea Baker', 'W', 'WWWWWWWWLL')])
def test_f2_future_opportunity_matches_the_plan_instruction(tmp_path, recommended, ann_record, bea_record):
    """Keep future evidence, contextualize the chosen send, retain other save warnings."""
    from playwright.sync_api import sync_playwright
    from tests.test_excel_war_room_formulas import _games, _payload
    from ui.ultimate_coach import render
    payload = _payload()
    payload['players'] = [p for p in payload['players'] if p['id'] in (1, 2, 10, 11)]
    for player in payload['players']:
        for history in player['team_history']:
            history['skill_level'] = 4
    payload['evidence'] = (_games(1, 10, ann_record) + _games(1, 11, 'WW')
                           + _games(2, 10, bea_record) + _games(2, 11, 'LLL'))
    for game in payload['evidence']:
        game['own_skill_level'] = game['opponent_skill_level'] = 4
    path = tmp_path / 'f2.html'
    path.write_text(render(payload, built_at='2026-10-20 13:00 UTC',
                           viewer_member_external_id='1001'), encoding='utf-8')
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        try:
            for width in (390, 1280):
                page = browser.new_page(viewport={'width': width, 'height': 850})
                errors = []
                page.on('pageerror', lambda e: errors.append(str(e)))
                page.goto(path.as_uri())
                page.select_option('#team-a', 'sharks-a|d1|Fall 2026')
                page.select_option('#team-b', 'falcons-a|d1|Fall 2026')
                page.locator('#next-send [data-next="10"]').click()
                answer = page.locator('#next-send .ns-plan-answer')
                assert answer.locator('b').inner_text() == recommended
                assert answer.locator('.mp-level').inner_text() == 'Strong'
                row = page.locator('#next-send .ns-medal').filter(
                    has=page.locator('[data-send-our="1"]'))
                assert row.count() == 1
                warning = row.locator('.ns-save')
                if recommended == 'Ann Archer':
                    assert 'consider saving' not in warning.inner_text()
                    assert 'Sending now spends' in warning.inner_text()
                    assert 'opportunity spent' in answer.inner_text()
                else:
                    assert 'consider saving' in warning.inner_text()
                assert 'Eve Ellis' in warning.inner_text()
                assert '2-0 in 2 meetings' in warning.inner_text()
                assert errors == []
                assert not page.evaluate('document.documentElement.scrollWidth > innerWidth')
                page.close()
        finally:
            browser.close()


def test_f2_paper_future_opportunity_matches_the_fresh_plan(tmp_path):
    from openpyxl import load_workbook
    from tests.excel_formula_eval import Workbook
    from tests.test_excel_war_room_formulas import _games, _payload
    from ui.export_excel_ultimate_coach import write_workbook
    payload = _payload()
    payload['players'] = [p for p in payload['players'] if p['id'] in (1, 2, 10, 11)]
    for player in payload['players']:
        for history in player['team_history']:
            history['skill_level'] = 4
    payload['evidence'] = (_games(1, 10, 'WWWWW') + _games(1, 11, 'WW')
                           + _games(2, 10, 'LLL') + _games(2, 11, 'LLL'))
    for game in payload['evidence']:
        game['own_skill_level'] = game['opponent_skill_level'] = 4
    wb = load_workbook(write_workbook(payload, tmp_path / 'f2.xlsx',
                                    built_at='2026-10-07 18:00 UTC', viewer_member_external_id='1001'))
    book = Workbook(wb)
    book.set('Command Center', 'C6', 'Cam Cole')
    shown = book.value('Engine', book.name('wr_NsCard').ref)
    assert 'Sending now spends' in shown and 'consider saving' not in shown
    answer = book.value('Engine', book.name('wr_PlanAnswer').ref)
    assert 'Ann Archer (Strong)' in answer and 'opportunity spent' in answer
    top = next(c.row for c in wb['Captain Packet']['A']
               if isinstance(c.value, str) and c.value.startswith('Recorded options preview'))
    packet = book.display('Captain Packet', f'C{top+1}')
    assert 'Sending now spends' in packet and 'consider saving' not in packet
    assert 'Eve Ellis' in packet and '2-0 in 2 meetings' in packet
    book.set('Lineup Lab', 'C7', 23)
    assert 're-plan on the phone' in book.value('Engine', book.name('wr_PlanAnswer').ref)
    assert 'consider saving' in book.value('Engine', book.name('wr_NsCard').ref)
    wb.close()



def test_f3_equal_strength_alternative_is_explained_even_when_naive_matches_plan():
    """A roster-first selected answer must still explain rejected equal-value alternatives."""
    from tests.test_captain_decision_review import plan, answer
    p, values = plan({'Alice': 4, 'Bob': 4}, {'Xavier': 4, 'Yvonne': 4},
                     {('Alice', 'Xavier'): (4, 6), ('Alice', 'Yvonne'): (0, 3),
                      ('Bob', 'Xavier'): (4, 6), ('Bob', 'Yvonne'): (6, 7)}, games=2)
    assert values[('Alice', 'Xavier')]['level'] == values[('Bob', 'Xavier')]['level'] == 'Solid'
    response = answer(p, 'Xavier')
    assert response['our'] == 'Alice'
    assert 'Bob' in response['cost_later'] and 'equally supported' in response['cost_later']
    assert 'for Yvonne (Strong)' in response['cost_later']
    names = ['Alice', 'Bob']
    naive = max(names, key=lambda n: (values[(n, 'Xavier')]['points'], -names.index(n)))
    assert naive == response['our']  # >= alone never reaches the proposed explanation.
    proposed_guard = (naive != response['our'] and
                      values[(naive, 'Xavier')]['points'] >= values[(response['our'], 'Xavier')]['points'])
    assert proposed_guard is False
    chosen_total = values[('Alice', 'Xavier')]['points'] + values[('Bob', 'Yvonne')]['points']
    alternative_total = values[('Bob', 'Xavier')]['points'] + values[('Alice', 'Yvonne')]['points']
    assert chosen_total - alternative_total == 9  # Existing utility, never match points or win estimates.


def test_f3_no_cost_note_for_interchangeable_answers_or_unusable_alternatives():
    from tests.test_match_night_engine import night
    from analytics.match_night_engine import match_plan
    ours, theirs, values = night([('A', 4), ('B', 4), ('C', 7)], [('X', 4), ('Y', 4)],
                                [['Solid', 'Strong'], ['Solid', 'Strong'], ['Strong', 'Strong']])
    p = match_plan(ours, theirs, values, games_left=2, cap=8)
    assert all(r['cost_later'] == '' for r in p['responses'])
    assert all(r['our'] != 'C' for r in p['responses'])


def test_f3_future_cost_is_visible_on_phone_and_desktop(tmp_path):
    from playwright.sync_api import sync_playwright
    from tests.test_excel_war_room_formulas import _games, _payload
    from ui.ultimate_coach import render
    payload = _payload()
    payload['players'] = [p for p in payload['players'] if p['id'] in (1, 2, 10, 11)]
    for player in payload['players']:
        for history in player['team_history']:
            history['skill_level'] = 4
    payload['evidence'] = (_games(1, 10, 'WWWWLL') + _games(1, 11, 'LLL')
                           + _games(2, 10, 'WWWWLL') + _games(2, 11, 'WWWWWWL'))
    for game in payload['evidence']:
        game['own_skill_level'] = game['opponent_skill_level'] = 4
    path = tmp_path / 'f3.html'
    path.write_text(render(payload, built_at='2026-10-07 18:00 UTC',
                           viewer_member_external_id='1001'), encoding='utf-8')
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        try:
            for width in (390, 1280):
                page = browser.new_page(viewport={'width': width, 'height': 900})
                page.goto(path.as_uri())
                page.select_option('#team-a', 'sharks-a|d1|Fall 2026')
                page.select_option('#team-b', 'falcons-a|d1|Fall 2026')
                page.locator('#next-send [data-next="10"]').click()
                answer = page.locator('#next-send .ns-plan-answer')
                assert answer.locator('b').inner_text() == 'Ann Archer'
                cost = answer.locator('.mp-cost').inner_text()
                assert 'equally supported' in cost and 'keep Bea Baker available for Eve Ellis (Strong)' in cost
                assert not page.evaluate('document.documentElement.scrollWidth > innerWidth')
                page.close()
        finally:
            browser.close()


def test_f1_designated_answers_are_conditional_send_instructions(tmp_path):
    from playwright.sync_api import sync_playwright
    from tests.test_excel_war_room_formulas import _games, _payload, _hist
    from ui.ultimate_coach import render
    payload = _payload()
    payload['players'] = [p for p in payload['players'] if p['id'] in (1, 2, 3, 10, 11, 12)]
    for player in payload['players']:
        if player['id'] == 12:
            player['team_history'] = [_hist('falcons-a', 'd1', 'EIGHT', 4)]
        for history in player['team_history']:
            history['skill_level'] = 4
    payload['evidence'] = (_games(1, 10, 'WWWWWW') + _games(1, 11, 'WWWWWW')
                           + _games(2, 12, 'WWWWWW') + _games(2, 11, 'LLL'))
    for game in payload['evidence']:
        game['own_skill_level'] = game['opponent_skill_level'] = 4
    path = tmp_path / 'f1.html'
    path.write_text(render(payload, built_at='2026-10-20 13:00 UTC',
                           viewer_member_external_id='1001'), encoding='utf-8')
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        try:
            for width in (390, 1280):
                page = browser.new_page(viewport={'width': width, 'height': 900})
                page.goto(path.as_uri())
                page.select_option('#team-a', 'sharks-a|d1|Fall 2026')
                page.select_option('#team-b', 'falcons-a|d1|Fall 2026')
                holds = page.locator('#match-plan .mp-hold')
                assert holds.count() == 1
                assert holds.locator('.mp-label').text_content() == 'Designated answers'
                assert 'send when that matchup is put up' in holds.inner_text()
                assert 'Ann Archer' in holds.inner_text() and 'Cam Cole' in holds.inner_text()
                assert not page.evaluate('document.documentElement.scrollWidth > innerWidth')
                page.close()
        finally:
            browser.close()
