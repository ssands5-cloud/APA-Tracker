"""Tests for ui/export_excel_team_matchup_engine.py."""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.lineup_lab import LineupLabResult, LineupSlot
from analytics.pairing_evidence import EvidenceLabel, PairingEvidence, PairingEvidenceMatrix
from analytics.team_matchup_engine import build_team_matchup_report
from ui.export_excel_team_matchup_engine import write_workbook


def _matrix() -> PairingEvidenceMatrix:
    pairing = PairingEvidence(
        player_id=1, player_external_id="P1", player_name="Ann", player_skill_level=5,
        opponent_id=2, opponent_external_id="P2", opponent_name="Bob", opponent_skill_level=4,
        format="8-Ball Open", session_name="Fall 2026",
        evidence_label=EvidenceLabel.DIRECT, observed_win_rate=0.75, direct_evidence_count=4,
        modeled_win_probability=0.7, model_source="analytics.head_to_head",
    )
    return PairingEvidenceMatrix(
        our_team_external_id="OUR1", opponent_team_external_id="OPP1",
        format="8-Ball Open", session_name="Fall 2026",
        expected_pairings=((1, 2),), pairings=(pairing,),
        counts={"DIRECT": 1, "INDIRECT": 0, "UNKNOWN": 0, "total_feasible_pairings": 1},
        our_roster_available=True, opponent_roster_available=True,
    )


class TestWriteWorkbook:
    def test_three_sheets_per_real_scope(self, tmp_path):
        report = build_team_matchup_report(_matrix(), "Mark It Up", "Corner Pockets")
        path = write_workbook([report], tmp_path / "team_matchup_engine.xlsx")

        workbook = load_workbook(path)
        titles = workbook.sheetnames
        assert any("Corner Pockets" in t for t in titles)
        assert any("Rank" in t for t in titles)
        assert any("Lineup" in t for t in titles)

    def test_roster_sheet_carries_real_names(self, tmp_path):
        report = build_team_matchup_report(_matrix(), "Mark It Up", "Corner Pockets")
        path = write_workbook([report], tmp_path / "out.xlsx")

        workbook = load_workbook(path)
        roster_sheet = workbook[workbook.sheetnames[0]]
        values = [cell.value for row in roster_sheet.iter_rows() for cell in row]
        assert "Ann" in values
        assert "Bob" in values

    def test_lineup_sheet_reflects_a_real_approved_lineup(self, tmp_path):
        lineup_result = LineupLabResult(
            assignments=(
                LineupSlot(
                    player_id=1, player_name="Ann", player_skill_level=5,
                    opponent_id=2, opponent_name="Bob", opponent_skill_level=4,
                    evidence_label=EvidenceLabel.DIRECT, observed_win_rate=0.75,
                    direct_evidence_count=4, modeled_win_probability=0.7,
                    model_source="analytics.head_to_head", lineup_score=0.75,
                    lineup_score_source="direct",
                ),
            ),
            unassigned_players=(), unassigned_opponents=(),
            total_score=0.75, skill_total=5, is_legal=True, blocked_reason=None,
        )
        report = build_team_matchup_report(
            _matrix(), "Mark It Up", "Corner Pockets", lineup_result=lineup_result,
        )
        path = write_workbook([report], tmp_path / "out.xlsx")

        workbook = load_workbook(path)
        lineup_sheet = next(workbook[t] for t in workbook.sheetnames if "Lineup" in t)
        assert lineup_sheet.cell(row=2, column=1).value == 1
        assert lineup_sheet.cell(row=2, column=2).value == "Ann"

    def test_an_empty_report_list_still_produces_an_openable_workbook(self, tmp_path):
        path = write_workbook([], tmp_path / "empty.xlsx")
        workbook = load_workbook(path)
        assert len(workbook.sheetnames) >= 1

    def test_a_long_scope_name_never_truncates_the_rank_or_lineup_suffix(self, tmp_path):
        """Regression: truncating a whole "<scope> Lineup" candidate to
        Excel's 31-char sheet-name limit from the right previously chopped
        the suffix itself down to an unrecognizable "Lin" for any
        real opponent name/format long enough to push past the limit."""
        report = build_team_matchup_report(_matrix(), "Mark It Up", "Corner Pockets")
        path = write_workbook([report], tmp_path / "out.xlsx")

        workbook = load_workbook(path)
        assert any(t.endswith("Rank") for t in workbook.sheetnames)
        assert any(t.endswith("Lineup") for t in workbook.sheetnames)
        assert all(len(t) <= 31 for t in workbook.sheetnames)

    def test_duplicate_sheet_titles_are_disambiguated(self, tmp_path):
        """Two real scopes with the same opponent+format truncated name
        must not silently collide on one sheet."""
        r1 = build_team_matchup_report(_matrix(), "Mark It Up", "Corner Pockets")
        r2 = build_team_matchup_report(_matrix(), "Mark It Up", "Corner Pockets")
        path = write_workbook([r1, r2], tmp_path / "out.xlsx")

        workbook = load_workbook(path)
        assert len(set(workbook.sheetnames)) == len(workbook.sheetnames)
