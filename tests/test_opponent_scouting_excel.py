"""Tests for the Opponent_Scouting Excel sheet (ui/export_excel.py)."""

from __future__ import annotations

import json

from openpyxl import Workbook, load_workbook

from ui.export_excel import (
    OPPONENT_SCOUTING_COLUMNS,
    OPPONENT_SCOUTING_SHEET,
    append_opponent_scouting_sheet,
    opponent_scouting_rows,
)


def payload():
    return {
        "opponent_scouting": [
            {
                "opponent_id": "O1", "opponent_name": "Bob",
                "opponent_team_id": "T2", "opponent_team_name": "Corner Pockets",
                "times_faced": 3, "avg_matchup_score": 62.5, "avg_win_probability": 0.35,
                "opponent_volatility": 0.72, "is_danger_matchup": True,
                "danger_reasons": ["our average win probability against them (35%) is below the 40% threshold"],
            },
            {
                "opponent_id": "O2", "opponent_name": "Carol",
                "opponent_team_id": "T2", "opponent_team_name": "Corner Pockets",
                "times_faced": 1, "avg_matchup_score": None, "avg_win_probability": None,
                "opponent_volatility": None, "is_danger_matchup": False, "danger_reasons": [],
            },
        ],
    }


class TestOpponentScoutingRows:
    def test_flattens_every_real_field(self):
        rows = opponent_scouting_rows(payload())
        bob = next(r for r in rows if r["Opponent Player"] == "Bob")
        assert bob["Opponent Team"] == "Corner Pockets"
        assert bob["Times Faced"] == 3
        assert bob["Avg Matchup Score"] == 62.5
        assert bob["Avg Win Probability"] == 0.35
        assert bob["Opponent Volatility"] == 0.72
        assert bob["Danger Matchup"] == "Yes"
        assert "40%" in bob["Danger Reasons"]

    def test_missing_signals_show_as_no_data_not_zero(self):
        rows = opponent_scouting_rows(payload())
        carol = next(r for r in rows if r["Opponent Player"] == "Carol")
        assert carol["Avg Matchup Score"] == "No data"
        assert carol["Avg Win Probability"] == "No data"
        assert carol["Opponent Volatility"] == "No data"
        assert carol["Danger Matchup"] == "No"

    def test_no_opponents_at_all_is_an_empty_list(self):
        assert opponent_scouting_rows({}) == []


class TestAppendOpponentScoutingSheet:
    def test_creates_a_frozen_filterable_sheet_with_every_real_column(self, tmp_path):
        workbook_path = tmp_path / "workbook.xlsx"
        document_path = tmp_path / "lineups.json"
        workbook = Workbook()
        workbook.active.title = "Existing"
        workbook.save(workbook_path)
        document_path.write_text(json.dumps(payload()), encoding="utf-8")

        assert append_opponent_scouting_sheet(workbook_path, document_path) == str(workbook_path)
        book = load_workbook(workbook_path)
        sheet = book[OPPONENT_SCOUTING_SHEET]
        assert [cell.value for cell in sheet[1]] == OPPONENT_SCOUTING_COLUMNS
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:H3"

        row = [cell.value for cell in sheet[2]]
        assert row[0:3] == ["Corner Pockets", "Bob", 3]
        assert sheet["E2"].number_format == "0%"

    def test_a_danger_matchup_is_conditionally_highlighted(self, tmp_path):
        workbook_path = tmp_path / "workbook.xlsx"
        document_path = tmp_path / "lineups.json"
        workbook = Workbook()
        workbook.active.title = "Existing"
        workbook.save(workbook_path)
        document_path.write_text(json.dumps(payload()), encoding="utf-8")

        append_opponent_scouting_sheet(workbook_path, document_path)
        sheet = load_workbook(workbook_path)[OPPONENT_SCOUTING_SHEET]
        formulas = [rule.formula[0] for rules in sheet.conditional_formatting for rule in rules.rules]
        assert '"Yes"' in formulas

    def test_is_idempotent_and_replaces_a_prior_sheet(self, tmp_path):
        workbook_path = tmp_path / "workbook.xlsx"
        document_path = tmp_path / "lineups.json"
        workbook = Workbook()
        workbook.active.title = "Existing"
        workbook.create_sheet(OPPONENT_SCOUTING_SHEET)
        workbook.save(workbook_path)
        document_path.write_text(json.dumps(payload()), encoding="utf-8")

        append_opponent_scouting_sheet(workbook_path, document_path)
        append_opponent_scouting_sheet(workbook_path, document_path)
        assert load_workbook(workbook_path).sheetnames == ["Existing", OPPONENT_SCOUTING_SHEET]

    def test_no_opponents_still_produces_a_real_header_only_sheet(self, tmp_path):
        workbook_path = tmp_path / "workbook.xlsx"
        document_path = tmp_path / "lineups.json"
        workbook = Workbook()
        workbook.active.title = "Existing"
        workbook.save(workbook_path)
        document_path.write_text(json.dumps({"opponent_scouting": []}), encoding="utf-8")

        append_opponent_scouting_sheet(workbook_path, document_path)
        sheet = load_workbook(workbook_path)[OPPONENT_SCOUTING_SHEET]
        assert sheet.max_row == 1
