"""Tests for scripts/build_captains_edge.py.

Seeds a real SQLite database through the actual ORM models -- not a
hand-rolled schema -- so a column rename in database/models.py breaks these
tests rather than silently producing artifacts with empty columns.
"""

from __future__ import annotations

import json
import sqlite3

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from database.models import Base, Match, Player, PlayerHeadToHead, PlayerMatchup, PlayerTeamHistory, Team
from scripts.build_captains_edge import (
    NoDatabaseError,
    build,
    build_payload,
    connect_read_only,
    fetch_head_to_head,
    fetch_matchups,
    fetch_players,
    render_html,
    resolve_db_path,
    tag_for,
)


@pytest.fixture
def db_path(tmp_path):
    """A populated database: one of our players, two opponents, scored
    pairings across both formats, and a real head-to-head game."""
    path = tmp_path / "apa.db"
    engine = create_engine(f"sqlite:///{path}")
    Base.metadata.create_all(engine)

    with Session(engine) as db:
        team = Team(external_id="T1", name="Chalk It Up")
        alice = Player(external_id="P1", name="Alice", skill_level=5, team=team,
                       matches_won=8, matches_played=10, win_pct=0.8, ppm=2.1, pa=0.7)
        bob = Player(external_id="P2", name="Bob", skill_level=4)
        carol = Player(external_id="P3", name="Carol", skill_level=6)
        db.add_all([team, alice, bob, carol])
        db.flush()

        match = Match(external_id="M1", week=7, match_date="2026-08-01",
                      home_team_id="T1", away_team_id="T2",
                      home_score=3, away_score=2)
        db.add(match)
        db.flush()

        db.add_all([
            PlayerMatchup(
                player_id=alice.id, opponent_id=bob.id, matches_played=6,
                win_rate=0.8333, avg_points_earned=5.5, avg_own_skill_level=5.0,
                avg_opponent_skill_level=4.0, sl_delta=-1.0, trend="stable",
                volatility=0, matchup_score=78, confidence_score=71,
                format="EIGHT_BALL", session_name="Summer 2026",
            ),
            PlayerMatchup(
                player_id=alice.id, opponent_id=carol.id, matches_played=2,
                win_rate=0.0, avg_points_earned=1.5, avg_own_skill_level=5.0,
                avg_opponent_skill_level=6.0, sl_delta=1.0, trend="down",
                volatility=2, matchup_score=31, confidence_score=44,
                format="NINE_BALL", session_name="Summer 2026",
            ),
        ])
        db.add(PlayerHeadToHead(
            player_id=alice.id, opponent_id=bob.id, match_id=match.id,
            own_skill_level=5, opponent_skill_level=4, result="W",
            format="EIGHT_BALL", session_name="Summer 2026",
        ))
        db.commit()

    engine.dispose()
    return path


@pytest.fixture
def connection(db_path):
    conn = connect_read_only(db_path)
    yield conn
    conn.close()


class TestTag:
    """Tags bucket the existing matchup_score; they never alter it, and the
    65/35 anchors are shared with scripts/render_demo_html.py."""

    def test_the_bands_match_the_demos_thresholds(self):
        assert tag_for(78) == "Favored"
        assert tag_for(65) == "Favored"
        assert tag_for(64) == "Even"
        assert tag_for(36) == "Even"
        assert tag_for(35) == "Avoid"
        assert tag_for(0) == "Avoid"

    def test_a_missing_score_is_not_quietly_called_even(self):
        """A pairing the engine could not score is unknown, not balanced."""
        assert tag_for(None) == "No data"


class TestReadOnly:
    def test_the_connection_refuses_writes(self, connection):
        """The captain's real database must survive any bug in this tool."""
        with pytest.raises(sqlite3.OperationalError):
            connection.execute("DELETE FROM player_matchups")

    def test_a_missing_database_names_every_path_it_tried(self, tmp_path):
        with pytest.raises(NoDatabaseError) as exc:
            resolve_db_path(str(tmp_path / "nope.db"))
        assert "nope.db" in str(exc.value)

    def test_an_explicit_path_wins(self, db_path):
        assert resolve_db_path(str(db_path)) == db_path


class TestQueries:
    def test_matchups_carry_both_names_and_every_scoring_field(self, connection):
        rows = {r["opponent_name"]: r for r in fetch_matchups(connection)}
        row = rows["Bob"]
        assert row["player_name"] == "Alice"
        assert row["matches_played"] == 6
        assert row["win_rate"] == pytest.approx(0.8333)
        assert row["matchup_score"] == 78
        assert row["confidence_score"] == 71
        assert row["sl_delta"] == -1.0
        assert row["avg_own_skill_level"] == 5.0
        assert row["avg_opponent_skill_level"] == 4.0
        assert row["format"] == "EIGHT_BALL"

    def test_the_tag_comes_from_the_stored_score(self, connection):
        rows = {r["opponent_name"]: r for r in fetch_matchups(connection)}
        assert rows["Bob"]["tag"] == "Favored"
        assert rows["Carol"]["tag"] == "Avoid"

    def test_only_players_with_pairings_are_offered(self, connection):
        """A roster name with no matchup rows would promise an answer the app
        cannot give, so it must not reach the selector."""
        names = {p["player_name"] for p in fetch_players(connection)}
        assert names == {"Alice", "Bob", "Carol"}

    def test_players_carry_their_roster_context(self, connection):
        alice = next(p for p in fetch_players(connection) if p["player_name"] == "Alice")
        assert alice["team_name"] == "Chalk It Up"
        assert alice["skill_level"] == 5
        assert alice["matches_played"] == 10

    def test_head_to_head_games_resolve_their_match(self, connection):
        games = fetch_head_to_head(connection)
        assert len(games) == 1
        assert games[0]["result"] == "W"
        assert games[0]["week"] == 7
        assert games[0]["match_date"] == "2026-08-01"
        assert games[0]["opponent_name"] == "Bob"

    def test_formats_and_sessions_come_from_real_rows(self, connection):
        payload = build_payload(connection)
        assert payload["formats"] == ["EIGHT_BALL", "NINE_BALL"]
        assert payload["sessions"] == ["Summer 2026"]


class TestScope:
    """A demo run scopes this builder to its own one real team pairing --
    see scripts/build_full_production_demo.py's build_documents(), which
    passes its own (our_team_id, opponent_team_id, format, session_name)
    scope dict through unchanged. Unrelated division pairings existing in
    the same real database must not appear in a scoped build's output."""

    SCOPE = {
        "our_team_id": "T1", "opponent_team_id": "T2",
        "format": "EIGHT_BALL", "session_name": "Summer 2026",
    }

    @pytest.fixture
    def scoped_db_path(self, tmp_path):
        path = tmp_path / "scoped.db"
        engine = create_engine(f"sqlite:///{path}")
        Base.metadata.create_all(engine)
        with Session(engine) as db:
            team_a = Team(external_id="T1", name="Chalk It Up")
            team_b = Team(external_id="T2", name="Corner Pockets")
            team_c = Team(external_id="T3", name="Unrelated C")
            team_d = Team(external_id="T4", name="Unrelated D")
            alice = Player(external_id="P1", name="Alice", skill_level=5, team=team_a)
            bob = Player(external_id="P2", name="Bob", skill_level=4, team=team_b)
            dave = Player(external_id="P3", name="Dave", skill_level=5, team=team_c)
            erin = Player(external_id="P4", name="Erin", skill_level=6, team=team_d)
            db.add_all([team_a, team_b, team_c, team_d, alice, bob, dave, erin])
            db.flush()

            match = Match(external_id="M1", week=7, match_date="2026-08-01",
                          home_team_id="T1", away_team_id="T2",
                          home_score=3, away_score=2)
            db.add(match)
            db.flush()

            # Real current-roster membership -- the authoritative source
            # this builder's scope filtering must use, never players.team_id
            # (see _scope_team_pair_clause's docstring: a real player can be
            # rostered on several teams at once for the same session).
            for player, team_external_id in (
                (alice, "T1"), (bob, "T2"), (dave, "T3"), (erin, "T4"),
            ):
                db.add(PlayerTeamHistory(
                    player_id=player.id, team_external_id=team_external_id,
                    session_name="Summer 2026", is_current=True,
                ))
            db.flush()

            db.add_all([
                PlayerMatchup(
                    player_id=alice.id, opponent_id=bob.id, matches_played=6,
                    win_rate=0.8333, matchup_score=78, confidence_score=71,
                    format="EIGHT_BALL", session_name="Summer 2026",
                ),
                # Unrelated real pairing on unrelated teams, same format/session.
                PlayerMatchup(
                    player_id=dave.id, opponent_id=erin.id, matches_played=3,
                    win_rate=0.5, matchup_score=50, confidence_score=50,
                    format="EIGHT_BALL", session_name="Summer 2026",
                ),
            ])
            db.add_all([
                PlayerHeadToHead(
                    player_id=alice.id, opponent_id=bob.id, match_id=match.id,
                    own_skill_level=5, opponent_skill_level=4, result="W",
                    format="EIGHT_BALL", session_name="Summer 2026",
                ),
                PlayerHeadToHead(
                    player_id=dave.id, opponent_id=erin.id, match_id=match.id,
                    own_skill_level=5, opponent_skill_level=6, result="L",
                    format="EIGHT_BALL", session_name="Summer 2026",
                ),
            ])
            db.commit()
        engine.dispose()
        return path

    def test_fetch_matchups_excludes_an_unrelated_team_pairing(self, scoped_db_path):
        conn = connect_read_only(scoped_db_path)
        try:
            scoped = fetch_matchups(conn, scope=self.SCOPE)
            unscoped = fetch_matchups(conn)
        finally:
            conn.close()
        assert len(scoped) == 1
        assert scoped[0]["player_name"] == "Alice"
        assert len(unscoped) == 2

    def test_fetch_players_excludes_players_from_other_teams(self, scoped_db_path):
        conn = connect_read_only(scoped_db_path)
        try:
            names = {p["player_name"] for p in fetch_players(conn, scope=self.SCOPE)}
        finally:
            conn.close()
        assert names == {"Alice", "Bob"}

    def test_fetch_head_to_head_excludes_an_unrelated_game(self, scoped_db_path):
        conn = connect_read_only(scoped_db_path)
        try:
            games = fetch_head_to_head(conn, scope=self.SCOPE)
        finally:
            conn.close()
        assert len(games) == 1
        assert games[0]["opponent_name"] == "Bob"

    def test_build_payload_is_scoped_end_to_end(self, scoped_db_path):
        conn = connect_read_only(scoped_db_path)
        try:
            payload = build_payload(conn, scope=self.SCOPE)
        finally:
            conn.close()
        assert len(payload["matchups"]) == 1
        assert {p["player_name"] for p in payload["players"]} == {"Alice", "Bob"}
        assert len(payload["head_to_head"]) == 1

    def test_scope_resolves_correctly_when_a_player_has_several_current_teams(self, tmp_path):
        """The real production bug this scoping is built against: a real
        player can be rostered on several current teams for the same
        session (different formats/divisions). players.team_id is a single
        scalar FK -- whichever team ingest last touched -- so filtering on
        it silently drops or misattributes a multi-team player. Scoping
        must use player_team_history instead, which records every
        membership."""
        path = tmp_path / "multi_team.db"
        engine = create_engine(f"sqlite:///{path}")
        Base.metadata.create_all(engine)
        with Session(engine) as db:
            team_a = Team(external_id="T1", name="Chalk It Up")
            team_b = Team(external_id="T2", name="Corner Pockets")
            team_other = Team(external_id="T9", name="A Whole Different Division")
            # Alice's own players.team_id (the single scalar FK) points at
            # the UNRELATED team -- exactly the real production shape found
            # on data/apa_tracker.db, where a player's team_id happened to
            # point at one of several concurrent teams, not the scoped one.
            alice = Player(external_id="P1", name="Alice", skill_level=5, team=team_other)
            bob = Player(external_id="P2", name="Bob", skill_level=4, team=team_b)
            db.add_all([team_a, team_b, team_other, alice, bob])
            db.flush()
            db.add_all([
                PlayerTeamHistory(
                    player_id=alice.id, team_external_id="T9",
                    session_name="Summer 2026", is_current=True,
                ),
                # Alice is ALSO, concurrently, on the scope's own team T1.
                PlayerTeamHistory(
                    player_id=alice.id, team_external_id="T1",
                    session_name="Summer 2026", is_current=True,
                ),
                PlayerTeamHistory(
                    player_id=bob.id, team_external_id="T2",
                    session_name="Summer 2026", is_current=True,
                ),
            ])
            db.add(PlayerMatchup(
                player_id=alice.id, opponent_id=bob.id, matches_played=6,
                win_rate=0.8333, matchup_score=78, confidence_score=71,
                format="EIGHT_BALL", session_name="Summer 2026",
            ))
            db.commit()
        engine.dispose()

        conn = connect_read_only(path)
        try:
            scoped = fetch_matchups(conn, scope=self.SCOPE)
        finally:
            conn.close()

        assert len(scoped) == 1
        assert scoped[0]["player_name"] == "Alice"

    def test_build_writes_a_scoped_decision_document(self, scoped_db_path, tmp_path):
        html_path, xlsx_path = build(str(scoped_db_path), str(tmp_path), scope=self.SCOPE)
        assert html_path.is_file()
        assert xlsx_path.is_file()
        decision = json.loads((tmp_path / "captains_edge.json").read_text(encoding="utf-8"))
        team_ids = {lineup["team_id"] for lineup in decision["lineups"]}
        # Internal team pks are opaque here, but only our own scoped team's
        # players should ever appear in a scoped decision document.
        all_players = {p["player_name"] for lineup in decision["lineups"] for p in lineup["players"]}
        assert "Dave" not in all_players
        assert "Erin" not in all_players


class TestDegradedDatabases:
    def test_a_schema_with_no_rows_yields_an_empty_payload_not_a_crash(self, tmp_path):
        path = tmp_path / "empty.db"
        engine = create_engine(f"sqlite:///{path}")
        Base.metadata.create_all(engine)
        engine.dispose()

        conn = connect_read_only(path)
        try:
            payload = build_payload(conn)
        finally:
            conn.close()
        assert payload["matchups"] == []
        assert payload["players"] == []

    def test_a_database_missing_the_matchup_table_entirely(self, tmp_path):
        """An older database predating the matchup engine has nothing to
        show, which is different from being broken."""
        path = tmp_path / "old.db"
        sqlite3.connect(path).close()

        conn = connect_read_only(path)
        try:
            assert fetch_matchups(conn) == []
            assert fetch_players(conn) == []
            assert fetch_head_to_head(conn) == []
        finally:
            conn.close()


class TestHtml:
    def test_build_writes_a_self_contained_page(self, db_path, tmp_path):
        html_path, _ = build(str(db_path), str(tmp_path))
        html = html_path.read_text(encoding="utf-8")

        assert html.startswith("<!DOCTYPE html>")
        # No network dependencies: it has to open from a file:// URL on a
        # captain's laptop, possibly with no internet.
        assert "http://" not in html
        assert "https://" not in html

    def test_the_data_is_embedded_and_parses_as_json(self, db_path, tmp_path):
        html_path, _ = build(str(db_path), str(tmp_path))
        html = html_path.read_text(encoding="utf-8")

        opener = '<script id="payload" type="application/json">'
        start = html.index(opener) + len(opener)
        payload = json.loads(html[start:html.index("</script>", start)].replace("<\\/", "</"))

        assert {p["player_name"] for p in payload["players"]} == {"Alice", "Bob", "Carol"}
        assert len(payload["matchups"]) == 2
        assert len(payload["head_to_head"]) == 1

    def test_a_closing_script_tag_in_data_cannot_break_the_page(self):
        """Escaping `</` is what stops embedded JSON from closing its own
        script tag early and blanking the app."""
        payload = {
            "generated_at": "now", "banner": "", "min_sample": 3,
            "players": [{"player_name": "</script><b>x"}],
            "matchups": [], "head_to_head": [], "formats": [], "sessions": [],
        }
        html = render_html(payload)
        assert "</script><b>" not in html
        assert "<\\/script>" in html

    def test_the_banner_reaches_the_page(self, db_path, tmp_path):
        html_path, _ = build(str(db_path), str(tmp_path), banner="DEMO DATA")
        assert "DEMO DATA" in html_path.read_text(encoding="utf-8")


class TestWorkbook:
    def test_it_has_the_three_captain_facing_sheets(self, db_path, tmp_path):
        _, xlsx_path = build(str(db_path), str(tmp_path))
        book = load_workbook(xlsx_path)
        assert book.sheetnames == ["Players", "Matchups", "H2H History"]

    def test_matchup_rows_carry_the_real_values(self, db_path, tmp_path):
        _, xlsx_path = build(str(db_path), str(tmp_path))
        sheet = load_workbook(xlsx_path)["Matchups"]
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        rows = [dict(zip(header, [c.value for c in r]))
                for r in sheet.iter_rows(min_row=2)]

        bob = next(r for r in rows if r["Opponent"] == "Bob")
        assert bob["Player"] == "Alice"
        assert bob["Tag"] == "Favored"
        assert bob["Matchup Score"] == 78
        assert bob["Confidence"] == 71
        assert bob["Matches Played"] == 6
        assert bob["Win Rate"] == pytest.approx(0.8333)
        assert bob["SL Delta"] == -1.0

    def test_headers_are_frozen_and_filterable_on_every_sheet(self, db_path, tmp_path):
        """Without these the sheet is a dump, not a lookup table."""
        _, xlsx_path = build(str(db_path), str(tmp_path))
        book = load_workbook(xlsx_path)
        for name in book.sheetnames:
            sheet = book[name]
            assert sheet.freeze_panes == "A2", name
            assert sheet.auto_filter.ref is not None, name

    def test_the_tag_column_is_conditionally_coloured(self, db_path, tmp_path):
        """Green/amber/red on Tag is what makes the sheet scannable without
        reading a single number."""
        _, xlsx_path = build(str(db_path), str(tmp_path))
        sheet = load_workbook(xlsx_path)["Matchups"]
        formulas = [
            rule.formula[0]
            for rules in sheet.conditional_formatting
            for rule in rules.rules
        ]
        assert '"Favored"' in formulas
        assert '"Even"' in formulas
        assert '"Avoid"' in formulas

    def test_h2h_history_carries_the_game_detail(self, db_path, tmp_path):
        _, xlsx_path = build(str(db_path), str(tmp_path))
        sheet = load_workbook(xlsx_path)["H2H History"]
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        row = dict(zip(header, [c.value for c in list(sheet.iter_rows(min_row=2))[0]]))

        assert row["Player"] == "Alice"
        assert row["Opponent"] == "Bob"
        assert row["Result"] == "W"
        assert row["Week"] == 7
        assert row["Own SL"] == 5
        assert row["Opp SL"] == 4

    def test_an_empty_database_still_produces_openable_artifacts(self, tmp_path):
        """A captain who runs this before the first sync should get an empty
        cheat-sheet, not a traceback or a corrupt file."""
        path = tmp_path / "empty.db"
        engine = create_engine(f"sqlite:///{path}")
        Base.metadata.create_all(engine)
        engine.dispose()

        html_path, xlsx_path = build(str(path), str(tmp_path / "out"))
        assert html_path.is_file() and xlsx_path.is_file()
        book = load_workbook(xlsx_path)
        assert book.sheetnames == ["Players", "Matchups", "H2H History"]
