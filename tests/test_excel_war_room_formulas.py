"""End-to-end Match Day / War Room / Lineup Lab / Scouting Cards / Captain
Packet / Coach Dashboard formula results, computed by tests/excel_formula_eval.py
(no Excel, COM, pywin32 or macros) on the workbook's real formulas and lookup
sheets.

The rosters are built so the matrix holds every evidence category:

    our \\ their     Cam Cole (SL6)          Eve Ellis (SL3)
    Ann (SL5)       2-0 direct  -> G        shared opp Zed only -> I
    Bea (SL4)       0-2 direct  -> R        nothing             -> X
    Dee (no SL)     1-1 direct  -> E        0-2 direct          -> R
"""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from analytics.ultimate_coach_match_day import build_match_day_section
from tests.excel_formula_eval import Workbook, serial
from ui.export_excel_ultimate_coach import write_workbook

MD, WR, LL, SC, CP, CD = "Match Day", "War Room", "Lineup Lab", "Scouting Cards", "Captain Packet", "Coach Dashboard"
SHARKS = "Sharks · Fall 2026 · 8-Ball"
SHARKS9 = "Sharks · Fall 2026 · 9-Ball"
FALCONS = "Falcons · Fall 2026 · 8-Ball"
OWLS = "Owls · Fall 2026 · 8-Ball"
ANN, BEA, DEE = "Ann Archer (APA record ID 1001)", "Bea Baker (APA record ID 1002)", "Dee Diaz (APA record ID 1003)"
ANN2 = "Ann Archer (APA record ID 1004)"
CAM, EVE, ZED = "Cam Cole (APA record ID 2001)", "Eve Ellis (APA record ID 2002)", "Zed Zane (APA record ID 2003)"


def _hist(team, div, fmt, sl, won=None, played=None):
    return {"team_external_id": team, "team_name": team.split("-")[0].capitalize(), "division_id": div,
            "session_name": "Fall 2026", "format": fmt, "is_current": True, "skill_level": sl,
            "matches_won": won, "matches_played": played}


def _player(pid, ext, name, history):
    return {"id": pid, "external_id": ext, "name": name, "current_skill_level": None, "current_matches_won": None,
            "current_matches_played": None, "team_history": history, "career_stats": []}


def _games(a, b, results, date="2026-09-20T19:00:00-06:00"):
    rows = []
    for n, result in enumerate(results):
        flip = "L" if result == "W" else "W"
        rows += [
            {"player_id": a, "opponent_id": b, "format": "EIGHT", "result": result, "match_date": date,
             "session_name": "Fall 2026", "own_skill_level": 5, "opponent_skill_level": 4, "points_earned": 3 - n},
            {"player_id": b, "opponent_id": a, "format": "EIGHT", "result": flip, "match_date": date,
             "session_name": "Fall 2026", "own_skill_level": 4, "opponent_skill_level": 5, "points_earned": None},
        ]
    return rows


def _fixture(mid, date, home, away, *, bye=False, status="UNPLAYED", scored=False, hs=None, as_=None):
    names = {"sharks-a": "Sharks", "falcons-a": "Falcons", "owls-a": "Owls", "bye": "BYE"}
    return {"match_id": mid, "match_external_id": str(mid), "match_date": date, "format": "EIGHT",
            "format_raw": "8-Ball Open", "session_name": "Fall 2026", "week": mid, "status": status, "location": None,
            "home_team_id": home, "home_team_name": names[home], "away_team_id": away, "away_team_name": names[away],
            "home_score": hs, "away_score": as_, "is_bye": bye, "is_scored": scored, "is_finalized": scored}


def _payload():
    players = [
        _player(1, "1001", "Ann Archer", [_hist("sharks-a", "d1", "EIGHT", 5, 6, 8), _hist("sharks-b", "d9", "NINE", 4)]),
        _player(2, "1002", "Bea Baker", [_hist("sharks-a", "d1", "EIGHT", 4, 3, 6)]),
        _player(3, "1003", "Dee Diaz", [_hist("sharks-a", "d1", "EIGHT", None)]),
        _player(4, "1004", "Ann Archer", []),
        _player(10, "2001", "Cam Cole", [_hist("falcons-a", "d1", "EIGHT", 6, 7, 8)]),
        _player(11, "2002", "Eve Ellis", [_hist("falcons-a", "d1", "EIGHT", 3, 2, 4)]),
        _player(12, "2003", "Zed Zane", [_hist("owls-a", "d1", "EIGHT", 4, 1, 2)]),
        _player(13, "2004", "Gus Gale", [_hist("owls-a", "d1", "EIGHT", 5, 2, 2)]),
    ]
    evidence = (_games(1, 10, "WW") + _games(2, 10, "LL") + _games(3, 10, "WL") + _games(1, 12, "W")
                + _games(11, 12, "L") + _games(3, 11, "LL"))
    payload = {
        "schema": "test", "probability_publication": "FORBIDDEN", "matchup_probability": None,
        "predictive_confidence": None, "requires_live_apa_login": False, "database_mutated": False,
        "name_matching_used": False, "trust": {}, "counts": {"players": len(players), "head_to_head_rows": len(evidence)},
        "players": players, "evidence": evidence,
    }
    payload["match_day"] = build_match_day_section([
        _fixture(1, "2026-10-11T11:00:00-06:00", "sharks-a", "falcons-a"),
        _fixture(2, "2026-10-18T11:00:00-06:00", "sharks-a", "bye", bye=True),
        _fixture(3, "2026-10-25T11:00:00-06:00", "sharks-a", "owls-a"),
        _fixture(4, "2026-10-25T19:00:00-06:00", "sharks-a", "falcons-a"),
        _fixture(5, "2026-09-27T11:00:00-06:00", "owls-a", "sharks-a", status="COMPLETED", scored=True, hs=3, as_=2),
    ], players)
    return payload


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    path = write_workbook(_payload(), tmp_path_factory.mktemp("wr") / "uc.xlsx", built_at="2026-10-07 18:00 UTC",
                          viewer_member_external_id="1001", viewer_card_number="80000001")
    return load_workbook(path)


@pytest.fixture()
def book(built):
    return Workbook(built)


def _row(wb, sheet, text, col="A"):
    for cell in wb[sheet][col]:
        if cell.value == text:
            return cell.row
    raise AssertionError(f"{text!r} not found on {sheet}")


def _roster(book, sheet, first, R=8):
    ours = [(book.display(sheet, f"A{r}"), book.display(sheet, f"C{r}"), book.display(sheet, f"D{r}"),
             book.display(sheet, f"E{r}")) for r in range(first, first + R)]
    theirs = [(book.display(sheet, f"G{r}"), book.display(sheet, f"I{r}"), book.display(sheet, f"J{r}"),
               book.display(sheet, f"L{r}")) for r in range(first, first + R)]
    return [x for x in ours if x[0]], [x for x in theirs if x[0]]


def _packet_roster(book, R=8):
    """Captain Packet page-1 rosters: ours (A, C, D, E); theirs (G, K = "SL n · W-L", L = played).
    Page 1 is decision first (best sends, risks), so the rosters start below them."""
    first = next(r for r in range(1, 80) if book.display(CP, f"A{r}") == "Player (APA record ID)") + 1
    ours = [(book.display(CP, f"A{r}"), book.display(CP, f"C{r}"), book.display(CP, f"D{r}"),
             book.display(CP, f"E{r}")) for r in range(first, first + R)]
    theirs = [(book.display(CP, f"G{r}"), book.display(CP, f"K{r}"), book.display(CP, f"L{r}"))
              for r in range(first, first + R)]
    return [x for x in ours if x[0]], [x for x in theirs if x[0]]


def _opportunities(book, wb):
    top = _row(wb, WR, "Top opportunities — best-supported sends (favorable direct first, then even, then indirect)")
    return [(book.display(WR, f"A{r}"), book.display(WR, f"C{r}")) for r in range(top + 1, top + 3)]


def _list_after(book, wb, title, n, sheet=WR):
    top = _row(wb, sheet, title)
    return [v for v in (book.display(sheet, f"A{r}") for r in range(top + 1, top + 1 + n)) if v]


def _no_bad_values(book, wb):
    bad = []
    for sheet in (MD, WR, LL, SC, CP, CD):
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    v = book.display(sheet, cell.coordinate)
                    if v == 0 or (isinstance(v, str) and v.startswith("#")):
                        bad.append((sheet, cell.coordinate, v))
    return bad


# ---------------------------------------------------------------------------

def test_match_day_defaults_to_the_viewers_next_fixture_and_auto_selects_one(book, built):
    assert book.display(MD, "B13") == ANN
    assert "League card #80000001" in book.display(MD, "C13") and "never used as identity" in book.display(MD, "C13")
    assert book.display(MD, "B14") == SHARKS and book.display(MD, "C14") == "Your choice."
    assert book.display(MD, "B15") == "8-Ball & 9-Ball"
    assert book.display(MD, "B16") == "Sun Oct 11, 2026"
    assert book.display(MD, "C16") == "Suggested: the earliest scheduled date on or after the build date."
    assert book.display(MD, "B17") == "Sun Oct 11, 2026 · 11:00 AM MDT (America/Denver) · Home vs Falcons · 8-Ball Open"
    assert book.display(MD, "C17") == "One fixture on this date — used automatically."
    assert book.display(MD, "B19") == f"Opponent roster: {FALCONS}"
    fresh = book.display(MD, "A4")
    assert "Built Wed Oct 7, 2026" in fresh and "latest recorded result Sun Sep 27, 2026" in fresh
    assert "never refreshes itself" in fresh
    assert _no_bad_values(book, built) == []


def test_every_interactive_tab_follows_match_day(book, built):
    status = f"Following Match Day: {SHARKS} vs {FALCONS} · Sun Oct 11, 2026 · 8-Ball"
    assert book.display(WR, "A4") == status
    assert book.display(SC, "A4") == status
    ours, theirs = _roster(book, WR, 15)
    assert ours == [(ANN, "5", "6-2", "Unknown · —"), (BEA, "4", "3-3", "Unknown · —"),
                    (DEE, "No data", "No data", "Unknown · —")]
    assert theirs == [(CAM, "6", "7-1", "—"), (EVE, "3", "2-2", "—")]
    assert book.display(SC, "A6") == f"{CAM} · SL 6"
    assert book.display(CP, "A2") == "Sun Oct 11, 2026 · 11:00 AM MDT (America/Denver) · Home vs Falcons · 8-Ball Open"
    assert _packet_roster(book) == (ours, [(n, f"SL {sl} · {wl}", pl) for n, sl, wl, pl in theirs])
    assert book.display(CD, "A4") == f"Following Match Day: {ANN} vs (choose Player B) · 8-Ball"
    assert book.display(LL, "D5") == "✓ Matches Match Day’s team."
    assert book.display(LL, "C9") == "Sun Oct 11, 2026 · 11:00 AM MDT · Home vs Falcons · match 1"  # marks belong to the exact default fixture
    assert book.display(LL, "D9").startswith("✓ Matches Match Day’s fixture")
    assert book.display(LL, "D6") == "✓ Matches Match Day’s opponent."


def test_matrix_shows_direct_records_indirect_evidence_and_gaps_with_colors(book, built):
    header = _row(built, WR, "Our player ↓ / opponent →")
    assert [book.display(WR, f"{c}{header}") for c in "BC"] == [CAM, EVE]
    cells = [[book.display(WR, f"{c}{header + i}") for c in "BC"] for i in (1, 2, 3)]
    assert cells == [["2-0 (2)", "≈ 1-0 vs 0-1 (1 shared)"], ["0-2 (2)", "No evidence"], ["1-1 (2)", "0-2 (2)"]]
    cats = [[book.display(WR, f"{c}{header + i}") for c in "OP"] for i in (1, 2, 3)]
    assert cats == [["G", "I"], ["R", "X"], ["E", "R"]]
    rules = built[WR].conditional_formatting
    ranges = {str(cf.sqref) for cf in rules}
    assert f"B{header + 1}:B{header + 8}" in ranges and f"C{header + 1}:C{header + 8}" in ranges


def test_matrix_captain_view_switches_text_not_evidence(book, built):
    header = _row(built, WR, "Our player ↓ / opponent →")
    assert book.display(WR, f"J{header - 1}") == "View" and book.display(WR, f"K{header - 1}") == "Evidence view"
    book.set(WR, f"K{header - 1}", "Captain view")
    cells = [[book.display(WR, f"{c}{header + i}") for c in "BC"] for i in (1, 2, 3)]
    assert cells == [["🟢 2-0", "🟡 ≈"], ["🔴 0-2", "⚪"], ["🟡 1-1", "🔴 0-2"]]     # same as the HTML Captain view
    assert [[book.display(WR, f"{c}{header + i}") for c in "OP"] for i in (1, 2, 3)] == [["G", "I"], ["R", "X"], ["E", "R"]]


def test_best_sends_risks_and_unique_options_use_the_disclosed_rules(book, built):
    assert _opportunities(book, built) == [
        (f"vs {CAM} · SL 6", f"1. {ANN} — 2-0 (2) · 2. {DEE} — 1-1 (2)"),
        (f"vs {EVE} · SL 3", f"≈ {ANN} — ≈ 1-0 vs 0-1 (1 shared)"),      # shared-only: never numbered
    ]
    threats = _list_after(book, built, "Opponents with winning records against our roster (by recorded direct results; unplayed only)", 3)
    assert threats == [f"{EVE} · SL 3 · 2-0 vs our roster (2 meetings, 1 of our players)"]
    concerning = _list_after(book, built, "Concerning pairings — more direct losses than wins (remaining players vs unplayed opponents)", 5)
    assert concerning == [f"{BEA} vs {CAM}: 0-2 direct (2 meetings)", f"{DEE} vs {EVE}: 0-2 direct (2 meetings)"]
    risks = _list_after(book, built, "Open risks — unplayed opponents with no favorable direct option left among our remaining players", 8)
    assert risks == [f"{EVE} — only even or indirect evidence left"]
    unique = _list_after(book, built, "Unique favorable options — consider saving (the only remaining favorable direct option vs an unplayed opponent)", 8)
    assert unique == [f"{ANN} — only favorable direct option vs {CAM}"]


def test_lineup_marks_change_candidates_but_never_the_evidence(book, built):
    book.set(LL, "C12", "Unavailable")   # Ann
    ours, _ = _roster(book, WR, 15)
    assert ours[0] == (ANN, "5", "6-2", "Unavailable · —") and ours[1][3] == "Unknown · —"
    assert _opportunities(book, built) == [
        (f"vs {CAM} · SL 6", f"1. {DEE} — 1-1 (2)"),
        (f"vs {EVE} · SL 3", "No favorable, even or indirect evidence among our remaining players."),
    ]
    risks = _list_after(book, built, "Open risks — unplayed opponents with no favorable direct option left among our remaining players", 8)
    assert risks == [f"{CAM} — only even or indirect evidence left", f"{EVE} — no evidence-backed option left"]
    snapshot = _row(built, WR, "Lineup Lab snapshot (your marks applied; edit them on Lineup Lab)")
    assert book.display(WR, f"C{snapshot + 1}") == ("2 of 3 (2 with unknown availability — still counted as remaining)")
    # The evidence itself is untouched: inspecting Cam still shows Ann's real 2-0.
    inspect = _row(built, WR, "Inspect opponent")
    book.set(WR, f"C{inspect}", CAM)
    assert [book.display(WR, f"{c}{inspect + 2}") for c in "ACE"] == [ANN, "1", "2-0 (2 meetings)"]


def test_lineup_marks_only_apply_to_the_team_they_were_made_for(book, built):
    book.set(LL, "C12", "Unavailable")
    book.set(LL, "C5", SHARKS9)
    ours, _ = _roster(book, WR, 15)
    assert ours[0][3] == "Unknown · —"
    end = 15 + 8
    assert book.display(WR, f"A{end}").startswith(f"⚠ Lineup Lab marks are for {SHARKS9} — not applied to {SHARKS}.")
    assert book.display(LL, "D5") == f"⚠ Match Day shows {SHARKS} — these marks are NOT applied there."
    assert book.display(LL, "B12") == "Ann Archer (APA record ID 1001)"  # rows list the named (9-Ball) team


def test_selected_lineup_and_remaining_totals_disclose_missing_skill_levels(book, built):
    snapshot = _row(built, WR, "Lineup Lab snapshot (your marks applied; edit them on Lineup Lab)")
    assert book.display(WR, f"C{snapshot + 2}") == (
        "9 known subtotal · 1 remaining player(s) without a captured SL — not a complete total")
    book.set(LL, "D12", "Played")
    book.set(LL, "D13", "Planned")
    assert book.display(WR, f"C{snapshot + 3}") == "Skill total 9 for 2 selected player(s) — not a lineup-legality check."
    book.set(LL, "D14", "Planned")
    book.set(LL, "C7", 23)
    assert book.display(WR, f"C{snapshot + 3}") == ("Known subtotal 9 · 1 selected player(s) without a captured SL · "
                                                    "your reference cap: 23 (user-entered, not verified) — not a lineup-legality check.")
    assert book.display(WR, f"C{snapshot + 1}") == "2 of 3 (2 with unknown availability — still counted as remaining)"


def test_opponent_played_marks_retire_that_opponent(book, built):
    book.set(LL, "C23", "Played")   # Cam
    assert _opportunities(book, built)[0] == (f"vs {CAM} · SL 6", "Already played.")
    risks = _list_after(book, built, "Open risks — unplayed opponents with no favorable direct option left among our remaining players", 8)
    assert risks == [f"{EVE} — only even or indirect evidence left"]
    unique = _list_after(book, built, "Unique favorable options — consider saving (the only remaining favorable direct option vs an unplayed opponent)", 8)
    assert unique == ["None right now."]


def test_changing_the_player_rebuilds_team_date_and_opponent_and_explains_stale_inputs(book, built):
    book.set(MD, "B6", EVE)
    assert book.display(MD, "B14") == FALCONS
    assert book.display(MD, "C14").startswith(f"⚠ “{SHARKS}” is not one of Eve Ellis’s current teams — ignored. ")
    assert book.display(MD, "B17") == "Sun Oct 11, 2026 · 11:00 AM MDT (America/Denver) · Away vs Sharks · 8-Ball Open"
    assert book.display(WR, "A4") == f"Following Match Day: {FALCONS} vs {SHARKS} · Sun Oct 11, 2026 · 8-Ball"
    ours, theirs = _roster(book, WR, 15)
    assert [o[0] for o in ours] == [CAM, EVE] and [t[0] for t in theirs] == [ANN, BEA, DEE]
    assert book.display(LL, "D5") == f"⚠ Match Day shows {FALCONS} — these marks are NOT applied there."
    book.set(MD, "B6", ANN2)  # same name, different identity, no team
    assert book.display(MD, "B13") == ANN2 and book.display(MD, "B14") == "—"
    assert book.display(MD, "C14") == "No current team captured for Ann Archer."
    assert _roster(book, WR, 15) == ([], [])  # nothing stale left on screen
    book.set(MD, "B6", "Ann Archer")  # a bare name is never an identity
    assert book.display(MD, "B13") == "⚠ No verified player matches “Ann Archer”."
    assert _no_bad_values(book, built) == []


def test_bye_multiple_fixtures_and_invalid_dates(book, built):
    book.set(MD, "B9", "Sun Oct 18, 2026")
    assert book.display(MD, "B19") == "Opponent roster: Not applicable (bye)."
    assert book.display(WR, "A7") == "No opponent roster for the selected fixture."
    assert _roster(book, WR, 15)[1] == []
    book.set(MD, "B9", "Sun Oct 25, 2026")
    assert book.display(MD, "C17") == "2 fixtures on this date — pick one in the Fixture cell (none is chosen for you)."
    assert book.display(MD, "B17") == "—" and _roster(book, WR, 15)[1] == []
    choice = book.value("Engine", book.name("uc_FixtureList")[1][1])
    assert choice == "2 · vs Falcons · 7:00 PM MDT · Home · 8-Ball Open"
    book.set(MD, "B10", choice)
    assert "7:00 PM MDT" in book.display(MD, "B17") and book.display(MD, "C17") == "Your choice of 2 fixtures on this date."
    assert [t[0] for t in _roster(book, WR, 15)[1]] == [CAM, EVE]
    book.set(MD, "B9", serial(date(2026, 10, 12)))
    assert book.display(MD, "B16") == "—"
    assert book.display(MD, "C16") == (f"⚠ “Monday, Oct 12, 2026” has no {SHARKS} fixture (8-Ball & 9-Ball) — ignored. "
                                       "Pick a date from the list.")
    assert book.display(MD, "B17") == "—" and _no_bad_values(book, built) == []


def test_format_and_team_changes_keep_valid_choices_and_explain_empty_schedules(book, built):
    book.set(MD, "B7", SHARKS9)
    assert book.display(MD, "B14") == SHARKS9 and book.display(MD, "C14") == "Your choice."
    assert book.display(MD, "C16") == ("No upcoming scheduled date was captured for this team and format — choose a date.")
    assert book.display(WR, "A7") == "No opponent roster for the selected fixture."
    book.set(MD, "B7", SHARKS)
    book.set(MD, "B8", "All recorded formats")
    assert book.display(MD, "B16") == "Sun Oct 11, 2026" and book.display(MD, "B15") == "All recorded formats"


def test_war_room_local_override_affects_only_that_tab_and_clearing_restores(book, built):
    book.set(WR, "C6", OWLS)
    assert book.display(WR, "A4") == f"Using local selections: {SHARKS} vs {OWLS} · 8-Ball"
    assert [t[0] for t in _roster(book, WR, 15)[1]] == ["Gus Gale (APA record ID 2004)", ZED]
    assert book.display(MD, "B17").endswith("Home vs Falcons · 8-Ball Open")       # Match Day untouched
    assert book.display(CD, "A4").startswith("Following Match Day")
    # Only the War Room changed: the packet, scouting cards and Coach Dashboard still follow Match Day.
    assert book.display(CP, "A2") == book.display(MD, "B17")
    book.set(WR, "C5", OWLS)  # a pairing with no scheduled fixture -> rosters only, said plainly
    assert "precomputed for teams scheduled to play each other" in book.display(WR, "A7")
    book.set(WR, "C5", "")
    book.set(WR, "C6", "")
    assert book.display(WR, "A4").startswith("Following Match Day:")
    book.set(WR, "C6", "Not A Team")
    assert book.display(WR, "A7").startswith("⚠ “Not A Team” is not a team label — ignored.")
    assert book.display(WR, "A4").startswith("Following Match Day:")


def test_coach_dashboard_defaults_to_the_viewer_and_offers_tonights_opponents(book, built):
    choices = [book.value(s, r) for s, r in book.name("cd_PlayerBList")]
    assert choices[:2] == [CAM, EVE] and choices[-1] == "All players…"
    assert book.display(CD, "C12") == "Choose Player B above."
    book.set(CD, "C6", CAM)
    assert [book.display(CD, f"C{r}") for r in (19, 20, 21)] == ["2-0", 2, "A: 5 · B: 6"]
    book.set(CD, "C6", "All players…")
    book.set(CD, "C7", ZED)
    assert book.display(CD, "C19") == "1-0"
    book.set(CD, "C5", BEA)
    assert book.display(CD, "A4") == f"Using local selections: {BEA} vs {ZED} · 8-Ball"
    assert book.display(CD, "C19") == "No recorded direct meeting in this format"
    book.set(CD, "C5", "")
    assert book.display(CD, "A4").startswith("Following Match Day")


def test_inspect_views_rank_candidates_and_show_one_players_evidence_vs_all(book, built):
    inspect = _row(built, WR, "Inspect opponent")
    book.set(WR, f"C{inspect}", CAM)
    rows = [[book.display(WR, f"{c}{inspect + 1 + k}") for c in "ACEJ"] for k in (1, 2, 3)]
    assert rows == [[ANN, "1", "2-0 (2 meetings)", "Direct record"], [DEE, "2", "1-1 (2 meetings)", "Direct record"],
                    [BEA, "3", "0-2 (2 meetings)", "Direct record"]]
    ours = _row(built, WR, "Inspect our player")
    book.set(WR, f"C{ours}", DEE)
    got = [[book.display(WR, f"{c}{ours + 1 + j}") for c in "ACEJ"] for j in (1, 2)]
    assert got == [[CAM, "Even direct", "1-1 (2 meetings)", "2"], [EVE, "Concerning direct", "0-2 (2 meetings)", "1"]]


def test_scouting_card_and_meetings_show_facts_samples_and_missing_information(book, built):
    card = {book.display(SC, f"A{r}"): book.display(SC, f"C{r}") for r in range(7, 19)}
    assert card["Quick read"] == ("Even: 3-3 vs our roster (6 meetings) · Best answer on record: Ann Archer (2-0) · "
                                  "Avoid: Bea Baker (0-2)")
    assert card["Team record"] == "7-1 (this team, Fall 2026)"
    assert card["Vs our roster"] == "3-3 in 6 meetings with 3 of our 3 players"
    assert card["Meetings with our players"] == f"vs {ANN}: 0-2 · vs {BEA}: 2-0 · vs {DEE}: 1-1"
    assert card["Record by opponent SL"] == "vs SL5 3-3"
    assert card["Coach observations"] == "(add notes on Coach Notes or Lineup Lab)"
    book.set(LL, "D23", "Breaks hard; slow safeties")
    assert book.display(SC, "C18") == "Breaks hard; slow safeties"
    top = _row(built, WR, "Direct meetings between the rosters (newest first)")
    first = [book.display(WR, f"{c}{top + 2}") for c in "ABEI"]
    assert first[0] == "Sun Sep 20, 2026" and first[3] in ("W", "L")
    meetings = [book.display(WR, f"B{r}") for r in range(top + 2, top + 2 + 40)]
    assert sum(1 for m in meetings if m) == 8  # 2+2+2 vs Cam, 2 vs Eve


def test_lineup_marks_belong_to_one_fixture_and_notes_follow_the_player(book, built):
    """GPT audit #84 repro: Played on Oct 11 must not leak into the Oct 25 Falcons fixture."""
    book.set(LL, "D12", "Played")        # Ann played on Oct 11
    book.set(LL, "C23", "Played")        # Cam played on Oct 11
    book.set(LL, "D23", "Slow safeties")
    assert _roster(book, WR, 15)[0][0][3] == "Unknown · Played"
    book.set(MD, "B9", "Sun Oct 25, 2026")
    book.set(MD, "B10", book.value("Engine", book.name("uc_FixtureList")[1][1]))   # the Falcons fixture
    assert "vs Falcons" in book.display(MD, "B17")
    ours, theirs = _roster(book, WR, 15)
    assert ours[0][3] == "Unknown · —" and theirs[0][3] == "—"          # nothing leaked
    assert book.display(WR, "A23").startswith("⚠ Lineup Lab marks are for Sun Oct 11, 2026 · 11:00 AM MDT · Home vs Falcons · match 1 — not applied to this fixture.")
    assert book.display(LL, "D9").startswith("⚠ Match Day’s fixture is Sun Oct 25, 2026 · 7:00 PM MDT · Home vs Falcons")
    assert book.display(SC, "C18") == "Slow safeties"                    # notes describe the player
    book.set(LL, "C9", book.value("Engine", book.name("uc_PlanKeyList").ref))  # re-plan for this fixture
    assert _roster(book, WR, 15)[0][0][3] == "Unknown · Played"
    book.set(MD, "B9", "Sun Oct 11, 2026")                               # back: Oct 25 marks don't apply
    assert _roster(book, WR, 15)[0][0][3] == "Unknown · —"
    # War Room exploring by hand (no fixture): marks apply only with the date blank.
    book.set(WR, "C6", FALCONS)
    assert _roster(book, WR, 15)[0][0][3] == "Unknown · —"
    book.set(LL, "C9", "")
    assert _roster(book, WR, 15)[0][0][3] == "Unknown · Played"
    assert _no_bad_values(book, built) == []



def test_two_matches_on_the_same_day_never_share_marks(tmp_path):
    """GPT audit #84 repro: two Falcons fixtures on Oct 25 (7 PM and 9 PM). Marks made for the first must
    not apply to the second; switching back restores them."""
    payload = _payload()
    from tests.test_excel_war_room_formulas import _fixture as fx
    fixtures = payload["match_day"]["fixtures"]
    raw = [fx(1, "2026-10-11T11:00:00-06:00", "sharks-a", "falcons-a"),
           fx(4, "2026-10-25T19:00:00-06:00", "sharks-a", "falcons-a"),
           fx(6, "2026-10-25T21:00:00-06:00", "sharks-a", "falcons-a")]
    payload["match_day"] = build_match_day_section(raw, payload["players"])
    wb = load_workbook(write_workbook(payload, tmp_path / "two.xlsx", built_at="2026-10-07 18:00 UTC",
                                      viewer_member_external_id="1001"))
    book = Workbook(wb)
    book.set(MD, "B9", "Sun Oct 25, 2026")
    first, second = (book.value(s, r) for s, r in book.name("uc_FixtureList")[:2])
    assert "7:00 PM" in first and "9:00 PM" in second
    book.set(MD, "B10", first)
    book.set(LL, "C9", book.value("Engine", book.name("uc_PlanKeyList").ref))
    assert book.display(LL, "C9").endswith("· match 4")
    book.set(LL, "D12", "Played")          # Ann
    book.set(LL, "C23", "Played")          # Cam
    ours, theirs = _roster(book, WR, 15)
    assert ours[0][3] == "Unknown · Played" and theirs[0][3] == "Played"
    book.set(MD, "B10", second)            # same day, same teams, another match
    ours, theirs = _roster(book, WR, 15)
    assert ours[0][3] == "Unknown · —" and theirs[0][3] == "—"
    assert book.display(LL, "D9").startswith("⚠ Match Day’s fixture is Sun Oct 25, 2026 · 9:00 PM MDT")
    assert "· match 4 — not applied to this fixture." in book.display(WR, "A23")
    assert _packet_roster(book)[0][0][3] == "Unknown · —"     # the packet agrees
    book.set(MD, "B10", first)             # back to the first match: its marks return
    assert _roster(book, WR, 15)[0][0][3] == "Unknown · Played"
    book.set(LL, "C9", "")                 # no fixture named: nothing applies (fail-closed)
    assert _roster(book, WR, 15)[0][0][3] == "Unknown · —"


def test_war_room_overrides_never_change_other_tabs(book, built):
    """GPT audit #84 repro: a War Room override must not change the Coach Dashboard's Player B pool or format,
    the scouting cards or the Captain Packet -- and the packet never mixes one fixture's heading with
    another opponent's roster."""
    pool = lambda: [book.value(s, r) for s, r in book.name("cd_PlayerBList")][:2]
    packet_heading, packet_rosters = book.display(CP, "A2"), _packet_roster(book)
    card = book.display(SC, "A6")
    assert pool() == [CAM, EVE]
    book.set(WR, "C6", OWLS)
    assert [t[0] for t in _roster(book, WR, 15)[1]] == ["Gus Gale (APA record ID 2004)", ZED]   # War Room changed
    assert pool() == [CAM, EVE]                                   # Coach Dashboard did not
    assert book.display(CP, "A2") == packet_heading and _packet_roster(book) == packet_rosters
    assert book.display(SC, "A6") == card and book.display(SC, "A4").startswith("Following Match Day")
    book.set(WR, "C5", SHARKS9)
    assert book.display(CD, "A4") == f"Following Match Day: {ANN} vs (choose Player B) · 8-Ball"
    assert book.display(CP, "A2") == packet_heading and _packet_roster(book) == packet_rosters
    # The Coach Dashboard's own override still works and clears back.
    book.set(CD, "C8", "9-Ball")
    assert book.display(CD, "A4").startswith("Using local selections") and book.display(CD, "A4").endswith("· 9-Ball")
    book.set(CD, "C8", "")
    assert book.display(CD, "A4").startswith("Following Match Day")
    book.set(WR, "C5", ""); book.set(WR, "C6", "")
    assert _no_bad_values(book, built) == []


def test_packet_evidence_is_packed_and_every_line_names_its_opponent(book, built):
    ws = built[CP]
    first = next(c.row for c in ws["A"] if c.value == "Opponent (APA record ID)") + 1
    lines = [[book.display(CP, f"{c}{r}") for c in "ACDHK"] for r in range(first, first + 8)]
    assert lines[:6] == [
        [f"vs {CAM}", "1", ANN, "2-0 (2)", "Favorable direct"],
        [f"vs {CAM}", "2", DEE, "1-1 (2)", "Even direct"],
        [f"vs {CAM}", "3", BEA, "0-2 (2)", "Concerning direct"],
        [f"vs {EVE}", "1", DEE, "0-2 (2)", "Concerning direct"],
        [f"vs {EVE}", "≈", ANN, "≈ 1-0 vs 0-1 (1 shared)", "Indirect only"],
        [f"vs {EVE}", "—", BEA, "No evidence", "No evidence"],
    ]
    assert lines[6:] == [["", "", "", "", ""]] * 2                     # packed: nothing after the last line
    # Meeting history uses real columns; the opponent card shows notes with the missing-information line.
    book.set(LL, "D23", "Breaks hard")
    cards_row = next(c.row for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("Scouting cards")) + 1
    texts = [book.display(CP, f"B{r}") for r in range(cards_row, cards_row + 5)]
    assert any("Notes: Breaks hard" in t for t in texts if isinstance(t, str))


def test_packet_meeting_page_shows_every_record_two_across_with_dividers(book, built):
    """Paul (#84 UAT): larger meeting text, two columns, numbered newest first, a count so nothing is silently
    dropped, and black row dividers / light shading only on rows that hold a meeting."""
    ws = built[CP]
    top = next(c.row for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("Meeting history"))
    assert book.display(CP, f"A{top + 1}") == "Showing 8 of 8 recorded meeting(s)."
    first = top + 3
    one, two, three = book.display(CP, f"A{first}"), book.display(CP, f"G{first}"), book.display(CP, f"A{first + 1}")
    assert one.startswith("1. Sun Sep 20, 2026\n") and two.startswith("2. ") and three.startswith("3. ")
    assert " · SL " in one and "\nvs " in book.display(CP, f"B{first}")
    assert book.display(CP, f"A{first + 4}") == "" and book.display(CP, f"G{first + 4}") == ""   # 8 meetings = 4 rows
    assert ws[f"A{first}"].font.sz == 10.5 and ws.row_dimensions[first].height == 29
    rules = [r for cf in ws.conditional_formatting for r in cf.rules
             if str(cf.sqref).startswith(f"A{first}") or str(cf.sqref).startswith(f"G{first}")]
    assert rules and all(r.dxf.border.bottom.style == "thin" and r.dxf.border.bottom.color.rgb.endswith("000000") for r in rules)
    assert all('<>""' in r.formula[0] for r in rules)          # empty slots stay blank: no ruled empty table


def test_start_here_explains_the_workbook_and_names_the_build(built):
    ws = built["START HERE"]
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    for needle in ("1 · What Ultimate Coach does", "2 · Quick start", "Step 1 · Go to Match Day", "Step 8 · Print the Captain Packet",
                   "3 · Workbook tour", "4 · Match night workflow", "5 · Important limitations",
                   "Historical records are not predictions", "No validated win-probability model", "6 · Mobile match night (phone)", "https://ssands5-cloud.github.io/APA-Tracker/",
                   "Add to Home Screen", "7 · Build information",
                   "Workbook version: ", "Build date: Wed Oct 7, 2026", "Data freshness: latest recorded result Sun Sep 27, 2026"):
        assert needle in text, needle
    links = {c.hyperlink.location for row in ws.iter_rows() for c in row if c.hyperlink}
    for sheet in ("Match Day", "Command Center", "War Room", "Lineup Lab", "Coach Notes", "Captain Packet"):
        assert f"'{sheet}'!A1" in links and sheet in built.sheetnames, sheet


def _next_send(book):
    # The whole answer is one wrapped cell, a line per item (CHAR(10)), then the note row; compare line by line.
    lines = [line for r in (8, 9) for line in str(book.display("Command Center", f"B{r}") or "").split("\n")]
    return [x for x in lines if x not in ("", "None")]


def test_command_center_next_send_answers_the_player_they_put_up(book, built):
    assert book.display("Command Center", "B6") == "They put up:" and book.display("Command Center", "B7") == "WHO SHOULD I SEND NEXT?"
    assert _next_send(book) == ["Pick the opponent player they put up (cell C6)."]
    book.set("Command Center", "C6", CAM)
    assert _next_send(book) == [
        "Medals = ordered direct records among our remaining players (same evidence = same medal). Recorded results only — not odds.",
        f"🥇 {ANN} — 2-0 direct record (2 meetings) — favorable · availability unknown",
        f"🥈 {DEE} — 1-1 direct record (2 meetings) — even · availability unknown",
        f"⚠ Avoid: {BEA} — 0-2 (2 meetings)",
    ]
    book.set("Command Center", "C6", EVE)        # shared-opponent only: never medalled
    assert _next_send(book) == [
        "No direct record to order — shared-opponent candidates only (≈, not ordered)",
        f"≈ Not ordered (shared-opponent results only): {ANN} (availability unknown)",
        f"⚠ Avoid: {DEE} — 0-2 (2 meetings)",
        f"❓ Unknown (no evidence, not weak): {BEA}",
    ]
    book.set("Command Center", "C6", CAM)
    book.set(LL, "C12", "Unavailable")           # Ann out: Dee is the only medal left
    assert _next_send(book)[1:] == [f"🥇 {DEE} — 1-1 direct record (2 meetings) — even · availability unknown",
                                    f"⚠ Avoid: {BEA} — 0-2 (2 meetings)"]
    book.set(LL, "C14", "Available")             # Dee marked Available: the caveat goes, nothing else changes
    assert _next_send(book)[1] == f"🥇 {DEE} — 1-1 direct record (2 meetings) — even"
    book.set(LL, "C23", "Played")                # Cam already played
    assert _next_send(book) == [f"{CAM} has already played (Lineup Lab)."]


def test_command_center_summarises_tonight_from_match_day(book, built):
    cc = "Command Center"
    vals = [book.display(cc, f"{c}{r}") for r in range(1, 40) for c in "BFJ"]
    text = "\n".join(str(v) for v in vals if v not in ("", None, 0))
    assert "Sun Oct 11, 2026 · 11:00 AM MDT (America/Denver) · Home vs Falcons · 8-Ball Open" in text
    assert "Available: 0" in text and "Unknown: 3 (not the same as unavailable)" in text and "Already used: 0" in text
    assert f"{CAM} · SL 6" in text and f"{EVE} · SL 3" in text
    assert "Favorable direct record (any sample size): 1" in text and "Concerning (more direct losses than wins): 2" in text
    assert "Limited evidence: 1 even direct · 1 shared-opponent only" in text
    assert "Insufficient evidence (nothing recorded): 1" in text
    assert f"vs {CAM}: best-supported send: {ANN} — reason: 2-0 direct record (2 meetings) — favorable" in text
    book.set(LL, "C12", "Unavailable")
    after = "\n".join(str(book.display(cc, f"{c}{r}")) for r in range(1, 40) for c in "BFJ")
    assert "Unavailable: 1" in after and f"vs {CAM}: best-supported send: {DEE} — reason: 1-1 direct record (2 meetings) — even" in after
    book.set(WR, "C6", OWLS)          # a War Room override never changes the Command Center
    assert "\n".join(str(book.display(cc, f"{c}{r}")) for r in range(1, 40) for c in "BFJ") == after


def test_coach_notes_reach_the_card_marked_as_opinion(book, built):
    ws = built["Coach Notes"]
    assert "opinions, not APA facts" in ws["A1"].value
    book.set("Coach Notes", "A5", CAM)
    book.set("Coach Notes", "B5", "Slow shooter")
    book.set("Coach Notes", "C5", "Strong safety player")
    book.set("Coach Notes", "D5", "Plays the long game")
    assert book.display(SC, "C18") == "Coach: Slow shooter · Strong safety player: Plays the long game"
    book.set(LL, "D23", "Breaks hard")       # tonight's note joins, still separated from evidence
    assert book.display(SC, "C18") == "Coach: Slow shooter · Strong safety player: Plays the long game · Breaks hard"
    book.set("Coach Notes", "A6", EVE)           # a threat's note is shown with the threat, labelled as opinion
    book.set("Coach Notes", "D6", "Runs out from anywhere")
    threats = _list_after(book, built, "Opponents with winning records against our roster (by recorded direct results; unplayed only)", 3)
    assert threats == [f"{EVE} · SL 3 · 2-0 vs our roster (2 meetings, 1 of our players) · 📝 Coach: Runs out from anywhere "
                       "(your opinion, not APA facts)"]


def test_start_here_example_and_tour_are_true_to_the_build(built):
    ws = built["START HERE"]
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    # Worked example comes from this build's own data: configured viewer -> next fixture. Selections only.
    for needle in (f"1 · Player: {ANN}", f"2 · Team: {SHARKS}", "3 · Format: 8-Ball & 9-Ball (the default)",
                   "4 · Scheduled date: Sun Oct 11, 2026", f"5 · Fixture: Home vs {FALCONS} (the only fixture that day"):
        assert needle in text, needle
    assert "War Room → Inspect" in text and "direct record and number of meetings" in text
    assert "Coach Dashboard" in text and "shared opponents)" not in text      # no longer claims shared-opponent evidence


def test_lineup_lab_best_sends_state_their_reason(book, built):
    ws = built[LL]
    top = next(c.row for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("Best remaining send per unplayed opponent"))
    lines = [book.display(LL, f"A{r}") for r in range(top + 1, top + 4)]
    assert lines[:2] == [
        f"vs {CAM}: best-supported send: {ANN} — reason: 2-0 direct record (2 meetings) — favorable",
        f"vs {EVE}: shared-opponent candidate, not ordered (the only one): {ANN} — reason: shared-opponent results only: "
        "ours 1-0 vs theirs 0-1 across "
        "1 shared opponent (no direct meetings)",
    ]
    book.set(LL, "C12", "Unavailable")
    assert book.display(LL, f"A{top + 2}") == f"vs {EVE}: no evidence-backed option left among our remaining players"



def test_inspect_is_blank_not_an_error_when_nothing_is_picked_and_after_clearing(book, built):
    """Real-Excel UAT: Excel evaluates every OR() argument, so blank Inspect showed #VALUE!. The evaluator now
    follows Excel (MATCH of an empty cell is #N/A), and the guards are error-free."""
    inspect = _row(built, WR, "Inspect opponent")
    cells = [f"{c}{inspect + 2 + k}" for k in range(3) for c in "ACDEGJ"]
    assert all(book.display(WR, ref) == "" for ref in cells)
    book.set(WR, f"C{inspect}", CAM)
    assert book.display(WR, f"A{inspect + 2}") == ANN
    book.set(WR, f"C{inspect}", "")
    assert all(book.display(WR, ref) == "" for ref in cells)
    assert _no_bad_values(book, built) == []


def test_helper_cells_are_hidden_from_the_captain(built):
    ws = built[WR]
    hidden = [c for c, d in ws.column_dimensions.items() if d.hidden]
    assert hidden and all(c >= "O" for c in hidden)                 # the matrix category helper grid
    cd = built[CD]
    assert cd.row_dimensions[15].hidden and cd.row_dimensions[16].hidden
    # START HERE: every literal sentence fits its merged cell (GPT saw truncated tour/quick-start text).
    from openpyxl.utils import get_column_letter
    start = built["START HERE"]
    spans = {(m.min_row, m.min_col): m.max_col for m in start.merged_cells.ranges}
    for row in start.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                width = sum(start.column_dimensions[get_column_letter(k)].width or 8.43
                            for k in range(c.column, spans.get((c.row, c.column), c.column) + 1))
                size = c.font.sz or 11
                lines = -(-len(c.value) // max(10, int(width * 11 / size * 1.05)))
                assert (start.row_dimensions[c.row].height or 15) >= lines * size * 1.2, (c.coordinate, c.value[:40])


def _wrapped_lines(text, width_units, size):
    """Word-wrapped line count for text in a cell this wide (Excel breaks at spaces, then inside long words)."""
    per_line = max(6, int(width_units * 10 / size))   # real Excel showed ~10 chars per 9 width units at 9 pt
    lines = 0
    for paragraph in text.split("\n"):
        lines, used = lines + 1, 0
        for word in paragraph.split(" "):
            while len(word) > per_line:
                lines, used, word = lines + (used > 0), 0, word[per_line:]
            if used and used + 1 + len(word) > per_line:
                lines, used = lines + 1, len(word)
            else:
                used += (1 if used else 0) + len(word)
    return lines


def test_formula_rows_fit_their_worst_case_text(built):
    # Real-Excel UAT 2a67d78: the Next Send "≈ Not ordered" list showed 2 of its 3 lines (the last candidate was
    # cut off), and narrow War Room matrix columns cut "(14 shared)" and the opponent header's record ID.
    # Excel never auto-fits formula rows, so each must be tall enough for its longest possible text.
    from openpyxl.utils import get_column_letter
    cc = built["Command Center"]
    width = lambda ws, c1, c2: sum(ws.column_dimensions[get_column_letter(k)].width or 8.43 for k in range(c1, c2 + 1))
    players = built["Players"]
    head = [c.value for c in players[1]]
    col = head.index("Player Label") + 1
    name = max((str(v) for (v,) in players.iter_rows(min_row=2, min_col=col, max_col=col, values_only=True) if v), key=len)
    wr = built[WR]
    header = _row(built, WR, "Our player ↓ / opponent →")
    roster = sum(1 for c in wr[header] if str(c.value).startswith("=INDEX(wr_OppLabels,"))
    # The whole Next Send answer is ONE cell right under the title (no empty medal rows in between), naming each
    # remaining player at most once: size it for k medal lines plus the other roster-k names in the lists.
    assert str(cc["B8"].value).endswith("_NsCard") and "NsOppNotes" not in str(cc["B8"].value)
    assert not any("_NsMedalText" in str(cc.cell(row=r, column=2).value) for r in range(8, 30))
    medal = (f"🥇 {name} — 12-0 direct record (12 meetings) — favorable · tied (same evidence) · consider saving — "
             "our only favorable direct option vs another unplayed opponent · availability unknown")
    headline = "Medals = ordered direct records among our remaining players (same evidence = same medal). Recorded results only — not odds."
    for k in range(0, 5):
        parts = [headline] + [medal] * k + ["+ more direct candidates — see War Room Inspect"]
        if roster - k:
            parts.append("≈ Not ordered (shared-opponent results only): "
                         + "; ".join([f"{name} (availability unknown)"] * max(roster - k - 2, 1)))
            parts += [f"⚠ Avoid: {name} — 0-12 (12 meetings)", f"❓ Unknown (no evidence, not weak): {name}"]
        need = _wrapped_lines("\n".join(parts), width(cc, 2, 11), 11) * 11 * 1.2
        assert cc.row_dimensions[8].height >= need, (k, cc.row_dimensions[8].height, need)
    # Printing fits the width only: a one-page squeeze made the sheet unreadable on paper (native print preview).
    assert cc.page_setup.fitToWidth == 1 and cc.page_setup.fitToHeight == 0
    cols = [c.column for c in wr[header] if str(c.value).startswith("=INDEX(wr_OppLabels,")]
    narrow = min(width(wr, c, c) for c in cols)
    assert (wr.row_dimensions[header].height or 15) >= _wrapped_lines(name, narrow, 8) * 8 * 1.2
    for r in range(header + 1, header + roster + 1):
        assert (wr.row_dimensions[r].height or 15) >= _wrapped_lines("≈ 41-53 vs 81-58 (61 shared)", narrow, 9) * 9 * 1.2, r
    inspect = _row(built, WR, "Inspect opponent")
    basis = ("Shared-opponent results only (no direct meetings) — not ordered against other indirect candidates; "
             "compare ours vs theirs")
    for r in range(inspect + 2, inspect + 2 + roster):
        assert (wr.row_dimensions[r].height or 15) >= _wrapped_lines(basis, width(wr, 10, 12), 9) * 9 * 1.2, r


def test_tall_rows_are_top_aligned_so_values_stay_with_their_row(built):
    # Real-Excel UAT 787f6d7: Inspect's rank and SL were bottom-aligned, so in the taller rows they sat beside the
    # NEXT player's name. Every filled cell in a row taller than two lines must read from the top.
    for sheet in (WR, "Command Center"):
        ws = built[sheet]
        merged_tails = {(r, c) for m in ws.merged_cells.ranges for r in range(m.min_row, m.max_row + 1)
                        for c in range(m.min_col, m.max_col + 1) if (r, c) != (m.min_row, m.min_col)}
        bad = [c.coordinate for row in ws.iter_rows() for c in row
               if c.value not in (None, "") and (c.row, c.column) not in merged_tails and c.column < 15
               and (ws.row_dimensions[c.row].height or 15) > 32 and c.alignment.vertical != "top"]
        assert bad == [], (sheet, bad[:20])


def test_captain_packet_page_one_is_decision_first(book, built):
    sends = _row(built, CP, "Best sends — top opportunities per opponent (favorable direct first, then even, then indirect)")
    risks = _row(built, CP, "Top risks")
    roster = _row(built, CP, "OUR TEAM")
    assert sends < risks < roster
    assert book.display(CP, f"C{sends + 1}") == f"1. {ANN} — 2-0 (2) · 2. {DEE} — 1-1 (2)"


FAY = "Fay Fox (APA record ID 1005)"


def _tied_payload():
    """The fixture plus Fay, whose 2-0 vs Cam is the same evidence as Ann's: a real tie for the first send."""
    payload = _payload()
    payload["players"].append(_player(5, "1005", "Fay Fox", [_hist("sharks-a", "d1", "EIGHT", 4, 2, 5)]))
    payload["evidence"] = payload["evidence"] + _games(5, 10, "WW")
    payload["counts"] = {"players": len(payload["players"]), "head_to_head_rows": len(payload["evidence"])}
    return payload


def test_tied_and_shared_only_sends_are_never_ranked_apart_in_excel(tmp_path):
    """GPT audit #84: outside the HTML Tonight panel, Excel numbered tied and shared-only picks 1., 2., 3. and called a
    shared-only pick "best-supported". Equal evidence now shares a number ("1="), shared-only is "≈", and the single
    best-send sentence says when it is tied. The texts equal the shared Python functions."""
    from analytics.ultimate_coach_war_room import SENDABLE, best_send_text, send_list_text
    from tests.test_ultimate_coach_war_room_browser import OURS, THEIRS, _python_war_room_for
    payload = _tied_payload()
    wb = load_workbook(write_workbook(payload, tmp_path / "tied.xlsx", built_at="2026-10-07 18:00 UTC",
                                      viewer_member_external_id="1001", viewer_card_number="80000001"))
    book = Workbook(wb)
    py = _python_war_room_for(payload)
    rows = [[r for r in b["rows"] if r["category"] in SENDABLE] for b in py["blocks"]]
    assert send_list_text(rows[0]) == f"1= {ANN} — 2-0 (2) · 1= {FAY} — 2-0 (2) · 2. {DEE} — 1-1 (2)"
    assert best_send_text(rows[0]) == (f"best-supported send (tied with 1 other, same evidence): {ANN} — reason: "
                                       "2-0 direct record (2 meetings) — favorable")
    # War Room top opportunities
    assert [x[1] for x in _opportunities(book, wb)] == [send_list_text(rows[0]), send_list_text(rows[1])]
    # Command Center
    cc = "\n".join(str(book.display("Command Center", f"{c}{r}")) for r in range(1, 60) for c in "BFJ")
    assert f"vs {CAM}: {best_send_text(rows[0])}" in cc
    assert f"vs {EVE}: {best_send_text(rows[1])}" in cc and "shared-opponent candidate, not ordered (the only one)" in cc
    # Lineup Lab
    top = next(c.row for c in wb[LL]["A"] if isinstance(c.value, str) and c.value.startswith("Best remaining send per unplayed opponent"))
    assert book.display(LL, f"A{top + 1}") == f"vs {CAM}: {best_send_text(rows[0])}"
    # Captain Packet page 1 (two sends)
    sends = _row(wb, CP, "Best sends — top opportunities per opponent (favorable direct first, then even, then indirect)")
    assert book.display(CP, f"C{sends + 1}") == send_list_text(rows[0], limit=2)
    # Ann unavailable: Fay alone in the top group -> numbered "1." again, no tie wording.
    book.set(LL, "C12", "Unavailable")
    assert _opportunities(book, wb)[0][1] == f"1. {FAY} — 2-0 (2) · 2. {DEE} — 1-1 (2)"
    assert book.display(LL, f"A{top + 1}") == f"vs {CAM}: best-supported send: {FAY} — reason: 2-0 direct record (2 meetings) — favorable"
