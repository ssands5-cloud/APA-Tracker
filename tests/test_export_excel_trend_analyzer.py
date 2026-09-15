"""Tests for ui/export_excel_trend_analyzer.py."""

from __future__ import annotations

from openpyxl import load_workbook

from analytics.trend_analyzer import build_report
from ui.export_excel_trend_analyzer import (
    HISTORY_SHEET,
    SUMMARY_SHEET,
    build_workbook,
    write_workbook,
)


def trend_row(external_id="P1"):
    return {
        "player_id": 1, "player_external_id": external_id, "player_name": "Alice",
        "format": "8-ball", "session_name": "Fall 2026", "sample_size": 5,
        "current_skill_level": 6, "regression_slope": 0.1, "volatility": 0.2,
        "sl_stability": 0.833333, "hot_cold_flag": "HOT",
        "projected_sl_change_probability": 0.3,
    }


class TestBuildWorkbook:
    def test_summary_sheet_present_history_absent_when_empty(self):
        report = build_report("T1", "Chalk It Up", "Fall 2026", [trend_row()], [])
        workbook = build_workbook(report)
        assert SUMMARY_SHEET in workbook.sheetnames
        assert HISTORY_SHEET not in workbook.sheetnames

    def test_history_sheet_present_when_real_observations_exist(self):
        history = [{
            "player_id": 1, "player_external_id": "P1", "match_id": "M1",
            "match_order": 1, "match_date": "2026-09-01", "format": "8-ball",
            "session_name": "Fall 2026", "skill_level": 5,
        }]
        report = build_report("T1", "Chalk It Up", "Fall 2026", [trend_row()], history)
        workbook = build_workbook(report)
        assert HISTORY_SHEET in workbook.sheetnames
        assert workbook[HISTORY_SHEET].max_row == 2  # header + 1 observation

    def test_row_count_matches_measured_players(self):
        report = build_report(
            "T1", "Chalk It Up", "Fall 2026",
            [trend_row("P1"), trend_row("P2")], [],
        )
        sheet = build_workbook(report)[SUMMARY_SHEET]
        assert sheet.max_row == 3  # header + 2 players

    def test_write_workbook_produces_a_real_readable_file(self, tmp_path):
        report = build_report("T1", "Chalk It Up", "Fall 2026", [trend_row()], [])
        out_path = write_workbook(report, tmp_path / "trend_analyzer.xlsx")

        assert out_path.exists()
        reopened = load_workbook(out_path)
        assert SUMMARY_SHEET in reopened.sheetnames
