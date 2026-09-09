"""Tests for the workbook's post-built Lineup Optimizer sheet."""

from __future__ import annotations

import json

from openpyxl import load_workbook
from openpyxl import Workbook

from ui.export_excel import (
    LINEUP_COLUMNS,
    LINEUP_RISK_COLUMNS,
    LINEUP_RISK_TITLE,
    LINEUP_SHEET,
    append_lineup_optimizer_sheet,
    lineup_optimizer_rows,
    lineup_risk_rows,
)


def payload():
    return {
        "lineups": [{
            "team_id": "T1",
            "team_name": "Chalk It Up",
            "opponent_team_id": "T2",
            "opponent_team_name": "Corner Pockets",
            "format": "8-ball",
            "session_name": "Summer 2026",
            "assignments": [
                {
                    "player_name": "Alice",
                    "opponent_name": "Bob",
                    "matchup_score_raw": 80,
                    "win_probability": 0.8,
                    "confidence": None,
                    "risk_factor": 0.1,
                    "final_score": 0.72,
                    "lineup_rank": 1,
                    "source_pairing": True,
                    "rationale": "High win probability.",
                },
            ],
        }],
    }


def test_rows_keep_raw_score_and_render_missing_signals_as_no_data():
    rows = lineup_optimizer_rows(payload())
    assert rows[0]["Matchup Score"] == 80
    assert rows[0]["Win Probability"] == 0.8
    assert rows[0]["Confidence"] == "No data"
    assert rows[0]["Source Pairing"] == "Yes"


def test_append_creates_a_frozen_filterable_sheet(tmp_path):
    workbook_path = tmp_path / "workbook.xlsx"
    document_path = tmp_path / "lineups.json"
    workbook = Workbook()
    workbook.active.title = "Existing"
    workbook.save(workbook_path)
    document_path.write_text(json.dumps(payload()), encoding="utf-8")

    assert append_lineup_optimizer_sheet(workbook_path, document_path) == str(workbook_path)
    book = load_workbook(workbook_path)
    sheet = book[LINEUP_SHEET]
    assert [cell.value for cell in sheet[1]] == LINEUP_COLUMNS
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:N2"
    row = [cell.value for cell in sheet[2]]
    assert row[0:6] == ["Chalk It Up", "Corner Pockets", "8-ball", "Summer 2026", "Alice", "Bob"]
    assert row[6:10] == [80, 0.8, "No data", 0.1]
    assert sheet["H2"].number_format == "0%"


def test_append_is_idempotent_and_replaces_a_prior_sheet(tmp_path):
    workbook_path = tmp_path / "workbook.xlsx"
    document_path = tmp_path / "lineups.json"
    workbook = Workbook()
    workbook.active.title = "Existing"
    workbook.create_sheet(LINEUP_SHEET)
    workbook.save(workbook_path)
    document_path.write_text(json.dumps(payload()), encoding="utf-8")

    append_lineup_optimizer_sheet(workbook_path, document_path)
    append_lineup_optimizer_sheet(workbook_path, document_path)
    assert load_workbook(workbook_path).sheetnames == ["Existing", LINEUP_SHEET]


def payload_with_risk():
    document = payload()
    document["lineups"][0]["lineup_risk"] = {
        "upset_risk_index": 0.14,
        "anchor_stability_score": 0.7,
        "anchor_player_name": "Alice",
        "lineup_volatility_load": 0.2,
        "danger_matchup_count": 1,
        "lineup_risk_score": 0.36,
    }
    return document


class TestLineupRiskFooter:
    """analytics.lineup_risk's team-level summary, rendered BELOW the
    per-pairing table on the same sheet -- no new sheet, no new artifact."""

    def test_rows_flatten_each_lineups_real_risk_block(self):
        rows = lineup_risk_rows(payload_with_risk())
        assert len(rows) == 1
        assert rows[0]["Team"] == "Chalk It Up"
        assert rows[0]["Lineup Risk Score"] == 0.36
        assert rows[0]["Anchor"] == "Alice"
        assert rows[0]["Danger Matchups"] == 1

    def test_a_lineup_without_a_risk_block_is_skipped_not_blank_filled(self):
        """An older document written before lineup risk existed -- absence
        is a real fact about the document, not an unknown-risk lineup."""
        assert lineup_risk_rows(payload()) == []

    def test_a_missing_anchor_shows_as_no_data_not_an_invented_name(self):
        document = payload_with_risk()
        document["lineups"][0]["lineup_risk"]["anchor_player_name"] = None
        document["lineups"][0]["lineup_risk"]["anchor_stability_score"] = None
        rows = lineup_risk_rows(document)
        assert rows[0]["Anchor"] == "No data"
        assert rows[0]["Anchor Stability"] == "No data"

    def test_the_footer_is_written_below_the_table_and_outside_the_autofilter(self, tmp_path):
        workbook_path = tmp_path / "workbook.xlsx"
        document_path = tmp_path / "lineups.json"
        workbook = Workbook()
        workbook.active.title = "Existing"
        workbook.save(workbook_path)
        document_path.write_text(json.dumps(payload_with_risk()), encoding="utf-8")

        append_lineup_optimizer_sheet(workbook_path, document_path)
        sheet = load_workbook(workbook_path)[LINEUP_SHEET]

        # The filter still covers ONLY the header + the one real data row --
        # a captain filtering the pairing table must never filter away the
        # summary rows that explain it.
        assert sheet.auto_filter.ref == "A1:N2"

        values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        flattened = [row[0] for row in values]
        assert LINEUP_RISK_TITLE in flattened
        title_index = flattened.index(LINEUP_RISK_TITLE)
        assert title_index > 1  # strictly below the header + data rows

        header = [cell for cell in values[title_index + 1] if cell is not None]
        assert header == LINEUP_RISK_COLUMNS
        summary = values[title_index + 2]
        assert summary[0] == "Chalk It Up"
        assert summary[4] == 0.36  # Lineup Risk Score
        assert summary[6] == "Alice"  # Anchor

    def test_no_footer_at_all_when_no_lineup_carries_risk(self, tmp_path):
        workbook_path = tmp_path / "workbook.xlsx"
        document_path = tmp_path / "lineups.json"
        workbook = Workbook()
        workbook.active.title = "Existing"
        workbook.save(workbook_path)
        document_path.write_text(json.dumps(payload()), encoding="utf-8")

        append_lineup_optimizer_sheet(workbook_path, document_path)
        sheet = load_workbook(workbook_path)[LINEUP_SHEET]
        flattened = [row[0].value for row in sheet.iter_rows()]
        assert LINEUP_RISK_TITLE not in flattened
        assert sheet.max_row == 2  # header + the single data row, nothing more
