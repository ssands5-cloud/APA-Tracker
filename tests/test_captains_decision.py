"""Tests for the Captain's Decision Engine.

Covers every case the spec's test list names: matchup_score normalisation,
risk_factor boundaries, confidence clamping, lineup ordering, stale pruning,
UI headers and sorting, and Excel formatting.

Named test_captains_decision to avoid colliding with
tests/test_build_captains_edge.py, which covers the older cheat-sheet
reporter in the same builder file.
"""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from analytics.captains_edge import (
    CONFIDENCE_BASE,
    Bounds,
    ScoreBounds,
    build_rationale,
    compute_bounds,
    confidence,
    lineup_recommendation,
    matchup_score,
    normalize,
    risk_factor,
)
from ui.tabs.captains_edge import (
    COLUMNS,
    HIGH_CONFIDENCE,
    HIGH_RISK,
    NO_DATA,
    load_rows,
    render,
    row_class,
)


def pairing(player="P1", opponent="O1", win=0.6, points=2.0, balls=None,
            player_name="Alice", opponent_name="Bob", team_pk=1):
    return {
        "player_id": player, "player_name": player_name,
        "opponent_id": opponent, "opponent_name": opponent_name,
        "win_probability": win, "expected_points": points,
        "expected_balls": balls, "team_pk": team_pk,
    }


FULL_BOUNDS = ScoreBounds(expected_points=Bounds(0.0, 4.0),
                          expected_balls=Bounds(0.0, 50.0))


# =============================================================================
# Normalisation
# =============================================================================

class TestNormalize:
    def test_it_maps_to_zero_one_against_population_bounds(self):
        assert normalize(0.0, Bounds(0.0, 4.0)) == 0.0
        assert normalize(4.0, Bounds(0.0, 4.0)) == 1.0
        assert normalize(1.0, Bounds(0.0, 4.0)) == 0.25

    def test_a_missing_value_is_none(self):
        assert normalize(None, Bounds(0.0, 4.0)) is None

    def test_unusable_bounds_yield_none(self):
        """A component nothing in the population had cannot be normalised."""
        assert normalize(2.0, Bounds()) is None

    def test_a_degenerate_range_is_full_marks_not_a_division_by_zero(self):
        """If every candidate expects the same output, none is disadvantaged
        by that component."""
        assert normalize(3.0, Bounds(3.0, 3.0)) == 1.0


class TestComputeBounds:
    def test_it_takes_min_and_max_across_the_population(self):
        bounds = compute_bounds([
            pairing(points=1.0, balls=10.0),
            pairing(points=4.0, balls=30.0),
        ])
        assert (bounds.expected_points.low, bounds.expected_points.high) == (1.0, 4.0)
        assert (bounds.expected_balls.low, bounds.expected_balls.high) == (10.0, 30.0)

    def test_an_all_null_component_is_unusable(self):
        bounds = compute_bounds([pairing(balls=None), pairing(balls=None)])
        assert not bounds.expected_balls.usable


# =============================================================================
# matchup_score
# =============================================================================

class TestMatchupScore:
    def test_it_applies_the_specified_weights(self):
        score = matchup_score(1.0, 4.0, 50.0, FULL_BOUNDS)
        # every component at full marks -> 1.0 regardless of weighting
        assert score == pytest.approx(1.0)

    def test_it_matches_the_formula_longhand(self):
        win, points, balls = 0.5, 1.0, 10.0
        expected = (0.6 * win
                    + 0.3 * normalize(points, FULL_BOUNDS.expected_points)
                    + 0.1 * normalize(balls, FULL_BOUNDS.expected_balls))
        assert matchup_score(win, points, balls, FULL_BOUNDS) == pytest.approx(
            expected, abs=1e-6
        )

    def test_win_probability_dominates(self):
        """0.6 of the weight. A better win chance must beat better output."""
        strong_win = matchup_score(0.9, 0.0, 0.0, FULL_BOUNDS)
        strong_output = matchup_score(0.1, 4.0, 50.0, FULL_BOUNDS)
        assert strong_win > strong_output

    def test_a_null_component_is_skipped_and_weights_renormalise(self):
        """A 9-ball pairing has no expected_points by construction. Without
        renormalising it would score up to 0.3 lower purely for lacking a
        field that does not apply to it."""
        with_balls_only = matchup_score(0.8, None, 50.0, FULL_BOUNDS)
        assert with_balls_only == pytest.approx(
            (0.6 * 0.8 + 0.1 * 1.0) / 0.7, abs=1e-6
        )

    def test_win_probability_alone_returns_itself(self):
        assert matchup_score(0.73, None, None, FULL_BOUNDS) == pytest.approx(0.73)

    def test_an_all_null_pairing_has_no_score(self):
        """No evidence is not a score of zero."""
        assert matchup_score(None, None, None, FULL_BOUNDS) is None

    def test_the_score_stays_within_zero_and_one(self):
        assert 0.0 <= matchup_score(1.0, 4.0, 50.0, FULL_BOUNDS) <= 1.0
        assert 0.0 <= matchup_score(0.0, 0.0, 0.0, FULL_BOUNDS) <= 1.0


# =============================================================================
# risk_factor
# =============================================================================

class TestRiskFactor:
    def test_it_matches_the_formula(self):
        assert risk_factor(0.5, 0.5) == pytest.approx(0.25)

    def test_a_perfectly_stable_player_has_no_risk(self):
        assert risk_factor(0.0, 1.0) == 0.0

    def test_it_clamps_at_one(self):
        assert risk_factor(5.0, 0.1) == 1.0

    def test_it_clamps_at_zero(self):
        """Stability above 1 would drive the product negative."""
        assert risk_factor(0.5, 1.5) == 0.0

    def test_more_volatility_means_more_risk(self):
        assert risk_factor(0.8, 0.5) > risk_factor(0.2, 0.5)

    def test_missing_inputs_yield_none(self):
        assert risk_factor(None, 0.5) is None
        assert risk_factor(0.5, None) is None


# =============================================================================
# confidence
# =============================================================================

class TestConfidence:
    def test_the_flag_sets_the_base(self):
        assert confidence(0.0, "HOT") == pytest.approx(CONFIDENCE_BASE["HOT"])
        assert confidence(0.0, "NEUTRAL") == pytest.approx(CONFIDENCE_BASE["NEUTRAL"])
        assert confidence(0.0, "COLD") == pytest.approx(CONFIDENCE_BASE["COLD"])

    def test_the_slope_nudges_it(self):
        assert confidence(0.1, "NEUTRAL") == pytest.approx(0.60)
        assert confidence(-0.1, "NEUTRAL") == pytest.approx(0.40)

    def test_it_clamps_at_one(self):
        assert confidence(5.0, "HOT") == 1.0

    def test_it_clamps_at_zero(self):
        assert confidence(-5.0, "COLD") == 0.0

    def test_a_cold_player_climbing_is_not_written_off(self):
        assert confidence(0.2, "COLD") > confidence(0.0, "COLD")

    def test_missing_inputs_yield_none(self):
        assert confidence(None, "HOT") is None
        assert confidence(0.1, None) is None

    def test_an_unrecognised_flag_is_none_not_guessed(self):
        """An unknown flag is a bug upstream; inventing a base would hide it."""
        assert confidence(0.1, "WARM") is None

    def test_the_flag_is_matched_case_insensitively(self):
        assert confidence(0.0, "hot") == pytest.approx(CONFIDENCE_BASE["HOT"])


# =============================================================================
# lineup_recommendation
# =============================================================================

class TestLineupOrdering:
    def test_players_are_ranked_by_matchup_score_descending(self):
        entries = lineup_recommendation([
            pairing(player="P1", player_name="Low", win=0.2),
            pairing(player="P2", player_name="High", win=0.9),
            pairing(player="P3", player_name="Mid", win=0.5),
        ], trends={})
        assert [e.player_name for e in entries] == ["High", "Mid", "Low"]
        assert [e.recommended_order for e in entries] == [1, 2, 3]

    def test_order_runs_one_to_n_and_is_not_padded_to_five(self):
        """A team with three scored players has three recommendations.
        Inventing two more would be fabrication."""
        entries = lineup_recommendation([
            pairing(player=f"P{i}", player_name=f"Player {i}", win=0.5)
            for i in range(3)
        ], trends={})
        assert [e.recommended_order for e in entries] == [1, 2, 3]

    def test_one_entry_per_player_keeping_their_best_pairing(self):
        """A player has many possible opponents; the lineup answers which one
        they should play."""
        entries = lineup_recommendation([
            pairing(player="P1", opponent="O1", opponent_name="Weak", win=0.9),
            pairing(player="P1", opponent="O2", opponent_name="Strong", win=0.2),
        ], trends={})
        assert len(entries) == 1
        assert entries[0].opponent_name == "Weak"

    def test_an_unscored_player_is_listed_but_unranked(self):
        """Ordering a player with no score would imply a judgement the
        evidence does not support."""
        entries = lineup_recommendation([
            pairing(player="P1", player_name="Scored", win=0.7),
            pairing(player="P2", player_name="Unscored", win=None,
                    points=None, balls=None),
        ], trends={})
        by_name = {e.player_name: e for e in entries}
        assert by_name["Scored"].recommended_order == 1
        assert by_name["Unscored"].recommended_order is None
        assert by_name["Unscored"].matchup_score is None

    def test_trends_supply_risk_and_confidence(self):
        entries = lineup_recommendation(
            [pairing(player="P1")],
            trends={"P1": {"volatility": 0.5, "sl_stability": 0.5,
                           "regression_slope": 0.1, "hot_cold_flag": "HOT"}},
        )
        assert entries[0].risk_factor == pytest.approx(0.25)
        assert entries[0].confidence == pytest.approx(0.85)

    def test_a_player_with_no_trend_row_still_ranks(self):
        """Missing form data must not remove a player from the lineup."""
        entries = lineup_recommendation([pairing(player="P1")], trends={})
        assert entries[0].recommended_order == 1
        assert entries[0].risk_factor is None
        assert entries[0].confidence is None

    def test_an_empty_pairing_list_yields_no_lineup(self):
        assert lineup_recommendation([], trends={}) == []


class TestRationale:
    def test_it_names_the_opponent_and_the_score(self):
        text = build_rationale(0.75, 0.1, 0.8, "Bob")
        assert "Bob" in text and "0.75" in text

    def test_it_says_so_when_form_is_unknown(self):
        assert "form unknown" in build_rationale(0.5, 0.1, None, "Bob")

    def test_it_flags_high_volatility(self):
        assert "volatility" in build_rationale(0.5, 0.9, 0.5, "Bob")

    def test_an_unscored_pairing_explains_itself(self):
        assert "not enough history" in build_rationale(None, None, None, "Bob")


# =============================================================================
# Builder: document shape, idempotency, stale pruning
# =============================================================================

class TestDecisionDocument:
    def _document(self, tmp_path):
        from scripts.build_captains_edge import build_decision_document, connect_read_only
        from scripts.build_captains_edge import resolve_db_path

        connection = connect_read_only(resolve_db_path())
        try:
            return build_decision_document(connection)
        finally:
            connection.close()

    def test_it_reports_how_the_roster_was_resolved(self, tmp_path):
        """Name-matching must never be mistaken for an id join."""
        document = self._document(tmp_path)
        assert "roster_resolution" in document
        assert document["roster_resolution"]

    def test_it_records_the_normalisation_bounds(self, tmp_path):
        """The scores are population-relative; without the bounds a reader
        cannot tell what a 0.7 was measured against."""
        document = self._document(tmp_path)
        assert set(document["normalization"]) == {"expected_points", "expected_balls"}

    def test_lineups_carry_ranked_players(self, tmp_path):
        document = self._document(tmp_path)
        assert document["lineups"]
        players = document["lineups"][0]["players"]
        ranked = [p for p in players if p["recommended_order"] is not None]
        assert [p["recommended_order"] for p in ranked] == list(
            range(1, len(ranked) + 1)
        )

    def test_rebuilding_is_idempotent(self, tmp_path):
        """The document is derived entirely from current rows."""
        first = self._document(tmp_path)
        second = self._document(tmp_path)
        assert first["lineups"] == second["lineups"]

    def test_the_file_is_rewritten_whole_so_stale_lineups_cannot_survive(self, tmp_path):
        """Pruning by full replacement: a team that no longer has pairings
        must not linger from a previous build."""
        from scripts.build_captains_edge import resolve_db_path, write_decision_json

        path = tmp_path / "captains_edge.json"
        path.write_text(json.dumps({
            "lineups": [{"team_id": "GONE", "team_name": "Disbanded", "players": []}]
        }), encoding="utf-8")

        write_decision_json(str(resolve_db_path()), path)
        rebuilt = json.loads(path.read_text(encoding="utf-8"))
        assert all(l["team_id"] != "GONE" for l in rebuilt["lineups"])


# =============================================================================
# UI
# =============================================================================

class TestCaptainsEdgeTab:
    def test_the_headers_match_the_spec_order(self):
        assert [label for _, label, _ in COLUMNS] == [
            "Player", "Opponent", "Matchup Score", "Risk", "Confidence",
            "Recommended Order", "Rationale",
        ]

    def test_high_confidence_is_highlighted(self):
        assert "high-confidence" in row_class({"confidence": HIGH_CONFIDENCE})

    def test_high_risk_is_highlighted(self):
        assert "high-risk" in row_class({"risk_factor": HIGH_RISK})

    def test_both_can_apply_at_once(self):
        """A strong run on an unsettled skill level is worth flagging twice."""
        classes = row_class({"confidence": 0.9, "risk_factor": 0.9})
        assert "high-confidence" in classes and "high-risk" in classes

    def test_an_ordinary_row_is_not_highlighted(self):
        assert row_class({"confidence": 0.5, "risk_factor": 0.1}) == ""

    def test_missing_values_never_trigger_a_highlight(self):
        assert row_class({"confidence": None, "risk_factor": None}) == ""

    def test_rows_are_sorted_by_recommended_order_with_unranked_last(self):
        rows = load_rows({"lineups": [{"team_id": "1", "team_name": "T", "players": [
            {"player_name": "Third", "recommended_order": 3},
            {"player_name": "Unranked", "recommended_order": None},
            {"player_name": "First", "recommended_order": 1},
        ]}]})
        assert [r["player_name"] for r in rows] == ["First", "Third", "Unranked"]

    def test_it_can_filter_to_one_team(self):
        document = {"lineups": [
            {"team_id": "1", "team_name": "A", "players": [{"player_name": "Alice"}]},
            {"team_id": "2", "team_name": "B", "players": [{"player_name": "Bob"}]},
        ]}
        assert [r["player_name"] for r in load_rows(document, team_id="2")] == ["Bob"]

    def test_nulls_render_as_no_data(self):
        html = render([{"player_name": "Alice", "opponent_name": None,
                        "matchup_score": None, "risk_factor": None,
                        "confidence": None, "recommended_order": None,
                        "rationale": ""}])
        assert NO_DATA in html

    def test_it_warns_when_the_roster_was_matched_by_name(self):
        html = render([{"player_name": "Alice", "matchup_score": 0.5,
                        "recommended_order": 1, "rationale": ""}],
                      resolution="player_name (id spaces do not join)")
        assert "NAME" in html

    def test_it_carries_no_external_resources(self):
        html = render([{"player_name": "Alice", "matchup_score": 0.5,
                        "recommended_order": 1, "rationale": ""}])
        assert "http://" not in html and "https://" not in html

    def test_an_empty_tab_explains_what_to_run(self):
        assert "build_captains_edge" in render([])


# =============================================================================
# Excel
# =============================================================================

class TestCaptainsEdgeSheet:
    def _sheet(self, tmp_path):
        from sqlalchemy.orm import Session

        from database.engine import create_db_engine
        from scheduler.graphql_sync import load_config
        from ui.export_excel import export_to_excel

        config = load_config("apa_config.yaml")
        with Session(create_db_engine(config)) as db:
            path = export_to_excel(
                db, {"export": {"excel_output_path": str(tmp_path / "wb.xlsx")}}
            )
        return load_workbook(path)["Captain's Edge"]

    def test_headers_match_the_spec_order(self, tmp_path):
        sheet = self._sheet(tmp_path)
        assert [c.value for c in next(sheet.iter_rows(max_row=1))] == [
            "Player", "Opponent", "Matchup Score", "Risk", "Confidence",
            "Recommended Order", "Rationale",
        ]

    def test_the_header_is_frozen_and_filtered(self, tmp_path):
        sheet = self._sheet(tmp_path)
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None

    def test_confidence_and_risk_are_conditionally_formatted(self, tmp_path):
        sheet = self._sheet(tmp_path)
        formulas = [rule.formula[0]
                    for rules in sheet.conditional_formatting
                    for rule in rules.rules]
        assert str(HIGH_CONFIDENCE) in formulas
        assert str(HIGH_RISK) in formulas

    def test_the_other_sheets_still_ship(self, tmp_path):
        from sqlalchemy.orm import Session

        from database.engine import create_db_engine
        from scheduler.graphql_sync import load_config
        from ui.export_excel import export_to_excel

        config = load_config("apa_config.yaml")
        with Session(create_db_engine(config)) as db:
            path = export_to_excel(
                db, {"export": {"excel_output_path": str(tmp_path / "wb.xlsx")}}
            )
        names = load_workbook(path).sheetnames
        assert "Matchups" in names
        assert "Head-to-Head" in names
        assert "Player Trends" in names
