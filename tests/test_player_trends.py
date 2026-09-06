"""Tests for the Player Trend Analyzer.

Every assertion here traces to the finalized governing spec: the ddof=1
volatility, the 1/(1+sigma) stability, the absolute hot/cold thresholds, the
probability formula and its constants, and the minimum-evidence table. Where
a value is checked numerically it is cross-checked against an independent
computation (statistics.stdev, or the formula written out longhand) rather
than against a number this implementation happened to produce.
"""

from __future__ import annotations

import math
import statistics

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from analytics.player_trends import (
    COLD_SLOPE_MAX,
    HOT_SLOPE_MIN,
    MIN_OBSERVATIONS_FLAG,
    MIN_OBSERVATIONS_PROBABILITY,
    PROBABILITY_SAMPLE_CAP,
    STABLE_VOLATILITY_MAX,
    VOLATILITY_WINDOW,
    evaluate,
    hot_cold_flag,
    projected_sl_change_probability,
    rolling_window,
    sl_stability,
    trend_slope,
    trend_strength,
    volatility,
)
from database.ingest import ingest_player_trends
from database.models import Base, Match, Player, PlayerMatch, PlayerTrend
from scripts.build_player_trends import build_rows, grouped_history
from ui.tabs.trends import NO_DATA, load_rows, render, row_class


# --- spec §1: volatility -----------------------------------------------------

class TestVolatility:
    def test_it_is_the_SAMPLE_standard_deviation(self):
        """ddof=1, cross-checked against the standard library rather than
        against our own output."""
        levels = [4, 4, 5, 5, 6]
        assert volatility(levels) == pytest.approx(statistics.stdev(levels), abs=1e-4)

    def test_it_is_not_the_population_standard_deviation(self):
        """The two differ, and picking the wrong one would silently shift
        every stability and probability downstream."""
        levels = [4, 6]
        assert volatility(levels) != pytest.approx(statistics.pstdev(levels), abs=1e-4)

    def test_an_unchanged_skill_level_has_zero_spread(self):
        assert volatility([5, 5, 5]) == 0.0

    def test_one_reading_is_null_not_zero(self):
        """Spec: do not fabricate zeros. One match has no spread, which is a
        different fact from zero spread."""
        assert volatility([5]) is None
        assert volatility([]) is None

    def test_missing_skill_levels_are_excluded_not_treated_as_zero(self):
        assert volatility([5, None, 5]) == 0.0


# --- spec §2: stability ------------------------------------------------------

class TestStability:
    def test_it_is_the_inverse_of_one_plus_volatility(self):
        assert sl_stability(0.0) == 1.0
        assert sl_stability(1.0) == 0.5
        assert sl_stability(0.5477) == pytest.approx(1 / 1.5477, abs=1e-4)

    def test_it_falls_as_volatility_rises(self):
        assert sl_stability(0.2) > sl_stability(0.9)

    def test_no_volatility_means_no_stability(self):
        """Stability derived from no observed spread would assert steadiness
        that was never measured."""
        assert sl_stability(None) is None


# --- spec §5: regression -----------------------------------------------------

class TestTrendSlope:
    def test_a_flat_history_has_zero_slope(self):
        assert trend_slope([5, 5, 5, 5]) == 0.0

    def test_a_rising_skill_level_has_a_positive_slope(self):
        assert trend_slope([4, 5, 6]) == pytest.approx(1.0)

    def test_a_falling_skill_level_has_a_negative_slope(self):
        assert trend_slope([6, 5, 4]) == pytest.approx(-1.0)

    def test_it_matches_least_squares_computed_longhand(self):
        ys = [4, 4, 5, 5, 6, 6]
        n = len(ys)
        xs = list(range(1, n + 1))
        expected = ((n * sum(x * y for x, y in zip(xs, ys)) - sum(xs) * sum(ys))
                    / (n * sum(x * x for x in xs) - sum(xs) ** 2))
        assert trend_slope(ys) == pytest.approx(expected, abs=1e-6)

    def test_two_observations_are_enough(self):
        """Spec §6 sets the minimum at 2, not 3."""
        assert trend_slope([4, 5]) is not None

    def test_one_observation_is_null(self):
        assert trend_slope([5]) is None
        assert trend_slope([]) is None

    def test_it_is_not_capped_at_the_volatility_window(self):
        """Spec §5: the slope uses ALL matches, while volatility uses the
        last 20. A long steady history followed by a climb must not read the
        same as the climb alone."""
        long_history = [3] * 30 + [7]
        assert trend_slope(long_history) != trend_slope(rolling_window(long_history))


class TestTrendStrength:
    def test_it_is_zero_for_a_flat_trend(self):
        assert trend_strength(0.0) == 0.0

    def test_it_is_direction_agnostic(self):
        assert trend_strength(0.3) == trend_strength(-0.3)

    def test_it_stays_within_zero_and_one(self):
        assert 0.0 <= trend_strength(50.0) <= 1.0

    def test_it_is_null_without_a_slope(self):
        assert trend_strength(None) is None


# --- spec §3: hot / cold -----------------------------------------------------

class TestHotColdFlag:
    def test_a_steady_climb_is_hot(self):
        assert hot_cold_flag(HOT_SLOPE_MIN, 0.2, 10) == "hot"

    def test_a_steady_decline_is_cold(self):
        assert hot_cold_flag(COLD_SLOPE_MAX, 0.2, 10) == "cold"

    def test_a_climb_with_erratic_skill_levels_is_not_hot(self):
        """Both conditions are required: a steep slope through a volatile
        skill level is noise, not a trend."""
        assert hot_cold_flag(0.5, STABLE_VOLATILITY_MAX + 0.01, 10) == "neutral"

    def test_a_shallow_slope_is_neutral(self):
        assert hot_cold_flag(0.01, 0.1, 10) == "neutral"

    def test_thresholds_are_inclusive(self):
        assert hot_cold_flag(HOT_SLOPE_MIN, STABLE_VOLATILITY_MAX, 10) == "hot"
        assert hot_cold_flag(COLD_SLOPE_MAX, STABLE_VOLATILITY_MAX, 10) == "cold"

    def test_too_little_evidence_is_null_not_neutral(self):
        """Spec §6: below 5 observations there is no flag. "Neutral" would
        claim an observation that was never made."""
        assert hot_cold_flag(0.5, 0.1, MIN_OBSERVATIONS_FLAG - 1) is None

    def test_no_volatility_means_no_flag(self):
        assert hot_cold_flag(0.5, None, 10) is None


# --- spec §4: projected probability ------------------------------------------

class TestProjectedProbability:
    def test_it_matches_the_spec_formula_longhand(self):
        slope, sigma, n = 0.3, 0.2, 10
        expected = 0.5 * math.tanh(4 * slope) * (1 - sigma) * (n / 20)
        assert projected_sl_change_probability(slope, sigma, n) == pytest.approx(
            expected, abs=1e-4
        )

    def test_a_bigger_sample_raises_it(self):
        few = projected_sl_change_probability(0.3, 0.2, 5)
        many = projected_sl_change_probability(0.3, 0.2, 20)
        assert many > few

    def test_higher_volatility_damps_it(self):
        calm = projected_sl_change_probability(0.3, 0.1, 10)
        erratic = projected_sl_change_probability(0.3, 0.6, 10)
        assert erratic < calm

    def test_the_sample_term_is_capped_at_twenty(self):
        at_cap = projected_sl_change_probability(0.3, 0.2, PROBABILITY_SAMPLE_CAP)
        beyond = projected_sl_change_probability(0.3, 0.2, PROBABILITY_SAMPLE_CAP * 5)
        assert at_cap == beyond

    def test_a_downward_trend_clamps_to_zero(self):
        """The heuristic describes upward pressure; direction lives in
        trend_slope, which is stored beside it."""
        assert projected_sl_change_probability(-0.5, 0.2, 10) == 0.0

    def test_volatility_above_one_clamps_to_zero_not_negative(self):
        assert projected_sl_change_probability(0.5, 1.8, 10) == 0.0

    def test_it_stays_within_zero_and_one(self):
        value = projected_sl_change_probability(50.0, 0.0, 20)
        assert 0.0 <= value <= 1.0

    def test_too_little_evidence_is_null(self):
        assert projected_sl_change_probability(
            0.5, 0.1, MIN_OBSERVATIONS_PROBABILITY - 1
        ) is None

    def test_no_volatility_means_no_probability(self):
        assert projected_sl_change_probability(0.5, None, 10) is None


# --- spec §6: minimum evidence, end to end -----------------------------------

class TestEvaluate:
    def test_a_single_match_yields_only_what_one_match_supports(self):
        result = evaluate(points=[2.0], skill_levels=[5], player_id="P1")
        assert result.matches_considered == 1
        assert result.avg_points_last_20 == 2.0
        assert result.volatility_last_20 is None
        assert result.trend_slope is None
        assert result.sl_stability is None
        assert result.hot_cold_flag is None
        assert result.projected_sl_change_probability is None

    def test_four_matches_give_a_slope_but_no_flag_or_probability(self):
        """The exact shape of this project's real data: enough for a slope,
        short of the 5 the flag and probability require."""
        result = evaluate(points=[1.0] * 4, skill_levels=[4, 4, 5, 5], player_id="P1")
        assert result.trend_slope is not None
        assert result.volatility_last_20 is not None
        assert result.hot_cold_flag is None
        assert result.projected_sl_change_probability is None

    def test_five_matches_unlock_the_gated_metrics(self):
        result = evaluate(points=[1.0] * 5, skill_levels=[4, 4, 5, 5, 5],
                          player_id="P1", format_="8-Ball Open")
        assert result.hot_cold_flag is not None
        assert result.projected_sl_change_probability is not None

    def test_volatility_is_windowed_but_the_slope_is_not(self):
        """Spec §1 vs §5. A history longer than the window must show the two
        spans diverging."""
        levels = [3] * 25 + [7] * 3
        result = evaluate(points=[1.0] * 28, skill_levels=levels, player_id="P1")
        assert result.matches_considered == VOLATILITY_WINDOW
        assert result.trend_slope == trend_slope(levels)


# --- database ----------------------------------------------------------------

@pytest.fixture
def db(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'trends.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session


@pytest.fixture
def seeded(db):
    """One player with five 8-ball matches and a climbing skill level."""
    player = Player(external_id="P1", name="Alice", skill_level=5)
    db.add(player)
    db.flush()

    for week, level in enumerate([4, 4, 5, 5, 5], start=1):
        match = Match(external_id=f"M{week}", week=week, format="8-Ball Open",
                      session_name="Fall 2026", home_team_id="T1", away_team_id="T2")
        db.add(match)
        db.flush()
        db.add(PlayerMatch(player_id=player.id, match_id=match.id,
                           skill_level=level, points_earned=2.0, result="W"))
    db.commit()
    return db


class TestDatabaseTable:
    def test_the_table_has_exactly_the_specified_columns(self):
        columns = {c.name for c in PlayerTrend.__table__.columns}
        assert columns == {
            "id", "player_id", "format", "matches_considered",
            "avg_points_last_20", "volatility_last_20", "trend_slope",
            "trend_strength", "sl_stability", "hot_cold_flag",
            "projected_sl_change_probability",
        }

    def test_nulls_are_stored_as_nulls_not_zeros(self, db):
        player = Player(external_id="P9", name="Thin", skill_level=5)
        db.add(player)
        db.commit()
        ingest_player_trends(db, [{
            "player_id": "P9", "format": "8-Ball Open", "matches_considered": 1,
            "avg_points_last_20": 2.0, "volatility_last_20": None,
            "trend_slope": None, "trend_strength": None, "sl_stability": None,
            "hot_cold_flag": None, "projected_sl_change_probability": None,
        }])
        row = db.query(PlayerTrend).one()
        assert row.volatility_last_20 is None
        assert row.hot_cold_flag is None
        assert row.projected_sl_change_probability is None

    def test_it_upserts_rather_than_duplicating(self, seeded):
        rows = build_rows(seeded)
        ingest_player_trends(seeded, rows)
        ingest_player_trends(seeded, rows)
        assert seeded.query(PlayerTrend).count() == len(rows)

    def test_an_unknown_player_is_skipped_not_guessed(self, seeded):
        assert ingest_player_trends(seeded, [{
            "player_id": "GHOST", "format": "8-Ball Open", "matches_considered": 5,
        }]) == 0


class TestBuilder:
    def test_it_groups_by_player_and_format(self, seeded):
        groups = grouped_history(seeded)
        assert len(groups) == 1
        assert list(groups)[0][1] == "8-Ball Open"

    def test_it_orders_history_chronologically(self, seeded):
        """Order is load-bearing: the regression reads against match order."""
        matches = list(grouped_history(seeded).values())[0]
        assert [m.skill_level for m in matches] == [4, 4, 5, 5, 5]

    def test_it_builds_one_row_per_player_format(self, seeded):
        rows = build_rows(seeded)
        assert len(rows) == 1
        assert rows[0]["player_id"] == "P1"
        assert rows[0]["matches_considered"] == 5
        assert rows[0]["trend_slope"] == pytest.approx(0.3, abs=1e-6)

    def test_a_match_with_no_format_is_skipped_not_guessed(self, db):
        player = Player(external_id="P2", name="Bob")
        db.add(player)
        db.flush()
        match = Match(external_id="MX", week=1, home_team_id="T1", away_team_id="T2")
        db.add(match)
        db.flush()
        db.add(PlayerMatch(player_id=player.id, match_id=match.id,
                           skill_level=5, points_earned=2.0))
        db.commit()
        assert build_rows(db) == []

    def test_an_empty_database_builds_nothing_rather_than_failing(self, db):
        assert build_rows(db) == []


# --- UI ----------------------------------------------------------------------

class TestTrendsTab:
    def test_hot_and_cold_rows_are_highlighted(self):
        assert row_class({"hot_cold_flag": "hot"}) == "hot"
        assert row_class({"hot_cold_flag": "cold"}) == "cold"

    def test_neutral_and_unmeasured_rows_are_not(self):
        assert row_class({"hot_cold_flag": "neutral"}) == ""
        assert row_class({"hot_cold_flag": None}) == ""

    def test_nulls_render_as_no_data_never_as_zero(self, seeded):
        ingest_player_trends(seeded, build_rows(seeded))
        html = render(load_rows(seeded))
        assert NO_DATA in html

    def test_it_renders_the_player_and_a_table(self, seeded):
        ingest_player_trends(seeded, build_rows(seeded))
        html = render(load_rows(seeded))
        assert "Alice" in html
        assert "<table" in html

    def test_it_carries_no_external_resources(self, seeded):
        """It has to open at a venue with no internet."""
        ingest_player_trends(seeded, build_rows(seeded))
        html = render(load_rows(seeded))
        assert "http://" not in html
        assert "https://" not in html

    def test_an_empty_tab_explains_what_to_run(self):
        assert "build_player_trends" in render([])


# --- Excel -------------------------------------------------------------------

class TestExcelSheet:
    def _workbook(self, db, tmp_path):
        from ui.export_excel import export_to_excel

        return load_workbook(export_to_excel(
            db, {"export": {"excel_output_path": str(tmp_path / "wb.xlsx")}}
        ))

    def test_the_sheet_exists_with_every_field(self, seeded, tmp_path):
        ingest_player_trends(seeded, build_rows(seeded))
        sheet = self._workbook(seeded, tmp_path)["Player Trends"]
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        assert header == ["Player", "Format", "Matches", "Avg Points",
                          "Slope (SL/match)", "Strength", "Volatility",
                          "SL Stability", "Trend", "SL Change Probability"]

    def test_headers_freeze_and_filter(self, seeded, tmp_path):
        ingest_player_trends(seeded, build_rows(seeded))
        sheet = self._workbook(seeded, tmp_path)["Player Trends"]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None

    def test_hot_and_cold_are_colour_coded(self, seeded, tmp_path):
        ingest_player_trends(seeded, build_rows(seeded))
        sheet = self._workbook(seeded, tmp_path)["Player Trends"]
        formulas = [rule.formula[0]
                    for rules in sheet.conditional_formatting
                    for rule in rules.rules]
        assert '"hot"' in formulas
        assert '"cold"' in formulas

    def test_slope_and_volatility_stay_numeric(self, seeded, tmp_path):
        """Spec §7: numeric, not strings -- otherwise the sheet cannot sort
        or chart them."""
        ingest_player_trends(seeded, build_rows(seeded))
        sheet = self._workbook(seeded, tmp_path)["Player Trends"]
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        row = dict(zip(header, [c.value for c in list(sheet.iter_rows(min_row=2))[0]]))
        assert isinstance(row["Slope (SL/match)"], (int, float))
        assert isinstance(row["Volatility"], (int, float))

    def test_the_other_sheets_still_ship(self, seeded, tmp_path):
        """Adding a sheet must not disturb the existing engines'."""
        names = self._workbook(seeded, tmp_path).sheetnames
        assert "Matchups" in names
        assert "Head-to-Head" in names
