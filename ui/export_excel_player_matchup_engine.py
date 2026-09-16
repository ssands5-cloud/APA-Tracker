"""Player Matchup Engine (Coach Mode) Excel: one filterable sheet over
already-built ``analytics.player_matchup_engine.PlayerMatchupReport``
objects. A REPORTER -- no computation happens here.
"""

from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analytics.player_matchup_engine import PlayerMatchupReport

COLUMNS = [
    "Our Team", "Opponent Team", "Player", "Player SL", "Player Trend (whole history)",
    "Player Volatility", "Opponent", "Opponent SL", "Opponent Trend (whole history)",
    "Opponent Volatility", "Format", "Session", "Evidence", "Observed Win Rate",
    "Direct W-L", "Direct Matches", "Modeled Probability", "Model Source", "Summary",
]


def _rows(reports: Sequence[PlayerMatchupReport]) -> list[list]:
    return [
        [
            r.our_team_external_id, r.opponent_team_external_id,
            r.player_name, r.player_skill_level, r.player_trend.trend, r.player_trend.volatility,
            r.opponent_name, r.opponent_skill_level, r.opponent_trend.trend, r.opponent_trend.volatility,
            r.format, r.session_name, r.evidence_label.value, r.observed_win_rate,
            f"{r.direct_wins}-{r.direct_losses}" if r.direct_wins is not None else "",
            r.direct_evidence_count, r.modeled_win_probability, r.model_source or "", r.summary,
        ]
        for r in reports
    ]


def write_workbook(reports: Sequence[PlayerMatchupReport], path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Player Matchup Engine"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")

    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    rows = _rows(reports)
    for row in rows:
        sheet.append(row)

    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"

    for index, name in enumerate(COLUMNS, start=1):
        width = max(len(str(name)) + 4, 12)
        for row in rows[:200]:
            width = max(width, min(len(str(row[index - 1])) + 2, 60))
        sheet.column_dimensions[get_column_letter(index)].width = width

    # Observed Win Rate and Modeled Probability read as percentages.
    for col in (14, 17):
        for row in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                cell.number_format = "0%"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
