"""Season Projection Excel export: a standalone workbook, per
docs/season_projection.md's "Excel layout" -- never a patch to the
existing apa_stats.xlsx.

Three values-only sheets: Season_Projection (one scope-level summary row),
Remaining_Matches (one row per real remaining match), and
Standings_History (one row per deduplicated real record change).
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analytics.season_projection import RemainingMatchProjection, SeasonProjection, StandingsPoint

SUMMARY_SHEET = "Season_Projection"
MATCHES_SHEET = "Remaining_Matches"
HISTORY_SHEET = "Standings_History"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")

MATCH_COLUMNS = (
    "Week", "Match ID", "Opponent Team ID", "Opponent", "Win Probability",
    "Probability Source", "Upset Likelihood",
)
HISTORY_COLUMNS = ("Captured At", "Wins", "Losses", "Win Rate")


def _header(sheet, columns: Sequence[str]) -> None:
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def build_workbook(
    projection: SeasonProjection,
    standings_points: Sequence[StandingsPoint],
    team_external_id: str,
    team_name: str,
    session_name: str,
    actual_wins: Optional[int],
    actual_losses: Optional[int],
    capture_time: Optional[str],
) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = SUMMARY_SHEET
    projected_final_wins = (
        round(actual_wins + projection.expected_wins, 2) if actual_wins is not None else None
    )
    projected_final_losses = (
        round(actual_losses + projection.expected_losses, 2) if actual_losses is not None else None
    )
    summary.append(["Field", "Value"])
    for cell in summary[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for field, value in (
        ("Team External ID", team_external_id),
        ("Team Name", team_name),
        ("Session", session_name),
        ("Actual Wins", actual_wins),
        ("Actual Losses", actual_losses),
        ("Remaining Matches", len(projection.matches)),
        ("Coverage", projection.coverage),
        ("Expected Remaining Wins", projection.expected_wins),
        ("Expected Remaining Losses", projection.expected_losses),
        ("Projected Final Wins", projected_final_wins),
        ("Projected Final Losses", projected_final_losses),
        ("Capture Time", capture_time),
    ):
        summary.append([field, value])
    for row in summary.iter_rows(min_row=2, max_row=summary.max_row, min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0%" if 0 <= cell.value <= 1 else "0.00"
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 24

    matches_sheet = workbook.create_sheet(MATCHES_SHEET)
    _header(matches_sheet, MATCH_COLUMNS)
    for m in projection.matches:
        matches_sheet.append([
            m.week, m.match_id, m.opponent_team_id, m.opponent_team_name,
            m.win_probability, m.probability_source.value, m.upset_likelihood,
        ])
    for index, width in enumerate((10, 14, 18, 22, 16, 18, 16), start=1):
        matches_sheet.column_dimensions[get_column_letter(index)].width = width
    last_column = get_column_letter(len(MATCH_COLUMNS))
    matches_sheet.auto_filter.ref = f"A1:{last_column}{max(matches_sheet.max_row, 1)}"

    if standings_points:
        history_sheet = workbook.create_sheet(HISTORY_SHEET)
        _header(history_sheet, HISTORY_COLUMNS)
        for p in standings_points:
            history_sheet.append([p.captured_at, p.wins, p.losses, p.win_rate])
        for index, width in enumerate((22, 10, 10, 12), start=1):
            history_sheet.column_dimensions[get_column_letter(index)].width = width

    return workbook


def write_workbook(
    projection: SeasonProjection,
    standings_points: Sequence[StandingsPoint],
    team_external_id: str,
    team_name: str,
    session_name: str,
    actual_wins: Optional[int],
    actual_losses: Optional[int],
    capture_time: Optional[str],
    path: str | Path,
) -> Path:
    workbook = build_workbook(
        projection, standings_points, team_external_id, team_name, session_name,
        actual_wins, actual_losses, capture_time,
    )
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
