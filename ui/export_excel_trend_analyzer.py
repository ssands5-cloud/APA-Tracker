"""Trend Analyzer Excel export: a standalone workbook, per
docs/trend_analyzer.md's "Excel layout" -- the dedicated production artifact
alongside the existing general-workbook ``Player Trends`` sheet (unchanged).

Two values-only sheets: Trend_Analyzer (one row per player/format/session)
and Trend_History (one row per real skill observation).
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from analytics.trend_analyzer import TrendAnalyzerReport

SUMMARY_SHEET = "Trend_Analyzer"
HISTORY_SHEET = "Trend_History"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")

ROW_COLUMNS = (
    "Player External ID", "Player Name", "Format", "Session", "Sample Size",
    "Current SL", "Regression Slope", "Volatility", "SL Stability", "Trend Score",
    "Hot/Cold Indicator", "Projected SL Change Probability",
)
HISTORY_COLUMNS = (
    "Player External ID", "Match ID", "Match Order", "Match Date", "Format",
    "Session", "Skill Level",
)


def _header(sheet, columns) -> None:
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def build_workbook(report: TrendAnalyzerReport) -> Workbook:
    workbook = Workbook()
    rows_sheet = workbook.active
    rows_sheet.title = SUMMARY_SHEET
    _header(rows_sheet, ROW_COLUMNS)
    for r in report.rows:
        rows_sheet.append([
            r.player_external_id, r.player_name, r.format, r.session_name, r.sample_size,
            r.current_skill_level, r.regression_slope, r.volatility, r.sl_stability,
            r.trend_score, r.hot_cold_indicator, r.projected_sl_change_probability,
        ])
    for index, width in enumerate((18, 22, 12, 14, 12, 12, 16, 14, 14, 14, 16, 24), start=1):
        rows_sheet.column_dimensions[get_column_letter(index)].width = width
    last_column = get_column_letter(len(ROW_COLUMNS))
    rows_sheet.auto_filter.ref = f"A1:{last_column}{max(rows_sheet.max_row, 1)}"

    if report.history:
        history_sheet = workbook.create_sheet(HISTORY_SHEET)
        _header(history_sheet, HISTORY_COLUMNS)
        for h in report.history:
            history_sheet.append([
                h.player_external_id, h.match_id, h.match_order, h.match_date,
                h.format, h.session_name, h.skill_level,
            ])
        for index, width in enumerate((18, 14, 12, 16, 12, 14, 12), start=1):
            history_sheet.column_dimensions[get_column_letter(index)].width = width

    return workbook


def write_workbook(report: TrendAnalyzerReport, path: str | Path) -> Path:
    workbook = build_workbook(report)
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
