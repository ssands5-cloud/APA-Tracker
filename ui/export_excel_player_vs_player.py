"""Player vs Player Excel export: a standalone workbook, never a patch to
the existing apa_stats.xlsx.

Per docs/player_vs_player_exports.md: "A later integration may copy the
same materialized rows into apa_stats.xlsx, but the production demo should
not patch the existing workbook after generation." This module only ever
builds a fresh, standalone `Workbook()` -- it never opens or modifies
`ui/export_excel.py`'s output.

Two sheets, real data only:

    Player_vs_Player   one row per feasible pairing (DIRECT/INDIRECT/UNKNOWN
                        alike), Stage 1 evidence plus the pair's own
                        real summary -- see analytics.player_vs_player_matrix.
    PvP_Game_History    one row per real recognized game, only when at
                        least one real game exists anywhere in the export.
                        Per the "no empty placeholder sheet" rule, this
                        sheet is OMITTED entirely (not written with headers
                        and no rows) when there is nothing real to put in
                        it.

Deterministic formatting: fixed column widths (never computed from cell
content), fixed column order, no conditional formatting, no volatile
formulas, no timestamps sourced from the local clock.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from analytics.player_vs_player_matrix import PlayerVsPlayerExportRow

SUMMARY_SHEET = "Player_vs_Player"
HISTORY_SHEET = "PvP_Game_History"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")

SUMMARY_COLUMNS = (
    "Session", "Format", "Our Team ID", "Opponent Team ID",
    "Player", "Player External ID", "Player SL",
    "Opponent", "Opponent External ID", "Opponent SL",
    "Evidence Label", "Direct Matches", "Observed Win Rate",
    "Total Games", "Wins", "Losses",
    "History Reliability (games)",
    "Last Recorded Skill-Gap Probability",
    "Modeled Win Probability (experimental)",
    "Trend (whole history)", "Trend (recent)",
    "Forward-Looking Pair Projection (experimental)",
)

SUMMARY_WIDTHS = {
    "Session": 14, "Format": 12, "Our Team ID": 16, "Opponent Team ID": 16,
    "Player": 20, "Player External ID": 14, "Player SL": 10,
    "Opponent": 20, "Opponent External ID": 14, "Opponent SL": 12,
    "Evidence Label": 12, "Direct Matches": 12, "Observed Win Rate": 14,
    "Total Games": 10, "Wins": 8, "Losses": 8,
    "History Reliability (games)": 18,
    "Last Recorded Skill-Gap Probability": 22,
    "Modeled Win Probability (experimental)": 22,
    "Trend (whole history)": 16, "Trend (recent)": 14,
    "Forward-Looking Pair Projection (experimental)": 24,
}

HISTORY_COLUMNS = (
    "Player External ID", "Opponent External ID", "Match ID", "Date",
    "Result", "Own SL", "Opponent SL", "Points Earned",
    "9-Ball Balls", "Format", "Session",
)

HISTORY_WIDTHS = {
    "Player External ID": 14, "Opponent External ID": 14, "Match ID": 10,
    "Date": 12, "Result": 8, "Own SL": 8, "Opponent SL": 10,
    "Points Earned": 12, "9-Ball Balls": 12, "Format": 12, "Session": 14,
}


def _style_header(sheet: Worksheet, columns: Sequence[str]) -> None:
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _apply_widths(sheet: Worksheet, columns: Sequence[str], widths: dict[str, int]) -> None:
    for index, name in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(name, 14)


def _summary_row(row: PlayerVsPlayerExportRow) -> list:
    summary = row.summary
    return [
        row.session_name, row.format, row.our_team_external_id, row.opponent_team_external_id,
        row.player_name, row.player_external_id, row.player_skill_level,
        row.opponent_name, row.opponent_external_id, row.opponent_skill_level,
        row.evidence_label.value, row.direct_matches, row.observed_win_rate,
        summary.total_games, summary.wins, summary.losses,
        summary.reliability, summary.skill_only_probability,
        summary.modeled_win_probability, summary.trend, summary.recent_trend,
        summary.next_match_projection,
    ]


def _history_rows(rows: Sequence[PlayerVsPlayerExportRow]) -> list[list]:
    history = []
    for row in rows:
        for game in row.summary.games:
            history.append([
                row.player_external_id, row.opponent_external_id,
                game.match_id, game.match_date, game.result,
                game.own_skill_level, game.opponent_skill_level,
                game.points_earned, game.nine_ball_points,
                game.format, game.session_name,
            ])
    return history


def build_workbook(rows: Sequence[PlayerVsPlayerExportRow]) -> Optional[Workbook]:
    """A fresh, standalone workbook for these real export rows, or ``None``
    when there are no feasible pairings at all -- never an empty
    placeholder workbook (see the module docstring's "no empty placeholder
    sheet" rule)."""
    if not rows:
        return None

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = SUMMARY_SHEET
    _style_header(summary_sheet, SUMMARY_COLUMNS)
    for row in rows:
        summary_sheet.append(_summary_row(row))
    _apply_widths(summary_sheet, SUMMARY_COLUMNS, SUMMARY_WIDTHS)
    last_column = get_column_letter(len(SUMMARY_COLUMNS))
    summary_sheet.auto_filter.ref = f"A1:{last_column}{max(summary_sheet.max_row, 1)}"

    history_rows = _history_rows(rows)
    if history_rows:
        history_sheet = workbook.create_sheet(HISTORY_SHEET)
        _style_header(history_sheet, HISTORY_COLUMNS)
        for game_row in history_rows:
            history_sheet.append(game_row)
        _apply_widths(history_sheet, HISTORY_COLUMNS, HISTORY_WIDTHS)
        last_history_column = get_column_letter(len(HISTORY_COLUMNS))
        history_sheet.auto_filter.ref = f"A1:{last_history_column}{max(history_sheet.max_row, 1)}"

    return workbook


def write_workbook(rows: Sequence[PlayerVsPlayerExportRow], path: str | Path) -> Optional[Path]:
    """Write the workbook to ``path``, or write nothing and return ``None``
    when there are no feasible pairings -- callers must not create an
    empty file either."""
    workbook = build_workbook(rows)
    if workbook is None:
        return None
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
