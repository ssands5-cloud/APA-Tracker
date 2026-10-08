"""Excel: the coach-facing half of the Ultimate Coach workbook.

Match Day is the control panel; everything else follows it:

    Match Day        Player -> Team -> Format -> Scheduled date -> Fixture
                     (yellow inputs), then the EFFECTIVE matchup (calculated).
    War Room         Captain's War Room for tonight: rosters, best-supported
                     sends, risks, Lineup Lab snapshot, color matrix, inspect
                     views, direct meetings, team comparison. Optional local
                     overrides (blank = follow Match Day).
    Lineup Lab       The captain's own planning marks (availability, lineup,
                     played, notes) + instant results. Marks are inputs, never
                     evidence.
    Scouting Cards   One card per opponent player.
    Captain Packet   Printable: page 1 summary, page 2+ detail.
    Coach Dashboard  Player A = you, Player B = an opponent from tonight's
                     roster or an explicit "All players..." choice.
    Engine           Every calculated selection as a named cell, kept apart
                     from inputs so exploring can't break the links.

Excel constraints honoured: ordinary formulas, tables, defined names and
data-validation dropdowns only -- no VBA/macros, no COM, no dynamic-array or
volatile functions. Where Excel can't do what the HTML does (erase a stale
input, move inputs when a roster changes, list every shared opponent), the
closest reliable workflow is used and the limitation is stated on-sheet.

Evidence rules come from analytics.ultimate_coach_matchup_evidence (ranking)
and analytics.ultimate_coach_war_room (categories, best sends, risks, cards);
nothing here re-scores or blends evidence.
"""

from __future__ import annotations

from collections import defaultdict
from typing import Any

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.pagebreak import Break

import ui.export_excel_ultimate_coach as base
from analytics.ultimate_coach_match_day import DEFAULT_MATCH_DAY_TIMEZONE, excel_date_serial
from analytics.ultimate_coach_matchup_evidence import (
    build_pair_index,
    fixture_scope_pairs,
    format_label,
    members_by_scope,
)
from analytics.ultimate_coach_war_room import (
    CATEGORY_LABELS,
    build_local_date,
    date_label,
    default_matchup,
    freshness,
    meetings_index,
    scope_dates,
    sl_bucket_index,
    suggested_date,
    war_room_pair,
    COACH_TAGS,
    MATCH_NIGHT_GUIDE,
    ONBOARDING_LIMITS,
    ONBOARDING_WHAT,
)

ME_COLUMNS = ["Key", "Rank", "Player", "SL", "Direct Record", "Shared-Opponent Evidence", "Basis", "Rows",
              "Player ID", "Category", "Cell", "Reason", "Captain Cell", "Rank Group"]
TC_COLUMNS = [
    "Pair Key", "Format", "Our Team", "Opponent Team", "Our Rostered", "Opp Rostered", "Our Captured SL",
    "Opp Captured SL", "Our SL Total", "Opp SL Total", "Our Games", "Opp Games", "Our No-Game Players",
    "Opp No-Game Players", "Direct Meetings", "Opponents Met", "Pairings",
]
THREATS_COLUMNS = ["Key", "Opp Player ID", "Threat"]
CONCERNING_COLUMNS = ["Key", "Our Player ID", "Opp Player ID", "Pairing"]
MEETINGS_COLUMNS = ["Key", "Date", "Our Player", "Opponent", "Result", "Our SL", "Their SL", "Session", "Points",
                    "Pair Key"]
SCOUTING_COLUMNS = ["Key", "Opponent", "SL", "Team Record", "Lifetime", "Sample", "Vs Our Roster", "Met List",
                    "Shared Summary", "By SL", "Winning SL", "Losing SL", "Missing", "Quick Read"]
DATEKEYS_COLUMNS = ["Key", "Group Key", "Serial", "Date Display"]
SUGGESTED_COLUMNS = ["Group Key", "Serial", "Date Display"]

ALL_PLAYERS_CHOICE = "All players…"
AVAILABILITY_CHOICES = ["Available", "Unavailable", "Unknown"]
LINEUP_CHOICES = ["Planned", "Played"]
CONCERNING_SLOTS = 20
MEETING_SLOTS = 40

def _cf_fill(rgb: str) -> PatternFill:
    """A fill for CONDITIONAL formatting. Excel paints a conditional (dxf) solid fill from its background
    colour, so both colours are set -- with only fgColor the rule fires but nothing is painted (real-Excel
    UAT: uncoloured matrix, invisible packet card headers)."""
    return PatternFill(fill_type="solid", fgColor=rgb, bgColor=rgb)


CAT_FILLS = {code: _cf_fill(rgb) for code, rgb in
             (("G", "CFE8D4"), ("R", "F4CCCC"), ("E", "FFF0B3"), ("I", "FFF0B3"), ("X", "E3E3E3"))}
CARD_HEAD_FILL_RGB = "E8DCCB"   # light rail tint; the name is dark text so it prints readably without any fill
CALC_FILL = PatternFill("solid", fgColor="EEF3EF")
BIG_LINK = Font(color="1F5C99", underline="single", bold=True, size=12)
SMALL = Font(size=9)
WARN_FONT = Font(bold=True, color="8A5A00")


# ---------------------------------------------------------------------------
# Lookup tables (static, precomputed per directed team pairing with a fixture)
# ---------------------------------------------------------------------------

def _text(value: Any, empty: str = "No data") -> str:
    return empty if value is None or value == "" else str(value)


def write_lookup_tables(wb, *, payload: dict[str, Any], match_day: dict[str, Any], rosters: list[dict[str, Any]],
                        teams: list[dict[str, Any]], pairs: list[dict[str, Any]], built_at: str) -> dict[str, Any]:
    players = payload.get("players") or []
    tz_name = match_day.get("display_timezone") or DEFAULT_MATCH_DAY_TIMEZONE
    index = build_pair_index(pairs)
    members = members_by_scope(rosters, players)
    roster_ids = {r["player_id"] for r in rosters}
    sl_index = sl_bucket_index(payload.get("evidence") or [], roster_ids)
    games = meetings_index(payload.get("evidence") or [], roster_ids, tz_name)
    players_by_id = {p["id"]: p for p in players}
    label_by_scope = {t["team_scope_key"]: t["team_label"] for t in teams}
    session_by_scope = {t["team_scope_key"]: t["session_name"] for t in teams}

    me_rows: list[list[Any]] = []
    tc_rows: list[list[Any]] = []
    threat_rows: list[list[Any]] = []
    concern_rows: list[list[Any]] = []
    meeting_rows: list[list[Any]] = []
    card_rows: list[list[Any]] = []
    max_threats = max_concern = max_meetings = 0
    for our, opp, fmt in fixture_scope_pairs(match_day):
        if our not in members or opp not in members:
            continue
        wr = war_room_pair(members[our], members[opp], index, fmt, sl_index=sl_index, meetings=games,
                           players_by_id=players_by_id, session_name=session_by_scope.get(opp, ""))
        key = f"{our}|{opp}|{fmt}"
        c = wr["comparison"]
        tc_rows.append([
            key, c["format"], label_by_scope.get(our, our), label_by_scope.get(opp, opp),
            c["ours"]["rostered"], c["theirs"]["rostered"], c["ours"]["captured_sl"], c["theirs"]["captured_sl"],
            c["ours"]["sl_total"], c["theirs"]["sl_total"], c["ours"]["games"], c["theirs"]["games"],
            c["ours"]["no_games"], c["theirs"]["no_games"], c["direct_meetings"], c["opponents_met"], c["pairings"],
        ])
        for j, block in enumerate(wr["blocks"], start=1):
            me_rows.append([f"{key}|{j}", "vs", block["opponent_label"], base._sl_display(block["opponent"]),
                            block["opponent_sample"], "—", block["note"], len(block["rows"]),
                            block["opponent"]["id"], "—", "—", "—", "—", "—"])
            for row in block["rows"]:
                me_rows.append([None, row["rank"], row["player"], base._sl_display(row["member"]), row["direct_text"],
                                row["shared_text"], row["basis"], row["position"], row["member"]["id"],
                                row["category"], row["cell"], row["reason"], row["captain"],
                                # Next Send medals: players with the same evidence share a rank group (and a medal).
                                int(row["rank"].rstrip("=")) if row["category"] in ("G", "E") else "—"])
        for k, threat in enumerate(wr["threats"], start=1):
            threat_rows.append([f"{key}|{k}", threat["opponent"]["id"], threat["threat_text"]])
        max_threats = max(max_threats, len(wr["threats"]))
        for k, item in enumerate(wr["concerning"][:CONCERNING_SLOTS], start=1):
            concern_rows.append([f"{key}|{k}", item["our"]["id"], item["opp"]["id"], item["text"]])
        max_concern = max(max_concern, min(len(wr["concerning"]), CONCERNING_SLOTS))
        for k, game in enumerate(wr["meetings"], start=1):
            meeting_rows.append([f"{key}|{k}", game["date"], base._player_label(game["our"]["name"], game["our"]["external_id"]),
                                 base._player_label(game["opp"]["name"], game["opp"]["external_id"]),
                                 _text(game["result"]), _text(game["own_sl"]), _text(game["opp_sl"]),
                                 _text(game["session"]), _text(game["points"]), key])
        max_meetings = max(max_meetings, len(wr["meetings"]))
        for j, card in enumerate(wr["cards"], start=1):
            card_rows.append([f"{key}|{j}", card["label"], card["sl"], card["team_record"], card["lifetime"],
                              card["sample"], card["vs_ours"], card["met_list"], card["shared_summary"],
                              card["by_sl"], card["winning_sl"], card["losing_sl"], card["missing"],
                              card["quick_read"]])

    def table(name: str, sheet_name: str, columns: list[str], rows: list[list[Any]], widths: dict[str, int]) -> None:
        ws = wb.create_sheet(sheet_name)
        base._style_header(ws, columns)
        for row in rows:
            ws.append(row)
        base._apply_widths(ws, columns, widths)
        base._autotable(ws, name, columns, len(rows))

    table("TeamComparison_Table", "Team Comparison", TC_COLUMNS, tc_rows,
          {"Pair Key": 40, "Our Team": 32, "Opponent Team": 32, "Direct Meetings": 28, "Pairings": 44})
    table("MatchupEvidence_Table", "Matchup Evidence", ME_COLUMNS, me_rows,
          {"Key": 44, "Rank": 6, "Player": 40, "SL": 8, "Direct Record": 20, "Shared-Opponent Evidence": 60,
           "Basis": 70, "Rows": 6, "Player ID": 10, "Category": 9, "Cell": 26, "Reason": 70})
    table("Threats_Table", "Threats", THREATS_COLUMNS, threat_rows, {"Key": 44, "Threat": 90})
    table("Concerning_Table", "Concerning", CONCERNING_COLUMNS, concern_rows, {"Key": 44, "Pairing": 90})
    table("Meetings_Table", "Meetings", MEETINGS_COLUMNS, meeting_rows,
          {"Key": 44, "Date": 17, "Our Player": 38, "Opponent": 38})
    table("Scouting_Table", "Scouting", SCOUTING_COLUMNS, card_rows,
          {"Key": 44, "Opponent": 38, "Vs Our Roster": 50, "Met List": 70, "By SL": 50, "Missing": 50, "Quick Read": 90})

    build_local = build_local_date(built_at, tz_name)
    dates = scope_dates(match_day)
    date_rows, suggested_rows = [], []
    max_dates = 0
    for (scope, code), isos in sorted(dates.items()):
        group = f"{scope}|{code}"
        for k, iso in enumerate(isos, start=1):
            date_rows.append([f"{group}|{k}", group, excel_date_serial(iso), date_label(iso)])
        max_dates = max(max_dates, len(isos))
        suggestion = suggested_date(isos, build_local)
        if suggestion:
            suggested_rows.append([group, excel_date_serial(suggestion), date_label(suggestion)])
    table("DateKeys_Table", "Date Keys", DATEKEYS_COLUMNS, date_rows, {"Key": 50, "Group Key": 44, "Date Display": 18})
    table("SuggestedDates_Table", "Suggested Dates", SUGGESTED_COLUMNS, suggested_rows,
          {"Group Key": 44, "Date Display": 18})
    return {
        "pairs": len(tc_rows), "evidence_rows": len(me_rows), "max_threats": max_threats,
        "max_concerning": max_concern, "max_meetings": max_meetings, "max_dates": max_dates,
        "build_local": build_local, "freshness": freshness(match_day, build_local), "tz": tz_name,
    }


# ---------------------------------------------------------------------------
# Engine: every calculated selection as a named cell
# ---------------------------------------------------------------------------

class Engine:
    def __init__(self, wb):
        self.wb = wb
        self.ws = wb.create_sheet("Engine")
        self.ws["A1"] = "Engine — calculated selections (do not edit)"
        self.ws["A1"].font = base.TITLE_FONT
        self.ws["A2"] = ("Every interactive tab reads these named cells. They are formulas derived from the yellow "
                         "input cells on Match Day, War Room, Lineup Lab and Coach Dashboard — editing here would break "
                         "the links, so change the inputs instead.")
        self.ws["A2"].font = base.MUTED_FONT
        for letter, width in zip("ABCDEFGHIJ", (34, 46, 30, 16, 16, 16, 16, 16, 16, 16)):
            self.ws.column_dimensions[letter].width = width
        self.row = 4

    def name(self, name: str, ref: str) -> None:
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)

    def cell(self, name: str, label: str, formula: str) -> str:
        r = self.row
        self.row += 1
        self.ws.cell(row=r, column=1, value=label).font = base.HELPER_FONT
        self.ws.cell(row=r, column=2, value=formula)
        self.ws.cell(row=r, column=3, value=name).font = base.HELPER_FONT
        self.name(name, f"Engine!$B${r}")
        return name

    def section(self, title: str) -> None:
        self.row += 1
        self.ws.cell(row=self.row, column=1, value=title).font = base.SUBHEAD_FONT
        self.row += 1

    def block(self, title: str, n: int, columns: list[tuple[str, Any]]) -> dict[str, tuple[str, int, int]]:
        """n rows; columns: (name, formula_fn(k, row) -> str). Registers a range name per column."""
        self.section(title)
        first = self.row
        out = {}
        for c, (name, _) in enumerate(columns, start=2):
            self.ws.cell(row=first - 1, column=c, value=name).font = base.HELPER_FONT
        for k in range(1, n + 1):
            r = first + k - 1
            self.ws.cell(row=r, column=1, value=f"{title} {k}").font = base.HELPER_FONT
            for c, (name, fn) in enumerate(columns, start=2):
                self.ws.cell(row=r, column=c, value=fn(k, r))
        last = first + n - 1
        for c, (name, _) in enumerate(columns, start=2):
            letter = get_column_letter(c)
            self.name(name, f"Engine!${letter}${first}:${letter}${last}")
            out[name] = (letter, first, last)
        self.row = last + 1
        return out


def _fmt_label_formula(code_ref: str) -> str:
    return (f'IF({code_ref}="EIGHT","8-Ball",IF({code_ref}="NINE","9-Ball",IF({code_ref}="MASTERS","Masters",'
            f'IF({code_ref}="MASTERS ALT","Masters Alt",{code_ref}))))')


def build_engine(wb, *, slots: dict[str, int], tz_name: str) -> dict[str, Any]:
    e = Engine(wb)
    R, T, D, F = slots["roster"], slots["teams"], slots["dates"], slots["fixtures"]
    md = "'Match Day'!"

    def guard(ref: str) -> str:
        return f'IF({ref}="","",{ref})'

    # ---- Match Day: player -> team -> format -> date -> fixture ----
    e.section("Match Day — player")
    e.cell("uc_ViewerKey", "Player record ID",
           f'=IF({md}$B$6="","",IF(ISNUMBER({md}$B$6),TEXT({md}$B$6,"0"),IFERROR(INDEX(Players_Table[External ID],'
           f'MATCH(TRIM({md}$B$6),Players_Table[Player Label],0)),TRIM({md}$B$6))))')
    e.cell("uc_ViewerName", "Player name",
           '=IF(uc_ViewerKey="","",IFERROR(INDEX(Players_Table[Player Name],MATCH(uc_ViewerKey,Players_Table[External ID],0)),""))')
    e.cell("uc_ViewerLabel", "Player label", '=IF(uc_ViewerName="","",uc_ViewerName&" (APA record ID "&uc_ViewerKey&")")')
    e.cell("uc_TeamCount", "Current team scopes",
           '=IF(uc_ViewerName="",0,COUNTIF(PlayerTeams_Table[APA Record ID],uc_ViewerKey))')
    e.cell("uc_FormatInput", "Format input", "=" + guard(f"{md}$B$8"))
    e.cell("uc_Code", "Format filter code",
           '=IF(OR(uc_FormatInput="",uc_FormatInput="8-Ball & 9-Ball"),"89",IF(uc_FormatInput="All recorded formats","ALL","F:"&uc_FormatInput))')
    e.cell("uc_FormatLabel", "Format label",
           '=IF(uc_Code="89","8-Ball & 9-Ball",IF(uc_Code="ALL","All recorded formats",uc_FormatInput))')
    teams = e.block("Team slot", T, [
        ("uc_TeamList", lambda k, r: f'=IF(uc_ViewerName="","",IFERROR(INDEX(PlayerTeams_Table[Team],'
                                     f'MATCH(uc_ViewerKey&"|{k}",PlayerTeams_Table[Viewer Slot Key],0)),""))'),
        ("uc_TeamScopes", lambda k, r: f'=IF(B{r}="","",IFERROR(INDEX(PlayerTeams_Table[Team Scope Key],'
                                       f'MATCH(uc_ViewerKey&"|{k}",PlayerTeams_Table[Viewer Slot Key],0)),""))'),
        ("uc_TeamNextSerial", lambda k, r: f'=IF(C{r}="","",IFERROR(INDEX(SuggestedDates_Table[Serial],'
                                           f'MATCH(C{r}&"|"&uc_Code,SuggestedDates_Table[Group Key],0)),""))'),
        ("uc_TeamNextDisplay", lambda k, r: f'=IF(C{r}="","",IFERROR(INDEX(SuggestedDates_Table[Date Display],'
                                            f'MATCH(C{r}&"|"&uc_Code,SuggestedDates_Table[Group Key],0)),"No upcoming fixture captured"))'),
    ])
    e.section("Match Day — team")
    e.cell("uc_TeamInput", "Team input", "=" + guard(f"{md}$B$7"))
    e.cell("uc_TeamInputValid", "Team input valid", '=AND(uc_TeamInput<>"",COUNTIF(uc_TeamList,uc_TeamInput)>0)')
    e.cell("uc_EarliestSerial", "Earliest upcoming date among the player's teams", "=MIN(uc_TeamNextSerial)")
    e.cell("uc_EarliestSlot", "Team slot with the earliest upcoming date",
           '=IF(uc_EarliestSerial=0,"",MATCH(uc_EarliestSerial,uc_TeamNextSerial,0))')
    e.cell("uc_TeamLabel", "Effective team",
           '=IF(uc_TeamInputValid,uc_TeamInput,IF(uc_TeamCount=0,"",IF(uc_TeamCount=1,INDEX(uc_TeamList,1),'
           'IF(uc_EarliestSlot<>"",INDEX(uc_TeamList,uc_EarliestSlot),""))))')
    e.cell("uc_TeamMsg", "Team note",
           '=IF(uc_ViewerName="","",IF(uc_TeamCount=0,"No current team captured for "&uc_ViewerName&".",'
           'IF(uc_TeamInputValid,"Your choice.",IF(uc_TeamInput<>"","⚠ “"&uc_TeamInput&"” is not one of "&uc_ViewerName&'
           '"’s current teams — ignored. ","")&IF(uc_TeamLabel="","Choose a team.",IF(uc_TeamCount=1,"Their only current team.",'
           '"Suggested: the team with the earliest scheduled date on or after the build date.")))))')
    e.cell("uc_TeamScope", "Effective team scope key",
           '=IF(uc_TeamLabel="","",IFERROR(INDEX(Teams_Table[Scope Key],MATCH(uc_TeamLabel,Teams_Table[Team],0)),""))')
    e.cell("uc_DateGroup", "Date group key", '=IF(uc_TeamScope="","",uc_TeamScope&"|"&uc_Code)')
    e.cell("uc_DateCount", "Scheduled dates", '=IF(uc_DateGroup="",0,COUNTIF(DateKeys_Table[Group Key],uc_DateGroup))')
    dates = e.block("Date slot", D, [
        ("uc_DateList", lambda k, r: f'=IF({k}>uc_DateCount,"",IFERROR(INDEX(DateKeys_Table[Date Display],'
                                     f'MATCH(uc_DateGroup&"|{k}",DateKeys_Table[Key],0)),""))'),
        ("uc_DateSerials", lambda k, r: f'=IF(B{r}="","",INDEX(DateKeys_Table[Serial],'
                                        f'MATCH(uc_DateGroup&"|{k}",DateKeys_Table[Key],0)))'),
    ])
    e.section("Match Day — date")
    e.cell("uc_SuggestedSerial", "Suggested date (serial)",
           '=IF(uc_DateGroup="","",IFERROR(INDEX(SuggestedDates_Table[Serial],MATCH(uc_DateGroup,SuggestedDates_Table[Group Key],0)),""))')
    e.cell("uc_DateInput", "Date input", "=" + guard(f"{md}$B$9"))
    e.cell("uc_DateInputSerial", "Date input (serial)",
           '=IF(uc_DateInput="","",IF(ISNUMBER(uc_DateInput),INT(uc_DateInput),IFERROR(INDEX(uc_DateSerials,'
           'MATCH(uc_DateInput,uc_DateList,0)),IFERROR(INT(DATEVALUE(uc_DateInput)),"bad"))))')
    e.cell("uc_DateInputValid", "Date input valid",
           '=IF(ISNUMBER(uc_DateInputSerial),COUNTIF(uc_DateSerials,uc_DateInputSerial)>0,FALSE)')
    e.cell("uc_DateSerial", "Effective date (serial)",
           '=IF(uc_DateInput="",uc_SuggestedSerial,IF(uc_DateInputValid,uc_DateInputSerial,""))')
    e.cell("uc_DateDisplay", "Effective date",
           '=IF(uc_DateSerial="","",IFERROR(INDEX(uc_DateList,MATCH(uc_DateSerial,uc_DateSerials,0)),""))')
    e.cell("uc_DateMsg", "Date note",
           '=IF(uc_TeamScope="","",IF(uc_DateInput="",IF(uc_SuggestedSerial="","No upcoming scheduled date was captured '
           'for this team and format — choose a date.","Suggested: the earliest scheduled date on or after the build date."),'
           'IF(uc_DateInputValid,"Your choice.","⚠ “"&IF(ISNUMBER(uc_DateInput),TEXT(uc_DateInput,"dddd, mmm d, yyyy"),'
           'uc_DateInput)&"” has no "&uc_TeamLabel&" fixture ("&uc_FormatLabel&") — ignored. Pick a date from the list.")))')
    e.cell("uc_FixGroup", "Fixture group key",
           '=IF(OR(uc_TeamScope="",uc_DateSerial=""),"",uc_TeamScope&"|"&uc_DateSerial&"|"&uc_Code)')
    e.cell("uc_FixCount", "Fixtures on the date",
           '=IF(uc_FixGroup="",0,COUNTIF(ScheduleKeys_Table[Group Key],uc_FixGroup))')
    fixtures = e.block("Fixture slot", F, [
        ("uc_FixRows", lambda k, r: f'=IF({k}>uc_FixCount,"",IFERROR(INDEX(ScheduleKeys_Table[Schedule Row],'
                                    f'MATCH(uc_FixGroup&"|{k}",ScheduleKeys_Table[Slot Key],0)),""))'),
        ("uc_FixtureList", lambda k, r: f'=IF(B{r}="","","{k} · vs "&INDEX(Schedule_Table[Opponent],B{r})&" · "&'
                                        f'INDEX(Schedule_Table[Kickoff],B{r})&" · "&INDEX(Schedule_Table[Home/Away],B{r})&'
                                        f'" · "&INDEX(Schedule_Table[Format],B{r}))'),
    ])
    e.section("Match Day — fixture")
    e.cell("uc_FixInput", "Fixture input", "=" + guard(f"{md}$B$10"))
    # A blank input must not MATCH an empty ("") fixture slot -- Excel matches zero-length text.
    e.cell("uc_FixChoice", "Chosen fixture slot",
           '=IF(uc_FixCount=0,"",IF(uc_FixCount=1,1,IF(uc_FixInput="","",IFERROR(MATCH(uc_FixInput,uc_FixtureList,0),""))))')
    e.cell("uc_SchedRow", "Schedule row", '=IF(uc_FixChoice="","",INDEX(uc_FixRows,uc_FixChoice))')
    e.cell("uc_FixMsg", "Fixture note",
           '=IF(uc_DateSerial="","",IF(uc_FixCount=0,"No fixture on this date.",IF(uc_FixCount=1,"One fixture on this date — '
           'used automatically.",IF(uc_FixChoice="",IF(uc_FixInput<>"","⚠ The fixture you picked is not on this date — ignored. ",'
           '"")&uc_FixCount&" fixtures on this date — pick one in the Fixture cell (none is chosen for you).","Your choice of "&'
           'uc_FixCount&" fixtures on this date."))))')
    e.cell("uc_OppScope", "Opponent scope key",
           '=IF(uc_SchedRow="","",IF(INDEX(Schedule_Table[Opponent Scope Key],uc_SchedRow)="—","",'
           'INDEX(Schedule_Table[Opponent Scope Key],uc_SchedRow)))')
    e.cell("uc_OppLabel", "Opponent team", '=IF(uc_OppScope="","",INDEX(Schedule_Table[Opponent Team],uc_SchedRow))')
    e.cell("uc_Category", "Fixture format code", '=IF(uc_SchedRow="","",INDEX(Schedule_Table[Format Category],uc_SchedRow))')
    e.cell("uc_PairKey", "Match Day evidence pair key",
           '=IF(OR(uc_TeamScope="",uc_OppScope="",uc_Category=""),"",uc_TeamScope&"|"&uc_OppScope&"|"&uc_Category)')
    e.cell("uc_Fixture", "Fixture line",
           f'=IF(uc_SchedRow="","",INDEX(Schedule_Table[Date Display],uc_SchedRow)&" · "&INDEX(Schedule_Table[Kickoff],'
           f'uc_SchedRow)&" ({tz_name}) · "&INDEX(Schedule_Table[Home/Away],uc_SchedRow)&" vs "&INDEX(Schedule_Table[Opponent],'
           f'uc_SchedRow)&" · "&INDEX(Schedule_Table[Format],uc_SchedRow))')
    e.cell("uc_Venue", "Venue and status",
           '=IF(uc_SchedRow="","","Venue: "&INDEX(Schedule_Table[Venue],uc_SchedRow)&" · Status: "&INDEX(Schedule_Table[Status],'
           'uc_SchedRow)&" · Score (home–away): "&INDEX(Schedule_Table[Score (home–away)],uc_SchedRow)&" · Session: "&'
           'INDEX(Schedule_Table[Session],uc_SchedRow))')
    # The exact fixture a Lineup Lab plan belongs to: unique per fixture (it ends with the match id), so
    # two matches on the same day -- even against the same opponent -- never share marks (GPT audit #84).
    e.cell("uc_PlanKey", "Fixture plan key",
           '=IF(uc_SchedRow="","",INDEX(Schedule_Table[Date Display],uc_SchedRow)&" · "&INDEX(Schedule_Table[Kickoff],'
           'uc_SchedRow)&" · "&INDEX(Schedule_Table[Home/Away],uc_SchedRow)&" vs "&INDEX(Schedule_Table[Opponent],uc_SchedRow)&'
           '" · match "&INDEX(Schedule_Table[Match ID],uc_SchedRow))')
    e.name("uc_PlanKeyList", f"Engine!$B${e.row - 1}")
    e.cell("uc_OppMsg", "Opponent roster note",
           '=IF(uc_SchedRow="","",IF(uc_OppScope="","Opponent roster: "&INDEX(Schedule_Table[Opponent Roster],uc_SchedRow)&".",'
           '"Opponent roster: "&uc_OppLabel))')

    # ---- War Room effective (optional local overrides) ----
    wr = "'War Room'!"
    e.section("War Room — effective teams")
    e.cell("wr_OurInput", "Our team override", "=" + guard(f"{wr}$C$5"))
    e.cell("wr_OppInput", "Opponent override", "=" + guard(f"{wr}$C$6"))
    e.cell("wr_OurValid", "Our override valid", '=AND(wr_OurInput<>"",COUNTIF(Teams_Table[Team],wr_OurInput)>0)')
    e.cell("wr_OppValid", "Opponent override valid", '=AND(wr_OppInput<>"",COUNTIF(Teams_Table[Team],wr_OppInput)>0)')
    e.cell("wr_Local", "Using local selections", "=OR(wr_OurValid,wr_OppValid)")
    e.cell("wr_OurLabel", "Effective our team", "=IF(wr_OurValid,wr_OurInput,uc_TeamLabel)")
    e.cell("wr_OppLabel", "Effective opponent team", "=IF(wr_OppValid,wr_OppInput,uc_OppLabel)")
    e.cell("wr_OurScope", "Our scope key",
           '=IF(wr_OurLabel="","",IFERROR(INDEX(Teams_Table[Scope Key],MATCH(wr_OurLabel,Teams_Table[Team],0)),""))')
    e.cell("wr_OppScope", "Opponent scope key",
           '=IF(wr_OppLabel="","",IFERROR(INDEX(Teams_Table[Scope Key],MATCH(wr_OppLabel,Teams_Table[Team],0)),""))')
    e.cell("wr_Category", "Format code",
           '=IF(wr_Local,IFERROR(INDEX(Teams_Table[Format Category],MATCH(wr_OurLabel,Teams_Table[Team],0)),""),uc_Category)')
    e.cell("wr_FormatLabel", "Format", "=" + _fmt_label_formula("wr_Category"))
    e.cell("wr_PairKey", "Evidence pair key",
           '=IF(OR(wr_OurScope="",wr_OppScope="",wr_Category="",wr_Category="No data"),"",wr_OurScope&"|"&wr_OppScope&"|"&wr_Category)')
    e.cell("wr_HasEvidence", "Precomputed evidence exists",
           '=IF(wr_PairKey="",FALSE,ISNUMBER(MATCH(wr_PairKey&"|1",MatchupEvidence_Table[Key],0)))')
    e.cell("wr_Status", "Status line",
           '=IF(wr_Local,"Using local selections: ","Following Match Day: ")&IF(wr_OurLabel="","(no team)",wr_OurLabel)&'
           '" vs "&IF(wr_OppLabel="","(no opponent)",wr_OppLabel)&IF(wr_Local,""," · "&IF(uc_DateDisplay="","no date",'
           'uc_DateDisplay))&IF(wr_FormatLabel="",""," · "&wr_FormatLabel)')
    e.cell("wr_Msg", "War Room note",
           '=IF(AND(wr_OurInput<>"",NOT(wr_OurValid)),"⚠ “"&wr_OurInput&"” is not a team label — ignored. ","")&'
           'IF(AND(wr_OppInput<>"",NOT(wr_OppValid)),"⚠ “"&wr_OppInput&"” is not a team label — ignored. ","")&'
           'IF(wr_OurLabel="","Choose your team on Match Day.",IF(wr_OppLabel="","No opponent roster for the selected fixture.",'
           'IF(wr_HasEvidence,"","Evidence comparisons are precomputed for teams scheduled to play each other this session; '
           'this pairing has no scheduled fixture (or the formats differ), so only rosters are shown.")))')

    # ---- Lineup Lab guard ----
    ll = "'Lineup Lab'!"
    e.section("Lineup Lab — which teams the marks belong to")
    e.cell("ll_OurTeam", "Planning for our team", "=" + guard(f"{ll}$C$5"))
    e.cell("ll_OppTeam", "Planning for opponent", "=" + guard(f"{ll}$C$6"))
    # Availability / lineup / played describe ONE match night (GPT audit #84: team-only keys leaked
    # "Played" into the next fixture). They apply only when Lineup Lab's match date is the War Room's
    # Match Day date; a War Room local exploration (no fixture) applies them only with the date blank.
    # Coach notes describe a player and follow the opponent team regardless of date.
    e.cell("ll_PlanIn", "Planning for fixture", "=" + guard(f"{ll}$C$9"))
    e.cell("ll_FixOK", "Marks belong to the shown fixture",
           '=IF(wr_Local,ll_PlanIn="",AND(ll_PlanIn<>"",uc_PlanKey<>"",ll_PlanIn=uc_PlanKey))')
    e.cell("ll_OurTeamOK", "Our team matches", '=AND(ll_OurTeam<>"",ll_OurTeam=wr_OurLabel)')
    e.cell("ll_OppTeamOK", "Opponent team matches (notes)", '=AND(ll_OppTeam<>"",ll_OppTeam=wr_OppLabel)')
    e.cell("ll_OurOK", "Our marks apply", "=AND(ll_OurTeamOK,ll_FixOK)")
    e.cell("ll_OppOK", "Opponent marks apply", "=AND(ll_OppTeamOK,ll_FixOK)")
    e.cell("ll_OurScope", "Planning team scope",
           '=IF(ll_OurTeam="","",IFERROR(INDEX(Teams_Table[Scope Key],MATCH(ll_OurTeam,Teams_Table[Team],0)),""))')
    e.cell("ll_OppScope", "Planning opponent scope",
           '=IF(ll_OppTeam="","",IFERROR(INDEX(Teams_Table[Scope Key],MATCH(ll_OppTeam,Teams_Table[Team],0)),""))')
    e.cell("ll_Cap", "Reference cap (user-entered)", f'=IF(ISNUMBER({ll}$C$7),{ll}$C$7,"")')

    # ---- rosters for the effective pairing ----
    ll_our_first, ll_opp_first = slots["ll_our_first"], slots["ll_opp_first"]

    def roster_cols(prefix: str, scope: str, ll_first: int, ok: str, our: bool):
        cols = [
            (f"{prefix}Rows", lambda k, r: f'=IF({scope}="","",IFERROR(MATCH({scope}&"|{k}",TeamRosters_Table[Roster Slot Key],0),""))'),
            (f"{prefix}Pids", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[Player ID],B{r}))'),
            (f"{prefix}Labels", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[Player Label],B{r}))'),
            (f"{prefix}SL", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[SL Display],B{r}))'),
            (f"{prefix}WL", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[W-L Display],B{r}))'),
            (f"{prefix}SLNum", lambda k, r: f'=IF(B{r}="","",IF(INDEX(TeamRosters_Table[SL Display],B{r})="No data","",'
                                            f'INDEX(TeamRosters_Table[Skill Level],B{r})))'),
        ]
        if our:
            cols += [
                (f"{prefix}Avail", lambda k, r: f'=IF(B{r}="","",IF({ok},IF({ll}$C${ll_first + k - 1}="","Unknown",'
                                                f'{ll}$C${ll_first + k - 1}),"Unknown"))'),
                (f"{prefix}Lineup", lambda k, r: f'=IF(B{r}="","",IF({ok},IF({ll}$D${ll_first + k - 1}="","—",'
                                                 f'{ll}$D${ll_first + k - 1}),"—"))'),
                (f"{prefix}Remaining", lambda k, r: f'=AND(B{r}<>"",H{r}<>"Unavailable",I{r}<>"Played")'),
            ]
        else:
            cols += [
                (f"{prefix}Played", lambda k, r: f'=IF(B{r}="","",IF({ok},IF({ll}$C${ll_first + k - 1}="Played","Played","—"),"—"))'),
                (f"{prefix}Unplayed", lambda k, r: f'=AND(B{r}<>"",H{r}<>"Played")'),
                (f"{prefix}Notes", lambda k, r: f'=IF(B{r}="","",IFERROR(IF(INDEX(cn_Summary,MATCH(INDEX(TeamRosters_Table[Player Label],B{r}),'
                                                f'cn_Player,0))="","","Coach: "&INDEX(cn_Summary,MATCH(INDEX(TeamRosters_Table[Player Label],B{r}),'
                                                f'cn_Player,0))&IF(AND(ll_OppTeamOK,{ll}$D${ll_first + k - 1}<>"")," · ","")),"")&'
                                                f'IF(ll_OppTeamOK,IF({ll}$D${ll_first + k - 1}="","",'
                                                f'{ll}$D${ll_first + k - 1}),""))'),
                (f"{prefix}Start", lambda k, r: f'=IF(OR(B{r}="",NOT(wr_HasEvidence)),"",IFERROR(MATCH(wr_PairKey&"|{k}",'
                                                f'MatchupEvidence_Table[Key],0),""))'),
                (f"{prefix}Count", lambda k, r: f'=IF(K{r}="",0,INDEX(MatchupEvidence_Table[Rows],K{r}))'),
                (f"{prefix}Card", lambda k, r: f'=IF(OR(B{r}="",NOT(wr_HasEvidence)),"",IFERROR(MATCH(wr_PairKey&"|{k}",'
                                               f'Scouting_Table[Key],0),""))'),
            ]
        return cols

    ours = e.block("Our roster slot", R, roster_cols("wr_Our", "wr_OurScope", ll_our_first, "ll_OurOK", True))
    theirs = e.block("Opponent roster slot", R, roster_cols("wr_Opp", "wr_OppScope", ll_opp_first, "ll_OppOK", False))
    our_pid_col, our_first = ours["wr_OurPids"][0], ours["wr_OurPids"][1]
    opp_start_col, opp_count_col, opp_first = theirs["wr_OppStart"][0], theirs["wr_OppCount"][0], theirs["wr_OppStart"][1]

    # Lineup Lab display rosters (the planning team named on Lineup Lab).
    e.block("Planning our slot", R, [
        ("ll_OurLabels", lambda k, r: f'=IF(ll_OurScope="","",IFERROR(INDEX(TeamRosters_Table[Player Label],'
                                      f'MATCH(ll_OurScope&"|{k}",TeamRosters_Table[Roster Slot Key],0)),""))'),
        ("ll_OurSL", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[SL Display],MATCH(ll_OurScope&"|{k}",TeamRosters_Table[Roster Slot Key],0)))'),
        ("ll_OurWL", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[W-L Display],MATCH(ll_OurScope&"|{k}",TeamRosters_Table[Roster Slot Key],0)))'),
    ])
    e.block("Planning opponent slot", R, [
        ("ll_OppLabels", lambda k, r: f'=IF(ll_OppScope="","",IFERROR(INDEX(TeamRosters_Table[Player Label],'
                                      f'MATCH(ll_OppScope&"|{k}",TeamRosters_Table[Roster Slot Key],0)),""))'),
        ("ll_OppSL", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[SL Display],MATCH(ll_OppScope&"|{k}",TeamRosters_Table[Roster Slot Key],0)))'),
        ("ll_OppWL", lambda k, r: f'=IF(B{r}="","",INDEX(TeamRosters_Table[W-L Display],MATCH(ll_OppScope&"|{k}",TeamRosters_Table[Roster Slot Key],0)))'),
    ])

    # ---- matrix: position of our player i inside opponent j's ranked block ----
    e.section("Matrix (row = our slot, column = opponent slot)")
    pos_first = e.row
    for j in range(1, R + 1):
        e.ws.cell(row=pos_first - 1, column=1 + j, value=f"pos opp {j}").font = base.HELPER_FONT
    grid = {}
    for i in range(1, R + 1):
        r = pos_first + i - 1
        e.ws.cell(row=r, column=1, value=f"pos our {i}").font = base.HELPER_FONT
        for j in range(1, R + 1):
            start = f"Engine!${opp_start_col}${opp_first + j - 1}"
            count = f"Engine!${opp_count_col}${opp_first + j - 1}"
            pid = f"Engine!${our_pid_col}${our_first + i - 1}"
            e.ws.cell(row=r, column=1 + j, value=(
                f'=IF(OR({pid}="",{start}=""),"",IFERROR(MATCH({pid},INDEX(MatchupEvidence_Table[Player ID],{start}+1):'
                f'INDEX(MatchupEvidence_Table[Player ID],{start}+{count}),0),""))'))
    e.row = pos_first + R + 1
    cat_first = e.row + 1
    e.ws.cell(row=cat_first - 1, column=1, value="Category grid").font = base.SUBHEAD_FONT
    cell_first = cat_first + R + 1
    e.ws.cell(row=cell_first - 1, column=1, value="Cell text grid").font = base.SUBHEAD_FONT
    cap_first = cell_first + R + 1
    e.ws.cell(row=cap_first - 1, column=1, value="Captain View grid").font = base.SUBHEAD_FONT
    captain = {}
    for i in range(1, R + 1):
        for j in range(1, R + 1):
            col = get_column_letter(1 + j)
            pos = f"${col}${pos_first + i - 1}"
            start = f"Engine!${opp_start_col}${opp_first + j - 1}"
            e.ws.cell(row=cat_first + i - 1, column=1 + j,
                      value=f'=IF({pos}="","",INDEX(MatchupEvidence_Table[Category],{start}+{pos}))')
            e.ws.cell(row=cell_first + i - 1, column=1 + j,
                      value=f'=IF({pos}="","",INDEX(MatchupEvidence_Table[Cell],{start}+{pos}))')
            grid[(i, j)] = (f"Engine!${col}${cat_first + i - 1}", f"Engine!${col}${cell_first + i - 1}",
                            f"Engine!${col}${pos_first + i - 1}")
            e.ws.cell(row=cap_first + i - 1, column=1 + j,
                      value=f'=IF({pos}="","",INDEX(MatchupEvidence_Table[Captain Cell],{start}+{pos}))')
            captain[(i, j)] = f"Engine!${col}${cap_first + i - 1}"
    last_col = get_column_letter(1 + R)
    for j in range(1, R + 1):
        col = get_column_letter(1 + j)
        e.name(f"wr_CatCol{j}", f"Engine!${col}${cat_first}:${col}${cat_first + R - 1}")
    e.row = cap_first + R + 1

    # ---- per-opponent counts: favorable / sendable remaining ----
    remaining = "wr_OurRemaining"
    e.block("Opponent evidence counts", R, [
        ("wr_GreenLeft", lambda k, r: f'=COUNTIFS(wr_CatCol{k},"G",{remaining},TRUE)'),
        ("wr_SendLeft", lambda k, r: f'=COUNTIFS(wr_CatCol{k},"G",{remaining},TRUE)+COUNTIFS(wr_CatCol{k},"E",{remaining},TRUE)'
                                     f'+COUNTIFS(wr_CatCol{k},"I",{remaining},TRUE)'),
        ("wr_RiskOpen", lambda k, r: f'=AND(INDEX(wr_OppUnplayed,{k}),B{r}=0)'),
        ("wr_RiskRun", lambda k, r: f'={"0" if k == 1 else f"E{r - 1}"}+IF(D{r},1,0)'),
    ])
    e.cell("wr_RiskCount", "Open risks", "=MAX(wr_RiskRun)")
    # Captain Packet evidence list: where each opponent's rows start in one continuous, packed list.
    e.block("Packet evidence start", R, [
        ("wr_PkStart", lambda k, r: "=1" if k == 1 else f"=B{r - 1}+INDEX(wr_OppCount,{k - 1})"),
    ])
    e.cell("wr_PkTotal", "Packet evidence lines", "=SUM(wr_OppCount)")

    # ---- best-supported sends per opponent (ranking order, sendable + remaining) ----
    e.section("Best-supported sends (per opponent: candidate rows in ranking order)")
    # Per opponent j, one row per ranked candidate: pid, category, remaining, sendable (ok), running count of
    # sendable rows, evidence group key (MatchupEvidence_Table[Rank Group]), direct (G/E and ok), group number
    # among remaining direct candidates, last group key, and the display label: "1." / "1=" (tied) / "≈"
    # (shared-opponent only, never numbered). Mirrors analytics send_labels / best_send_text (GPT audit #84).
    bs_first = e.row + 1
    send_refs, send_rng = {}, {}
    heads = ("pid", "cat", "rem", "ok", "run", "key", "dir", "grp", "last", "lbl")
    for j in range(1, R + 1):
        base_col = 2 + (j - 1) * len(heads)
        for off, head in enumerate(heads):
            e.ws.cell(row=bs_first - 1, column=base_col + off, value=f"o{j} {head}").font = base.HELPER_FONT
        start = f"Engine!${opp_start_col}${opp_first + j - 1}"
        count = f"Engine!${opp_count_col}${opp_first + j - 1}"
        L = {h: get_column_letter(base_col + i) for i, h in enumerate(heads)}
        rng = {h: f"Engine!${L[h]}${bs_first}:${L[h]}${bs_first + R - 1}" for h in heads}
        for rr in range(1, R + 1):
            r = bs_first + rr - 1
            c = {h: f"{L[h]}{r}" for h in heads}
            p = {h: f"{L[h]}{r - 1}" for h in heads}
            formulas = {
                "pid": f'=IF(OR({start}="",{rr}>{count}),"",INDEX(MatchupEvidence_Table[Player ID],{start}+{rr}))',
                "cat": f'=IF({c["pid"]}="","",INDEX(MatchupEvidence_Table[Category],{start}+{rr}))',
                "rem": f'=IF({c["pid"]}="",FALSE,IFERROR(INDEX({remaining},MATCH({c["pid"]},wr_OurPids,0)),FALSE))',
                "ok": f'=AND({c["pid"]}<>"",OR({c["cat"]}="G",{c["cat"]}="E",{c["cat"]}="I"),{c["rem"]})',
                "run": f'=IF({c["ok"]},1,0)' if rr == 1 else f'={p["run"]}+IF({c["ok"]},1,0)',
                "key": f'=IF({c["pid"]}="","",INDEX(MatchupEvidence_Table[Rank Group],{start}+{rr}))',
                "dir": f'=AND({c["ok"]},OR({c["cat"]}="G",{c["cat"]}="E"))',
                "grp": (f'=IF({c["dir"]},1,0)' if rr == 1 else
                        f'=IF({c["dir"]},IF({p["grp"]}=0,1,IF({c["key"]}={p["last"]},{p["grp"]},{p["grp"]}+1)),{p["grp"]})'),
                "last": f'=IF({c["dir"]},{c["key"]},"")' if rr == 1 else f'=IF({c["dir"]},{c["key"]},{p["last"]})',
                "lbl": (f'=IF(NOT({c["ok"]}),"",IF({c["cat"]}="I","≈",{c["grp"]}&'
                        f'IF(COUNTIFS({rng["grp"]},{c["grp"]},{rng["dir"]},TRUE)>1,"=",".")))'),
            }
            for off, head in enumerate(heads):
                e.ws.cell(row=r, column=base_col + off, value=formulas[head])
        send_refs[j] = rng["run"]
        send_rng[j] = rng
    e.row = bs_first + R + 1

    def send_text(n: int, k: int) -> str:
        rng, m = send_rng[k], f"MATCH({n},{send_rng[k]['run']},0)"
        row = f"INDEX(wr_OppStart,{k})+{m}"
        return (f'=IF(OR(INDEX(wr_OppStart,{k})="",NOT(INDEX(wr_OppUnplayed,{k}))),"",IFERROR(INDEX({rng["lbl"]},{m})&" "&'
                f'INDEX(MatchupEvidence_Table[Player],{row})&" — "&INDEX(MatchupEvidence_Table[Cell],{row}),""))')

    def best_text(k: int) -> str:
        rng, m = send_rng[k], f"MATCH(1,{send_rng[k]['run']},0)"
        row = f"INDEX(wr_OppStart,{k})+{m}"
        n_i = f'COUNTIFS({rng["cat"]},"I",{rng["ok"]},TRUE)'
        tie = f'COUNTIFS({rng["grp"]},1,{rng["dir"]},TRUE)'
        tag = (f'IF(INDEX({rng["cat"]},{m})="I","shared-opponent candidate, not ordered ("&IF({n_i}=1,"the only one",'
               f'"one of "&{n_i})&"): ",IF({tie}>1,"best-supported send (tied with "&({tie}-1)&" other"&IF({tie}>2,"s","")&'
               f'", same evidence): ","best-supported send: "))')
        return (f'=IF(OR(INDEX(wr_OppStart,{k})="",NOT(INDEX(wr_OppUnplayed,{k}))),"",IFERROR({tag}&'
                f'INDEX(MatchupEvidence_Table[Player],{row})&" — reason: "&INDEX(MatchupEvidence_Table[Reason],{row}),""))')

    e.block("Best send", R, [
        (f"wr_Send{n}", (lambda n: lambda k, r: send_text(n, k))(n)) for n in (1, 2, 3)
    ] + [
        # The labelled list (War Room, packet) and the single best send with its reason (Lineup Lab, Command Center).
        ("wr_SendList", lambda k, r: f'=B{r}&IF(C{r}="",""," · "&C{r})&IF(D{r}="",""," · "&D{r})'),
        ("wr_BestText", lambda k, r: best_text(k)),
    ])

    # ---- threats, concerning pairings, unique favorable options ----
    e.block("Threat row", R, [
        ("wr_ThreatRow", lambda k, r: f'=IF(NOT(wr_HasEvidence),"",IFERROR(MATCH(wr_PairKey&"|{k}",Threats_Table[Key],0),""))'),
        ("wr_ThreatOk", lambda k, r: f'=IF(B{r}="",FALSE,IFERROR(INDEX(wr_OppUnplayed,MATCH(INDEX(Threats_Table[Opp Player ID],B{r}),wr_OppPids,0)),FALSE))'),
        ("wr_ThreatRun", lambda k, r: (f'=IF(C{r},1,0)' if k == 1 else f'=D{r - 1}+IF(C{r},1,0)')),
    ])
    e.block("Concerning row", CONCERNING_SLOTS, [
        ("wr_ConcernRow", lambda k, r: f'=IF(NOT(wr_HasEvidence),"",IFERROR(MATCH(wr_PairKey&"|{k}",Concerning_Table[Key],0),""))'),
        ("wr_ConcernOk", lambda k, r: (
            f'=IF(B{r}="",FALSE,AND(IFERROR(INDEX(wr_OurRemaining,MATCH(INDEX(Concerning_Table[Our Player ID],B{r}),wr_OurPids,0)),FALSE),'
            f'IFERROR(INDEX(wr_OppUnplayed,MATCH(INDEX(Concerning_Table[Opp Player ID],B{r}),wr_OppPids,0)),FALSE)))')),
        ("wr_ConcernRun", lambda k, r: (f'=IF(C{r},1,0)' if k == 1 else f'=D{r - 1}+IF(C{r},1,0)')),
    ])
    e.section("Unique favorable options (row = our slot, column = opponent slot)")
    u_first = e.row + 1
    for i in range(1, R + 1):
        r = u_first + i - 1
        for j in range(1, R + 1):
            cat = grid[(i, j)][0]
            e.ws.cell(row=r, column=1 + j, value=(
                f'=AND({cat}="G",INDEX(wr_OurRemaining,{i}),INDEX(wr_OppUnplayed,{j}),INDEX(wr_GreenLeft,{j})=1)'))
        e.ws.cell(row=r, column=2 + R, value=f'=COUNTIF(B{r}:{last_col}{r},TRUE)')
        e.ws.cell(row=r, column=3 + R, value=f'=IFERROR(MATCH(TRUE,B{r}:{last_col}{r},0),"")')
        e.ws.cell(row=r, column=4 + R, value=(f'=IF({get_column_letter(2 + R)}{r}>0,1,0)' if i == 1 else
                                              f'={get_column_letter(4 + R)}{r - 1}+IF({get_column_letter(2 + R)}{r}>0,1,0)'))
    for off, name in enumerate(("wr_UniqueCount", "wr_UniqueFirst", "wr_UniqueRun")):
        col = get_column_letter(2 + R + off)
        e.name(name, f"Engine!${col}${u_first}:${col}${u_first + R - 1}")
    e.row = u_first + R + 1

    # ---- Next Send (Command Center): "they put up this player -- who do I send?" ----
    # Mirrors analytics.next_send: medals only for ordered direct candidates (favorable, then even) among our
    # remaining players, same rank group = same medal; shared-only candidates unordered; Avoid worst first.
    cc, ME = "'Command Center'!", "MatchupEvidence_Table"
    e.section("Next Send")
    e.cell("wr_NsInput", "They put up", "=" + guard(f"{cc}$C$6"))
    e.cell("wr_NsJ", "Opponent slot", '=IF(wr_NsInput="","",IFERROR(MATCH(wr_NsInput,wr_OppLabels,0),""))')
    e.cell("wr_NsOpen", "Opponent unplayed", '=IF(wr_NsJ="",FALSE,INDEX(wr_OppUnplayed,wr_NsJ))')
    e.cell("wr_NsStart", "Ranked block start", '=IF(wr_NsJ="","",INDEX(wr_OppStart,wr_NsJ))')
    e.cell("wr_NsCount", "Ranked rows", '=IF(OR(wr_NsJ="",wr_NsStart=""),0,INDEX(wr_OppCount,wr_NsJ))')
    e.cell("wr_NsSelfUnique", "One favorable option left vs them", '=IF(wr_NsJ="",FALSE,INDEX(wr_GreenLeft,wr_NsJ)=1)')
    medal = 'IF(H{r}=1,"🥇",IF(H{r}=2,"🥈","🥉"))'

    def _join(k: int, r: int, col: str, cond: str, text: str) -> str:
        """Running ", "-joined list down a column; the last row holds the whole list (no TEXTJOIN needed)."""
        if k == 1:
            return f'=IF({cond},{text},"")'
        prev = f"{col}{r - 1}"
        return f'=IF({cond},{prev}&IF({prev}="","","; ")&{text},{prev})'

    e.block("Next Send ranked row", R, [
        ("wr_NsRow", lambda k, r: f'=IF(OR(NOT(wr_NsOpen),{k}>wr_NsCount),"",wr_NsStart+{k})'),
        ("wr_NsSlot", lambda k, r: f'=IF(B{r}="","",IFERROR(MATCH(INDEX({ME}[Player ID],B{r}),wr_OurPids,0),""))'),
        ("wr_NsCat", lambda k, r: f'=IF(B{r}="","",INDEX({ME}[Category],B{r}))'),
        ("wr_NsRem", lambda k, r: f'=IF(C{r}="",FALSE,INDEX(wr_OurRemaining,C{r}))'),
        ("wr_NsKey", lambda k, r: f'=IF(B{r}="","",INDEX({ME}[Rank Group],B{r}))'),
        ("wr_NsElig", lambda k, r: f'=AND(E{r},OR(D{r}="G",D{r}="E"))'),
        ("wr_NsGrp", lambda k, r: (f'=IF(G{r},1,0)' if k == 1 else
                                    f'=IF(G{r},IF(H{r - 1}=0,1,IF(F{r}=I{r - 1},H{r - 1},H{r - 1}+1)),H{r - 1})')),
        ("wr_NsLastKey", lambda k, r: (f'=IF(G{r},F{r},"")' if k == 1 else f'=IF(G{r},F{r},I{r - 1})')),
        ("wr_NsMedalRun", lambda k, r: (f'=IF(AND(G{r},H{r}<=3),1,0)' if k == 1 else f'=J{r - 1}+IF(AND(G{r},H{r}<=3),1,0)')),
        ("wr_NsMedalText", lambda k, r: (
            f'=IF(G{r},IF(H{r}<=3,{medal.format(r=r)}&" "&INDEX({ME}[Player],B{r})&" — "&INDEX({ME}[Reason],B{r})'
            f'&IF(COUNTIFS(wr_NsGrp,H{r},wr_NsElig,TRUE)>1," · tied (same evidence)","")'
            f'&IF(INDEX(wr_UniqueCount,C{r})-IF(AND(D{r}="G",wr_NsSelfUnique),1,0)>0,'
            f'" · consider saving — our only favorable direct option vs another unplayed opponent","")'
            f'&IF(INDEX(wr_OurAvail,C{r})="Unknown"," · availability unknown",""),""),"")')),
        ("wr_NsMore", lambda k, r: f'=AND(G{r},H{r}>3)'),
        ("wr_NsUnordRun", lambda k, r: (f'=IF(AND(E{r},D{r}="I"),1,0)' if k == 1 else f'=M{r - 1}+IF(AND(E{r},D{r}="I"),1,0)')),
        ("wr_NsUnordJoin", lambda k, r: _join(k, r, "N", f'AND(E{r},D{r}="I")', f'INDEX({ME}[Player],B{r})'
                                              f'&IF(INDEX(wr_OurAvail,C{r})="Unknown"," (availability unknown)","")')),
        ("wr_NsUnkJoin", lambda k, r: _join(k, r, "O", f'AND(E{r},D{r}="X")', f'INDEX({ME}[Player],B{r})')),
        ("wr_NsRevRow", lambda k, r: f'=IF(OR(NOT(wr_NsOpen),{k}>wr_NsCount),"",wr_NsStart+wr_NsCount+1-{k})'),
        ("wr_NsAvoidOk", lambda k, r: (f'=IF(P{r}="",FALSE,AND(INDEX({ME}[Category],P{r})="R",IFERROR(INDEX(wr_OurRemaining,'
                                        f'MATCH(INDEX({ME}[Player ID],P{r}),wr_OurPids,0)),FALSE)))')),
        ("wr_NsAvoidJoin", lambda k, r: _join(k, r, "R", f'Q{r}',
                                              f'INDEX({ME}[Player],P{r})&" — "&INDEX({ME}[Direct Record],P{r})')),
    ])
    e.cell("wr_NsHeadline", "Next Send headline", (
        '=IF(wr_NsInput="","Pick the opponent player they put up (cell C6).",IF(wr_NsJ="","That player is not on tonight’s opponent roster.",'
        'IF(NOT(wr_NsOpen),wr_NsInput&" has already played (Lineup Lab).",IF(MAX(wr_NsMedalRun)>0,"Medals = ordered direct records among our '
        'remaining players (same evidence = same medal). Recorded results only — not odds.",IF(MAX(wr_NsUnordRun)>0,"No direct record to order — '
        'shared-opponent candidates only (≈, not ordered)","No evidence-backed option left among our remaining players")))))'))
    # Each remaining player is in at most one of these lists, so together they name at most R players. One wrapped
    # cell (a line per list) can therefore be sized once for R names; three rows would each need room for all R
    # (real-Excel UAT 2a67d78 showed a 2-line row hiding the last "≈" candidate).
    e.cell("wr_NsUnordLine", "Next Send not-ordered line",
           f'=IF(INDEX(wr_NsUnordJoin,{R})="","","≈ Not ordered (shared-opponent results only): "&INDEX(wr_NsUnordJoin,{R}))')
    e.cell("wr_NsAvoidLine", "Next Send avoid line", f'=IF(INDEX(wr_NsAvoidJoin,{R})="","","⚠ Avoid: "&INDEX(wr_NsAvoidJoin,{R}))')
    e.cell("wr_NsUnkLine", "Next Send unknown line",
           f'=IF(INDEX(wr_NsUnkJoin,{R})="","","❓ Unknown (no evidence, not weak): "&INDEX(wr_NsUnkJoin,{R}))')
    e.cell("wr_NsLists", "Next Send lists (one line each)",
           '=wr_NsUnordLine&IF(AND(wr_NsUnordLine<>"",wr_NsAvoidLine&wr_NsUnkLine<>""),CHAR(10),"")'
           '&wr_NsAvoidLine&IF(AND(wr_NsAvoidLine<>"",wr_NsUnkLine<>""),CHAR(10),"")&wr_NsUnkLine')
    # The whole Next Send answer as ONE wrapped cell, a line per item: headline, up to four medal lines, the
    # "+ more" note, then the lists. Fixed rows per medal left a tall empty gap between the medal and the lists
    # on screen and on paper (real-Excel UAT 0c2e474); one cell keeps the answer contiguous and is sized once.
    for n in range(1, 5):
        e.cell(f"wr_NsMedal{n}", f"Next Send medal line {n}", f'=IFERROR(INDEX(wr_NsMedalText,MATCH({n},wr_NsMedalRun,0)),"")')
    e.cell("wr_NsMoreLine", "Next Send more-candidates line",
           '=IF(COUNTIF(wr_NsMedalRun,">4")+COUNTIF(wr_NsMore,TRUE)=0,"","+ more direct candidates — see War Room Inspect")')
    e.cell("wr_NsCard", "Next Send answer (one line each)", "=wr_NsHeadline" + "".join(
        f'&IF({x}="","",CHAR(10)&{x})' for x in ("wr_NsMedal1", "wr_NsMedal2", "wr_NsMedal3", "wr_NsMedal4",
                                                 "wr_NsMoreLine", "wr_NsLists")))

    # ---- Lineup Lab metrics ----
    e.section("Lineup Lab metrics")
    e.cell("wr_OurCount", "Our rostered", '=COUNTIF(wr_OurRows,">0")')
    e.cell("wr_RemainingCount", "Remaining players", "=COUNTIF(wr_OurRemaining,TRUE)")
    e.cell("wr_RemainingUnknown", "Remaining with unknown availability", '=COUNTIFS(wr_OurRemaining,TRUE,wr_OurAvail,"Unknown")')
    e.cell("wr_RemainingSL", "Remaining known SL total", "=SUMIFS(wr_OurSLNum,wr_OurRemaining,TRUE)")
    e.cell("wr_RemainingMissing", "Remaining without a captured SL", '=COUNTIFS(wr_OurRemaining,TRUE,wr_OurSLNum,"")')
    e.cell("wr_SelectedCount", "Selected lineup (Planned + Played)", '=COUNTIF(wr_OurLineup,"Planned")+COUNTIF(wr_OurLineup,"Played")')
    e.cell("wr_SelectedSL", "Selected known SL total",
           '=SUMIFS(wr_OurSLNum,wr_OurLineup,"Planned")+SUMIFS(wr_OurSLNum,wr_OurLineup,"Played")')
    e.cell("wr_SelectedMissing", "Selected without a captured SL",
           '=COUNTIFS(wr_OurLineup,"Planned",wr_OurSLNum,"")+COUNTIFS(wr_OurLineup,"Played",wr_OurSLNum,"")')
    e.cell("wr_UnplayedCount", "Opponents not yet played", "=COUNTIF(wr_OppUnplayed,TRUE)")
    e.cell("wr_GreenOpps", "Unplayed opponents with a favorable option left", '=COUNTIFS(wr_OppUnplayed,TRUE,wr_GreenLeft,">0")')
    e.cell("wr_SendOpps", "Unplayed opponents with any evidence-backed option left", '=COUNTIFS(wr_OppUnplayed,TRUE,wr_SendLeft,">0")')

    # ---- Coach Dashboard effective ----
    cd = "'Coach Dashboard'!"
    e.section("Coach Dashboard — effective players")
    e.block("Player B choice", R + 1, [
        ("cd_PlayerBList", lambda k, r: (f'=INDEX(wr_OppLabels,{k})' if k <= R else f'="{ALL_PLAYERS_CHOICE}"')),
    ])
    e.cell("cd_AInput", "Player A override", "=" + guard(f"{cd}$C$5"))
    e.cell("cd_BInput", "Player B choice", "=" + guard(f"{cd}$C$6"))
    e.cell("cd_AnyInput", "Any player", "=" + guard(f"{cd}$C$7"))
    e.cell("cd_FmtInput", "Format override", "=" + guard(f"{cd}$C$8"))
    e.cell("cd_Local", "Using local selections", '=OR(cd_AInput<>"",cd_FmtInput<>"")')
    e.cell("cd_ALabel", "Player A", '=IF(cd_AInput<>"",cd_AInput,uc_ViewerLabel)')
    e.cell("cd_BLabel", "Player B", f'=IF(cd_BInput="","",IF(cd_BInput="{ALL_PLAYERS_CHOICE}",cd_AnyInput,cd_BInput))')
    e.cell("cd_AID", "Player A id", '=IFERROR(INDEX(Players_Table[Player ID],MATCH(cd_ALabel,Players_Table[Player Label],0)),"")')
    e.cell("cd_BID", "Player B id", '=IFERROR(INDEX(Players_Table[Player ID],MATCH(cd_BLabel,Players_Table[Player Label],0)),"")')
    e.cell("cd_FmtLabel", "Format", '=IF(cd_FmtInput<>"",cd_FmtInput,IF(wr_FormatLabel="","8-Ball",wr_FormatLabel))')
    e.cell("cd_Status", "Status line",
           '=IF(cd_Local,"Using local selections","Following Match Day")&": "&IF(cd_ALabel="","(Player A)",cd_ALabel)&" vs "&'
           f'IF(cd_BLabel="","(choose Player B)",cd_BLabel)&" · "&cd_FmtLabel')

    return {"grid": grid, "captain": captain, "ours": ours, "theirs": theirs}


# ---------------------------------------------------------------------------
# Sheet helpers
# ---------------------------------------------------------------------------

def _dv(sheet, cell: str, source: str, prompt: str) -> None:
    v = DataValidation(type="list", formula1=source, allow_blank=True)
    v.promptTitle, v.prompt = "Choose", prompt
    sheet.add_data_validation(v)
    v.add(cell)


def _calc(cell, value: Any, *, bold: bool = False) -> None:
    cell.value = value
    cell.fill = CALC_FILL
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def _link(sheet, row: int, col: int, text: str, target: str, *, font: Font = base.LINK_FONT) -> None:
    c = sheet.cell(row=row, column=col, value=text)
    c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{target}'!A1", display=text)
    c.font = font


def _head(sheet, title: str, subtitle: str, *, last: int, current: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = title
    sheet["A1"].font = base.TITLE_FONT
    sheet.row_dimensions[1].height = 30
    sheet["A2"] = subtitle
    sheet["A2"].font = base.SUBTITLE_FONT
    sheet["A2"].alignment = base.WRAP_TOP
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
    nav = ["Match Day", "War Room", "Lineup Lab", "Scouting Cards", "Captain Packet", "Coach Dashboard"]
    sheet.cell(row=3, column=1, value="Go to:").font = base.MUTED_FONT
    col = 2
    for name in nav:
        if name == current:
            continue
        _link(sheet, 3, col, name, name)
        col += 1


def _span(sheet, row: int, c1: int, c2: int, value: Any, *, font: Font | None = None, fill=None, wrap: bool = True,
          height: float | None = None) -> None:
    cell = sheet.cell(row=row, column=c1, value=value)
    if font is not None:
        cell.font = font
    if fill is not None:
        for c in range(c1, c2 + 1):
            sheet.cell(row=row, column=c).fill = fill
    if wrap:
        cell.alignment = base.WRAP_TOP
    if c2 > c1:
        sheet.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    if height:
        sheet.row_dimensions[row].height = height


def _blank_if(ref: str, expr: str) -> str:
    return f'=IF({ref}="","",{expr})'


# ---------------------------------------------------------------------------
# Match Day (control panel)
# ---------------------------------------------------------------------------

def build_match_day(wb, *, slots: dict[str, int], viewer_label: str | None, viewer_key: str | None,
                    card_number: str | None, default_team: str | None, stats: dict[str, Any],
                    format_options_source: str) -> None:
    ws = wb.create_sheet("Match Day", 0)
    for letter, width in zip("ABCDEFGHIJKL", (24, 46, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14)):
        ws.column_dimensions[letter].width = width
    _head(ws, "Ultimate Coach — Match Day",
          "Set up tonight's matchup once — every tab follows it. Yellow cells are yours; green-gray cells are calculated.",
          last=12, current="Match Day")
    fresh = stats["freshness"]
    _span(ws, 4, 1, 12,
          f"Built {fresh['build_date']} · offline snapshot: latest recorded result {fresh['latest_result']}"
          + (f" ({fresh['unplayed_before_build']} earlier fixtures still show UNPLAYED — results after the snapshot are not "
             "included)" if fresh["unplayed_before_build"] else "")
          + f" · times in {stats['tz']} · this file never refreshes itself.", font=base.MUTED_FONT, height=30)
    base._section(ws, 5, "1 · Set up the matchup (yellow = your input)", last_col=12)
    rows = [
        (6, "Player", viewer_label or "", "Pick your name — every entry shows the APA record ID. Typing just the record ID also works.",
         "PlayerLabelList"),
        (7, "Team", default_team or "", "Only this player's current teams (team · session · format).", "uc_TeamList"),
        (8, "Format", "8-Ball & 9-Ball", "Default: every 8-Ball and 9-Ball variant. Or pick one recorded format.",
         format_options_source),
        (9, "Scheduled date", "", "Blank = suggested date (the earliest scheduled date on or after the build date). "
                                  "Or pick one of this team's real dates.", "uc_DateList"),
        (10, "Fixture", "", "Only needed when the date has more than one fixture — nothing is chosen for you.",
         "uc_FixtureList"),
    ]
    for r, label, value, help_text, source in rows:
        ws.cell(row=r, column=1, value=label).font = base.LABEL_FONT
        base._input(ws.cell(row=r, column=2), value)
        ws.cell(row=r, column=2).number_format = "@" if r in (6, 7) else "General"
        _dv(ws, f"B{r}", source, help_text)
        _span(ws, r, 3, 12, help_text, font=base.MUTED_FONT, height=30)
    _span(ws, 11, 1, 12, "Excel can't erase what you typed: when an input stops fitting (for example after you change the "
                         "player), it is ignored and the reason is shown in step 2 — no stale results are kept.",
          font=base.MUTED_FONT, height=30)
    base._section(ws, 12, "2 · Effective matchup (calculated — every tab uses this)", last_col=12)
    eff = [
        (13, "Player", '=IF(uc_ViewerLabel<>"",uc_ViewerLabel,IF(uc_ViewerKey="","Choose a player in step 1.",'
                       '"⚠ No verified player matches “"&\'Match Day\'!$B$6&"”."))',
         (f'=IF(uc_ViewerKey="{viewer_key}","League card #{card_number} was verified to this record when the build was '
          f'configured (display only; card numbers are per league and never used as identity).","")'
          if viewer_key and card_number else '=""')),
        (14, "Team", '=IF(uc_TeamLabel="","—",uc_TeamLabel)', "=uc_TeamMsg"),
        (15, "Format", "=uc_FormatLabel", '=""'),
        (16, "Date", '=IF(uc_DateDisplay="","—",uc_DateDisplay)', "=uc_DateMsg"),
        (17, "Fixture", '=IF(uc_Fixture="","—",uc_Fixture)', "=uc_FixMsg"),
        (18, "Venue · status", '=IF(uc_Venue="","—",uc_Venue)', '=""'),
        (19, "Opponent roster", '=IF(uc_OppMsg="","—",uc_OppMsg)', '=""'),
    ]
    for r, label, value, note in eff:
        ws.cell(row=r, column=1, value=label).font = base.LABEL_FONT
        _calc(ws.cell(row=r, column=2), value, bold=True)
        _span(ws, r, 3, 12, note, font=base.MUTED_FONT)
        ws.row_dimensions[r].height = 30
    links = [("→ Captain's War Room", "War Room"), ("→ Lineup Lab", "Lineup Lab"), ("→ Scouting Cards", "Scouting Cards"),
             ("→ Captain Packet (print)", "Captain Packet"), ("→ Coach Dashboard", "Coach Dashboard")]
    for k, (text, target) in enumerate(links):
        _link(ws, 21, 1 + k * 2, text, target, font=BIG_LINK)
    ws.row_dimensions[21].height = 22

    base._section(ws, 23, "Fixtures on the effective date", last_col=12)
    heads = ["Fixture", "Kickoff", "Home/Away", "Format", "Opponent", "Venue", "Status", "Score (home–away)", "Opponent Roster"]
    base._subheader(ws, 24, [(1, "Fixture"), (3, "Kickoff"), (4, "Home/Away"), (5, "Format"), (6, "Opponent"),
                             (8, "Venue"), (9, "Status"), (10, "Score (home–away)"), (11, "Opponent Roster")])
    for k in range(1, slots["fixtures"] + 1):
        r = 24 + k
        row_ref = f"INDEX(uc_FixRows,{k})"
        _span(ws, r, 1, 2, f'=IF({row_ref}="","",INDEX(uc_FixtureList,{k}))', wrap=True)
        for col, name, c2 in ((3, "Kickoff", 3), (4, "Home/Away", 4), (5, "Format", 5), (6, "Opponent", 7), (8, "Venue", 8),
                              (9, "Status", 9), (10, "Score (home–away)", 10), (11, "Opponent Roster", 12)):
            _span(ws, r, col, c2, f'=IF({row_ref}="","",INDEX(Schedule_Table[{name}],{row_ref}))', font=SMALL)
        for c in range(1, 13):
            ws.cell(row=r, column=c).border = base.CELL_BORDER
    top = 26 + slots["fixtures"]
    base._section(ws, top, "Your teams (dropdown source for step 1)", last_col=12)
    base._subheader(ws, top + 1, [(1, "Slot"), (2, "Team"), (3, "Next scheduled date (this format)")])
    for k in range(1, slots["teams"] + 1):
        r = top + 1 + k
        ws.cell(row=r, column=1, value=f'=IF(INDEX(uc_TeamList,{k})="","","Team {k}")')
        ws.cell(row=r, column=2, value=f'=INDEX(uc_TeamList,{k})')
        _span(ws, r, 3, 6, f'=IF(INDEX(uc_TeamList,{k})="","",INDEX(uc_TeamNextDisplay,{k}))', wrap=False)
    top = top + slots["teams"] + 3
    base._section(ws, top, "Scheduled dates for the effective team and format", last_col=12)
    for k in range(1, slots["dates"] + 1):
        r = top + 1 + (k - 1) // 4
        c = 1 + ((k - 1) % 4) * 3
        ws.cell(row=r, column=c, value=f'=INDEX(uc_DateList,{k})')
    top = top + 2 + (slots["dates"] - 1) // 4
    base._section(ws, top, "About this data", last_col=12)
    notes = [
        "Dates, weekdays and kickoff times are converted to " + stats["tz"] + ". A fixture stored as 2026-08-30T01:00:00Z "
        "is Saturday Aug 29 at 7:00 PM in Denver.",
        "APA record ID = APA's member record number (the identity used everywhere). It is not the league card number "
        "printed on a member card.",
        "Rosters are each team's CURRENT captured roster, not a reconstruction of who played on a past date.",
    ]
    for k, text in enumerate(notes, start=1):
        _span(ws, top + k, 1, 12, text, font=base.MUTED_FONT, height=30)
    ws.freeze_panes = "A4"


# ---------------------------------------------------------------------------
# War Room
# ---------------------------------------------------------------------------

def _roster_section(ws, top: int, R: int, *, title_ours: str, title_theirs: str) -> int:
    base._subheader(ws, top, [(1, "OUR TEAM"), (7, "OPPONENT")])
    ws.cell(row=top, column=1).fill = base.SECTION_FILL
    ws.cell(row=top, column=1).font = base.SECTION_FONT
    ws.cell(row=top, column=7).fill = base.OPPONENT_FILL
    ws.cell(row=top, column=7).font = base.SECTION_FONT
    _span(ws, top, 2, 5, title_ours, font=base.SUBHEAD_FONT)
    _span(ws, top, 8, 12, title_theirs, font=base.SUBHEAD_FONT)
    base._subheader(ws, top + 1, [(1, "Player (APA record ID)"), (3, "SL"), (4, "Team W-L"), (5, "Availability · lineup"),
                                  (7, "Player (APA record ID)"), (9, "SL"), (10, "Team W-L"), (11, "vs our roster"),
                                  (12, "Played")])
    for k in range(1, R + 1):
        r = top + 1 + k
        _span(ws, r, 1, 2, f'=INDEX(wr_OurLabels,{k})', font=SMALL)
        ws.cell(row=r, column=3, value=f'=INDEX(wr_OurSL,{k})')
        ws.cell(row=r, column=4, value=f'=INDEX(wr_OurWL,{k})')
        ws.cell(row=r, column=5, value=f'=IF(INDEX(wr_OurLabels,{k})="","",INDEX(wr_OurAvail,{k})&" · "&INDEX(wr_OurLineup,{k}))')
        _span(ws, r, 7, 8, f'=INDEX(wr_OppLabels,{k})', font=SMALL)
        ws.cell(row=r, column=9, value=f'=INDEX(wr_OppSL,{k})')
        ws.cell(row=r, column=10, value=f'=INDEX(wr_OppWL,{k})')
        ws.cell(row=r, column=11, value=f'=IF(INDEX(wr_OppCard,{k})="","",INDEX(Scouting_Table[Vs Our Roster],INDEX(wr_OppCard,{k})))')
        ws.cell(row=r, column=11).font = SMALL
        ws.cell(row=r, column=11).alignment = base.WRAP_TOP
        ws.cell(row=r, column=12, value=f'=INDEX(wr_OppPlayed,{k})')
        for c in range(1, 13):
            ws.cell(row=r, column=c).border = base.CELL_BORDER
        ws.row_dimensions[r].height = 28
    return top + 2 + R


def build_war_room(wb, *, slots: dict[str, int], engine: dict[str, Any]) -> dict[str, int]:
    R = slots["roster"]
    ws = wb.create_sheet("War Room", 1)
    for letter, width in zip("ABCDEFGHIJKLM", (30, 15, 9, 11, 22, 3, 15, 15, 9, 11, 30, 10, 3)):
        ws.column_dimensions[letter].width = width
    _head(ws, "Captain's War Room — Match Night",
          "Who should I put up next? Everything for tonight's fixture, from recorded evidence only — no odds, no guarantees.",
          last=12, current="War Room")
    _span(ws, 4, 1, 12, "=wr_Status", font=Font(bold=True, size=12, color=base.FELT_DEEP))
    for r, label, help_text in ((5, "Our team (optional override)", "Blank = follow Match Day. Enter another team to explore; clear it to return."),
                                (6, "Opponent (optional override)", "Blank = follow Match Day's fixture opponent.")):
        _span(ws, r, 1, 2, label, font=base.LABEL_FONT)
        base._input(ws.cell(row=r, column=3))
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        _dv(ws, f"C{r}", "TeamNameList", help_text)
        _span(ws, r, 7, 12, help_text, font=base.MUTED_FONT)
    _span(ws, 7, 1, 12, "=wr_Msg", font=WARN_FONT)
    base._section(ws, 8, "Match summary", last_col=12)
    _span(ws, 9, 1, 12, '=IF(wr_Local,"Local exploration — not tied to Match Day’s fixture.",IF(uc_Fixture="","No fixture selected on Match Day.",uc_Fixture))',
          font=Font(bold=True, size=12), height=22)
    _span(ws, 10, 1, 12, '=IF(wr_Local,"",uc_Venue)', font=base.MUTED_FONT)

    base._section(ws, 12, "Rosters (SL = this team's captured skill level for its format)", last_col=12)
    end = _roster_section(ws, 13, R, title_ours="=wr_OurLabel", title_theirs="=wr_OppLabel")
    _span(ws, end, 1, 12, '=IF(AND(ll_OurTeamOK,ll_OppTeamOK,NOT(ll_FixOK)),"⚠ Lineup Lab marks are for "&IF(ll_PlanIn="",'
                          '"no fixture",ll_PlanIn)&" — not applied to this fixture. ","")&'
                          'IF(AND(ll_OurTeam<>"",NOT(ll_OurTeamOK)),"⚠ Lineup Lab marks are for "&ll_OurTeam&" — not applied to "&'
                          'wr_OurLabel&". ","")&IF(AND(ll_OppTeam<>"",NOT(ll_OppOK)),"⚠ Opponent marks are for "&ll_OppTeam&'
                          '" — not applied. ","")&"Team W-L = this team’s captured record this session. “No data” = not captured. '
                          'Availability, lineup and played are your Lineup Lab marks, never evidence; Unknown is not Unavailable."',
          font=base.MUTED_FONT, height=30)

    top = end + 2
    base._section(ws, top, "Top opportunities — best-supported sends (favorable direct first, then even, then indirect)",
                  last_col=12)
    for k in range(1, R + 1):
        r = top + k
        _span(ws, r, 1, 2, f'=IF(INDEX(wr_OppLabels,{k})="","","vs "&INDEX(wr_OppLabels,{k})&" · SL "&INDEX(wr_OppSL,{k}))',
              font=SMALL)
        _span(ws, r, 3, 12, (f'=IF(INDEX(wr_OppLabels,{k})="","",IF(NOT(INDEX(wr_OppUnplayed,{k})),"Already played.",'
                             f'IF(INDEX(wr_Send1,{k})="",IF(wr_HasEvidence,"No favorable, even or indirect evidence among our remaining players.",""),'
                             f'INDEX(wr_SendList,{k}))))'), font=SMALL)
        ws.row_dimensions[r].height = 28
    top = top + R + 2
    base._section(ws, top, "Top risks", last_col=12)
    _span(ws, top + 1, 1, 12, "Opponents with winning records against our roster (by recorded direct results; unplayed only)",
          font=base.SUBHEAD_FONT, fill=base.SUBHEAD_FILL)
    for n in range(1, 4):
        row = f"INDEX(wr_ThreatRow,MATCH({n},wr_ThreatRun,0))"
        note = f"INDEX(wr_OppNotes,MATCH(INDEX(Threats_Table[Opp Player ID],{row}),wr_OppPids,0))"
        _span(ws, top + 1 + n, 1, 12, f'=IFERROR(INDEX(Threats_Table[Threat],{row})&IFERROR(IF({note}="",""," · 📝 "&{note}&'
                                      f'" (your opinion, not APA facts)"),""),'
                                      f'IF({n}=1,IF(wr_HasEvidence,"No unplayed opponent has a winning recorded record against our roster.",""),""))',
              font=SMALL)
    t2 = top + 5
    _span(ws, t2, 1, 12, "Concerning pairings — more direct losses than wins (remaining players vs unplayed opponents)",
          font=base.SUBHEAD_FONT, fill=base.SUBHEAD_FILL)
    for n in range(1, 6):
        _span(ws, t2 + n, 1, 12, f'=IFERROR(INDEX(Concerning_Table[Pairing],INDEX(wr_ConcernRow,MATCH({n},wr_ConcernRun,0))),'
                                 f'IF({n}=1,IF(wr_HasEvidence,"No concerning direct records among remaining pairings.",""),""))',
              font=SMALL)
    t3 = t2 + 7
    _span(ws, t3, 1, 12, "Open risks — unplayed opponents with no favorable direct option left among our remaining players",
          font=base.SUBHEAD_FONT, fill=base.SUBHEAD_FILL)
    for n in range(1, R + 1):
        _span(ws, t3 + n, 1, 12, f'=IFERROR(INDEX(wr_OppLabels,MATCH({n},wr_RiskRun,0))&" — "&IF(INDEX(wr_SendLeft,MATCH({n},wr_RiskRun,0))>0,'
                                 f'"only even or indirect evidence left","no evidence-backed option left"),IF({n}=1,IF(wr_HasEvidence,'
                                 f'"None — every unplayed opponent still has a favorable direct option.",""),""))', font=SMALL)

    top = t3 + R + 2
    base._section(ws, top, "Lineup Lab snapshot (your marks applied; edit them on Lineup Lab)", last_col=12)
    metrics = [
        ("Remaining players", '=IF(wr_OurCount=0,"—",wr_RemainingCount&" of "&wr_OurCount&IF(wr_RemainingUnknown>0," ("&'
                              'wr_RemainingUnknown&" with unknown availability — still counted as remaining)",""))'),
        ("Remaining skill total", '=IF(wr_OurCount=0,"—",IF(wr_RemainingMissing=0,wr_RemainingSL&" (all remaining players have a captured SL)",'
                                  'wr_RemainingSL&" known subtotal · "&wr_RemainingMissing&" remaining player(s) without a captured SL — not a complete total"))'),
        ("Selected lineup (Planned + Played)", '=IF(wr_SelectedCount=0,"No players marked Planned or Played yet.",IF(wr_SelectedMissing=0,'
                                               '"Skill total "&wr_SelectedSL&" for "&wr_SelectedCount&" selected player(s)","Known subtotal "&wr_SelectedSL&'
                                               '" · "&wr_SelectedMissing&" selected player(s) without a captured SL")&IF(ll_Cap="",""," · your reference cap: "&'
                                               'll_Cap&" (user-entered, not verified)")&" — not a lineup-legality check.")'),
        ("Roster flexibility", '=IF(wr_UnplayedCount=0,"—","Favorable direct option left vs "&wr_GreenOpps&" of "&wr_UnplayedCount&'
                               '" unplayed opponents · any evidence-backed option vs "&wr_SendOpps&" of "&wr_UnplayedCount)'),
    ]
    for k, (label, formula) in enumerate(metrics, start=1):
        _span(ws, top + k, 1, 2, label, font=base.LABEL_FONT)
        _span(ws, top + k, 3, 12, formula, font=SMALL)
    u_top = top + len(metrics) + 1
    _span(ws, u_top, 1, 12, "Unique favorable options — consider saving (the only remaining favorable direct option vs an unplayed opponent)",
          font=base.SUBHEAD_FONT, fill=base.SUBHEAD_FILL)
    for n in range(1, R + 1):
        _span(ws, u_top + n, 1, 12, (f'=IFERROR(INDEX(wr_OurLabels,MATCH({n},wr_UniqueRun,0))&" — only favorable direct option vs "&'
                                     f'INDEX(wr_OppLabels,INDEX(wr_UniqueFirst,MATCH({n},wr_UniqueRun,0)))&IF(INDEX(wr_UniqueCount,'
                                     f'MATCH({n},wr_UniqueRun,0))>1," (and "&(INDEX(wr_UniqueCount,MATCH({n},wr_UniqueRun,0))-1)&" more)",""),'
                                     f'IF({n}=1,IF(wr_HasEvidence,"None right now.",""),""))'), font=SMALL)

    top = u_top + R + 2
    base._section(ws, top, "Matchup matrix — our players (rows) vs their players (columns)", last_col=12)
    _span(ws, top + 1, 1, 9, "Green = more direct wins than losses · Red = more direct losses than wins · Yellow = even direct "
                             "record, or shared-opponent evidence only (≈ ours vs theirs) · Gray = no evidence. Numbers in () are "
                             "meetings (Evidence view). Captain view: 🟢🟡⚪🔴 + our direct record. Not odds or predictions.",
          font=base.MUTED_FONT, height=30)
    ws.cell(row=top + 1, column=10, value="View").font = base.LABEL_FONT
    base._input(ws.cell(row=top + 1, column=11), "Evidence view")
    ws.merge_cells(start_row=top + 1, start_column=11, end_row=top + 1, end_column=12)
    _dv(ws, f"K{top + 1}", '"Evidence view,Captain view"', "Captain view: 🟢🟡⚪🔴 and the record. Evidence view: sample sizes too.")
    view_ref = f"$K${top + 1}"
    header = top + 2
    ws.cell(row=header, column=1, value="Our player ↓ / opponent →").font = base.SUBHEAD_FONT
    ws.cell(row=header, column=1).alignment = base.WRAP_TOP
    matrix_cols = [2, 3, 4, 5, 7, 8, 9, 10, 11, 12][:R]
    for j, col in enumerate(matrix_cols, start=1):
        c = ws.cell(row=header, column=col, value=f'=INDEX(wr_OppLabels,{j})')
        c.font = Font(bold=True, size=8, color=base.FELT_DEEP)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    narrow = min(ws.column_dimensions[get_column_letter(col)].width for col in matrix_cols)
    ws.row_dimensions[header].height = _worst_height(_longest_label(wb), narrow, 8, 42)
    grid = engine["grid"]
    helper_col0 = 15  # O.. : a same-sheet copy of the category grid for conditional formatting
    for i in range(1, R + 1):
        r = header + i
        ws.cell(row=r, column=1, value=f'=INDEX(wr_OurLabels,{i})').font = SMALL
        ws.cell(row=r, column=1).alignment = base.WRAP_TOP
        for j, col in enumerate(matrix_cols, start=1):
            cat_ref, cell_ref, _ = grid[(i, j)]
            c = ws.cell(row=r, column=col, value=f'=IF({view_ref}="Captain view",{engine["captain"][(i, j)]},{cell_ref})')
            c.font = SMALL
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r, column=helper_col0 + j - 1, value=f"={cat_ref}").font = base.HELPER_FONT
        ws.row_dimensions[r].height = _worst_height(WORST_EVIDENCE, narrow, 9, 30)
    first_row, last_row = header + 1, header + R
    for j in range(R):
        ws.column_dimensions[get_column_letter(helper_col0 + j)].hidden = True
    for j, col in enumerate(matrix_cols, start=1):
        letter = get_column_letter(col)
        helper = get_column_letter(helper_col0 + j - 1)
        for code, fill in CAT_FILLS.items():
            ws.conditional_formatting.add(f"{letter}{first_row}:{letter}{last_row}",
                                          FormulaRule(formula=[f'${helper}{first_row}="{code}"'], fill=fill, stopIfTrue=True))

    top = last_row + 2
    base._section(ws, top, "Inspect — choose an opponent to see all our candidates, or one of our players to see all opponents",
                  last_col=12)
    _span(ws, top + 1, 1, 2, "Inspect opponent", font=base.LABEL_FONT)
    base._input(ws.cell(row=top + 1, column=3))
    ws.merge_cells(start_row=top + 1, start_column=3, end_row=top + 1, end_column=5)
    _dv(ws, f"C{top + 1}", "wr_OppLabels", "Pick an opponent player from tonight's roster.")
    insp = f"$C${top + 1}"
    _span(ws, top + 1, 7, 12, f'=IF({insp}="","Pick an opponent to rank our players against them (ranking rule below).","")',
          font=base.MUTED_FONT)
    j_ref = f'IFERROR(MATCH({insp},wr_OppLabels,0),"")'
    base._subheader(ws, top + 2, [(1, "Our player (APA record ID)"), (3, "Rank"), (4, "SL"), (5, "Direct record"),
                                  (7, "Shared-opponent evidence (samples)"), (10, "Basis · ties · missing evidence")])
    for rr in range(1, R + 1):
        r = top + 2 + rr
        # Excel evaluates EVERY argument of OR(), so each guard term must itself be error-free: with nothing
        # picked, INDEX(...,"") is #VALUE! (real-Excel UAT showed #VALUE! here; GPT #84 P1).
        s = f'IFERROR(INDEX(wr_OppStart,{j_ref}),"")'
        n = f'IFERROR(INDEX(wr_OppCount,{j_ref}),0)'
        cond = f'OR({insp}="",{j_ref}="",{s}="",{rr}>{n})'
        _span(ws, r, 1, 2, f'=IF({cond},"",INDEX(MatchupEvidence_Table[Player],{s}+{rr}))', font=SMALL)
        # Top-aligned like the rest of the row: in a tall row a bottom-aligned rank sat beside the NEXT player's
        # name (real-Excel UAT 787f6d7).
        ws.cell(row=r, column=3, value=f'=IF({cond},"",INDEX(MatchupEvidence_Table[Rank],{s}+{rr}))').alignment = base.WRAP_TOP
        ws.cell(row=r, column=4, value=f'=IF({cond},"",INDEX(MatchupEvidence_Table[SL],{s}+{rr}))').alignment = base.WRAP_TOP
        _span(ws, r, 5, 6, f'=IF({cond},"",INDEX(MatchupEvidence_Table[Direct Record],{s}+{rr}))', font=SMALL)
        _span(ws, r, 7, 9, f'=IF({cond},"",INDEX(MatchupEvidence_Table[Shared-Opponent Evidence],{s}+{rr}))', font=SMALL)
        _span(ws, r, 10, 12, f'=IF({cond},"",INDEX(MatchupEvidence_Table[Basis],{s}+{rr}))', font=SMALL)
        basis_width = sum(ws.column_dimensions[get_column_letter(c)].width for c in (10, 11, 12))
        ws.row_dimensions[r].height = _worst_height(WORST_BASIS, basis_width, 9, 30)
    top2 = top + 3 + R + 1
    _span(ws, top2, 1, 2, "Inspect our player", font=base.LABEL_FONT)
    base._input(ws.cell(row=top2, column=3))
    ws.merge_cells(start_row=top2, start_column=3, end_row=top2, end_column=5)
    _dv(ws, f"C{top2}", "wr_OurLabels", "Pick one of our players.")
    insp2 = f"$C${top2}"
    i_ref = f'IFERROR(MATCH({insp2},wr_OurLabels,0),"")'
    base._subheader(ws, top2 + 1, [(1, "Opponent (APA record ID)"), (3, "Evidence"), (5, "Direct record"),
                                   (7, "Shared-opponent evidence (samples)"), (10, "Rank vs this opponent")])
    for j in range(1, R + 1):
        r = top2 + 1 + j
        s = f"INDEX(wr_OppStart,{j})"
        cond = f'OR({insp2}="",{i_ref}="",{s}="")'
        # position of our player i inside opponent j's block: recompute with MATCH on the block's ids
        p = (f'MATCH(INDEX(wr_OurPids,{i_ref}),INDEX(MatchupEvidence_Table[Player ID],{s}+1):'
             f'INDEX(MatchupEvidence_Table[Player ID],{s}+INDEX(wr_OppCount,{j})),0)')
        _span(ws, r, 1, 2, f'=IF({cond},"",INDEX(wr_OppLabels,{j}))', font=SMALL)
        _span(ws, r, 3, 4, (f'=IF({cond},"",IFERROR(IF(INDEX(MatchupEvidence_Table[Category],{s}+{p})="G","Favorable direct",'
                            f'IF(INDEX(MatchupEvidence_Table[Category],{s}+{p})="R","Concerning direct",IF(INDEX(MatchupEvidence_Table[Category],'
                            f'{s}+{p})="E","Even direct",IF(INDEX(MatchupEvidence_Table[Category],{s}+{p})="I","Indirect only","Insufficient")))),""))'),
              font=SMALL)
        _span(ws, r, 5, 6, f'=IF({cond},"",IFERROR(INDEX(MatchupEvidence_Table[Direct Record],{s}+{p}),""))', font=SMALL)
        _span(ws, r, 7, 9, f'=IF({cond},"",IFERROR(INDEX(MatchupEvidence_Table[Shared-Opponent Evidence],{s}+{p}),""))', font=SMALL)
        _span(ws, r, 10, 12, f'=IF({cond},"",IFERROR(INDEX(MatchupEvidence_Table[Rank],{s}+{p}),""))', font=SMALL)
        ws.row_dimensions[r].height = 30

    top = top2 + 2 + R + 1
    base._section(ws, top, "Direct meetings between the rosters (newest first)", last_col=12)
    base._subheader(ws, top + 1, [(1, "Date"), (2, "Our player"), (5, "Opponent"), (9, "Result (ours)"), (10, "SL ours/theirs"),
                                  (11, "Session")])
    for k in range(1, MEETING_SLOTS + 1):
        r = top + 1 + k
        row = f'IFERROR(MATCH(wr_PairKey&"|{k}",Meetings_Table[Key],0),"")'
        cond = f'OR(NOT(wr_HasEvidence),{row}="")'
        ws.cell(row=r, column=1, value=f'=IF({cond},"",INDEX(Meetings_Table[Date],{row}))').font = SMALL
        _span(ws, r, 2, 4, f'=IF({cond},"",INDEX(Meetings_Table[Our Player],{row}))', font=SMALL, wrap=False)
        _span(ws, r, 5, 8, f'=IF({cond},"",INDEX(Meetings_Table[Opponent],{row}))', font=SMALL, wrap=False)
        ws.cell(row=r, column=9, value=f'=IF({cond},"",INDEX(Meetings_Table[Result],{row}))').font = SMALL
        ws.cell(row=r, column=10, value=f'=IF({cond},"",INDEX(Meetings_Table[Our SL],{row})&" / "&INDEX(Meetings_Table[Their SL],{row}))').font = SMALL
        _span(ws, r, 11, 12, f'=IF({cond},"",INDEX(Meetings_Table[Session],{row}))', font=SMALL, wrap=False)
    _span(ws, top + 2 + MEETING_SLOTS, 1, 12,
          f'=IF(AND(wr_HasEvidence,ISNUMBER(MATCH(wr_PairKey&"|{MEETING_SLOTS + 1}",Meetings_Table[Key],0))),"More meetings exist — '
          f'filter the Meetings sheet by this pairing.",IF(AND(wr_HasEvidence,NOT(ISNUMBER(MATCH(wr_PairKey&"|1",Meetings_Table[Key],0)))),'
          f'"No recorded direct meetings between these rosters.",""))', font=base.MUTED_FONT)

    top = top + 4 + MEETING_SLOTS
    base._section(ws, top, "Team comparison", last_col=12)
    t = f'IFERROR(MATCH(wr_PairKey,TeamComparison_Table[Pair Key],0),"")'
    comps = [("Players rostered", "Our Rostered", "Opp Rostered"), ("Players with a captured SL", "Our Captured SL", "Opp Captured SL"),
             ("Skill total (captured SLs only)", "Our SL Total", "Opp SL Total"), ("Recorded games (all opponents)", "Our Games", "Opp Games"),
             ("Players with no recorded games", "Our No-Game Players", "Opp No-Game Players")]
    for k, (label, a, b) in enumerate(comps, start=1):
        r = top + k
        _span(ws, r, 1, 2, label, font=base.LABEL_FONT)
        _span(ws, r, 3, 5, f'=IF({t}="","",INDEX(TeamComparison_Table[{a}],{t}))', wrap=False)
        _span(ws, r, 7, 9, f'=IF({t}="","",INDEX(TeamComparison_Table[{b}],{t}))', wrap=False)
    for k, (label, col) in enumerate((("Direct meetings between the rosters", "Direct Meetings"),
                                      ("Opponent players our roster has met directly", "Opponents Met"),
                                      ("Player pairings by evidence", "Pairings")), start=len(comps) + 1):
        r = top + k
        _span(ws, r, 1, 2, label, font=base.LABEL_FONT)
        _span(ws, r, 3, 12, f'=IF({t}="","",INDEX(TeamComparison_Table[{col}],{t}))', wrap=False)
    top = top + len(comps) + 5
    base._section(ws, top, "How to read this", last_col=12)
    for k, text in enumerate([
        "Ranking rule (reviewed, unchanged): players with direct meetings first (by observed direct record, then more meetings), "
        "then shared-opponent results only (by our record against opponents both players faced, then more shared opponents, "
        "then more games); no evidence = listed last, not ranked. Ties are shown as “2=” and named.",
        "Best-supported sends keep that order and drop only concerning (losing direct) and no-evidence pairings. A direct "
        "record ranks ahead of indirect evidence even when it is small — read the records and sample sizes.",
        "No recorded meetings with our roster does not mean a weak opponent — it means unknown.",
        "Nothing here is a win probability (none is calibrated), a confidence claim or a guaranteed lineup.",
    ], start=1):
        _span(ws, top + k, 1, 12, text, font=base.MUTED_FONT, height=30)
    ws.freeze_panes = "A5"
    return {"matrix_header": header}


# ---------------------------------------------------------------------------
# Lineup Lab
# ---------------------------------------------------------------------------

_MD_PREFIX = {"wr_": "md_", "ll_": "lm_"}
_MD_SHEETS = ("Lineup Lab", "Scouting Cards", "Captain Packet", "Coach Dashboard", "Command Center")


def _to_md(text: str, engine_title: str, md_title: str) -> str:
    import re
    text = re.sub(r"(?<![A-Za-z0-9_])(wr|ll)_", lambda m: _MD_PREFIX[m.group(1) + "_"], text)
    return text.replace(f"{engine_title}!", f"'{md_title}'!")


def split_match_day_engine(wb) -> None:
    """Local overrides affect only their own tab (Paul, GPT audit #84).

    The War Room's override cells feed the wr_* pairing. Every other output --
    Lineup Lab, Scouting Cards, Captain Packet, Coach Dashboard -- must keep
    following Match Day, so the whole pairing engine is copied once to "Engine MD"
    with wr_/ll_ renamed md_/lm_ and the War Room override inputs pinned blank:
    md_* is always Match Day's fixture. Those four sheets (and the Coach
    Dashboard's engine cells) are re-pointed at md_*/lm_*, so a War Room override
    can never put one fixture's heading over another opponent's roster.
    """
    import re
    src = wb["Engine"]
    md = wb.copy_worksheet(src)
    md.title = "Engine MD"
    for row in md.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = _to_md(c.value, "Engine", "Engine MD")
    for name in list(wb.defined_names):
        if name[:3] in _MD_PREFIX:
            target = _MD_PREFIX[name[:3]] + name[3:]
            text = wb.defined_names[name].attr_text.replace("Engine!", "'Engine MD'!")
            wb.defined_names[target] = DefinedName(target, attr_text=text)
    for name in ("md_OurInput", "md_OppInput"):
        ref = wb.defined_names[name].attr_text.split("!")[1].replace("$", "")
        md[ref].value = '=""'
    # The Coach Dashboard's engine cells (Player B pool, format) follow Match Day too.
    cd_rows = []
    for name in wb.defined_names:
        if name.startswith("cd_"):
            m = re.search(r"\$?[A-Z]+\$?(\d+)", wb.defined_names[name].attr_text.split("!")[1])
            cd_rows.append(int(m.group(1)))
    for row in src.iter_rows(min_row=min(cd_rows) - 2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = _to_md(c.value, "Engine", "Engine")
    for sheet in _MD_SHEETS:
        ws = wb[sheet]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = _to_md(c.value, "Engine", "Engine MD")
        for dv in ws.data_validations.dataValidation:
            if dv.formula1:
                dv.formula1 = _to_md(dv.formula1, "Engine", "Engine MD")
    md.sheet_state = "hidden" if src.sheet_state == "hidden" else src.sheet_state


def build_lineup_lab(wb, *, slots: dict[str, int], default_team: str | None, default_opp: str | None,
                     default_date: str | None = None) -> None:
    R = slots["roster"]
    ws = wb.create_sheet("Lineup Lab", 2)
    for letter, width in zip("ABCDEFGH", (6, 44, 18, 30, 10, 12, 3, 3)):
        ws.column_dimensions[letter].width = width
    _head(ws, "Lineup Lab", "Your planning marks for tonight — inputs, never evidence. Results update instantly on this "
                            "sheet and on the War Room.", last=6, current="Lineup Lab")
    rows = [(5, "Planning for our team", default_team or "", "TeamNameList",
             '=IF(ll_OurTeamOK,"✓ Matches Match Day’s team.",IF(ll_OurTeam="","Name the team you are planning for.",'
             '"⚠ Match Day shows "&IF(wr_OurLabel="","no team",wr_OurLabel)&" — these marks are NOT applied there."))'),
            (6, "Planning for opponent", default_opp or "", "TeamNameList",
             '=IF(ll_OppOK,"✓ Matches Match Day’s opponent.",IF(ll_OppTeam="","Name the opponent team.","⚠ Match Day’s opponent is "&'
             'IF(wr_OppLabel="","none",wr_OppLabel)&" — opponent marks are NOT applied there."))'),
            (7, "Reference skill cap (optional)", "", None,
             '="Your own reference number (e.g. your division rule). Never assumed, never applied automatically."')]
    for r, label, value, source, status in rows:
        _span(ws, r, 1, 2, label, font=base.LABEL_FONT)
        base._input(ws.cell(row=r, column=3), value)
        if source:
            _dv(ws, f"C{r}", source, "Pick a team.")
        _span(ws, r, 4, 6, status, font=base.MUTED_FONT)
        ws.row_dimensions[r].height = 30
    _span(ws, 8, 1, 6, "Marks are for ONE fixture (below, ending in its match number). Pick another fixture on Match Day "
                       "— even later the same day — and these marks stop applying (it says so). Clear them, pick the new "
                       "fixture in the cell below, and plan again. Coach notes describe the player and always follow the "
                       "opponent team.",
          font=base.MUTED_FONT, height=44)
    _span(ws, 9, 1, 2, "Planning for fixture", font=base.LABEL_FONT)
    base._input(ws.cell(row=9, column=3), default_date or "")
    _dv(ws, "C9", "uc_PlanKeyList", "Pick Match Day's current fixture (the list offers exactly that one).")
    _span(ws, 9, 4, 6, '=IF(ll_FixOK,"✓ Matches Match Day’s fixture — availability, lineup and played applied.",'
                       'IF(uc_PlanKey="","Choose a fixture on Match Day first.",'
                       '"⚠ Match Day’s fixture is "&uc_PlanKey&" — availability, lineup and played are NOT applied. Notes still are."))',
          font=base.MUTED_FONT)
    ws.row_dimensions[9].height = 30
    base._section(ws, 10, "Our roster — availability and lineup (blank availability = Unknown)", last_col=6)
    base._subheader(ws, 11, [(1, "#"), (2, "Player (APA record ID)"), (3, "Availability"), (4, "Lineup"), (5, "SL"), (6, "Team W-L")])
    our_first = 12
    avail = DataValidation(type="list", formula1='"' + ",".join(AVAILABILITY_CHOICES) + '"', allow_blank=True)
    lineup = DataValidation(type="list", formula1='"' + ",".join(LINEUP_CHOICES) + '"', allow_blank=True)
    played = DataValidation(type="list", formula1='"Played"', allow_blank=True)
    for dvx in (avail, lineup, played):
        ws.add_data_validation(dvx)
    for k in range(1, R + 1):
        r = our_first + k - 1
        ws.cell(row=r, column=1, value=f'=IF(INDEX(ll_OurLabels,{k})="","",{k})')
        ws.cell(row=r, column=2, value=f'=INDEX(ll_OurLabels,{k})')
        for c in (3, 4):
            base._input(ws.cell(row=r, column=c))
            ws.cell(row=r, column=c).font = Font(bold=True)
        avail.add(f"C{r}")
        lineup.add(f"D{r}")
        ws.cell(row=r, column=5, value=f'=INDEX(ll_OurSL,{k})')
        ws.cell(row=r, column=6, value=f'=INDEX(ll_OurWL,{k})')
    opp_top = our_first + R + 1
    base._section(ws, opp_top, "Opponent roster — who has played, and your scouting notes", last_col=6)
    base._subheader(ws, opp_top + 1, [(1, "#"), (2, "Player (APA record ID)"), (3, "Played"), (4, "Coach notes"), (5, "SL"),
                                      (6, "Team W-L")])
    opp_first = opp_top + 2
    for k in range(1, R + 1):
        r = opp_first + k - 1
        ws.cell(row=r, column=1, value=f'=IF(INDEX(ll_OppLabels,{k})="","",{k})')
        ws.cell(row=r, column=2, value=f'=INDEX(ll_OppLabels,{k})')
        for c in (3, 4):
            base._input(ws.cell(row=r, column=c))
            ws.cell(row=r, column=c).font = Font(bold=True)
        played.add(f"C{r}")
        ws.cell(row=r, column=4).alignment = base.WRAP_TOP
        ws.cell(row=r, column=5, value=f'=INDEX(ll_OppSL,{k})')
        ws.cell(row=r, column=6, value=f'=INDEX(ll_OppWL,{k})')
    assert our_first == slots["ll_our_first"] and opp_first == slots["ll_opp_first"]
    res = opp_first + R + 1
    base._section(ws, res, "Results (instant, for the War Room's matchup)", last_col=6)
    lines = [
        ("Remaining players", '=IF(wr_OurCount=0,"—",wr_RemainingCount&" of "&wr_OurCount&IF(wr_RemainingUnknown>0," ("&wr_RemainingUnknown&" unknown availability)",""))'),
        ("Remaining skill total", '=IF(wr_OurCount=0,"—",IF(wr_RemainingMissing=0,wr_RemainingSL,wr_RemainingSL&" known · "&wr_RemainingMissing&" missing SL"))'),
        ("Selected lineup total", '=IF(wr_SelectedCount=0,"—",IF(wr_SelectedMissing=0,wr_SelectedSL&" ("&wr_SelectedCount&" players)",'
                                  'wr_SelectedSL&" known · "&wr_SelectedMissing&" missing SL")&IF(ll_Cap="",""," vs your reference "&ll_Cap))'),
        ("Flexibility", '=IF(wr_UnplayedCount=0,"—","Favorable option left vs "&wr_GreenOpps&" of "&wr_UnplayedCount&" unplayed opponents")'),
        ("Open risks", '=IF(wr_UnplayedCount=0,"—",wr_RiskCount&" unplayed opponent(s) with no favorable direct option left")'),
    ]
    for k, (label, formula) in enumerate(lines, start=1):
        _span(ws, res + k, 1, 2, label, font=base.LABEL_FONT)
        _span(ws, res + k, 3, 6, formula, font=SMALL)
    bs = res + len(lines) + 1
    _span(ws, bs, 1, 6, "Best remaining send per unplayed opponent — and why (the recorded evidence, with its sample size)",
          font=base.SUBHEAD_FONT, fill=base.SUBHEAD_FILL)
    for k in range(1, R + 1):
        _span(ws, bs + k, 1, 6, f'=IF(OR(INDEX(wr_OppLabels,{k})="",NOT(INDEX(wr_OppUnplayed,{k}))),"","vs "&INDEX(wr_OppLabels,{k})&": "&'
                                f'IF(INDEX(wr_BestText,{k})="","no evidence-backed option left among our remaining players",'
                                f'INDEX(wr_BestText,{k})))', font=SMALL, height=30)
    ws.freeze_panes = "A4"


# ---------------------------------------------------------------------------
# Scouting Cards
# ---------------------------------------------------------------------------

CARD_FIELDS = [("Quick read", "Quick Read"), ("Team record", "Team Record"), ("League lifetime", "Lifetime"), ("Recorded games", "Sample"),
               ("Vs our roster", "Vs Our Roster"), ("Meetings with our players", "Met List"),
               ("Shared-opponent evidence", "Shared Summary"), ("Record by opponent SL", "By SL"),
               ("Winning records vs", "Winning SL"), ("Losing records vs", "Losing SL"), ("Missing information", "Missing")]


def _card_block(ws, top: int, k: int, *, compact: bool = False) -> int:
    card = f"INDEX(wr_OppCard,{k})"
    _span(ws, top, 1, 12, f'=IF(INDEX(wr_OppLabels,{k})="","",INDEX(wr_OppLabels,{k})&" · SL "&INDEX(wr_OppSL,{k})&'
                          f'IF(INDEX(wr_OppPlayed,{k})="Played"," · already played",""))',
          font=Font(bold=True, size=12, color="FFFFFF"), fill=base.OPPONENT_FILL)
    fields = CARD_FIELDS if not compact else [f for f in CARD_FIELDS if f[1] in ("Quick Read", "Vs Our Roster", "Met List", "By SL", "Missing")]
    for n, (label, col) in enumerate(fields, start=1):
        r = top + n
        _span(ws, r, 1, 2, f'=IF(INDEX(wr_OppLabels,{k})="","","{label}")', font=base.LABEL_FONT)
        _span(ws, r, 3, 12, f'=IF(OR(INDEX(wr_OppLabels,{k})="",{card}=""),IF(INDEX(wr_OppLabels,{k})="","","Not precomputed for this pairing"),'
                            f'INDEX(Scouting_Table[{col}],{card}))', font=SMALL)
    r = top + len(fields) + 1
    _span(ws, r, 1, 2, f'=IF(INDEX(wr_OppLabels,{k})="","","Coach observations")', font=base.LABEL_FONT)
    _span(ws, r, 3, 12, f'=IF(INDEX(wr_OppLabels,{k})="","",IF(INDEX(wr_OppNotes,{k})="","(add notes on Coach Notes or Lineup Lab)",INDEX(wr_OppNotes,{k})))',
          font=SMALL)
    return r + 1


def build_scouting_cards(wb, *, slots: dict[str, int]) -> None:
    R = slots["roster"]
    ws = wb.create_sheet("Scouting Cards", 3)
    for letter, width in zip("ABCDEFGHIJKL", (22, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12)):
        ws.column_dimensions[letter].width = width
    _head(ws, "Opponent Scouting Cards", "One card per opponent player — recorded facts, sample sizes, what's missing, "
                                         "and your notes.", last=12, current="Scouting Cards")
    _span(ws, 4, 1, 12, "=wr_Status", font=Font(bold=True, color=base.FELT_DEEP))
    top = 6
    for k in range(1, R + 1):
        top = _card_block(ws, top, k) + 1
    _span(ws, top, 1, 12, "“Winning/losing records vs SLx” only state the direction of recorded results against opponents "
                          "of that skill level, with the record shown — never a prediction. No meetings with our roster means "
                          "unknown, not weak.", font=base.MUTED_FONT, height=30)
    ws.freeze_panes = "A5"


# ---------------------------------------------------------------------------
# Captain Packet (print)
# ---------------------------------------------------------------------------

# Captain Packet layout constants (Excel width units, points). Paul's real-Excel UAT (#84): readable
# text, no clipped Basis, no orphan opponent headings, no empty bands, keep the five-page grouping.
PACKET_WIDTHS = (26, 22, 7, 11, 24, 2, 14, 10, 9, 10, 22, 10)        # A..L
PACKET_PRINTABLE = ((11 - 2 * 0.35) * 72, (8.5 - 2 * 0.45) * 72)     # landscape Letter, margins below
PACKET_TITLE_ROWS = (26, 20)                                          # rows 1-2 repeat on every page


def _width_pt(width: float) -> float:
    return int(width * 7 + 5) * 0.75


def build_captain_packet(wb, *, slots: dict[str, int], stats: dict[str, Any]) -> dict[str, Any]:
    """Printable Captain Packet that follows Match Day (md_* after split_match_day_engine).

    Every printed row has an explicit height and every column an explicit width, so the page geometry
    is known at build time; one print scale is then chosen so the tallest page fits, and each page uses
    the largest font its row heights allow. Variable-length sections are PACKED (no reserved blank slot
    bands): evidence is one continuous list in which every line names its opponent, so a page break can
    never orphan an opponent heading. Facts are unchanged; only presentation differs from the War Room.
    """
    R = slots["roster"]
    ws = wb.create_sheet("Captain Packet", 4)
    for letter, width in zip("ABCDEFGHIJKL", PACKET_WIDTHS):
        ws.column_dimensions[letter].width = width
    for letter in "PQR":
        ws.column_dimensions[letter].width = 6
    ws.sheet_view.showGridLines = False
    heights: dict[int, float] = {}

    def h(row: int, pt: float) -> None:
        ws.row_dimensions[row].height = pt
        heights[row] = pt

    def font(size: float, **kw) -> Font:
        return Font(size=size, **kw)

    white = "FFFFFF"
    ws["A1"] = "Captain Packet — Match Night"
    ws["A1"].font = font(16, bold=True, color=base.FELT_DEEP)
    ws.merge_cells("A1:L1")
    h(1, PACKET_TITLE_ROWS[0])
    _span(ws, 2, 1, 12, '=IF(uc_Fixture="","Choose a fixture on Match Day.",uc_Fixture)',
          font=font(12, bold=True, color=base.FELT_DEEP))
    h(2, PACKET_TITLE_ROWS[1])
    _span(ws, 3, 1, 12, '=uc_Venue', font=font(9.5, color="5B6A61"))
    h(3, 15)
    h(4, 4)

    # ---------------- page 1 (decision first): best sends, risks, then the rosters ----------------
    r = 4
    r += 1
    base._section(ws, r, "Best sends — top opportunities per opponent (favorable direct first, then even, then indirect)",
                  last_col=12)
    h(r, 18)
    for k in range(1, R + 1):
        r += 1
        _span(ws, r, 1, 2, f'=IF(INDEX(wr_OppLabels,{k})="","","vs "&INDEX(wr_OppLabels,{k}))', font=font(10, bold=True))
        _span(ws, r, 3, 12, f'=IF(INDEX(wr_OppLabels,{k})="","",IF(NOT(INDEX(wr_OppUnplayed,{k})),"Already played.",'
                            f'IF(INDEX(wr_Send1,{k})="","No evidence-backed option left among our remaining players.",'
                            f'INDEX(wr_Send1,{k})&IF(INDEX(wr_Send2,{k})="",""," · "&INDEX(wr_Send2,{k})))))',
              font=font(10))
        h(r, 27)
    r += 1
    base._section(ws, r, "Top risks", last_col=12)
    h(r, 18)
    risk_lines = [f'=IFERROR("Dangerous: "&INDEX(Threats_Table[Threat],INDEX(wr_ThreatRow,MATCH({n},wr_ThreatRun,0))),"")'
                  for n in range(1, 4)]
    risk_lines += [f'=IFERROR("Avoid: "&INDEX(Concerning_Table[Pairing],INDEX(wr_ConcernRow,MATCH({n},wr_ConcernRun,0))),"")'
                   for n in range(1, 4)]
    risk_lines.append('=IF(wr_UnplayedCount=0,"","Open risks: "&wr_RiskCount&" unplayed opponent(s) with no favorable '
                      'direct option left among our remaining players.")')
    for formula in risk_lines:
        r += 1
        _span(ws, r, 1, 12, formula, font=font(10.5), wrap=False)
        h(r, 15)
    r += 1
    for c1, c2, text, sub, fill in ((1, 1, "OUR TEAM", "=wr_OurLabel", base.SECTION_FILL),
                                    (7, 7, "OPPONENT", "=wr_OppLabel", base.OPPONENT_FILL)):
        ws.cell(row=r, column=c1, value=text).font = font(10.5, bold=True, color=white)
        ws.cell(row=r, column=c1).fill = fill
        _span(ws, r, c1 + 1, c1 + 4 if c1 == 1 else 12, sub, font=font(10.5, bold=True, color=base.FELT_DEEP))
    h(r, 18)
    r += 1
    for col, text in ((1, "Player (APA record ID)"), (3, "SL"), (4, "Team W-L"), (5, "Availability · lineup"),
                      (7, "Player (APA record ID)"), (11, "SL · Team W-L"), (12, "Played")):
        cell = ws.cell(row=r, column=col, value=text)
        cell.font = font(9, bold=True, color=base.FELT_DEEP)
        cell.fill = base.SUBHEAD_FILL
    h(r, 16)
    roster_first = r + 1
    for k in range(1, R + 1):
        r = roster_first + k - 1
        f11 = font(11)
        _span(ws, r, 1, 2, f'=INDEX(wr_OurLabels,{k})', font=f11, wrap=False)
        ws.cell(row=r, column=3, value=f'=INDEX(wr_OurSL,{k})').font = f11
        ws.cell(row=r, column=4, value=f'=INDEX(wr_OurWL,{k})').font = f11
        ws.cell(row=r, column=5, value=f'=IF(INDEX(wr_OurLabels,{k})="","",INDEX(wr_OurAvail,{k})&" · "&INDEX(wr_OurLineup,{k}))').font = f11
        _span(ws, r, 7, 10, f'=INDEX(wr_OppLabels,{k})', font=f11, wrap=False)
        ws.cell(row=r, column=11, value=f'=IF(INDEX(wr_OppLabels,{k})="","","SL "&INDEX(wr_OppSL,{k})&" · "&INDEX(wr_OppWL,{k}))').font = f11
        ws.cell(row=r, column=12, value=f'=INDEX(wr_OppPlayed,{k})').font = f11
        for c in range(1, 13):
            if c != 6:
                ws.cell(row=r, column=c).border = base.CELL_BORDER
        h(r, 16)
    r = roster_first + R
    _span(ws, r, 1, 12, '=IF(AND(ll_OurTeamOK,ll_OppTeamOK,NOT(ll_FixOK)),"⚠ Lineup Lab marks are for another fixture — not '
                        'applied here. ","")&"Team W-L = this team’s captured record this session · “No data” = not captured · '
                        'availability, lineup and played are your Lineup Lab marks for this fixture, never evidence; Unknown is '
                        'not Unavailable."', font=font(9, color="5B6A61"))
    h(r, 26)
    r += 1
    fresh = stats["freshness"]
    _span(ws, r, 1, 12, f"Built {fresh['build_date']} · offline snapshot (latest recorded result {fresh['latest_result']}) · "
                        f"times {stats['tz']} · rosters = current captured rosters · SL = team-scope, format-specific · "
                        "colors/ranks = recorded evidence only, NOT CALIBRATED odds.", font=font(9, color="5B6A61"))
    h(r, 24)
    pages = [(1, r)]

    # ---------------- page 2: scouting cards, two across ----------------
    p2 = r + 1
    ws.row_breaks.append(Break(id=p2 - 1))
    base._section(ws, p2, "Scouting cards — the full cards are on the Scouting Cards sheet", last_col=12)
    h(p2, 18)
    r = p2
    card_fields = [("Vs our roster", "Vs Our Roster", 24), ("Met", "Met List", 36), ("By opponent SL", "By SL", 24),
                   ("Missing · notes", None, 24)]
    sides = ((1, 1, 2, 5), (7, 8, 9, 12))            # (header from, label from, value from, to)
    header_rows = []
    for band in range((R + 1) // 2):
        r += 1
        header_rows.append(r)
        h(r, 18)
        for side, (c0, cl, cv, c9) in enumerate(sides):
            k = band * 2 + side + 1
            if k > R:
                continue
            lbl = f"INDEX(wr_OppLabels,{k})"
            _span(ws, r, c0, c9, f'=IF({lbl}="","",{lbl}&" · SL "&INDEX(wr_OppSL,{k})&IF(INDEX(wr_OppPlayed,{k})="Played",'
                                 f'" · already played",""))', font=font(10.5, bold=True, color="3B2410"), wrap=False)
        for i, (label, col, pt) in enumerate(card_fields, start=1):
            rr = r + i
            h(rr, pt)
            for side, (c0, cl, cv, c9) in enumerate(sides):
                k = band * 2 + side + 1
                if k > R:
                    continue
                lbl, card = f"INDEX(wr_OppLabels,{k})", f"INDEX(wr_OppCard,{k})"
                _span(ws, rr, cl, cv - 1, f'=IF({lbl}="","","{label}")', font=font(9, bold=True, color="5B6A61"))
                value = (f'INDEX(Scouting_Table[{col}],{card})' if col else
                         f'INDEX(Scouting_Table[Missing],{card})&IF(INDEX(wr_OppNotes,{k})="",""," · Notes: "&INDEX(wr_OppNotes,{k}))')
                _span(ws, rr, cv, c9, f'=IF(OR({lbl}="",{card}=""),IF({lbl}="","","Not precomputed for this pairing"),{value})',
                      font=font(9.5))
        r += len(card_fields)
    for hr in header_rows:
        for c0, c9 in ((1, 5), (7, 12)):
            first, last = get_column_letter(c0), get_column_letter(c9)
            ws.conditional_formatting.add(f"{first}{hr}:{last}{hr}",
                                          FormulaRule(formula=[f'${first}${hr}<>""'], fill=_cf_fill(CARD_HEAD_FILL_RGB)))
    pages.append((p2, r))

    # ---------------- pages 3-4: evidence, packed, opponent named on every line ----------------
    p3 = r + 1
    ws.row_breaks.append(Break(id=p3 - 1))
    base._section(ws, p3, "Evidence by opponent — every line names the opponent", last_col=12)
    h(p3, 18)
    _span(ws, p3 + 1, 1, 12, "Rank: direct meetings first (by record, then more meetings) · “≈” = shared-opponent evidence "
                             "only, NOT ordered among themselves · “—” = no evidence · “2=” = tied (every row with the same "
                             "“n=” under one opponent is tied). Evidence: W-L (meetings) = direct record; ≈ ours vs theirs "
                             "(n shared) = records against shared opponents. Basis = the evidence category.",
          font=font(9, color="5B6A61"))
    h(p3 + 1, 26)
    hdr = p3 + 2
    for col, text in ((1, "Opponent (APA record ID)"), (3, "Rank"), (4, "Our player (APA record ID)"), (8, "Evidence"),
                      (11, "Basis")):
        cell = ws.cell(row=hdr, column=col, value=text)
        cell.font = font(9, bold=True, color=base.FELT_DEEP)
    for c in range(1, 13):
        ws.cell(row=hdr, column=c).fill = base.SUBHEAD_FILL
    h(hdr, 16)
    ev_first = hdr + 1
    for n in range(1, R * R + 1):
        r = ev_first + n - 1
        ws.cell(row=r, column=16, value=f'=IF({n}>wr_PkTotal,"",MATCH({n},wr_PkStart,1))').font = base.HELPER_FONT
        ws.cell(row=r, column=17, value=f'=IF(P{r}="","",{n}-INDEX(wr_PkStart,P{r})+1)').font = base.HELPER_FONT
        ws.cell(row=r, column=18, value=f'=IF(P{r}="","",INDEX(wr_OppStart,P{r})+Q{r})').font = base.HELPER_FONT
        f10 = font(10.5)
        _span(ws, r, 1, 2, f'=IF(P{r}="","","vs "&INDEX(wr_OppLabels,P{r}))', font=f10, wrap=False)
        ws.cell(row=r, column=3, value=f'=IF(P{r}="","",INDEX(MatchupEvidence_Table[Rank],R{r}))').font = f10
        _span(ws, r, 4, 7, f'=IF(P{r}="","",INDEX(MatchupEvidence_Table[Player],R{r}))', font=f10, wrap=False)
        _span(ws, r, 8, 10, f'=IF(P{r}="","",INDEX(MatchupEvidence_Table[Cell],R{r}))', font=f10, wrap=False)
        cat = f'INDEX(MatchupEvidence_Table[Category],R{r})'
        _span(ws, r, 11, 12, f'=IF(P{r}="","",IF({cat}="G","Favorable direct",IF({cat}="R","Concerning direct",'
                             f'IF({cat}="E","Even direct",IF({cat}="I","Indirect only","No evidence")))))', font=f10, wrap=False)
        h(r, 14.5)
    last_ev = ev_first + R * R - 1
    ws.conditional_formatting.add(f"A{ev_first}:L{last_ev}",
                                  FormulaRule(formula=[f'AND($P{ev_first}<>"",$Q{ev_first}=1)'], fill=_cf_fill(base.FELT_SOFT),
                                              font=Font(bold=True)))
    pages.append((p3, last_ev))

    # ---------------- page 5: meeting history ----------------
    # Paul (#84 UAT): every recorded meeting readable on page 5, larger text, clear row dividers. Excel without
    # macros cannot size rows to how many meetings a fixture has, so the page holds a fixed 40 slots as 20 rows x
    # 2 meetings (numbered newest first, filled row by row so short histories stay together at the top). Each
    # meeting is one row of two lines -- "n. date / result · SL" and "our player / vs opponent" (names + IDs).
    # A count line says how many meetings exist; any beyond 40 are on the Meetings sheet, never silently dropped.
    p5 = last_ev + 1
    ws.row_breaks.append(Break(id=p5 - 1))
    base._section(ws, p5, "Meeting history — direct meetings between the rosters (newest first)", last_col=12)
    h(p5, 18)
    total = 'IF(wr_PairKey="",0,COUNTIF(Meetings_Table[Pair Key],wr_PairKey))'
    _span(ws, p5 + 1, 1, 12, f'=IF({total}=0,"No recorded direct meetings between these rosters in this format.",'
                             f'"Showing "&MIN({total},{MEETING_SLOTS})&" of "&{total}&" recorded meeting(s)"&IF({total}>{MEETING_SLOTS},'
                             f'" — the older "&({total}-{MEETING_SLOTS})&" are listed on the Meetings sheet.","."))',
          font=font(10.5, bold=True, color=base.FELT_DEEP), wrap=False)
    h(p5 + 1, 16)
    hdr = p5 + 2
    blocks = ((1, 1, 2, 5), (7, 8, 9, 12))     # (first col, date/result to, names from, names to)
    for c0, cd, cn, c9 in blocks:
        _span(ws, hdr, c0, cd, "No. Date / Result · SL ours/theirs", font=font(9, bold=True, color=base.FELT_DEEP))
        _span(ws, hdr, cn, c9, "Our player (APA record ID) / vs Opponent (APA record ID)", font=font(9, bold=True, color=base.FELT_DEEP))
    for c in range(1, 13):
        ws.cell(row=hdr, column=c).fill = base.SUBHEAD_FILL
    h(hdr, 16)
    rows_per_side = MEETING_SLOTS // 2
    meet_first = hdr + 1
    for i in range(rows_per_side):
        r = meet_first + i
        for side, (c0, cd, cn, c9) in enumerate(blocks):
            k = i * 2 + side + 1
            row = f'IFERROR(MATCH(wr_PairKey&"|{k}",Meetings_Table[Key],0),"")'
            cond = f'OR(NOT(wr_HasEvidence),{row}="")'
            f12 = font(10.5)
            _span(ws, r, c0, cd, f'=IF({cond},"","{k}. "&INDEX(Meetings_Table[Date],{row})&CHAR(10)&INDEX(Meetings_Table[Result],{row})&'
                                 f'" · SL "&INDEX(Meetings_Table[Our SL],{row})&"/"&INDEX(Meetings_Table[Their SL],{row}))', font=f12)
            _span(ws, r, cn, c9, f'=IF({cond},"",INDEX(Meetings_Table[Our Player],{row})&CHAR(10)&"vs "&'
                                 f'INDEX(Meetings_Table[Opponent],{row}))', font=f12)
        h(r, 29)
    meet_last = meet_first + rows_per_side - 1
    # Dividers and subtle alternating shading on rows that hold a meeting; black rules print in black & white.
    from openpyxl.styles import Border, Side
    rule = Border(top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    for c0, c9 in ((1, 5), (7, 12)):
        a, b = get_column_letter(c0), get_column_letter(c9)
        rng = f"{a}{meet_first}:{b}{meet_last}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(${a}{meet_first}<>"",MOD(ROW(),2)=0)'], border=rule,
                                                       fill=_cf_fill("F2F2F2"), stopIfTrue=True))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${a}{meet_first}<>""'], border=rule))
    last = meet_last + 1
    _span(ws, last, 1, 12, '="Shared-opponent totals per pairing are in the evidence pages; the per-opponent breakdown is in the '
                           'HTML cockpit. This page always holds 40 slots (Excel cannot resize rows to a fixture without macros)."',
          font=font(9, color="5B6A61"))
    h(last, 24)
    pages.append((p5, last))

    # ---------------- one print scale for the whole sheet ----------------
    width_pt = sum(_width_pt(w) for w in PACKET_WIDTHS)
    title_pt = sum(PACKET_TITLE_ROWS)
    page_heights = []
    for i, (a, b) in enumerate(pages):
        body = sum(heights.get(x, 15) for x in range(a, b + 1))
        page_heights.append(body if i == 0 else body + title_pt)
    # Evidence (pages 3-4) is one section spread over two printed pages.
    budget = [ph / (2 if i == 2 else 1) for i, ph in enumerate(page_heights)]
    scale = min(1.0, PACKET_PRINTABLE[0] / width_pt, *(PACKET_PRINTABLE[1] / b for b in budget))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.page_setup.scale = max(10, int(scale * 100))
    ws.print_area = f"A1:L{last}"
    ws.print_title_rows = "1:2"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.page_margins.top = ws.page_margins.bottom = 0.45
    ws.page_margins.header = ws.page_margins.footer = 0.2
    _link(ws, 1, 14, "← Match Day", "Match Day")
    _link(ws, 2, 14, "← War Room", "War Room")
    return {"pages": pages, "page_heights": page_heights, "width_pt": width_pt, "scale": ws.page_setup.scale,
            "roster_first": roster_first, "evidence_first": ev_first}


# ---------------------------------------------------------------------------
# Coach Dashboard
# ---------------------------------------------------------------------------

def build_coach_dashboard(wb, *, format_options: list[str]) -> None:
    ws = wb.create_sheet("Coach Dashboard", 5)
    for letter, width in zip("ABCDEFG", (30, 6, 44, 18, 18, 18, 18)):
        ws.column_dimensions[letter].width = width
    _head(ws, "Coach Dashboard — Player vs Player", "Player A defaults to you; Player B comes from tonight's opponent roster "
                                                    "(or “All players…”). Nothing is preselected.", last=7, current="Coach Dashboard")
    _span(ws, 4, 1, 7, "=cd_Status", font=Font(bold=True, color=base.FELT_DEEP))
    inputs = [(5, "Player A (optional override)", "PlayerLabelList", "Blank = you (Match Day's player)."),
              (6, "Player B", "cd_PlayerBList", "Pick an opponent from tonight's roster, or “All players…” and choose below."),
              (7, "Any player (Player B = All players…)", "PlayerLabelList", "Only used when Player B is “All players…”."),
              (8, "Format (optional override)", '"' + ",".join(base._format_label(f) for f in format_options) + '"',
               "Blank = tonight's format.")]
    for r, label, source, help_text in inputs:
        _span(ws, r, 1, 2, label, font=base.LABEL_FONT)
        base._input(ws.cell(row=r, column=3))
        _dv(ws, f"C{r}", source, help_text)
        _span(ws, r, 4, 7, help_text, font=base.MUTED_FONT)
    base._section(ws, 10, "Effective comparison (calculated)", last_col=7)
    for r, label, formula in ((11, "Player A", '=IF(cd_ALabel="","—",cd_ALabel)'),
                              (12, "Player B", '=IF(cd_BLabel="","Choose Player B above.",IF(cd_BID="","⚠ “"&cd_BLabel&"” is not a verified player label.",cd_BLabel))'),
                              (13, "Format", "=cd_FmtLabel")):
        _span(ws, r, 1, 2, label, font=base.LABEL_FONT)
        _calc(ws.cell(row=r, column=3), formula, bold=True)
    fmt_code = base._format_code_formula(format_options, "cd_FmtLabel")[1:]
    ws["A15"] = "Pair key"
    ws["C15"] = f'=IF(OR(cd_AID="",cd_BID=""),"",cd_AID&"|"&{fmt_code}&"|"&cd_BID)'
    ws["A16"] = "Pair row in Player vs Player"
    for helper_row in (15, 16):          # internal lookup keys: kept for the formulas, hidden from the captain
        ws.row_dimensions[helper_row].hidden = True
    ws["C16"] = '=IF(C15="","",IFERROR(MATCH(C15,PlayerVsPlayer_Table[Pair Key],0),""))'
    for ref in ("A15", "C15", "A16", "C16"):
        ws[ref].font = base.HELPER_FONT
    base._section(ws, 18, "Direct record (Player A's perspective)", last_col=7)
    for r, label, formula in (
        (19, "Record", '=IF(OR(cd_AID="",cd_BID=""),"—",IF(C16="","No recorded direct meeting in this format",'
                       'INDEX(PlayerVsPlayer_Table[Wins],C16)&"-"&INDEX(PlayerVsPlayer_Table[Losses],C16)))'),
        (20, "Recorded meetings", '=IF(OR(cd_AID="",cd_BID=""),"—",IF(C16="","0 — absence of evidence, not a 0-0 tie",'
                                  'INDEX(PlayerVsPlayer_Table[Games],C16)))'),
        (21, "Captured SL tonight", '=IF(cd_ALabel="","A: —",IFERROR("A: "&INDEX(wr_OurSL,MATCH(cd_ALabel,wr_OurLabels,0)),"A: not on tonight’s roster"))&'
                                    '" · "&IF(cd_BLabel="","B: —",IFERROR("B: "&INDEX(wr_OppSL,MATCH(cd_BLabel,wr_OppLabels,0)),"B: not on tonight’s opponent roster"))'),
    ):
        _span(ws, r, 1, 2, label, font=base.LABEL_FONT)
        _calc(ws.cell(row=r, column=3), formula)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    _span(ws, 23, 1, 7, "Records only — no win percentage or odds are shown (none is calibrated). For shared-opponent evidence "
                        "between tonight's rosters, use the War Room's Inspect views.", font=base.MUTED_FONT, height=30)
    ws.freeze_panes = "A5"


# ---------------------------------------------------------------------------
# START HERE + Captain Command Center (Paul, "Make this a captain's weapon")
# ---------------------------------------------------------------------------

def _fit_height(text: Any, width_units: float, size: float, minimum: float) -> float:
    """Row height (points) that shows all of a wrapped literal text; formulas keep the given minimum."""
    if not isinstance(text, str) or text.startswith("="):
        return minimum
    per_line = max(10, int(width_units * 11 / size * 1.05))
    lines = max(1, -(-len(text) // per_line))
    return max(minimum, round(lines * size * 1.35 + 4, 1))


def _worst_height(text: str, width_units: float, size: float, minimum: float) -> float:
    """Row height (points) for the longest text a FORMULA cell can show. Excel never auto-fits formula rows,
    so this word-wraps the worst case (real-Excel UAT 2a67d78: clipped Next Send list and matrix cells)."""
    per_line = max(6, int(width_units * 10 / size))   # measured in real Excel: ~10 chars per 9 units at 9 pt
    lines = 0
    for paragraph in text.split("\n"):
        lines, used = lines + 1, 0
        for word in paragraph.split(" "):
            while len(word) > per_line:
                lines, used, word = lines + (used > 0), 0, word[per_line:]
            if used and used + 1 + len(word) > per_line:
                lines, used = lines + 1, len(word)
            else:
                used += (1 if used else 0) + len(word)
    return max(minimum, round(lines * size * 1.4 + 6, 1))


def _next_send_worst(label: str, roster: int, medals: int) -> str:
    """The longest Next Send answer with `medals` medal lines and the remaining players in the lists."""
    headline = ("Medals = ordered direct records among our remaining players (same evidence = same medal). Recorded "
                "results only — not odds.")
    medal = (f"🥇 {label} — 12-0 direct record (12 meetings) — favorable · tied (same evidence) · consider saving — "
             "our only favorable direct option vs another unplayed opponent · availability unknown")
    rest = max(roster - medals, 0)
    lines = [headline] + [medal] * medals + ["+ more direct candidates — see War Room Inspect"]
    if rest:
        lines.append("≈ Not ordered (shared-opponent results only): "
                     + "; ".join([f"{label} (availability unknown)"] * max(rest - 2, 1)))
        lines += [f"⚠ Avoid: {label} — 0-12 (12 meetings)", f"❓ Unknown (no evidence, not weak): {label}"]
    return "\n".join(lines)


def _longest_label(wb) -> str:
    """The longest player label in this workbook (Players sheet), or a long generic one when there is none."""
    labels = []
    if "Players" in wb.sheetnames:
        ws = wb["Players"]
        head = [c.value for c in ws[1]]
        if "Player Label" in head:
            col = head.index("Player Label") + 1
            labels = [str(v) for (v,) in ws.iter_rows(min_row=2, min_col=col, max_col=col, values_only=True) if v]
    return max(labels, key=len) if labels else WORST_LABEL


# A long but realistic player label, and the longest matrix evidence text, for sizing formula rows.
WORST_LABEL = "Christopher Fitzgerald (APA record ID 3487149)"
WORST_EVIDENCE = "≈ 41-53 vs 81-58 (61 shared)"
WORST_BASIS = ("Shared-opponent results only (no direct meetings) — not ordered against other indirect candidates; "
               "compare ours vs theirs")


def _card(ws, top: int, c1: int, c2: int, title: str, lines: list[Any], *, size: float = 11, title_fill=None,
          line_height: float = 18) -> int:
    """A titled card: a coloured title bar and plain lines underneath, framed by a light border. A line given as
    (value, worst_case_text) is sized for that text, since a formula line's height can't be fitted from itself."""
    from openpyxl.styles import Border, Side
    edge = Side(style="thin", color="CFC6B4")
    _span(ws, top, c1, c2, title, font=Font(bold=True, size=12, color="FFFFFF"), fill=title_fill or base.SECTION_FILL,
          wrap=False, height=22)
    r = top
    width = sum(ws.column_dimensions[get_column_letter(c)].width or 8.43 for c in range(c1, c2 + 1))
    for line in lines:
        r += 1
        line, worst = line if isinstance(line, tuple) else (line, None)
        floor = _worst_height(worst, width, size, line_height) if worst else line_height
        _span(ws, r, c1, c2, line, font=Font(size=size),
              height=max(_fit_height(line, width, size, floor), ws.row_dimensions[r].height or 0))
    for rr in range(top, r + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=rr, column=c)
            cell.border = Border(left=edge if c == c1 else None, right=edge if c == c2 else None,
                                 top=edge if rr == top else None, bottom=edge if rr == r else None)
    return r


TAB_TOUR = [
    ("START HERE", "This page: what Ultimate Coach does and how to use it in three minutes."),
    ("Command Center", "Tonight at a glance: opponent, date, venue, availability counts, the opponent roster and the evidence summary."),
    ("Match Day", "THE control panel. Pick player → team → format → date → fixture once; every other tab follows it."),
    ("War Room", "The match-night overview: both rosters, best sends per opponent, dangerous opponents, avoid sends, open risks, the colour matrix, and Inspect (shared-opponent evidence for one opponent or one of our players). Its own optional overrides only change this tab."),
    ("Lineup Lab", "Mark our players Available / Unavailable / Unknown and Planned / Played, and which opponents already played. Marks belong to one fixture."),
    ("Scouting Cards", "One card per opponent: records, meetings with our players, record by skill level, missing information and your observations."),
    ("Coach Notes", "Your own observations per player (tags such as 'Slow shooter' plus free text). Opinions, not APA facts — kept apart from evidence."),
    ("Captain Packet", "The printable match-night packet: page 1 summary, then scouting cards, evidence and meeting history."),
    ("Coach Dashboard", "One player vs one player: their direct record and number of meetings, plus tonight's captured skill levels. Defaults to you vs tonight's opponents. (Shared-opponent evidence is in War Room → Inspect.)"),
    ("Schedule, Team Rosters, Players …", "The recorded data behind everything above — reference only."),
]


def build_start_here(wb, *, stats: dict[str, Any], version: str, example: list[str] | None = None) -> None:
    ws = wb.create_sheet("START HERE", 0)
    for letter, width in zip("ABCDEFGHIJ", (3, 26, 26, 3, 26, 26, 3, 26, 26, 3)):
        ws.column_dimensions[letter].width = width
    ws.sheet_view.showGridLines = False
    _span(ws, 1, 2, 9, "Ultimate Coach — Captain's War Room", font=Font(bold=True, size=22, color=base.FELT_DEEP), wrap=False,
          height=36)
    _span(ws, 2, 2, 9, "Who should I put up next? Recorded, identity-verified APA evidence — organised for match night.",
          font=Font(italic=True, size=12, color="5B6A61"), height=20)
    fresh = stats["freshness"]
    _span(ws, 3, 2, 9, f"Workbook version {version} · built {fresh['build_date']} · data current to the latest recorded result "
                       f"{fresh['latest_result']} · times in {stats['tz']} · this file never refreshes itself.",
          font=Font(size=10, color="5B6A61"), height=18)
    top = 5
    a = _card(ws, top, 2, 3, "1 · What Ultimate Coach does", list(ONBOARDING_WHAT))
    steps = ["Step 1 · Go to Match Day", "Step 2 · Select Player (you)", "Step 3 · Select Team (and Format)",
             "Step 4 · Select Match Date (and Fixture if two share a day)", "Step 5 · Review the Command Center",
             "Step 6 · Lineup Lab: mark who's here — marks belong to that one fixture", "Step 7 · Review the War Room matchups",
             "Step 8 · Print the Captain Packet"]
    b = _card(ws, top, 5, 6, "2 · Quick start (3 minutes)", steps, title_fill=base.OPPONENT_FILL)
    flow = ["Match Day decides the fixture.", "↓ Command Center · War Room · Lineup Lab", "↓ Scouting Cards · Coach Dashboard",
            "↓ Captain Packet (print)", "Change Match Day and every tab follows.",
            "A tab's own override changes only that tab — clear it to follow Match Day again."]
    c = _card(ws, top, 8, 9, "4 · Match night workflow", flow)
    r = max(a, b, c) + 2
    # A worked example taken from THIS build's own data (the configured viewer's next fixture) -- selections
    # only, never results, and never hard-coded in the source.
    r = _card(ws, r, 2, 9, "Worked example from this build — what Match Day picks for you",
              example or ["No viewer is configured for this build: pick yourself on Match Day and it fills in team, date "
                          "and fixture the same way."], title_fill=base.OPPONENT_FILL) + 2
    _span(ws, r, 2, 9, "3 · Workbook tour", font=Font(bold=True, size=12, color="FFFFFF"), fill=base.SECTION_FILL, wrap=False, height=22)
    for name, text in TAB_TOUR:
        r += 1
        cell = ws.cell(row=r, column=2, value=name)
        cell.font = Font(bold=True, size=11, color=base.FELT_DEEP)
        if name in wb.sheetnames or name in ("Command Center", "Coach Notes"):
            cell.hyperlink = Hyperlink(ref=f"B{r}", location=f"'{name}'!A1", display=name)
            cell.font = Font(bold=True, size=11, color="1F5C99", underline="single")
        cell.alignment = base.WRAP_TOP
        _span(ws, r, 3, 9, text, font=Font(size=11),
              height=max(_fit_height(text, sum(ws.column_dimensions[c].width for c in "CDEFGHI"), 11, 20),
                         _fit_height(name, ws.column_dimensions["B"].width, 11, 20)))
    r += 2
    lim = _card(ws, r, 2, 9, "5 · Important limitations", list(ONBOARDING_LIMITS), title_fill=PatternFill("solid", fgColor="8A5A00"))
    r = _card(ws, lim + 2, 2, 9, "6 · Mobile match night (phone)", list(MATCH_NIGHT_GUIDE)) + 2
    _card(ws, r, 2, 9, "7 · Build information",
          [f"Workbook version: {version}", f"Build date: {fresh['build_date']}",
           f"Data freshness: latest recorded result {fresh['latest_result']}"
           + (f" · {fresh['unplayed_before_build']} earlier fixtures still show UNPLAYED in this snapshot"
              if fresh["unplayed_before_build"] else "")])
    _link(ws, 4, 2, "→ Start: Match Day", "Match Day", font=BIG_LINK)
    _link(ws, 4, 5, "→ Command Center", "Command Center", font=BIG_LINK)
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_command_center(wb, *, slots: dict[str, int], stats: dict[str, Any]) -> None:
    """Tonight in 30 seconds. Reads the Match Day pairing (md_* after split_match_day_engine)."""
    R = slots["roster"]
    ws = wb.create_sheet("Command Center", 1)
    for letter, width in zip("ABCDEFGHIJKL", (3, 30, 12, 14, 3, 30, 12, 14, 3, 34, 14, 3)):
        ws.column_dimensions[letter].width = width
    ws.sheet_view.showGridLines = False
    _span(ws, 1, 2, 11, "Captain Command Center — tonight at a glance", font=Font(bold=True, size=20, color=base.FELT_DEEP),
          wrap=False, height=32)
    _span(ws, 2, 2, 11, '=IF(uc_Fixture="","Choose a fixture on Match Day (player → team → format → date → fixture).",uc_Fixture)',
          font=Font(bold=True, size=14, color=base.FELT_DEEP), height=24)
    _span(ws, 3, 2, 11, '=IF(wr_OppLabel="",uc_OppMsg,"Opponent: "&wr_OppLabel&" · "&uc_Venue)', font=Font(size=11, color="5B6A61"),
          height=34)
    fresh = stats["freshness"]
    _span(ws, 4, 2, 11, f"Data freshness: built {fresh['build_date']} · latest recorded result {fresh['latest_result']} · "
                        "the file never refreshes itself.", font=Font(size=10, color="5B6A61"), height=16)
    _link(ws, 5, 2, "→ Match Day (change matchup)", "Match Day")
    _link(ws, 5, 6, "→ Lineup Lab (mark availability)", "Lineup Lab")
    _link(ws, 5, 10, "→ Captain Packet (print)", "Captain Packet")
    # Decision first: who do I send against the player they just put up? (Engine "Next Send", mirrors next_send.)
    _span(ws, 6, 2, 2, "They put up:", font=Font(bold=True, size=12, color=base.FELT_DEEP))
    base._input(ws.cell(row=6, column=3))
    ws.merge_cells(start_row=6, start_column=3, end_row=6, end_column=8)
    _dv(ws, "C6", "wr_OppLabels", "Pick the opponent player they just put up.")
    _span(ws, 6, 10, 11, '=IF(wr_OppLabel="","","Mark sends Played on Lineup Lab — they drop out here.")',
          font=Font(size=10, color="5B6A61"), height=22)
    # The whole answer is ONE wrapped cell (wr_NsCard: headline, medals, "+ more", then the lists), so it reads top
    # to bottom with no empty medal rows in between. Each remaining player appears at most once in it, so size the
    # cell for the worst mix -- k medal lines plus the other R-k names in the lists -- with this workbook's
    # longest label, or Excel hides the last names.
    worst = max((_next_send_worst(_longest_label(wb), R, k) for k in range(0, 5)),
                key=lambda t: _worst_height(t, 166, 11, 0))
    ns_lines = [
        ("=wr_NsCard", worst),
        '=IF(wr_NsJ="","",IF(INDEX(wr_OppNotes,wr_NsJ)="","","📝 "&INDEX(wr_OppNotes,wr_NsJ)&" (your opinion, not APA facts)"))',
    ]
    top = _card(ws, 7, 2, 11, "WHO SHOULD I SEND NEXT?", ns_lines, size=11, line_height=30) + 2
    cnt = lambda v: f'COUNTIF(wr_OurAvail,"{v}")'
    _card(ws, top, 2, 4, "MY TEAM", [
        '=IF(wr_OurLabel="","Choose your team on Match Day.",wr_OurLabel)',
        f'=IF(wr_OurCount=0,"","Available: "&{cnt("Available")})',
        f'=IF(wr_OurCount=0,"","Unavailable: "&{cnt("Unavailable")})',
        f'=IF(wr_OurCount=0,"","Unknown: "&{cnt("Unknown")}&" (not the same as unavailable)")',
        '=IF(wr_OurCount=0,"","Already used: "&COUNTIF(wr_OurLineup,"Played")&" · planned: "&COUNTIF(wr_OurLineup,"Planned"))',
        '=IF(wr_OurCount=0,"","Remaining skill total: "&IF(wr_RemainingMissing=0,wr_RemainingSL,wr_RemainingSL&" known ('
        '"&wr_RemainingMissing&" without SL)"))',
    ])
    opp_lines = ['=IF(wr_OppLabel="","No opponent roster for the selected fixture.",wr_OppLabel)']
    for k in range(1, R + 1):
        opp_lines.append(f'=IF(INDEX(wr_OppLabels,{k})="","",INDEX(wr_OppLabels,{k})&" · SL "&INDEX(wr_OppSL,{k})&'
                         f'IF(INDEX(wr_OppPlayed,{k})="Played"," · played",""))')
    opp_end = _card(ws, top, 6, 8, "OPPONENT TEAM", opp_lines, size=10, title_fill=base.OPPONENT_FILL, line_height=16)
    _span(ws, opp_end + 1, 6, 8, '=IF(wr_OppLabel="","","Missing information: "&COUNTIF(wr_OppSL,"No data")&" player(s) without a '
                                 'captured SL · "&wr_UnplayedCount&" not yet played")', font=Font(size=10, color="5B6A61"), height=30)
    cats = lambda code: "+".join(f'COUNTIF(wr_CatCol{j},"{code}")' for j in range(1, R + 1))
    summary = [
        '=IF(wr_PairKey="","Evidence summary appears once Match Day names a fixture with an opponent roster.",'
        '"Pairings between the two rosters, by recorded evidence:")',
        f'=IF(wr_PairKey="","","Favorable direct record (any sample size): "&({cats("G")}))',
        f'=IF(wr_PairKey="","","Concerning (more direct losses than wins): "&({cats("R")}))',
        f'=IF(wr_PairKey="","","Limited evidence: "&({cats("E")})&" even direct · "&({cats("I")})&" shared-opponent only")',
        f'=IF(wr_PairKey="","","Insufficient evidence (nothing recorded): "&({cats("X")}))',
        '=IF(wr_PairKey="","","Open risks: "&wr_RiskCount&" unplayed opponent(s) with no favorable direct option left")',
    ]
    _card(ws, top, 10, 11, "COACHING SUMMARY", summary, size=10.5, line_height=30)
    r = max(opp_end + 2, top + len(summary) + 2) + 1
    send_lines = []
    for k in range(1, R + 1):
        send_lines.append(f'=IF(OR(INDEX(wr_OppLabels,{k})="",NOT(INDEX(wr_OppUnplayed,{k}))),"","vs "&INDEX(wr_OppLabels,{k})&": "&'
                          f'IF(INDEX(wr_BestText,{k})="","no evidence-backed option left",INDEX(wr_BestText,{k})))')
    _card(ws, r, 2, 11, "BEST-SUPPORTED REMAINING SEND PER UNPLAYED OPPONENT — and the recorded evidence behind it",
          send_lines, size=10.5, line_height=30)
    ws.freeze_panes = "A7"
    ws.page_setup.orientation = "landscape"
    # Fit the WIDTH only: squeezing the whole sheet onto one page shrank it below readable size (real-Excel print
    # preview, 0c2e474). Extra pages are fine; the Captain Packet is the one-page-per-topic print.
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


COACH_NOTE_ROWS = 200


def build_coach_notes(wb) -> None:
    """Durable coach observations per player -- opinions, kept apart from evidence. The first row naming a
    player feeds that player's scouting cards and packet ("Coach: ..."); APA records are never changed."""
    ws = wb.create_sheet("Coach Notes")
    for letter, width in zip("ABCDEFG", (44, 24, 24, 60, 14, 3, 70)):
        ws.column_dimensions[letter].width = width
    _span(ws, 1, 1, 5, "Coach Notes — your observations (opinions, not APA facts)", font=Font(bold=True, size=18, color=base.FELT_DEEP),
          wrap=False, height=30)
    _span(ws, 2, 1, 5, "One row per player. Pick the player (name + APA record ID), up to two tags, and write what you saw. "
                       "These notes appear on that player's scouting card and in the Captain Packet, always marked “Coach:”. "
                       "They never change any recorded result.", font=Font(size=10, color="5B6A61"), height=30)
    _link(ws, 1, 7, "← START HERE", "START HERE")
    heads = ["Player (APA record ID)", "Tag 1", "Tag 2", "Observation", "Date noted"]
    for c, text in enumerate(heads, start=1):
        cell = ws.cell(row=4, column=c, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = base.SECTION_FILL
    ws.cell(row=4, column=7, value="Summary used on cards (calculated)").font = base.HELPER_FONT
    first, last = 5, 4 + COACH_NOTE_ROWS
    player = DataValidation(type="list", formula1="PlayerLabelList", allow_blank=True)
    tags = DataValidation(type="list", formula1='"' + ",".join(COACH_TAGS) + '"', allow_blank=True)
    ws.add_data_validation(player)
    ws.add_data_validation(tags)
    for r in range(first, last + 1):
        for c in range(1, 6):
            base._input(ws.cell(row=r, column=c))
        player.add(f"A{r}")
        tags.add(f"B{r}")
        tags.add(f"C{r}")
        ws.cell(row=r, column=4).alignment = base.WRAP_TOP
        ws.cell(row=r, column=7, value=f'=IF(A{r}="","",TRIM(B{r}&IF(AND(B{r}<>"",C{r}<>"")," · ","")&C{r}&'
                                       f'IF(AND(OR(B{r}<>"",C{r}<>""),D{r}<>""),": ","")&D{r}))').font = base.HELPER_FONT
    wb.defined_names["cn_Player"] = DefinedName("cn_Player", attr_text=f"'Coach Notes'!$A${first}:$A${last}")
    wb.defined_names["cn_Summary"] = DefinedName("cn_Summary", attr_text=f"'Coach Notes'!$G${first}:$G${last}")
    ws.freeze_panes = "A5"
