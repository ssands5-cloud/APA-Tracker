"""Full Production Demo Builder: docs/full_production_demo_builder.md.

Turns one authorized data source into one verified, immutable demo bundle.
This is a COORDINATION boundary: it calls the existing acquisition,
analytics, and renderer entry points and copies none of their logic. Every
number in the bundle is computed by the module that owns it.

Phases, in order, each stopping the build on failure:

    1. preflight   canonical repository root/boundary id, argument
                   exclusivity, contained and new output directory
    2. acquire     fixture: rebuild data/demo_coherent.db through
                   scripts/build_coherent_demo.py
                   live:    require APA_ACCESS_TOKEN and stage into
                            data/apa_tracker_regenerated.db
    3. lock        record the database SHA-256 and open it read-only; a hash
                   change during the build fails the run
    4. documents   run every analytics builder against that one database and
                   one scope, into a run-scoped temporary staging directory
    5. render      place staged artifacts into html/ and excel/ and write a
                   fresh unified index.html -- never the stale
                   exports/demo_dashboard.html
    6. verify      reconcile pair keys, evidence counts, lineup counts,
                   roster identities, scope, and artifact shape
    7. finalize    write manifest and checksums, re-read and verify them from
                   disk, then write READY last

A failed phase leaves no READY marker, so no partial run is ever
promotable or presentable. The builder never opens a browser or serves
anything; presentation belongs to scripts/run_production_demo.py.

CLI note: this exposes `--mode fixture|live` as the operator-facing source
selector. docs/full_production_demo_builder.md sketches `--fixtures PATH` /
`--live` instead; the mode flag is the form this project asked for, and the
document's actual contract -- phase order, snapshot lock, manifest content,
checksums, READY-last, containment, exit codes -- is implemented as written.

Usage:
    python scripts/build_full_production_demo.py --mode fixture
    python scripts/build_full_production_demo.py --mode live --promote
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import os
import shutil
import sqlite3
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Optional

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from analytics.pairing_evidence import build_pairing_evidence_matrix
from database.models import Match
from database.queries import CanonicalRosterError, canonical_current_roster
from scripts.build_captains_edge import connect_read_only

logger = logging.getLogger("build_full_production_demo")

PROJECT_ROOT = Path(__file__).resolve().parent.parent
BOUNDARY_FILE = PROJECT_ROOT / ".repo-boundary-id"
BOUNDARY_ID = "apa-tracker"
DEFAULT_RUN_ROOT = PROJECT_ROOT / "demo-runs"

MANIFEST_SCHEMA_VERSION = "demo-manifest-v1"
EVENT_SCHEMA_VERSION = "demo-event-v1"
MANIFEST_NAME = "demo_manifest.json"
CHECKSUMS_NAME = "checksums.sha256"
READY_NAME = "READY"
INDEX_NAME = "index.html"

FIXTURE_DB = PROJECT_ROOT / "data" / "demo_coherent.db"
LIVE_STAGING_DB = PROJECT_ROOT / "data" / "apa_tracker_regenerated.db"
LIVE_PRODUCTION_DB = PROJECT_ROOT / "data" / "apa_tracker.db"
TOKEN_ENV = "APA_ACCESS_TOKEN"

# The one coherent rehearsal scope scripts/build_coherent_demo.py creates.
FIXTURE_SCOPE = {
    "our_team_id": "90301",
    "opponent_team_id": "90302",
    "session_name": "2026 Rehearsal Session",
    "format": "8-Ball Open",
}
# Fixture-mode-only invariants. Live data has its own real shape and must
# never be asserted against these.
FIXTURE_EXPECTATIONS = {
    "our_roster": 5,
    "opponent_roster": 5,
    "total_feasible_pairings": 25,
    "DIRECT": 5,
    "INDIRECT": 20,
    "UNKNOWN": 0,
}

HTML_ARTIFACTS = {
    "team_strength.html",
    "season_projection.html",
    "trend_analyzer.html",
    "opponent_volatility.html",
    "player_vs_player.html",
    "player_matchup_explorer.html",
    "data_coverage.html",
    "captains_edge.html",
    "captain_first_edge.html",
}
EXCEL_ARTIFACTS = {
    "team_strength.xlsx",
    "season_projection.xlsx",
    "trend_analyzer.xlsx",
    "opponent_volatility.xlsx",
    "player_vs_player.xlsx",
    "data_coverage.xlsx",
    "captains_edge.xlsx",
}
JSON_ARTIFACTS = {"captains_edge.json", "lineups.json"}

# Declared-but-unimplemented demo features. The manifest and index say so in
# those words and link nothing, rather than shipping a plausible placeholder
# artifact (docs/full_production_demo_builder.md: the manifest "may mark them
# not_implemented; it may not emit a plausible placeholder artifact").
FEATURE_STATUS = {
    "team_strength": "available",
    "season_projection": "available",
    "trend_analyzer": "available",
    "opponent_volatility": "available",
    "player_vs_player": "available",
    "player_matchup_explorer": "available",
    "data_coverage": "available",
    "captain_first_edge": "available",
    "lineup_lab": "available",
    "captains_edge": "available",
    "match_difficulty_heatmap": "not_implemented",
    "captains_live_assistant": "not_implemented",
}

EXIT_OK = 0
EXIT_PREFLIGHT = 2
EXIT_AUTH = 3
EXIT_ACQUIRE = 4
EXIT_DOCUMENTS = 10
EXIT_RENDER = 11
EXIT_FINALIZE = 12
EXIT_INTERRUPT = 130


class BuildError(RuntimeError):
    """A phase failure carrying the exit category it must terminate with."""

    def __init__(self, message: str, code: int):
        super().__init__(message)
        self.code = code


@dataclass
class PhaseResult:
    name: str
    status: str
    detail: str = ""
    artifacts: list[str] = field(default_factory=list)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def resolve_contained(path: Path, root: Path) -> Path:
    """Resolve ``path`` and require it to stay under ``root``.

    Symlinks are resolved first, so a link pointing outside the repository
    is rejected on its target rather than its name.
    """
    resolved = Path(path).resolve()
    root_resolved = Path(root).resolve()
    if resolved != root_resolved and root_resolved not in resolved.parents:
        raise BuildError(
            f"path escapes the repository boundary: {path}", EXIT_PREFLIGHT
        )
    return resolved


def _git_revision() -> Optional[str]:
    head = PROJECT_ROOT / ".git" / "HEAD"
    if not head.is_file():
        return None
    try:
        content = head.read_text(encoding="utf-8").strip()
        if content.startswith("ref: "):
            ref = PROJECT_ROOT / ".git" / content[5:]
            return ref.read_text(encoding="utf-8").strip() if ref.is_file() else None
        return content
    except OSError:
        return None


def preflight(run_dir: Path, run_root: Path) -> PhaseResult:
    if not BOUNDARY_FILE.is_file() or BOUNDARY_FILE.read_text(encoding="utf-8").strip() != BOUNDARY_ID:
        raise BuildError(
            f"not the canonical APA-Tracker repository: {BOUNDARY_FILE} must contain {BOUNDARY_ID!r}",
            EXIT_PREFLIGHT,
        )
    resolve_contained(run_root, PROJECT_ROOT)
    resolved_run = resolve_contained(run_dir, run_root)
    if resolved_run.exists() and any(resolved_run.iterdir()):
        raise BuildError(f"run directory is not empty: {resolved_run}", EXIT_PREFLIGHT)
    return PhaseResult("preflight", "ok", f"run directory {resolved_run.name}")


def acquire_fixture() -> tuple[Path, PhaseResult]:
    """Rebuild the coherent rehearsal database through its own builder."""
    from scripts.build_coherent_demo import build as build_coherent

    db_file = build_coherent(str(FIXTURE_DB))
    if not db_file.is_file() or db_file.stat().st_size == 0:
        raise BuildError(f"fixture database was not produced: {db_file}", EXIT_ACQUIRE)
    return db_file, PhaseResult("acquire", "ok", "rebuilt the coherent fixture database")


def acquire_live(resume: bool = False) -> tuple[Path, PhaseResult]:
    """Stage a fresh authenticated scrape into the staging database.

    Requires a bearer token in the environment. There is deliberately no
    username/password path here: that mechanism belongs to a separate frozen
    scraper contract and is never used as a fallback. The token value is read
    directly by the scraper layer and is never logged, echoed, or written to
    the manifest.

    Ingests every accessible team, current roster, scheduled match, and
    completed scoresheet in each division the account's own teams belong to
    (scheduler.graphql_sync.run_division_wide) -- not only the account's own
    teams, which is as far as run_all_teams alone goes. A non-empty
    coverage_gaps result is a hard refusal: this staging database is never
    handed to the rest of the build, let alone promoted.

    ``resume=True`` preserves an existing staging database and continues an
    interrupted acquisition. A normal non-resume game-night refresh now
    starts by copying the verified production database into staging when
    production exists. That production DB is the permanent career archive:
    current GraphQL data is applied on top and run_division_wide receives
    resume=True so already-captured scoresheets are skipped. This keeps
    future refreshes short and prevents a current-season-only rebuild from
    erasing historical career rows.

    If no production database exists yet, non-resume mode still starts from
    an empty staging database and performs the original full current sync.
    A real APA access token is short-lived, so explicit --resume remains
    available for interrupted first-time acquisitions.
    """
    if not os.environ.get(TOKEN_ENV):
        raise BuildError(
            f"{TOKEN_ENV} is not set. Complete the APA /authorize consent flow, "
            f"then set {TOKEN_ENV} in this shell and re-run. Live mode never "
            "falls back to username/password or the legacy scraper.",
            EXIT_AUTH,
        )

    from scheduler.graphql_sync import run_division_wide

    sync_resume = resume
    seeded_from_production = False
    if not resume:
        if LIVE_STAGING_DB.exists():
            LIVE_STAGING_DB.unlink()
        if LIVE_PRODUCTION_DB.is_file() and LIVE_PRODUCTION_DB.stat().st_size:
            LIVE_STAGING_DB.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(LIVE_PRODUCTION_DB, LIVE_STAGING_DB)
            seeded_from_production = True
            sync_resume = True

    try:
        result = run_division_wide(
            str(PROJECT_ROOT / "apa_config.yaml"), export=False,
            db_path=str(LIVE_STAGING_DB), resume=sync_resume,
        )
    except Exception as exc:  # noqa: BLE001 - category, never the payload
        raise BuildError(f"live acquisition failed: {type(exc).__name__}", EXIT_ACQUIRE) from None

    if result["coverage_gaps"]:
        raise BuildError(
            "division-wide coverage is incomplete, refusing to stage this run: "
            + "; ".join(result["coverage_gaps"]),
            EXIT_ACQUIRE,
        )

    if not LIVE_STAGING_DB.is_file() or LIVE_STAGING_DB.stat().st_size == 0:
        raise BuildError(
            f"live acquisition produced no staging database at {LIVE_STAGING_DB}",
            EXIT_ACQUIRE,
        )

    # player_h2h_advantage (what scripts/build_lineups.py and Captain's
    # Edge's Decision Engine section both read) is populated by a SEPARATE
    # step, scripts/build_head_to_head.py, that neither run_all_teams() nor
    # run_division_wide() calls -- they only rebuild player_matchups (via
    # analytics.matchup_builder.build_matchups). build_coherent_demo.py
    # already does this for the fixture path; a real live run surfaced that
    # nothing did it for live.
    from database.engine import create_db_engine
    from database.ingest import ingest_h2h_advantage
    from scripts.build_head_to_head import build_rows as build_h2h_advantage_rows

    # ingest_h2h_advantage writes, so this connection is opened writable --
    # not through connect_read_only's mode=ro like every analytics builder
    # elsewhere in this module.
    write_engine = create_db_engine({"database": {"path": str(LIVE_STAGING_DB)}})
    with Session(write_engine) as db:
        rows = build_h2h_advantage_rows(db)
        if rows:
            ingest_h2h_advantage(db, rows)
    write_engine.dispose()

    totals = result["division_wide"]
    return LIVE_STAGING_DB, PhaseResult(
        "acquire", "ok",
        ("incremental seed from career-complete production; " if seeded_from_production else "")
        + f"{totals['teams_ingested']} team(s), {totals['matches_ingested']} match(es), "
        f"{totals['scored_matches_with_scoresheet']} scoresheet(s); coverage complete",
    )


def acquire_existing(source_db: Path) -> tuple[Path, PhaseResult]:
    """Use an already-acquired database as-is: never touches the network,
    never requires a token, and never rescrapes.

    For building a fresh verified run from data that was already ingested
    and, separately, already repaired (e.g. by
    scripts/backfill_player_identity.py) -- when the goal is a new manifest
    and checksums over the CURRENT contents of that database, not a new
    scrape. The file is used read-only exactly like any other acquired
    database; nothing here mutates it.
    """
    if not source_db.is_file() or source_db.stat().st_size == 0:
        raise BuildError(f"no database at {source_db}", EXIT_ACQUIRE)
    return source_db, PhaseResult(
        "acquire", "ok", f"using already-acquired database {source_db.name} (no rescrape)"
    )


def build_documents(db: Session, db_path: Path, scope: dict, staging: Path) -> PhaseResult:
    """Run every analytics builder against one database and one scope."""
    from scripts import (
        build_captain_first_edge,
        build_captains_edge,
        build_data_coverage,
        build_lineups,
        build_opponent_volatility,
        build_player_matchup_explorer,
        build_player_vs_player_export,
        build_season_projection,
        build_team_strength,
        build_trend_analyzer,
    )

    our = scope["our_team_id"]
    opponent = scope["opponent_team_id"]
    session_name = scope["session_name"]
    format_name = scope["format"]

    staging.mkdir(parents=True, exist_ok=True)
    built: list[str] = []
    try:
        build_team_strength.build(db, our, session_name, staging)
        build_season_projection.build(db, our, staging)
        build_trend_analyzer.build(db, our, session_name, staging, format_name)
        build_opponent_volatility.build(db, opponent, session_name, staging, format_name)
        build_player_vs_player_export.build(db, our, opponent, format_name, session_name, staging)
        # Whole-session scope on purpose: any two captured players, including
        # two who are both on other teams.
        build_player_matchup_explorer.build(db, session_name, staging, format_=format_name)
        build_data_coverage.build(db, our, opponent, format_name, session_name, staging)
        build_captain_first_edge.build(db, our, staging)
    except Exception as exc:  # noqa: BLE001 - phase category
        raise BuildError(f"analytics document construction failed: {exc}", EXIT_DOCUMENTS) from exc

    # These two own their own read-only connections by design. Scoped to
    # this run's exact team pairing -- like every other builder above --
    # rather than the whole division: a captain reviewing this run's
    # artifacts should see this run's own opponent, not every other real
    # team pairing division-wide sync happens to have ingested.
    try:
        build_captains_edge.build(str(db_path), str(staging), scope=scope)
        build_lineups.build(str(db_path), str(staging), scope=scope)
    except Exception as exc:  # noqa: BLE001 - phase category
        raise BuildError(f"legacy document construction failed: {exc}", EXIT_DOCUMENTS) from exc

    built = sorted(p.name for p in staging.iterdir() if p.is_file())
    return PhaseResult("documents", "ok", f"{len(built)} staged artifact(s)", built)


def _reconciliation(db: Session, scope: dict) -> dict:
    """Authoritative counts, computed once from the analytics layer."""
    our_roster = canonical_current_roster(db, scope["our_team_id"], scope["session_name"])
    opponent_roster = canonical_current_roster(db, scope["opponent_team_id"], scope["session_name"])
    matrix = build_pairing_evidence_matrix(
        db,
        our_team_external_id=scope["our_team_id"],
        opponent_team_external_id=scope["opponent_team_id"],
        format=scope["format"],
        session_name=scope["session_name"],
    )
    counts = dict(matrix.counts)
    return {
        "our_roster_count": len(our_roster),
        "opponent_roster_count": len(opponent_roster),
        "our_roster_player_ids": sorted(r.player_id for r in our_roster),
        "opponent_roster_player_ids": sorted(r.player_id for r in opponent_roster),
        "total_feasible_pairings": counts.get("total_feasible_pairings", 0),
        "DIRECT": counts.get("DIRECT", 0),
        "INDIRECT": counts.get("INDIRECT", 0),
        "UNKNOWN": counts.get("UNKNOWN", 0),
    }


def _require_nonempty(directory: Path, expected: set[str], label: str) -> list[str]:
    """Exact inventory as non-empty regular files.

    Shares scripts/reproducible_build.py's validator rather than restating
    the same four checks, so both build paths agree on what "the artifacts
    are present" means.
    """
    from scripts.reproducible_build import validated_artifacts

    try:
        return validated_artifacts(directory, expected)
    except (OSError, ValueError) as exc:
        raise BuildError(f"{label} artifact verification failed: {exc}", EXIT_RENDER) from None


def render(
    staging: Path,
    run_dir: Path,
    db_path: Path,
    scope: dict,
    reconciliation: dict,
    mode: str,
    database_sha256: str,
) -> tuple[PhaseResult, dict]:
    """Place staged artifacts into the bundle and write a fresh index."""
    html_dir = run_dir / "html"
    excel_dir = run_dir / "excel"
    json_dir = run_dir / "json"
    data_dir = run_dir / "data"
    for directory in (html_dir, excel_dir, json_dir, data_dir):
        directory.mkdir(parents=True, exist_ok=True)

    for item in sorted(staging.iterdir()):
        if not item.is_file():
            continue
        if item.suffix == ".html":
            destination = html_dir
        elif item.suffix == ".xlsx":
            destination = excel_dir
        elif item.suffix == ".json":
            destination = json_dir
        else:
            continue
        shutil.copy2(item, destination / item.name)

    shutil.copy2(db_path, data_dir / db_path.name)

    html_present = _require_nonempty(html_dir, HTML_ARTIFACTS, "html")
    excel_present = _require_nonempty(excel_dir, EXCEL_ARTIFACTS, "excel")
    json_present = _require_nonempty(json_dir, JSON_ARTIFACTS, "json")

    index_path = run_dir / INDEX_NAME
    index_path.write_text(
        _render_index(
            scope, reconciliation, html_present, excel_present, json_present,
            mode=mode, run_id=run_dir.name, database_sha256=database_sha256,
        ),
        encoding="utf-8",
    )

    artifacts = (
        [f"html/{name}" for name in html_present]
        + [f"excel/{name}" for name in excel_present]
        + [f"json/{name}" for name in json_present]
        + [f"data/{db_path.name}", INDEX_NAME]
    )
    return PhaseResult("render", "ok", f"{len(artifacts)} artifact(s)", artifacts), {
        "html": html_present,
        "excel": excel_present,
        "json": json_present,
    }


def _render_index(
    scope: dict,
    reconciliation: dict,
    html_names,
    excel_names,
    json_names,
    mode: str,
    run_id: str,
    database_sha256: str,
) -> str:
    """A fresh unified entrypoint linking every artifact in this run.

    Deliberately NOT exports/demo_dashboard.html, which is rendered from a
    different, older fixture set and links none of these.
    """
    def links(prefix: str, names) -> str:
        return "".join(
            f'<li><a href="{escape(prefix)}/{escape(name)}">{escape(name)}</a></li>'
            for name in names
        ) or "<li>None.</li>"

    source_label = (
        "VERIFIED FIXTURE (rehearsal data -- not live APA production evidence)"
        if mode == "fixture"
        else "LIVE APA DATA"
    )
    built_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%SZ")
    unimplemented = "".join(
        f"<li>{escape(name.replace('_', ' ').title())}: not implemented</li>"
        for name, status in sorted(FEATURE_STATUS.items())
        if status != "available"
    )

    rows = "".join(
        f"<tr><th>{escape(str(key).replace('_', ' ').title())}</th>"
        f"<td>{escape(str(value))}</td></tr>"
        for key, value in (
            ("our team", scope["our_team_id"]),
            ("opponent team", scope["opponent_team_id"]),
            ("session", scope["session_name"]),
            ("format", scope["format"]),
            ("our roster", reconciliation["our_roster_count"]),
            ("opponent roster", reconciliation["opponent_roster_count"]),
            ("feasible pairings", reconciliation["total_feasible_pairings"]),
            ("direct", reconciliation["DIRECT"]),
            ("indirect", reconciliation["INDIRECT"]),
            ("unknown", reconciliation["UNKNOWN"]),
        )
    )

    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<title>APA Tracker -- Demo Run</title>
<style>
body {{ font-family: system-ui, sans-serif; margin: 24px; color: #1c1f24; }}
.status-banner {{ border: 2px solid #1F3864; background: #eef2fa; padding: 12px 16px;
  margin: 0 0 18px; font-size: 14px; }}
.status-banner b {{ font-size: 15px; }}
table {{ border-collapse: collapse; }}
th, td {{ padding: 6px 10px; border-bottom: 1px solid #e2e5ea; text-align: left; }}
.unimplemented {{ color: #666e7a; }}
</style></head><body>
<h1>APA Tracker demo run</h1>

<div class="status-banner">
<b>Source: {escape(source_label)}</b><br>
Run <code>{escape(run_id)}</code> &middot; built {escape(built_at)} &middot;
database SHA-256 <code>{escape(database_sha256[:16])}&hellip;</code><br>
Status: verified -- manifest, checksums, and READY all written after every gate passed.
</div>

<p>Every artifact below was built from one database and one scope in a single
verified run. Counts come from the analytics layer, not from these pages.</p>

<h2>Scope and reconciliation</h2>
<table border="1" cellpadding="6" cellspacing="0"><tbody>{rows}</tbody></table>

<h2>Reports</h2>
<ul>{links("html", html_names)}</ul>

<h2>Workbooks</h2>
<ul>{links("excel", excel_names)}</ul>

<h2>JSON</h2>
<ul>{links("json", json_names)}</ul>

<h2>Not implemented</h2>
<p>These are declared demo features that do not exist yet. Nothing is linked
for them and no placeholder artifact was produced.</p>
<ul class="unimplemented">{unimplemented}</ul>

<p><small>Built by scripts/build_full_production_demo.py. This index replaces
the older exports/demo_dashboard.html, which is rendered from a different
fixture set and links none of these artifacts.</small></p>
</body></html>"""


def verify(run_dir: Path, db_path: Path, locked_hash: str, mode: str, reconciliation: dict) -> PhaseResult:
    """Reconcile the bundle against the authoritative analytics figures."""
    problems: list[str] = []

    if sha256_file(db_path) != locked_hash:
        problems.append("database hash changed during the build")

    counts_sum = reconciliation["DIRECT"] + reconciliation["INDIRECT"] + reconciliation["UNKNOWN"]
    if counts_sum != reconciliation["total_feasible_pairings"]:
        problems.append(
            f"evidence counts {counts_sum} do not sum to "
            f"{reconciliation['total_feasible_pairings']} feasible pairings"
        )

    expected_pairings = reconciliation["our_roster_count"] * reconciliation["opponent_roster_count"]
    if expected_pairings != reconciliation["total_feasible_pairings"]:
        problems.append(
            f"{reconciliation['total_feasible_pairings']} feasible pairings does not match "
            f"the {expected_pairings}-cell roster cross-product"
        )

    if mode == "fixture":
        for key, expected in FIXTURE_EXPECTATIONS.items():
            actual = reconciliation.get(
                {"our_roster": "our_roster_count", "opponent_roster": "opponent_roster_count"}.get(key, key)
            )
            if actual != expected:
                problems.append(f"fixture expectation {key}={expected} but found {actual}")

    pair_rows = _workbook_row_count(run_dir / "excel" / "player_vs_player.xlsx", "Player_vs_Player")
    if pair_rows != reconciliation["total_feasible_pairings"]:
        problems.append(
            f"player_vs_player.xlsx has {pair_rows} pair row(s) but the matrix has "
            f"{reconciliation['total_feasible_pairings']}"
        )

    lineups_path = run_dir / "json" / "lineups.json"
    try:
        lineups = json.loads(lineups_path.read_text(encoding="utf-8"))
        assignments = sum(len(entry.get("assignments", [])) for entry in lineups.get("lineups", []))
        unavailable = lineups.get("lineups_unavailable", [])
        # Zero assignments is only a real bug when nothing explains it. A
        # declared unavailability (the run's own scope exceeded the exact-
        # search guard) is a real, reported state -- see
        # scripts/build_lineups.py's "available": False groups -- not a
        # silent failure this build should hide behind a hard error.
        if assignments == 0 and not unavailable:
            problems.append("lineups.json contains no lineup assignments")
    except (OSError, ValueError) as exc:
        problems.append(f"lineups.json unreadable: {exc}")

    index_html = (run_dir / INDEX_NAME).read_text(encoding="utf-8")
    for token in ("http://", "https://"):
        if token in index_html:
            problems.append(f"index.html references an external resource ({token})")

    if problems:
        raise BuildError("; ".join(problems), EXIT_RENDER)
    return PhaseResult("verify", "ok", f"{reconciliation['total_feasible_pairings']} pairings reconciled")


def _workbook_row_count(path: Path, sheet_name: str) -> int:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return -1
        return max(workbook[sheet_name].max_row - 1, 0)
    finally:
        workbook.close()


def finalize(run_dir: Path, manifest: dict) -> PhaseResult:
    """Manifest, then checksums, then re-read both, then READY last."""
    manifest_path = run_dir / MANIFEST_NAME
    checksums_path = run_dir / CHECKSUMS_NAME

    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")

    lines = []
    for relative in sorted(manifest["artifacts"]):
        target = run_dir / relative
        lines.append(f"{sha256_file(target)}  {relative}")
    lines.append(f"{sha256_file(manifest_path)}  {MANIFEST_NAME}")
    checksums_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    # Re-read both from disk and verify before anything is marked ready.
    reloaded = json.loads(manifest_path.read_text(encoding="utf-8"))
    if reloaded != manifest:
        raise BuildError("manifest did not survive the round trip to disk", EXIT_FINALIZE)
    for line in checksums_path.read_text(encoding="utf-8").splitlines():
        digest, relative = line.split("  ", 1)
        if sha256_file(run_dir / relative) != digest:
            raise BuildError(f"checksum mismatch for {relative}", EXIT_FINALIZE)

    (run_dir / READY_NAME).write_text(
        json.dumps(
            {
                "schema_version": MANIFEST_SCHEMA_VERSION,
                "run_id": manifest["run_id"],
                "manifest_sha256": sha256_file(manifest_path),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    return PhaseResult("finalize", "ok", "manifest, checksums, READY")


class EventSink:
    """Versioned, redacted JSONL events for automation.

    The launcher reads the completion event from here rather than parsing
    stdout, so a run directory is never inferred from log text. Only
    non-secret fields are ever written: phase, status, counts, and paths
    relative to the repository root.
    """

    def __init__(self, path: Optional[Path]):
        self.path = path
        if path is not None:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("", encoding="utf-8")

    def emit(self, phase: str, status: str, **fields) -> None:
        if self.path is None:
            return
        event = {
            "schema_version": EVENT_SCHEMA_VERSION,
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "phase": phase,
            "status": status,
            **fields,
        }
        with self.path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(event, sort_keys=True) + "\n")


def run_build(
    mode: str,
    run_dir: Path,
    run_root: Path,
    promote: bool = False,
    events: Optional[EventSink] = None,
    resume: bool = False,
    source_db: Optional[Path] = None,
    opponent_team_id: Optional[str] = None,
    session_name: Optional[str] = None,
    format_name: Optional[str] = None,
) -> Path:
    events = events or EventSink(None)
    phases: list[PhaseResult] = [preflight(run_dir, run_root)]
    events.emit("preflight", "ok")

    if source_db is not None:
        db_path, acquire_result = acquire_existing(source_db)
    else:
        db_path, acquire_result = acquire_fixture() if mode == "fixture" else acquire_live(resume=resume)
    phases.append(acquire_result)

    locked_hash = sha256_file(db_path)
    phases.append(PhaseResult("lock", "ok", f"database sha256 {locked_hash[:12]}"))

    run_dir.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix="staging-", dir=run_dir))

    engine = create_engine("sqlite://", creator=lambda: connect_read_only(db_path))
    db = Session(bind=engine)
    try:
        # Live scope must be chosen from what was actually ingested, so it
        # needs the open database -- fixture scope is a fixed constant and
        # needs nothing from it.
        scope = (
            dict(FIXTURE_SCOPE)
            if mode == "fixture"
            else _live_scope(db, opponent_team_id, session_name, format_name)
        )
        phases.append(build_documents(db, db_path, scope, staging))
        reconciliation = _reconciliation(db, scope)
    finally:
        db.close()
        engine.dispose()

    render_result, _inventory = render(
        staging, run_dir, db_path, scope, reconciliation,
        mode=mode, database_sha256=locked_hash,
    )
    phases.append(render_result)
    shutil.rmtree(staging, ignore_errors=True)

    phases.append(verify(run_dir, db_path, locked_hash, mode, reconciliation))

    manifest = {
        "schema_version": MANIFEST_SCHEMA_VERSION,
        "run_id": run_dir.name,
        "mode": mode,
        "built_at": datetime.now(timezone.utc).isoformat(),
        "repository_revision": _git_revision(),
        "database_sha256": locked_hash,
        "database_file": db_path.name,
        "scope": scope,
        "source_label": (
            "verified_fixture" if mode == "fixture" else "live"
        ),
        "reconciliation": reconciliation,
        "feature_status": dict(FEATURE_STATUS),
        "formula_versions": _formula_versions(),
        "artifacts": sorted(render_result.artifacts),
        "artifact_hashes": {
            relative: sha256_file(run_dir / relative) for relative in sorted(render_result.artifacts)
        },
        "phases": [
            {"name": p.name, "status": p.status, "detail": p.detail} for p in phases
        ],
        "promotable": True,
    }
    phases.append(finalize(run_dir, manifest))

    if promote and mode == "live":
        promote_live_database(db_path)

    events.emit(
        "complete",
        "ok",
        run_dir=str(run_dir.resolve().relative_to(PROJECT_ROOT).as_posix()),
        run_id=manifest["run_id"],
        manifest_sha256=sha256_file(run_dir / MANIFEST_NAME),
        artifact_count=len(manifest["artifacts"]),
    )
    return run_dir


def _live_scope(
    db: Session,
    opponent_team_id: Optional[str] = None,
    session_name: Optional[str] = None,
    format_name: Optional[str] = None,
) -> dict:
    """A real (opponent, session, format) scope, chosen from what was
    actually ingested -- never a fixture constant, and never a guess.

    With no override, candidates are every real, non-bye match our
    configured team appears in, earliest scheduled week first. A candidate
    is only usable when BOTH sides resolve a real canonical current roster
    (database.queries.canonical_current_roster) for that match's own
    session -- the same identity guard every analytics builder in this
    project already requires, so a scope this function picks is guaranteed
    buildable, not merely plausible. The first such candidate wins; ties
    are broken by match external id for a stable, reproducible choice.

    When ``opponent_team_id``, ``session_name`` and ``format_name`` are ALL
    given, that exact scope is used instead of auto-selection -- for
    pinning a specific, already-verified real matchup (e.g. one an earlier
    export was built and checked against) rather than whichever real match
    happens to sort first. The same roster-resolution guard still applies:
    a pinned scope with no real match, or with either side unable to
    resolve a canonical current roster, is refused rather than guessed.
    """
    from scripts.build_captain_first_edge import _configured_our_team_id

    our_team_id = _configured_our_team_id()
    if not our_team_id:
        raise BuildError(
            "live mode needs team.team_id in apa_config.yaml to select a scope",
            EXIT_PREFLIGHT,
        )

    overrides_given = (opponent_team_id, session_name, format_name)
    if any(overrides_given) and not all(overrides_given):
        raise BuildError(
            "--opponent-team-id, --session and --format must be given together, or not at all",
            EXIT_PREFLIGHT,
        )
    if all(overrides_given):
        match = (
            db.query(Match)
            .filter(
                Match.is_bye.is_(False),
                Match.session_name == session_name,
                Match.format == format_name,
                (
                    (Match.home_team_id == our_team_id) & (Match.away_team_id == opponent_team_id)
                )
                | (
                    (Match.away_team_id == our_team_id) & (Match.home_team_id == opponent_team_id)
                ),
            )
            .first()
        )
        if match is None:
            raise BuildError(
                f"no real match between team {our_team_id} and {opponent_team_id} "
                f"in session {session_name!r}, format {format_name!r}",
                EXIT_DOCUMENTS,
            )
        try:
            our_roster = canonical_current_roster(db, our_team_id, session_name)
            opponent_roster = canonical_current_roster(db, opponent_team_id, session_name)
        except CanonicalRosterError as exc:
            raise BuildError(f"pinned scope has no canonical current roster: {exc}", EXIT_DOCUMENTS) from None
        if not our_roster or not opponent_roster:
            raise BuildError("pinned scope resolved an empty canonical current roster", EXIT_DOCUMENTS)
        return {
            "our_team_id": our_team_id,
            "opponent_team_id": opponent_team_id,
            "session_name": session_name,
            "format": format_name,
        }

    candidates = (
        db.query(Match)
        .filter(
            Match.is_bye.is_(False),
            (Match.home_team_id == our_team_id) | (Match.away_team_id == our_team_id),
        )
        .order_by(Match.week, Match.external_id)
        .all()
    )

    tried: list[str] = []
    for match in candidates:
        opponent_id = match.away_team_id if match.home_team_id == our_team_id else match.home_team_id
        if not opponent_id or not match.session_name or not match.format:
            continue
        tried.append(f"{opponent_id}/{match.session_name}/{match.format}")
        try:
            our_roster = canonical_current_roster(db, our_team_id, match.session_name)
            opponent_roster = canonical_current_roster(db, opponent_id, match.session_name)
        except CanonicalRosterError:
            continue  # a real identity problem on this candidate; try the next one
        if our_roster and opponent_roster:
            return {
                "our_team_id": our_team_id,
                "opponent_team_id": opponent_id,
                "session_name": match.session_name,
                "format": match.format,
            }

    raise BuildError(
        "no usable live scope: no real match gave both sides a canonical current "
        "roster" + (f" (tried {len(tried)} candidate(s))" if tried else " (no real matches found)"),
        EXIT_DOCUMENTS,
    )


def _formula_versions() -> dict:
    from analytics.opponent_volatility import FORMULA_VERSION as VOLATILITY
    from analytics.team_strength import FORMULA_VERSION as TEAM_STRENGTH
    from analytics.trend_analyzer import FORMULA_VERSION as TREND

    return {
        "team_strength": TEAM_STRENGTH,
        "trend_analyzer": TREND,
        "opponent_volatility": VOLATILITY,
    }


def promote_live_database(staged: Path) -> None:
    """Replace the production database after upstream verification succeeds.

    This is public so thin operator launchers can reuse the exact same backup-and-
    promote contract instead of copying it. Callers remain responsible for
    completing their own verification gates before invoking it.
    """
    backup = LIVE_PRODUCTION_DB.with_suffix(
        f".backup-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}.db"
    )
    if LIVE_PRODUCTION_DB.exists():
        shutil.copy2(LIVE_PRODUCTION_DB, backup)
        logger.info("Backed up the existing production database to %s", backup.name)
    shutil.copy2(staged, LIVE_PRODUCTION_DB)
    logger.info("Promoted %s to %s", staged.name, LIVE_PRODUCTION_DB.name)


def main(argv: Optional[list[str]] = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description="Build one verified demo bundle.")
    parser.add_argument("--mode", choices=("fixture", "live"), required=True)
    parser.add_argument("--out", help="run directory (default: demo-runs/<mode>-<utc timestamp>)")
    parser.add_argument("--run-root", default=str(DEFAULT_RUN_ROOT))
    parser.add_argument(
        "--promote",
        action="store_true",
        help="live mode only: replace data/apa_tracker.db after a successful build",
    )
    parser.add_argument(
        "--resume",
        action="store_true",
        help=(
            "live mode only: resume an interrupted acquisition instead of deleting "
            "the staging database and starting over -- skips scoresheets already "
            "fetched, since a real APA token can expire before a full division-wide "
            "sync finishes"
        ),
    )
    parser.add_argument("--events", help="write versioned redacted JSONL events here")
    parser.add_argument(
        "--source-db",
        help=(
            "live mode only: build from this already-acquired database instead of "
            "a fresh scrape -- never touches the network, never requires a token, "
            "and never rescrapes. For a new verified run over data that was already "
            "ingested and separately repaired in place."
        ),
    )
    parser.add_argument(
        "--opponent-team-id",
        help="live mode only: pin the scope's opponent team instead of auto-selecting one "
             "(must be given together with --session and --format)",
    )
    parser.add_argument(
        "--session",
        help="live mode only: pin the scope's session name (must be given with --opponent-team-id and --format)",
    )
    parser.add_argument(
        "--format",
        dest="format_name",
        help="live mode only: pin the scope's format (must be given with --opponent-team-id and --session)",
    )
    args = parser.parse_args(argv)

    if args.resume and args.mode != "live":
        parser.error("--resume is only meaningful with --mode live")
    if args.source_db and args.mode != "live":
        parser.error("--source-db is only meaningful with --mode live")
    if args.source_db and args.resume:
        parser.error("--source-db and --resume are mutually exclusive: one names an "
                      "existing database, the other resumes a fresh scrape")
    if args.source_db and args.promote:
        parser.error("--source-db and --promote are mutually exclusive: the source "
                      "database is not a freshly-acquired staging file to promote")
    scope_overrides = (args.opponent_team_id, args.session, args.format_name)
    if any(scope_overrides) and args.mode != "live":
        parser.error("--opponent-team-id/--session/--format are only meaningful with --mode live")
    if any(scope_overrides) and not all(scope_overrides):
        parser.error("--opponent-team-id, --session and --format must be given together, or not at all")

    run_root = Path(args.run_root)
    if args.out:
        run_dir = Path(args.out)
    else:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        run_dir = run_root / f"{args.mode}-{stamp}"

    events = EventSink(Path(args.events) if args.events else None)
    try:
        completed = run_build(
            args.mode, run_dir, run_root, promote=args.promote, events=events, resume=args.resume,
            source_db=Path(args.source_db) if args.source_db else None,
            opponent_team_id=args.opponent_team_id,
            session_name=args.session,
            format_name=args.format_name,
        )
    except BuildError as exc:
        logger.error("BUILD FAILED: %s", exc)
        events.emit("failed", "error", code=exc.code)
        return exc.code
    except KeyboardInterrupt:
        logger.error("interrupted")
        return EXIT_INTERRUPT

    logger.info("Run ready: %s", completed)
    logger.info("Open: %s", completed / INDEX_NAME)
    return EXIT_OK


if __name__ == "__main__":
    raise SystemExit(main())
