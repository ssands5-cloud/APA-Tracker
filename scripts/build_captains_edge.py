"""Build the Captain's Edge: a lineup cheat-sheet in two forms.

Writes two self-contained artifacts from the existing SQLite database:

    exports/captains_edge.html   an interactive app, opens with no server
    exports/captains_edge.xlsx   the same data, filterable in Excel

This is a REPORTER, not an engine. Every number it shows -- win rate,
matchup score, confidence, volatility, SL delta -- was computed by
analytics/matchups.py (P0/P1/P2) and persisted to `player_matchups` by the
sync. Nothing here recomputes any of it, so this view can never disagree
with the workbook or the demo about the same pairing. The only thing this
script derives is the Favored/Even/Avoid tag, which buckets the existing
matchup_score for display and never alters it.

The database is opened read-only at the SQLite level (``mode=ro``), so a bug
here cannot write to, migrate, or create the real data file. That is also
why this script does NOT use database.engine.create_db_engine(): that helper
calls ``Base.metadata.create_all()``, which writes.

Usage:
    python scripts/build_captains_edge.py
    python scripts/build_captains_edge.py --db path/to/apa.db
    python scripts/build_captains_edge.py --out-dir exports
"""

from __future__ import annotations

import argparse
import json
import logging
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUT_DIR = PROJECT_ROOT / "exports"
HTML_NAME = "captains_edge.html"
XLSX_NAME = "captains_edge.xlsx"

# Where the database might live, in the order we look. The configured path
# wins; the rest are fallbacks for a checkout whose apa_config.yaml has not
# been customised, and for the `apa_data/` layout used by some setups.
FALLBACK_DB_PATHS = (
    "data/apa_tracker.db",
    "apa_data/apa.db",
    "apa.db",
)

# The 65/35 anchors are the same ones scripts/render_demo_html.py already
# uses for "recommended opponents" and "opponents to be cautious of", so the
# two views cannot tag the same pairing differently.
FAVORED_AT = 65
AVOID_AT = 35

# A pairing with almost no history can post an extreme win rate that means
# nothing. The app still shows these; this is the default for its "hide
# low-sample" filter, not a hard exclusion.
DEFAULT_MIN_SAMPLE = 3


class NoDatabaseError(RuntimeError):
    """No readable SQLite database was found to build from."""


def tag_for(matchup_score: Optional[int]) -> str:
    """Bucket a 0-100 matchup_score into Favored / Even / Avoid.

    Returns "No data" for a missing score rather than guessing a bucket -- a
    pairing the engine could not score is not an "Even" one.
    """
    if matchup_score is None:
        return "No data"
    if matchup_score >= FAVORED_AT:
        return "Favored"
    if matchup_score <= AVOID_AT:
        return "Avoid"
    return "Even"


def _configured_db_path() -> Optional[Path]:
    """database.path from apa_config.yaml, if present and parseable. A
    malformed or absent config is not fatal -- the caller still has the
    fallbacks and --db."""
    config_path = PROJECT_ROOT / "apa_config.yaml"
    if not config_path.is_file():
        return None
    try:
        import yaml  # only needed to read the configured path

        config = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    except Exception:  # pragma: no cover - config is advisory, never required
        logger.debug("Could not read %s; falling back to default paths", config_path)
        return None
    configured = (config.get("database") or {}).get("path")
    return Path(configured) if configured else None


def resolve_db_path(explicit: Optional[str] = None) -> Path:
    """Find the database: an explicit path, then the configured one, then the
    known fallbacks.

    Raises NoDatabaseError naming every location tried -- "database not
    found" with no path in the message is the least actionable error a
    captain could be handed.
    """
    tried: list[Path] = []

    if explicit:
        path = Path(explicit)
        if path.is_file():
            return path
        tried.append(path)

    configured = _configured_db_path()
    if configured is not None:
        path = configured if configured.is_absolute() else PROJECT_ROOT / configured
        if path.is_file():
            return path
        tried.append(path)

    for candidate in FALLBACK_DB_PATHS:
        path = PROJECT_ROOT / candidate
        if path.is_file():
            return path
        tried.append(path)

    listed = "\n".join(f"  - {p}" for p in tried)
    raise NoDatabaseError(
        "No APA database found. Looked in:\n"
        f"{listed}\n\n"
        "Populate one first (for example: python -m scheduler.graphql_sync), "
        "or point this script at an existing file:\n"
        "  python scripts/build_captains_edge.py --db path/to/apa.db"
    )


def connect_read_only(db_path: Path) -> sqlite3.Connection:
    """Open the database read-only, so this tool cannot mutate real data."""
    connection = sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    return connection


def _table_exists(connection: sqlite3.Connection, name: str) -> bool:
    row = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (name,)
    ).fetchone()
    return row is not None


def fetch_matchups(connection: sqlite3.Connection) -> list[dict[str, Any]]:
    """Every scored pairing, straight out of `player_matchups` -- the table
    analytics/matchup_builder writes. No arithmetic beyond the display tag.
    """
    if not _table_exists(connection, "player_matchups"):
        return []
    rows = connection.execute(
        """
        SELECT
            p.external_id   AS player_id,
            p.name          AS player_name,
            o.external_id   AS opponent_id,
            o.name          AS opponent_name,
            m.matches_played,
            m.win_rate,
            m.avg_points_earned,
            m.avg_own_skill_level,
            m.avg_opponent_skill_level,
            m.sl_delta,
            m.trend,
            m.volatility,
            m.matchup_score,
            m.confidence_score,
            m.format,
            m.session_name
        FROM player_matchups m
        JOIN players p ON p.id = m.player_id
        JOIN players o ON o.id = m.opponent_id
        ORDER BY p.name, m.matchup_score DESC
        """
    ).fetchall()

    matchups = []
    for row in rows:
        record = dict(row)
        record["tag"] = tag_for(record["matchup_score"])
        matchups.append(record)
    return matchups


def fetch_players(connection: sqlite3.Connection) -> list[dict[str, Any]]:
    """Roster rows for everyone who appears in at least one scored pairing.

    Deliberately not every row in `players`: the roster carries people with
    no matchup rows at all, and offering them in the selector would promise
    an answer the app cannot give.
    """
    if not _table_exists(connection, "player_matchups"):
        return []
    rows = connection.execute(
        """
        SELECT DISTINCT
            p.external_id AS player_id,
            p.name        AS player_name,
            p.skill_level,
            p.matches_won,
            p.matches_played,
            p.win_pct,
            p.ppm,
            p.pa,
            t.name        AS team_name
        FROM players p
        LEFT JOIN teams t ON t.id = p.team_id
        WHERE p.id IN (SELECT player_id FROM player_matchups)
           OR p.id IN (SELECT opponent_id FROM player_matchups)
        ORDER BY p.name
        """
    ).fetchall()
    return [dict(row) for row in rows]


def fetch_head_to_head(connection: sqlite3.Connection) -> list[dict[str, Any]]:
    """The individual games behind each pairing, newest first.

    `player_matchups` says what the record is; this says which games made it,
    which is what a captain actually argues about. A missing table or no rows
    is normal on a database synced before head-to-head ingestion existed.
    """
    if not _table_exists(connection, "player_head_to_head"):
        return []
    rows = connection.execute(
        """
        SELECT
            p.external_id AS player_id,
            p.name        AS player_name,
            o.external_id AS opponent_id,
            o.name        AS opponent_name,
            h.own_skill_level,
            h.opponent_skill_level,
            h.result,
            h.format,
            h.session_name,
            mt.week,
            mt.match_date,
            mt.home_score,
            mt.away_score
        FROM player_head_to_head h
        JOIN players p ON p.id = h.player_id
        JOIN players o ON o.id = h.opponent_id
        LEFT JOIN matches mt ON mt.id = h.match_id
        ORDER BY mt.week DESC, mt.match_date DESC
        """
    ).fetchall()
    return [dict(row) for row in rows]


def _distinct(matchups: list[dict[str, Any]], key: str) -> list[str]:
    return sorted({m[key] for m in matchups if m.get(key)})


def build_payload(connection: sqlite3.Connection, banner: str = "") -> dict[str, Any]:
    """Everything both artifacts need, in one JSON-serialisable dict."""
    matchups = fetch_matchups(connection)
    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "banner": banner,
        "min_sample": DEFAULT_MIN_SAMPLE,
        "players": fetch_players(connection),
        "matchups": matchups,
        "head_to_head": fetch_head_to_head(connection),
        "formats": _distinct(matchups, "format"),
        "sessions": _distinct(matchups, "session_name"),
    }


def render_html(payload: dict[str, Any]) -> str:
    """Inline the payload into the app shell.

    A literal `</script>` inside the JSON would close the tag early and blank
    the page, so `</` is escaped -- the standard hazard of embedding JSON in
    HTML, worth guarding even though no player name should contain it.
    """
    data = json.dumps(payload, default=str).replace("</", "<\\/")
    return _APP_TEMPLATE.replace("__DATA__", data)


# --- Excel -------------------------------------------------------------------

def write_workbook(payload: dict[str, Any], path: Path) -> Path:
    """Write the three-sheet Excel cheat-sheet.

    Every sheet gets frozen headers and column filters; the Matchups sheet
    additionally colours its Tag column green/amber/red so a captain can scan
    it without reading a single number.
    """
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")

    def add_sheet(title: str, columns: list[str], rows: list[list[Any]]):
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            sheet.append(row)

        # Freeze the header and filter every column, so the sheet is usable
        # as a lookup table rather than a static dump.
        sheet.freeze_panes = "A2"
        last_column = get_column_letter(len(columns))
        sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"

        for index, name in enumerate(columns, start=1):
            width = max(len(str(name)) + 4, 12)
            for row in rows[:200]:
                width = max(width, min(len(str(row[index - 1])) + 2, 42))
            sheet.column_dimensions[get_column_letter(index)].width = width
        return sheet

    # Players
    add_sheet(
        "Players",
        ["Player", "Team", "Skill Level", "Matches Won", "Matches Played",
         "Win %", "PPM", "PA"],
        [
            [p["player_name"], p.get("team_name") or "", p.get("skill_level"),
             p.get("matches_won"), p.get("matches_played"), p.get("win_pct"),
             p.get("ppm"), p.get("pa")]
            for p in payload["players"]
        ],
    )

    # Matchups
    matchup_columns = [
        "Player", "Opponent", "Format", "Session", "Tag", "Matchup Score",
        "Confidence", "Matches Played", "Win Rate", "Avg Points",
        "Avg Own SL", "Avg Opp SL", "SL Delta", "Volatility", "Trend",
    ]
    matchup_rows = [
        [m["player_name"], m["opponent_name"], m.get("format") or "",
         m.get("session_name") or "", m["tag"], m.get("matchup_score"),
         m.get("confidence_score"), m.get("matches_played"), m.get("win_rate"),
         m.get("avg_points_earned"), m.get("avg_own_skill_level"),
         m.get("avg_opponent_skill_level"), m.get("sl_delta"),
         m.get("volatility"), m.get("trend") or ""]
        for m in payload["matchups"]
    ]
    sheet = add_sheet("Matchups", matchup_columns, matchup_rows)

    if matchup_rows:
        last = sheet.max_row
        tag_range = f"E2:E{last}"
        for text, fill_colour, font_colour in (
            ("Favored", "C6EFCE", "006100"),
            ("Even", "FFEB9C", "9C5700"),
            ("Avoid", "FFC7CE", "9C0006"),
        ):
            sheet.conditional_formatting.add(
                tag_range,
                CellIsRule(
                    operator="equal", formula=[f'"{text}"'],
                    fill=PatternFill("solid", fgColor=fill_colour),
                    font=Font(color=font_colour, bold=True),
                ),
            )
        # Win Rate reads as a percentage, not a bare 0.83.
        for row in sheet.iter_rows(min_row=2, min_col=9, max_col=9):
            for cell in row:
                cell.number_format = "0%"

    # H2H History
    add_sheet(
        "H2H History",
        ["Player", "Opponent", "Format", "Session", "Week", "Date", "Result",
         "Own SL", "Opp SL", "Home Score", "Away Score"],
        [
            [g.get("player_name") or "", g.get("opponent_name") or "",
             g.get("format") or "", g.get("session_name") or "", g.get("week"),
             g.get("match_date") or "", g.get("result") or "",
             g.get("own_skill_level"), g.get("opponent_skill_level"),
             g.get("home_score"), g.get("away_score")]
            for g in payload["head_to_head"]
        ],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build(db_path: Optional[str] = None, out_dir: Optional[str] = None,
          banner: str = "") -> tuple[Path, Path]:
    """Read the database and write both artifacts. Returns (html, xlsx)."""
    resolved = resolve_db_path(db_path)
    logger.info("Reading %s", resolved)

    connection = connect_read_only(resolved)
    try:
        payload = build_payload(connection, banner)
    finally:
        connection.close()

    directory = Path(out_dir) if out_dir else DEFAULT_OUT_DIR
    directory.mkdir(parents=True, exist_ok=True)

    html_path = directory / HTML_NAME
    html_path.write_text(render_html(payload), encoding="utf-8")

    xlsx_path = write_workbook(payload, directory / XLSX_NAME)

    logger.info(
        "Wrote %s and %s (%d players, %d pairings, %d games)",
        html_path.name, xlsx_path.name, len(payload["players"]),
        len(payload["matchups"]), len(payload["head_to_head"]),
    )
    return html_path, xlsx_path


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description="Build the Captain's Edge cheat-sheet.")
    parser.add_argument("--db", help="path to the SQLite database")
    parser.add_argument("--out-dir", help=f"output directory (default: {DEFAULT_OUT_DIR})")
    parser.add_argument("--banner", default="", help="notice shown in the app header")
    args = parser.parse_args()

    try:
        html_path, xlsx_path = build(args.db, args.out_dir, args.banner)
    except NoDatabaseError as exc:
        print(f"\n{exc}\n")
        return 1

    print(f"\nOpen in a browser:  {html_path}\nOpen in Excel:      {xlsx_path}\n")
    return 0


_APP_TEMPLATE = r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Captain's Edge</title>
<style>
  :root {
    --bg:#f5f6f8; --panel:#fff; --ink:#15171c; --muted:#666e7a; --line:#e2e5ea;
    --accent:#1d4ed8; --accent-soft:#eef2ff;
    --good:#15803d; --good-bg:#e6f6ec; --bad:#b91c1c; --bad-bg:#fdecec;
    --even:#92400e; --even-bg:#fdf3e3;
  }
  @media (prefers-color-scheme: dark) {
    :root {
      --bg:#131519; --panel:#1b1e24; --ink:#e9ecf1; --muted:#98a1ae; --line:#2b3039;
      --accent:#7ea2ff; --accent-soft:#222a3d;
      --good:#6ee7a0; --good-bg:#153020; --bad:#ff9b9b; --bad-bg:#331a1a;
      --even:#f0c580; --even-bg:#332715;
    }
  }
  * { box-sizing:border-box; }
  body { margin:0; background:var(--bg); color:var(--ink);
         font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif; }
  header { padding:16px 22px; border-bottom:1px solid var(--line); background:var(--panel); }
  h1 { margin:0; font-size:20px; letter-spacing:-.01em; }
  .sub { color:var(--muted); font-size:13px; margin-top:3px; }
  .banner { margin-top:10px; padding:8px 11px; border-radius:7px; background:var(--even-bg);
            color:var(--even); font-size:13px; border:1px solid var(--line); }
  .layout { display:grid; grid-template-columns:280px 1fr; gap:18px; padding:18px 22px;
            align-items:start; }
  @media (max-width:960px){ .layout{ grid-template-columns:1fr; } }
  .panel { background:var(--panel); border:1px solid var(--line); border-radius:10px; padding:15px; }
  .panel + .panel { margin-top:16px; }
  label { display:block; font-size:11.5px; font-weight:600; text-transform:uppercase;
          letter-spacing:.04em; color:var(--muted); margin:14px 0 5px; }
  label:first-child { margin-top:0; }
  select, input { width:100%; padding:8px 9px; font:inherit; color:var(--ink);
                  background:var(--bg); border:1px solid var(--line); border-radius:7px; }
  .check { display:flex; align-items:center; gap:8px; margin-top:10px; font-size:13.5px;
           color:var(--ink); }
  .check input { width:auto; }
  h2 { font-size:16px; margin:0 0 4px; }
  .h2note { color:var(--muted); font-size:12.5px; margin-bottom:12px; }
  .table-wrap { overflow-x:auto; }
  table { border-collapse:collapse; width:100%; font-size:13.5px; }
  th,td { padding:8px 9px; text-align:left; border-bottom:1px solid var(--line); white-space:nowrap; }
  th { font-size:11.5px; text-transform:uppercase; letter-spacing:.04em; color:var(--muted);
       cursor:pointer; user-select:none; }
  th:hover { color:var(--accent); }
  th.sorted::after { content:" \25BC"; font-size:9px; }
  th.sorted.asc::after { content:" \25B2"; }
  td.num, th.num { text-align:right; font-variant-numeric:tabular-nums; }
  tbody tr { cursor:pointer; }
  tbody tr:hover { background:var(--accent-soft); }
  tbody tr.active { background:var(--accent-soft); }
  .tag { display:inline-block; padding:2px 10px; border-radius:999px; font-size:12px; font-weight:650; }
  .tag.good{background:var(--good-bg);color:var(--good);}
  .tag.bad{background:var(--bad-bg);color:var(--bad);}
  .tag.even{background:var(--even-bg);color:var(--even);}
  .tag.none{background:var(--bg);color:var(--muted);}
  .rec td { padding-top:10px; padding-bottom:10px; }
  .rec .score { font-size:16px; font-weight:700; }
  .empty { color:var(--muted); padding:20px 4px; text-align:center; }
  .tiles { display:grid; grid-template-columns:repeat(auto-fit,minmax(120px,1fr)); gap:10px;
           margin:12px 0; }
  .tile { background:var(--bg); border:1px solid var(--line); border-radius:8px; padding:9px 11px; }
  .tile .lbl { font-size:10.5px; text-transform:uppercase; letter-spacing:.04em; color:var(--muted); }
  .tile .val { font-size:17px; font-weight:650; margin-top:2px; }
  .log-row { display:flex; gap:12px; align-items:center; font-size:13px; padding:6px 0;
             border-bottom:1px solid var(--line); }
  .w{color:var(--good);font-weight:700;} .l{color:var(--bad);font-weight:700;}
  .foot { padding:0 22px 26px; color:var(--muted); font-size:12px; }
  .close { float:right; cursor:pointer; color:var(--muted); font-size:13px; }
  .close:hover { color:var(--accent); }
</style>
</head>
<body>
<header>
  <h1>Captain's Edge</h1>
  <div class="sub" id="subtitle"></div>
  <div class="banner" id="banner" hidden></div>
</header>

<div class="layout">
  <aside>
    <div class="panel">
      <label for="playerA">Player</label>
      <input id="filter" type="search" placeholder="Search players…" style="margin-bottom:6px">
      <select id="playerA"></select>

      <label for="format">Format</label>
      <select id="format"></select>

      <label for="session">Session</label>
      <select id="session"></select>

      <label for="tagFilter">Show</label>
      <select id="tagFilter">
        <option value="__all__">Everyone</option>
        <option value="Favored">Favored only</option>
        <option value="Avoid">Avoid only</option>
      </select>

      <div class="check">
        <input type="checkbox" id="hideLow">
        <label for="hideLow" style="margin:0;text-transform:none;font-size:13.5px;font-weight:400;letter-spacing:0;color:var(--ink);">
          Hide low-sample (&lt; <span id="minSample"></span> games)
        </label>
      </div>
    </div>
  </aside>

  <main>
    <div class="panel">
      <h2>Recommended opponents</h2>
      <div class="h2note" id="rec-note"></div>
      <div class="table-wrap"><table class="rec">
        <thead><tr>
          <th>Opponent</th><th>Tag</th><th class="num">Score</th><th class="num">Confidence</th>
          <th class="num">Games</th><th class="num">Win rate</th>
          <th class="num">SL delta</th><th class="num">Volatility</th>
        </tr></thead>
        <tbody id="rec-body"></tbody>
      </table></div>
      <div class="empty" id="rec-empty" hidden></div>
    </div>

    <div class="panel" id="detail-panel" hidden></div>

    <div class="panel">
      <h2>All opponents</h2>
      <div class="h2note">Click any row for the game-by-game history. Click a column to sort.</div>
      <div class="table-wrap"><table>
        <thead><tr>
          <th data-k="opponent_name">Opponent</th>
          <th data-k="tag">Tag</th>
          <th class="num" data-k="matchup_score">Score</th>
          <th class="num" data-k="confidence_score">Confidence</th>
          <th class="num" data-k="matches_played">Games</th>
          <th class="num" data-k="win_rate">Win rate</th>
          <th class="num" data-k="avg_own_skill_level">Avg own SL</th>
          <th class="num" data-k="avg_opponent_skill_level">Avg opp SL</th>
          <th class="num" data-k="sl_delta">SL delta</th>
          <th class="num" data-k="volatility">Volatility</th>
          <th data-k="trend">Trend</th>
          <th data-k="format">Format</th>
          <th data-k="session_name">Session</th>
        </tr></thead>
        <tbody id="all-body"></tbody>
      </table></div>
      <div class="empty" id="all-empty" hidden></div>
    </div>
  </main>
</div>

<div class="foot" id="foot"></div>

<script id="payload" type="application/json">__DATA__</script>
<script>
(function () {
  "use strict";
  var DATA = JSON.parse(document.getElementById("payload").textContent);
  var ALL = "__all__";
  var selected = null;

  var el = function (id) { return document.getElementById(id); };
  var esc = function (s) {
    return String(s === null || s === undefined ? "" : s).replace(/[&<>"']/g, function (c) {
      return { "&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;" }[c];
    });
  };
  var pct = function (v){ return v==null ? "—" : Math.round(v*100)+"%"; };
  var num = function (v,d){ if(v==null) return "—"; return d?Number(v).toFixed(d):String(v); };
  var signed = function (v){ if(v==null) return "—"; return (v>0?"+":"")+Number(v).toFixed(2); };
  // Volatility is a raw count on older databases and a normalized rate after
  // P2. Render whichever is stored rather than assuming a scale.
  var vol = function (v){ if(v==null) return "—"; return Number.isInteger(v)?String(v):Number(v).toFixed(2); };
  var tagClass = function (t){
    return t==="Favored" ? "good" : t==="Avoid" ? "bad" : t==="No data" ? "none" : "even";
  };
  var tagCell = function (t){ return '<span class="tag '+tagClass(t)+'">'+esc(t)+"</span>"; };

  function fill(select, options, allLabel) {
    select.innerHTML = "";
    if (allLabel) {
      var o = document.createElement("option");
      o.value = ALL; o.textContent = allLabel; select.appendChild(o);
    }
    options.forEach(function (opt) {
      var o = document.createElement("option");
      o.value = opt.value; o.textContent = opt.label; select.appendChild(o);
    });
  }

  function scoped() {
    var f = el("format").value, s = el("session").value;
    return DATA.matchups.filter(function (m) {
      return (f===ALL || m.format===f) && (s===ALL || m.session_name===s);
    });
  }

  function rowsForPlayer() {
    var a = el("playerA").value;
    var rows = scoped().filter(function (m){ return m.player_id === a; });
    var tag = el("tagFilter").value;
    if (tag !== ALL) rows = rows.filter(function (m){ return m.tag === tag; });
    if (el("hideLow").checked) {
      rows = rows.filter(function (m){ return (m.matches_played||0) >= DATA.min_sample; });
    }
    return rows;
  }

  function renderRecommended(rows) {
    var best = rows.slice()
      .filter(function (m){ return m.matchup_score != null; })
      .sort(function (a,b){ return b.matchup_score - a.matchup_score; })
      .slice(0, 5);
    var body = el("rec-body"), empty = el("rec-empty");
    body.innerHTML = best.map(function (m) {
      return '<tr data-opp="'+esc(m.opponent_id)+'">' +
        "<td><strong>"+esc(m.opponent_name)+"</strong></td>" +
        "<td>"+tagCell(m.tag)+"</td>" +
        '<td class="num score">'+num(m.matchup_score)+"</td>" +
        '<td class="num">'+num(m.confidence_score)+"</td>" +
        '<td class="num">'+num(m.matches_played)+"</td>" +
        '<td class="num">'+pct(m.win_rate)+"</td>" +
        '<td class="num">'+signed(m.sl_delta)+"</td>" +
        '<td class="num">'+vol(m.volatility)+"</td></tr>";
    }).join("");
    empty.hidden = best.length > 0;
    if (!best.length) empty.textContent = "No opponents match these filters.";
    el("rec-note").textContent = best.length
      ? "Best " + best.length + " of " + rows.length + " by matchup score. Click a row for history."
      : "";
  }

  var sortKey = "matchup_score", sortAsc = false;

  function renderAll(rows) {
    var sorted = rows.slice().sort(function (a,b) {
      var x=a[sortKey], y=b[sortKey];
      if (x==null) return 1;
      if (y==null) return -1;
      if (typeof x === "string") return sortAsc ? x.localeCompare(y) : y.localeCompare(x);
      return sortAsc ? x-y : y-x;
    });
    el("all-body").innerHTML = sorted.map(function (m) {
      return '<tr data-opp="'+esc(m.opponent_id)+'"'+(selected===m.opponent_id?' class="active"':'')+">" +
        "<td>"+esc(m.opponent_name)+"</td>" +
        "<td>"+tagCell(m.tag)+"</td>" +
        '<td class="num">'+num(m.matchup_score)+"</td>" +
        '<td class="num">'+num(m.confidence_score)+"</td>" +
        '<td class="num">'+num(m.matches_played)+"</td>" +
        '<td class="num">'+pct(m.win_rate)+"</td>" +
        '<td class="num">'+num(m.avg_own_skill_level,2)+"</td>" +
        '<td class="num">'+num(m.avg_opponent_skill_level,2)+"</td>" +
        '<td class="num">'+signed(m.sl_delta)+"</td>" +
        '<td class="num">'+vol(m.volatility)+"</td>" +
        "<td>"+esc(m.trend||"—")+"</td>" +
        "<td>"+esc(m.format||"—")+"</td>" +
        "<td>"+esc(m.session_name||"—")+"</td></tr>";
    }).join("");
    el("all-empty").hidden = sorted.length > 0;
    if (!sorted.length) el("all-empty").textContent = "No opponents match these filters.";
    document.querySelectorAll("th[data-k]").forEach(function (th) {
      th.classList.toggle("sorted", th.dataset.k === sortKey);
      th.classList.toggle("asc", th.dataset.k === sortKey && sortAsc);
    });
  }

  function renderDetail(opponentId) {
    var panel = el("detail-panel");
    var row = rowsForPlayer().find(function (m){ return m.opponent_id === opponentId; });
    if (!row) { panel.hidden = true; return; }

    var games = DATA.head_to_head.filter(function (g) {
      return g.player_id === row.player_id && g.opponent_id === row.opponent_id &&
             (row.format ? g.format === row.format : true) &&
             (row.session_name ? g.session_name === row.session_name : true);
    });

    var log = games.length ? games.map(function (g) {
      var r = String(g.result||"").toUpperCase();
      var mark = r==="W" ? '<span class="w">W</span>' : r==="L" ? '<span class="l">L</span>' : esc(r||"?");
      var score = (g.home_score!=null && g.away_score!=null)
        ? g.home_score+"–"+g.away_score : "";
      return '<div class="log-row">'+mark+
        "<span>Week "+esc(g.week==null?"—":g.week)+"</span>"+
        '<span style="color:var(--muted)">'+esc(g.match_date||"")+"</span>"+
        '<span style="color:var(--muted)">SL '+esc(g.own_skill_level)+" vs "+esc(g.opponent_skill_level)+"</span>"+
        (score?'<span style="color:var(--muted)">'+esc(score)+"</span>":"")+
        "</div>";
    }).join("") : '<div class="empty">No individual games recorded for this pairing.</div>';

    panel.hidden = false;
    panel.innerHTML =
      '<span class="close" id="close-detail">close ✕</span>' +
      "<h2>"+esc(row.player_name)+" vs "+esc(row.opponent_name)+" "+tagCell(row.tag)+"</h2>" +
      '<div class="h2note">'+esc(row.format||"")+" · "+esc(row.session_name||"")+"</div>" +
      '<div class="tiles">' +
        tile("Matchup score", num(row.matchup_score)) +
        tile("Confidence", num(row.confidence_score)) +
        tile("Win rate", pct(row.win_rate)) +
        tile("Games", num(row.matches_played)) +
        tile("Avg own SL", num(row.avg_own_skill_level,2)) +
        tile("Avg opp SL", num(row.avg_opponent_skill_level,2)) +
        tile("SL delta", signed(row.sl_delta)) +
        tile("Volatility", vol(row.volatility)) +
      "</div>" + log;
    el("close-detail").addEventListener("click", function () {
      selected = null; panel.hidden = true; render();
    });
  }

  function tile(label, value) {
    return '<div class="tile"><div class="lbl">'+esc(label)+'</div><div class="val">'+value+"</div></div>";
  }

  function render() {
    var rows = rowsForPlayer();
    renderRecommended(rows);
    renderAll(rows);
    if (selected) renderDetail(selected); else el("detail-panel").hidden = true;
  }

  function refreshPlayers() {
    var rows = scoped();
    var needle = el("filter").value.trim().toLowerCase();
    var seen = {};
    rows.forEach(function (m){ seen[m.player_id] = m.player_name; });
    var people = Object.keys(seen).map(function (id){ return { value:id, label:seen[id] }; })
      .filter(function (p){ return !needle || p.label.toLowerCase().indexOf(needle) >= 0; })
      .sort(function (a,b){ return a.label.localeCompare(b.label); });

    var prev = el("playerA").value;
    fill(el("playerA"), people, null);
    if (people.some(function (p){ return p.value === prev; })) el("playerA").value = prev;
    selected = null;
    render();
  }

  document.addEventListener("click", function (event) {
    var tr = event.target.closest ? event.target.closest("tr[data-opp]") : null;
    if (!tr) return;
    selected = tr.dataset.opp;
    render();
    el("detail-panel").scrollIntoView({ behavior:"smooth", block:"nearest" });
  });

  fill(el("format"), DATA.formats.map(function (f){ return {value:f,label:f}; }), "All formats");
  fill(el("session"), DATA.sessions.map(function (s){ return {value:s,label:s}; }), "All sessions");
  el("minSample").textContent = DATA.min_sample;

  ["format","session","tagFilter"].forEach(function (id) {
    el(id).addEventListener("change", id === "tagFilter" ? render : refreshPlayers);
  });
  el("filter").addEventListener("input", refreshPlayers);
  el("playerA").addEventListener("change", function (){ selected = null; render(); });
  el("hideLow").addEventListener("change", render);
  document.querySelectorAll("th[data-k]").forEach(function (th) {
    th.addEventListener("click", function () {
      var k = th.dataset.k;
      if (k === sortKey) sortAsc = !sortAsc; else { sortKey = k; sortAsc = false; }
      render();
    });
  });

  if (DATA.banner) { el("banner").hidden = false; el("banner").textContent = DATA.banner; }
  el("subtitle").textContent = DATA.matchups.length
    ? DATA.players.length+" players · "+DATA.matchups.length+" scored pairings · "+
      DATA.head_to_head.length+" games"
    : "No scored pairings in this database yet.";
  el("foot").textContent = "Generated "+DATA.generated_at+
    " · read-only view of player_matchups · every score comes from the matchup engine, not from this page.";

  if (!DATA.matchups.length) {
    el("rec-empty").hidden = false;
    el("rec-empty").textContent = "This database has no scored pairings yet. Run a sync that builds matchups, then rebuild.";
    el("all-empty").hidden = false;
    el("all-empty").textContent = "Nothing to show yet.";
  } else {
    refreshPlayers();
  }
})();
</script>
</body>
</html>
"""


if __name__ == "__main__":
    raise SystemExit(main())
